"""Workbook I/O: read the vendor inventory, write the student cells, prove nothing else changed.

Optiv's input (``data/input/Meridian_Vendor_Input.xlsx``) has a "Vendor Inventory" sheet with headers in row 4,
provided (white / blue-grey) columns B-K, student (cream) columns L-V, the fictional worked example V-000 in
row 5 and the real vendors below it. Design: docs/design.md "Traceability & workbook I/O" and "2.11".

- ``read_workbook`` finds the sheet, header row and columns (alias table + rapidfuzz), reads the vendor profiles
  and reports problems as ``ValidationIssue``s instead of raising.
- ``write_workbook`` writes ``StudentCells`` into a copy and appends extra sheets. It checks every rule first and
  raises ``WorkbookWriteError`` (saving nothing) rather than bend one.
- ``check_fidelity`` is the semantic diff behind the fidelity gate.

Writer rules: only cream (FFFFF6E0) student cells of real vendor rows, never the example row or a provided cell;
values are stored as plain strings, so "=..." text never becomes a formula; ``column_dimensions`` is never
indexed by letter, because L-O and R-T are single ``<col>`` spans stored under L and R, and touching M-O / S-T
adds an overlapping ``<col>`` that Excel "repairs" on open. Workbooks are opened with rich text kept: the
example's evidence cell P5 has bold runs that a plain load flattens.

Issue codes (``ValidationIssue.code``):
  E01 no vendor sheet, or not an .xlsx file        E02 no "Vendor ID" header in rows 1-10
  E03 header not found (names the closest match)   E04 duplicate vendor ID
  E05 no vendor rows                               E06 student cell already filled (an error when writing)
  E07 vendor has no website (warning)              E08 unknown operational dependency label (warning)
  E09 write target is not a cream student cell     E10 value not allowed, or over its length budget
  E11 characters an .xlsx cannot store removed: control characters, lone surrogates, U+FFFE/U+FFFF (warning)
  E12 vendor ID is not a writable row (unknown, or the worked example)
  E13 output refused: sheet title taken or invalid, or destination is the source
"""

from __future__ import annotations

import datetime as dt
import io
import math
import os
import re
import tempfile
import unicodedata
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell, MergedCell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field
from rapidfuzz import fuzz, process

from footprint.models import STUDENT_FIELDS, SheetSpec, StudentCells, ValidationIssue, VendorProfile, WorkbookData

WorkbookSource = str | Path | bytes | BinaryIO
"""A workbook as a path, its raw bytes, or a seekable binary stream (e.g. a Streamlit upload)."""

CREAM = "FFFFF6E0"
"""ARGB fill of the student cells. The writer only writes cells with this fill."""

SHEET_NAME = "Vendor Inventory"
HEADER_SCAN_ROWS = 10  # the "Vendor ID" header must be in rows 1-10
FUZZY_MIN_SCORE = 85  # rapidfuzz token_sort_ratio needed to accept a header variant
FUZZY_MARGIN = 5  # ... and by how much it must beat the header's next-best field

PROVIDED_FIELDS: list[str] = [
    "vendor_id", "name", "description", "service", "category",
    "website", "business_process", "operational_dependency", "data_accessed", "data_volume",
]
"""VendorProfile fields read from the provided columns B-K."""

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "vendor_id": ("Vendor ID",),
    "name": ("Vendor Name",),
    "description": ("Vendor Description",),
    "service": ("Service / Product Provided", "Service / Product Provided to Meridian"),
    "category": ("Category",),
    "website": ("Public Website",),
    "business_process": ("Business Process Supported",),
    "operational_dependency": ("Operational Dependency",),
    "data_accessed": ("Data Classification Accessed",),
    "data_volume": ("Data Volume (annual)",),
    "criticality_tier": ("Criticality Tier",),
    "criticality_rationale": ("Criticality Rationale",),
    "assessment_depth": ("Assessment Depth Applied",),
    "ai_usage_detected": ("AI Usage Detected (Y/N)", "AI Usage Detected"),
    "evidence": ("Evidence Description along with OSINT Source(s) / Evidence Link", "Evidence Excerpt"),
    "how_ai_used": ("How Vendor May Be Using AI",),
    "ai_subprocessors": ("Named AI Sub-processors / Model Providers",),
    "ai_risk_class": ("AI Security Risk Classification",),
    "risk_rationale": ("Risk Rationale",),
    "recommended_action": ("Recommended Action",),
    "assessed_by": ("Assessed By / Date", "Assessedd By / Date"),
}
"""Field -> accepted header texts, in column order B-V. The first alias is the name used in messages."""

DEPENDENCY_LABELS: tuple[str, ...] = ("Total", "Critical", "High", "Moderate", "Medium", "Low")
"""Operational dependency labels the criticality rubric understands (compared case-insensitively)."""

ALLOWED_VALUES: dict[str, tuple[str, ...]] = {
    "criticality_tier": ("Critical", "High", "Medium", "Low"),
    "ai_usage_detected": ("Yes", "No", "Inconclusive"),
    "ai_risk_class": ("Critical", "High", "Medium", "Low", "None identified"),
}
"""Columns L, O and S take exactly one of these values."""

DEFAULT_LENGTH_BUDGETS: dict[str, int] = {
    **{field: 1050 for field in STUDENT_FIELDS},
    "evidence": 1390,
    "how_ai_used": 1670,
    "recommended_action": 1480,
}
"""Characters per student cell: what fits the column width at Excel's 409 pt row-height limit."""

MAX_ROW_HEIGHT = 409.0  # Excel's limit, in points
LINE_HEIGHT = 12.75  # one wrapped line of Arial 10, in points
ROW_PADDING = 8.0  # cell margins, in points

_NOTE_FONT = Font(name="Arial", size=10, italic=True)
_HEADER_FONT = Font(name="Arial", size=10, bold=True)
_BODY_FONT = Font(name="Arial", size=10)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor=CREAM)  # cream: blue-grey would mean "provided by Optiv"
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_BAD_TITLE_CHARS = re.compile(r"[\\/?*\[\]:]")
_RESERVED_TITLES = frozenset({"history"})  # Excel reserves "History" (case-insensitive)
_UNSTORABLE = re.compile(r"[\ud800-\udfff\ufffe\uffff]")  # not valid in XML; ILLEGAL_CHARACTERS_RE misses them
_LITERAL_ESCAPE = re.compile(r"_(x[0-9A-Fa-f]{4}_)")  # Excel decodes '_x0041_' as 'A' unless '_' is escaped
_STYLE_PARTS = ("number_format", "font", "fill", "border", "alignment", "protection")


class WriteReport(BaseModel):
    """What write_workbook changed."""

    written: list[str] = Field(default_factory=list, description="student cells written, e.g. 'L6'")
    skipped: list[str] = Field(default_factory=list, description="student cells of vendor rows left untouched")
    warnings: list[ValidationIssue] = Field(default_factory=list)
    row_heights: dict[int, float] = Field(default_factory=dict, description="raised rows -> new height in points")
    sheets_added: list[str] = Field(default_factory=list)


class WorkbookWriteError(ValueError):
    """A write would break a workbook rule. Lists every problem found; nothing was saved."""

    def __init__(self, issues: Sequence[ValidationIssue]) -> None:
        self.issues = list(issues)
        super().__init__("; ".join(f"{issue.code} {issue.message}" for issue in self.issues))

    @property
    def code(self) -> str:
        """Code of the first problem, e.g. 'E09'."""
        return self.issues[0].code


# --------------------------------------------------------------------------- reading


def read_workbook(source: WorkbookSource) -> WorkbookData:
    """Read the vendor profiles. Problems with the workbook are reported in ``WorkbookData.issues`` (E01-E08)
    rather than raised; only a missing file raises (FileNotFoundError)."""
    try:
        wb = _load(source)
    except (BadZipFile, InvalidFileException, KeyError) as exc:
        return _unreadable(ValidationIssue(code="E01", message=f"Not a readable .xlsx workbook ({exc})."))
    ws = _find_sheet(wb)
    if ws is None:
        message = f'No "{SHEET_NAME}" sheet, and no sheet with a "Vendor ID" header in rows 1-{HEADER_SCAN_ROWS}.'
        return _unreadable(ValidationIssue(code="E01", message=message))
    header_row = _find_header_row(ws)
    if header_row is None:
        message = f'Sheet "{ws.title}" has no "Vendor ID" header in rows 1-{HEADER_SCAN_ROWS}.'
        return _unreadable(ValidationIssue(code="E02", message=message), ws.title)
    column_map, header_issues = _map_headers(ws, header_row)
    vendors, example, row_issues = _read_rows(ws, header_row, column_map)
    return WorkbookData(
        sheet_name=ws.title,
        header_row=header_row,
        column_map=column_map,
        vendors=vendors,
        example=example,
        issues=header_issues + row_issues,
    )


def _load(source: WorkbookSource) -> Workbook:
    """Open a workbook with rich text kept, so that saving it leaves the provided cells unchanged."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    return load_workbook(source, rich_text=True)


def _unreadable(issue: ValidationIssue, sheet_name: str = "") -> WorkbookData:
    return WorkbookData(sheet_name=sheet_name, header_row=0, column_map={}, vendors=[], issues=[issue])


def _find_sheet(wb: Workbook) -> Worksheet | None:
    """The "Vendor Inventory" sheet, else the first sheet with a "Vendor ID" header in rows 1-10."""
    for ws in wb.worksheets:
        if _compact(ws.title) == _compact(SHEET_NAME):
            return ws
    return next((ws for ws in wb.worksheets if _find_header_row(ws) is not None), None)


def _find_header_row(ws: Worksheet) -> int | None:
    """Row (1-10) of the first cell that reads "Vendor ID"."""
    wanted = {_compact(alias) for alias in HEADER_ALIASES["vendor_id"]}
    for row in ws.iter_rows(min_row=1, max_row=min(HEADER_SCAN_ROWS, ws.max_row)):
        for cell in row:
            if _compact(_text(cell.value)) in wanted:
                return cell.row
    return None


def _map_headers(ws: Worksheet, header_row: int) -> tuple[dict[str, str], list[ValidationIssue]]:
    """Assign header cells to fields, best match first: exact aliases, then rapidfuzz scores >= 85.

    Each field takes at most one column and each column at most one field. A field left without a column is
    an E03 error naming the closest unassigned header.
    """
    headers = {get_column_letter(cell.column): _text(cell.value) for cell in ws[header_row] if _text(cell.value)}
    candidates = []
    for letter, text in headers.items():
        matches = {field: _header_match(field, text) for field in HEADER_ALIASES}
        for order, (field, (exact, score)) in enumerate(matches.items()):
            runner_up = max((s for f, (_, s) in matches.items() if f != field), default=0.0)
            if exact or (score >= FUZZY_MIN_SCORE and score - runner_up >= FUZZY_MARGIN):
                candidates.append((not exact, -score, column_index_from_string(letter), order, field, letter))
    assigned: dict[str, str] = {}
    for *_, field, letter in sorted(candidates):
        if field not in assigned and letter not in assigned.values():
            assigned[field] = letter
    column_map = {field: assigned[field] for field in HEADER_ALIASES if field in assigned}
    unassigned = {letter: text for letter, text in headers.items() if letter not in assigned.values()}
    missing = [field for field in HEADER_ALIASES if field not in assigned]
    return column_map, [_missing_header(ws.title, header_row, field, unassigned) for field in missing]


def _header_match(field: str, text: str) -> tuple[bool, float]:
    """(exact, score) of a header against a field's aliases. Exact ignores case, spacing and punctuation."""
    aliases = HEADER_ALIASES[field]
    if _compact(text) in {_compact(alias) for alias in aliases}:
        return True, 100.0
    # token_sort_ratio, not token_set_ratio: a subset ("Risk", "Criticality") must not score 100 against a field.
    best = process.extractOne(_norm(text), [_norm(alias) for alias in aliases], scorer=fuzz.token_sort_ratio)
    return False, best[1] if best else 0.0


def _missing_header(sheet: str, header_row: int, field: str, unassigned: Mapping[str, str]) -> ValidationIssue:
    wanted = HEADER_ALIASES[field][0]
    where = f'Header "{wanted}" not found in row {header_row} of "{sheet}"'
    choices = {letter: _norm(text) for letter, text in unassigned.items()}
    best = process.extractOne(_norm(wanted), choices, scorer=fuzz.token_set_ratio)
    if best is None:
        return ValidationIssue(code="E03", message=f"{where} (closest match: none; every header there is in use).")
    _, score, letter = best
    hint = f'closest match: "{unassigned[letter]}" in {letter}{header_row}, {score:.0f}% similar'
    return ValidationIssue(code="E03", cell=f"{letter}{header_row}", message=f"{where} ({hint}).")


def _read_rows(
    ws: Worksheet, header_row: int, column_map: Mapping[str, str]
) -> tuple[list[VendorProfile], VendorProfile | None, list[ValidationIssue]]:
    """Profiles from every row below the header with a Vendor ID; the example is set aside.

    Blank rows are skipped, and an E05 warning names them when vendors follow; a row with data but no Vendor ID
    is an E05 warning too, so no vendor is dropped silently.
    """
    vendors: list[VendorProfile] = []
    example: VendorProfile | None = None
    issues: list[ValidationIssue] = []
    first_row_of: dict[str, int] = {}
    gap: list[int] = []
    for row in range(header_row + 1, ws.max_row + 1):
        values = {
            field: _text(ws[f"{column_map[field]}{row}"].value) if field in column_map else ""
            for field in PROVIDED_FIELDS
        }
        if not values["vendor_id"]:
            if any(values.values()):
                message = f"Row {row} has profile data but no Vendor ID, so it was not assessed."
                cell = f"{column_map.get('vendor_id', 'B')}{row}"
                issues.append(ValidationIssue(code="E05", severity="warning", cell=cell, message=message))
            else:
                gap.append(row)
            continue
        if gap and (vendors or example):
            rows = f"{gap[0]}" if len(gap) == 1 else f"{gap[0]}-{gap[-1]}"
            message = f"Blank row(s) {rows} skipped; vendors below them were still read (from row {row})."
            issues.append(ValidationIssue(code="E05", severity="warning", message=message))
        gap = []
        profile = VendorProfile(row=row, **values, is_example=_is_example(values["vendor_id"], values["name"]))
        if profile.is_example:
            example = example or profile
            continue
        key = profile.vendor_id.casefold()
        if key in first_row_of:
            message = f'Vendor ID "{profile.vendor_id}" is used in rows {first_row_of[key]} and {row}.'
            issues.append(ValidationIssue(code="E04", cell=f"{column_map['vendor_id']}{row}", message=message))
        first_row_of.setdefault(key, row)
        vendors.append(profile)
        issues += _vendor_warnings(ws, profile, column_map)
    if not vendors:
        message = f'No vendor rows below the header in row {header_row} of "{ws.title}".'
        issues.append(ValidationIssue(code="E05", message=message))
    return vendors, example, issues


def _vendor_warnings(ws: Worksheet, vendor: VendorProfile, column_map: Mapping[str, str]) -> list[ValidationIssue]:
    """E06 student cells already filled, E07 no website, E08 unknown dependency label (all warnings)."""
    warnings = []
    filled = [
        f"{column_map[field]}{vendor.row}"
        for field in STUDENT_FIELDS
        if field in column_map and _text(ws[f"{column_map[field]}{vendor.row}"].value)
    ]
    if filled:
        message = f"{vendor.vendor_id}: {', '.join(filled)} already have values; writing them again needs overwrite."
        warnings.append(ValidationIssue(code="E06", severity="warning", cell=filled[0], message=message))
    if "website" in column_map and not vendor.website:
        message = f"{vendor.vendor_id} has no public website: criticality can still be rated, OSINT cannot start."
        cell = f"{column_map['website']}{vendor.row}"
        warnings.append(ValidationIssue(code="E07", severity="warning", cell=cell, message=message))
    known = {label.casefold() for label in DEPENDENCY_LABELS}
    if "operational_dependency" in column_map and vendor.operational_dependency.casefold() not in known:
        message = (
            f'{vendor.vendor_id}: operational dependency "{vendor.operational_dependency}" '
            f"is not one of {', '.join(DEPENDENCY_LABELS)}."
        )
        cell = f"{column_map['operational_dependency']}{vendor.row}"
        warnings.append(ValidationIssue(code="E08", severity="warning", cell=cell, message=message))
    return warnings


def _is_example(vendor_id: str, name: str) -> bool:
    """The fictional worked example: ID V-000, or a name marked fictional / do not edit / EXAMPLE..."""
    lowered = name.casefold()
    marked = "fictional" in lowered or "do not edit" in lowered or lowered.startswith("example")
    return vendor_id.upper() == "V-000" or marked


def _text(value: object) -> str:
    """A cell value as trimmed text ('' when empty). Otherwise verbatim: quotes, dashes and case are kept."""
    return "" if value is None else str(value).strip()


def _norm(text: str) -> str:
    """Header text as lower-case words: 'AI Usage Detected (Y/N)' -> 'ai usage detected y n'."""
    return " ".join(re.sub(r"[\W_]+", " ", unicodedata.normalize("NFKC", text).casefold()).split())


def _compact(text: str) -> str:
    """_norm without spaces, so that 'Public Web Site' equals 'Public Website'."""
    return _norm(text).replace(" ", "")


# --------------------------------------------------------------------------- writing


def write_workbook(
    source: WorkbookSource,
    destination: str | Path | BinaryIO,
    cells: Mapping[str, StudentCells],
    sheets: Sequence[SheetSpec] = (),
    *,
    overwrite: bool = False,
    last_modified_by: str | None = None,
    length_budgets: Mapping[str, int] | None = None,
    modified: dt.datetime | None = None,
) -> WriteReport:
    """Write student cells into a copy of ``source``, append ``sheets``, and save it to ``destination``.

    ``cells`` maps vendor ID -> values; None leaves a cell untouched and "" clears it. Every rule is checked
    before anything is written: all violations are raised together as WorkbookWriteError and nothing is saved.
    Rows that receive text are made taller to fit it, never shorter. ``modified`` (or SOURCE_DATE_EPOCH) pins
    the file's modified time, so the same inputs give byte-identical output.
    """
    _refuse_same_file(source, destination)
    budgets = _length_budgets(length_budgets)
    data = read_workbook(source)
    if not data.ok:
        raise WorkbookWriteError([issue for issue in data.issues if issue.severity == "error"])

    wb = _load(source)  # a fresh copy: reading may have created empty cells while scanning
    ws = wb[data.sheet_name]
    plan, problems, warnings = _plan_cells(ws, data, cells, budgets, overwrite)
    problems += _sheet_title_problems(wb, sheets)
    if problems:
        raise WorkbookWriteError(problems)

    for coordinate, text in plan.items():
        _put_text(ws[coordinate], text)
    student_columns = [data.column_map[field] for field in STUDENT_FIELDS]
    row_heights = _raise_row_heights(ws, {ws[coordinate].row for coordinate in plan}, student_columns)
    for spec in sheets:
        warnings += _append_sheet(wb, spec)
    if last_modified_by is not None:
        wb.properties.lastModifiedBy = last_modified_by
    _save(wb, destination, _fixed_time(modified))

    every_cell = [f"{column}{vendor.row}" for vendor in data.vendors for column in student_columns]
    return WriteReport(
        written=list(plan),
        skipped=[coordinate for coordinate in every_cell if coordinate not in plan],
        warnings=warnings,
        row_heights=row_heights,
        sheets_added=[spec.title for spec in sheets],
    )


def _refuse_same_file(source: WorkbookSource, destination: str | Path | BinaryIO) -> None:
    """E13 when the destination is the source (any spelling of the path), so the input is never overwritten."""
    same = source is destination
    if not same and isinstance(source, (str, os.PathLike)) and isinstance(destination, (str, os.PathLike)):
        src, dst = Path(source), Path(destination)
        same = os.path.normcase(src.resolve()) == os.path.normcase(dst.resolve())
        if not same and src.exists() and dst.exists():
            same = os.path.samefile(src, dst)
    if same:
        message = f"Destination {destination} is the source workbook; write the output to a new file."
        raise WorkbookWriteError([ValidationIssue(code="E13", message=message)])


def _length_budgets(overrides: Mapping[str, int] | None) -> dict[str, int]:
    """Default budgets updated with ``overrides``. Unknown field names are a programming error."""
    unknown = sorted(set(overrides or {}) - set(STUDENT_FIELDS))
    if unknown:
        raise ValueError(f"length_budgets names unknown student fields: {', '.join(unknown)}")
    return {**DEFAULT_LENGTH_BUDGETS, **(overrides or {})}


def _plan_cells(
    ws: Worksheet,
    data: WorkbookData,
    cells: Mapping[str, StudentCells],
    budgets: Mapping[str, int],
    overwrite: bool,
) -> tuple[dict[str, str], list[ValidationIssue], list[ValidationIssue]]:
    """Check every requested value against the rules. Returns ({coordinate: text}, errors, warnings)."""
    rows = {vendor.vendor_id: vendor.row for vendor in data.vendors}
    example_id = data.example.vendor_id if data.example else None
    plan: dict[str, str] = {}
    errors: list[ValidationIssue] = []
    warnings: list[ValidationIssue] = []
    for vendor_id, values in cells.items():
        row = rows.get(vendor_id)
        if row is None:
            why = "is the worked example, never edited" if vendor_id == example_id else "is not a vendor row"
            errors.append(ValidationIssue(code="E12", message=f'Vendor ID "{vendor_id}" {why}.'))
            continue
        for field in STUDENT_FIELDS:
            value = getattr(values, field)
            if value is None:
                continue
            cell = ws[f"{data.column_map[field]}{row}"]
            text, removed = _storable(value)
            if removed:
                warnings.append(_removed_characters(cell.coordinate, removed))
            errors += _target_problems(cell, overwrite) + _value_problems(cell.coordinate, field, text, budgets[field])
            if overwrite and _text(cell.value):
                message = f"{cell.coordinate}: the existing value was replaced (overwrite)."
                warnings.append(ValidationIssue(code="E06", severity="warning", cell=cell.coordinate, message=message))
            plan[cell.coordinate] = text
    return plan, errors, warnings


def _target_problems(cell: Cell | MergedCell, overwrite: bool) -> list[ValidationIssue]:
    """E09 unless the cell is a cream student cell; E06 if it already has a value and overwrite is off."""
    where = cell.coordinate
    if isinstance(cell, MergedCell) or not _is_cream(cell):
        message = f"{where} is not a cream student cell, so it is never written."
        return [ValidationIssue(code="E09", cell=where, message=message)]
    if _text(cell.value) and not overwrite:
        message = f"{where} already has a value; pass overwrite to replace it."
        return [ValidationIssue(code="E06", cell=where, message=message)]
    return []


def _value_problems(where: str, field: str, text: str, budget: int) -> list[ValidationIssue]:
    """E10 when the text is not an allowed value (L, O, S) or is over its budget. Text is never truncated."""
    problems = []
    allowed = ALLOWED_VALUES.get(field)
    if allowed and text and text not in allowed:
        message = f"{where} ({field}) must be one of {', '.join(allowed)}, not {_preview(text)}."
        problems.append(ValidationIssue(code="E10", cell=where, message=message))
    if len(text) > budget:
        message = f"{where} ({field}) has {len(text):,} characters; its budget is {budget:,}."
        problems.append(ValidationIssue(code="E10", cell=where, message=message))
    return problems


def _is_cream(cell: Cell) -> bool:
    """True for a solid fill of colour FFF6E0 (alpha ignored): the template's mark for a student cell."""
    fill = cell.fill
    color = getattr(fill, "fgColor", None)
    return (
        getattr(fill, "patternType", None) == "solid"
        and color is not None
        and color.type == "rgb"
        and str(color.rgb).upper()[-6:] == CREAM[-6:]
    )


def _storable(text: str) -> tuple[str, int]:
    """``text`` without the characters an .xlsx cannot store, and how many were removed."""
    text, control = ILLEGAL_CHARACTERS_RE.subn("", text)
    text, other = _UNSTORABLE.subn("", text)
    return text, control + other


def _put_text(cell: Cell, text: str) -> None:
    """Store text as a plain string; '' clears the cell.

    openpyxl types '=...' as a formula and '#N/A' as an error value, so the type is forced back to string.
    A literal '_xHHHH_' is escaped as '_x005F_xHHHH_' so that Excel shows it as written.
    """
    cell.value = _LITERAL_ESCAPE.sub(r"_x005F_\1", text) or None
    if text:
        cell.data_type = "s"


def _removed_characters(where: str, count: int) -> ValidationIssue:
    message = f"{where}: removed {count} character(s) that an .xlsx file cannot store."
    return ValidationIssue(code="E11", severity="warning", cell=where, message=message)


def _raise_row_heights(ws: Worksheet, rows: Iterable[int], columns: Sequence[str]) -> dict[int, float]:
    """Make each row tall enough for its longest student cell; never lower it, never raise it above 409 pt."""
    default = ws.sheet_format.defaultRowHeight or 15.0
    raised: dict[int, float] = {}
    for row in sorted(rows):
        lines = max(_estimate_lines(ws[f"{column}{row}"].value, _column_width(ws, column)) for column in columns)
        needed = min(MAX_ROW_HEIGHT, lines * LINE_HEIGHT + ROW_PADDING)
        dimension = ws.row_dimensions[row]
        if needed > (dimension.height or default):
            dimension.height = needed
            raised[row] = needed
    return raised


def _estimate_lines(value: object, width: float) -> int:
    """Wrapped lines of Arial 10 text in a column ``width`` units wide (about 1.15 characters per unit).

    Every paragraph takes at least one line, so blank lines between paragraphs count too.
    """
    text = "" if value is None else str(value)
    if not text:
        return 0
    per_line = max(10, int(width * 1.15))
    return sum(max(1, math.ceil(len(paragraph) / per_line)) for paragraph in text.split("\n"))


def _column_width(ws: Worksheet, letter: str) -> float:
    """Width of the <col> span that contains column ``letter``.

    Scans the existing spans instead of indexing ``column_dimensions[letter]``: L-O and R-T are single spans
    stored under L and R, and indexing M would create a second, overlapping <col> (Excel repair prompt).
    """
    index = column_index_from_string(letter)
    for span in ws.column_dimensions.values():
        if span.min is not None and span.max is not None and span.min <= index <= span.max:
            return span.width
    return ws.sheet_format.defaultColWidth or 8.43


def _sheet_title_problems(wb: Workbook, sheets: Sequence[SheetSpec]) -> list[ValidationIssue]:
    """E13 for an appended sheet whose title is invalid in Excel or already taken (case-insensitive)."""
    taken = {name.casefold() for name in wb.sheetnames}
    problems = []
    for spec in sheets:
        title = spec.title
        if not 1 <= len(title) <= 31 or _BAD_TITLE_CHARS.search(title) or title.startswith("'") or title.endswith("'"):
            message = f'Sheet title "{title}" is not valid in Excel (1-31 characters, none of \\ / ? * [ ] :).'
            problems.append(ValidationIssue(code="E13", message=message))
        elif title.casefold() in _RESERVED_TITLES:
            message = f'Sheet title "{title}" is reserved by Excel.'
            problems.append(ValidationIssue(code="E13", message=message))
        elif title.casefold() in taken:
            message = f'Sheet "{title}" already exists; sheets are never replaced.'
            problems.append(ValidationIssue(code="E13", message=message))
        taken.add(title.casefold())
    return problems


def _append_sheet(wb: Workbook, spec: SheetSpec) -> list[ValidationIssue]:
    """Add ``spec`` as the last sheet: italic note, bold cream header, wrapped rows, frozen and filtered header."""
    ws = wb.create_sheet(spec.title)
    header_row = 2 if spec.note else 1
    warnings: list[ValidationIssue | None] = []
    if spec.note:
        warnings.append(_put_value(ws["A1"], spec.note))
        ws["A1"].font = _NOTE_FONT
    for column, header in enumerate(spec.headers, start=1):
        cell = ws.cell(header_row, column)
        warnings.append(_put_value(cell, header))
        cell.font, cell.fill, cell.alignment = _HEADER_FONT, _HEADER_FILL, _WRAP_TOP
    for row, values in enumerate(spec.rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row, column)
            warnings.append(_put_value(cell, value))
            cell.font, cell.alignment = _BODY_FONT, _WRAP_TOP
    for column, header in enumerate(spec.headers, start=1):
        given = spec.column_widths[column - 1] if column <= len(spec.column_widths) else None
        values = [row_values[column - 1] for row_values in spec.rows if len(row_values) >= column]
        ws.column_dimensions[get_column_letter(column)].width = given or _default_width(header, values)
    if spec.headers:
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(spec.headers))}{header_row + len(spec.rows)}"
    return [warning for warning in warnings if warning]


def _put_value(cell: Cell, value: str | int | float | None) -> ValidationIssue | None:
    """Write an appended-sheet value: finite numbers stay numbers, anything else is stored as plain text."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        cell.value = value
        return None
    text, removed = _storable(str(value))
    _put_text(cell, text)
    return _removed_characters(f"'{cell.parent.title}'!{cell.coordinate}", removed) if removed else None


def _default_width(header: str, values: Iterable[object]) -> float:
    """A width that fits the longest line of the header or values, kept between 10 and 60 characters."""
    lines = [line for item in (header, *values) if item is not None for line in str(item).split("\n")]
    return float(min(60, max(10, max((len(line) for line in lines), default=0) + 2)))


def _save(wb: Workbook, destination: str | Path | BinaryIO, modified: dt.datetime | None) -> None:
    """Serialise in memory first (a failed save then opens no file), then write the stream, or the path through
    a temporary file that is renamed into place. ``modified`` pins docProps/core.xml for reproducible output."""
    buffer = io.BytesIO()
    wb.save(buffer)
    data = _pin_modified(buffer.getvalue(), modified) if modified is not None else buffer.getvalue()
    if not isinstance(destination, (str, os.PathLike)):
        destination.write(data)
        return
    target = Path(destination)
    handle, temporary = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=target.parent)
    try:
        with os.fdopen(handle, "wb") as stream:
            stream.write(data)
        os.replace(temporary, target)
    except BaseException:
        try:
            Path(temporary).unlink(missing_ok=True)
        except OSError:
            pass  # never hide the original error
        raise


def _pin_modified(data: bytes, modified: dt.datetime) -> bytes:
    """The .xlsx ``data`` with dcterms:modified and every zip entry time set to ``modified`` (UTC)."""
    stamp = modified.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    stamp_parts = max(modified.astimezone(dt.timezone.utc).timetuple()[:6], (1980, 1, 1, 0, 0, 0))
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            part = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                part = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + stamp.encode() + rb"\g<2>", part)
            info.date_time = stamp_parts  # entry timestamps too, or the zip still differs
            target.writestr(info, part)
    return output.getvalue()


def _fixed_time(modified: dt.datetime | None) -> dt.datetime | None:
    """``modified``, else SOURCE_DATE_EPOCH (reproducible builds) when it is set, else None (wall clock)."""
    if modified is not None:
        return modified
    epoch = os.environ.get("SOURCE_DATE_EPOCH", "").strip()
    return dt.datetime.fromtimestamp(int(epoch), dt.timezone.utc) if epoch.isdigit() else None


# --------------------------------------------------------------------------- fidelity


def check_fidelity(original: WorkbookSource, output: WorkbookSource) -> list[str]:
    """Semantic diff of ``output`` against ``original``: [] means that only allowed changes were made.

    Allowed: the values of student cells (L-V) on real vendor rows, taller real vendor rows, and sheets appended
    after the original ones (nothing is allowed if the original's inventory has read errors). Everything else is
    compared: sheet order and active sheet; per original sheet the cell values, number formats, fonts, fills,
    borders, alignment, protection and hyperlinks, merged ranges, <col> spans, row heights, freeze panes, sheet
    views, data validations, visibility and auto-filter.
    """
    before, after = _load(original), _load(output)
    inventory = read_workbook(original)
    vendor_rows = {vendor.row for vendor in inventory.vendors} if inventory.ok else set()
    student_cells = {f"{inventory.column_map[field]}{row}" for row in vendor_rows for field in STUDENT_FIELDS}

    diffs: list[str] = []
    if after.sheetnames[: len(before.sheetnames)] != before.sheetnames:
        diffs.append(f"Sheet order changed from {before.sheetnames} to {after.sheetnames}.")
    if after.active is None or after.active.title != before.active.title:
        now = after.active.title if after.active is not None else None
        diffs.append(f'The active sheet changed from "{before.active.title}" to "{now}".')
    for sheet in before.worksheets:
        if sheet.title not in after.sheetnames:
            diffs.append(f'Sheet "{sheet.title}" is missing.')
            continue
        is_inventory = sheet.title == inventory.sheet_name
        diffs += _sheet_diffs(
            sheet,
            after[sheet.title],
            student_cells if is_inventory else set(),
            vendor_rows if is_inventory else set(),
        )
    return diffs


def _sheet_diffs(a: Worksheet, b: Worksheet, student_cells: set[str], vendor_rows: set[int]) -> list[str]:
    """Differences between an original sheet and the same sheet in the output."""
    where = f'"{a.title}"'
    diffs: list[str] = []
    for row in range(1, max(a.max_row, b.max_row) + 1):
        for column in range(1, max(a.max_column, b.max_column) + 1):
            first = a.cell(row, column)
            diffs += _cell_diffs(first, b.cell(row, column), value_may_change=first.coordinate in student_cells)
    diffs += _row_diffs(a, b, vendor_rows)
    diffs += _column_diffs(a, b)
    merged_a, merged_b = sorted(map(str, a.merged_cells.ranges)), sorted(map(str, b.merged_cells.ranges))
    if merged_a != merged_b:
        diffs.append(f"{where}: merged ranges changed from {merged_a} to {merged_b}.")
    if a.freeze_panes != b.freeze_panes:
        diffs.append(f"{where}: freeze panes changed from {a.freeze_panes} to {b.freeze_panes}.")
    if repr(a.views) != repr(b.views):
        view_a, view_b = a.sheet_view, b.sheet_view
        diffs.append(
            f"{where}: sheet view changed (zoom {view_a.zoomScale} -> {view_b.zoomScale}, "
            f"top-left cell {view_a.topLeftCell} -> {view_b.topLeftCell})."
        )
    rules_a = [repr(rule) for rule in a.data_validations.dataValidation]
    rules_b = [repr(rule) for rule in b.data_validations.dataValidation]
    if rules_a != rules_b:
        diffs.append(f"{where}: data validations changed.")
    if a.sheet_state != b.sheet_state:
        diffs.append(f"{where}: visibility changed from {a.sheet_state} to {b.sheet_state}.")
    if a.auto_filter.ref != b.auto_filter.ref:
        diffs.append(f"{where}: auto-filter changed from {a.auto_filter.ref} to {b.auto_filter.ref}.")
    return diffs


def _cell_diffs(a: Cell | MergedCell, b: Cell | MergedCell, value_may_change: bool) -> list[str]:
    where = f"'{a.parent.title}'!{a.coordinate}"
    diffs = []
    if not value_may_change and (a.data_type, repr(a.value)) != (b.data_type, repr(b.value)):
        diffs.append(f"{where}: value changed from {_preview(a.value)} to {_preview(b.value)}.")
    if value_may_change and b.data_type in ("f", "e"):
        diffs.append(f"{where}: student cell holds a formula or error value, not plain text.")
    diffs += [f"{where}: {part} changed." for part in _STYLE_PARTS if repr(getattr(a, part)) != repr(getattr(b, part))]
    if _hyperlink(a) != _hyperlink(b):
        diffs.append(f"{where}: hyperlink changed from {_hyperlink(a)} to {_hyperlink(b)}.")
    return diffs


def _row_diffs(a: Worksheet, b: Worksheet, growable: set[int]) -> list[str]:
    """Row heights and visibility must match, except that the rows in ``growable`` may become taller."""
    diffs = []
    for row in sorted(set(a.row_dimensions) | set(b.row_dimensions)):
        (height_a, hidden_a), (height_b, hidden_b) = _row_shape(a, row), _row_shape(b, row)
        if (height_a, hidden_a) == (height_b, hidden_b):
            continue
        if row in growable and hidden_a == hidden_b and height_b >= height_a:
            continue
        diffs.append(
            f'"{a.title}": row {row} changed from {height_a:g} pt{" (hidden)" if hidden_a else ""} '
            f'to {height_b:g} pt{" (hidden)" if hidden_b else ""}.'
        )
    return diffs


def _row_shape(ws: Worksheet, row: int) -> tuple[float, bool]:
    default = ws.sheet_format.defaultRowHeight or 15.0
    dimension = ws.row_dimensions.get(row)  # .get: never create a row definition while checking
    if dimension is None:
        return default, False
    return (default if dimension.height is None else dimension.height), bool(dimension.hidden)


def _column_diffs(a: Worksheet, b: Worksheet) -> list[str]:
    """<col> spans must match one for one. Only existing keys are looked up, so no definition is created."""
    diffs = []
    for key, span in a.column_dimensions.items():
        other = b.column_dimensions.get(key)
        if other is None or _span(span) != _span(other):
            now = _span(other) if other is not None else None
            diffs.append(f'"{a.title}": column {key} definition changed from {_span(span)} to {now}.')
    for key in sorted(b.column_dimensions.keys() - a.column_dimensions.keys(), key=column_index_from_string):
        diffs.append(f'"{a.title}": column {key} gained a <col> definition {_span(b.column_dimensions.get(key))}.')
    return diffs


def _span(dimension: ColumnDimension) -> tuple[int | None, int | None, float, bool]:
    """(min, max, width, hidden) of a <col> definition."""
    return dimension.min, dimension.max, dimension.width, bool(dimension.hidden)


def _hyperlink(cell: Cell | MergedCell) -> str | None:
    link = cell.hyperlink
    if link is None:
        return None
    return (link.target or "") + (f"#{link.location}" if link.location else "")


def _preview(value: object, limit: int = 50) -> str:
    """Short quoted form of a cell value for messages."""
    if value is None:
        return "empty"
    text = str(value)
    quoted = repr(text if len(text) <= limit else text[:limit] + "...")
    return quoted + (" (rich text)" if isinstance(value, CellRichText) else "")
