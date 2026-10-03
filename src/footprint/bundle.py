"""Colab bundle and notebook support (design: Demo; docs/contracts_p3.md section 14).

``build_bundle`` packs everything ``notebooks/footprint_demo.ipynb`` needs to replay the assessment on Google Colab,
or on any machine without the private repository: the footprint wheel (``uv build``), the source tree, config,
prompts, seeds, the input workbook, the frozen evidence pack (index, raw blobs, text store, screenshots, LLM cache,
AI-policy decisions), the review stores and the run records.

- The zip is deterministic: entries sorted by path, fixed timestamps and permissions, deflate (gzip blobs, the
  wheel and .xlsx files are stored as they are, being compressed already). Its ``BUNDLE_MANIFEST.json`` maps
  every file to its SHA-256.
- It never holds ``.env``, other secret files, git or virtual-env folders, caches, ``scratch/`` or ``submission/``.
  From a ``.streamlit`` folder it takes ``config.toml`` only (``app/.streamlit/config.toml`` keeps the app on
  localhost with no usage statistics wherever it runs); ``secrets.toml`` and everything else there stays out.
  A secret scan refuses the build when any file (gzip blobs and zip members included) holds an API-key-like token
  or the value of a key from the environment or ``.env``. The error names files, never values. Only reviewed
  public captures listed in ``allow_tokens_in`` may hold a token (a vendor page's own Maps key); never our values.
- Colab replay installs the build machine's versions of the text-extraction and robots packages and of pydantic
  (``REPLAY_PACKAGES``, recorded in the manifest), and asks for a session restart if the kernel had already
  imported an older one. pydantic matters because the LLM cache keys hash the response schemas it generates; the
  manifest also records those schema hashes (``llm_schemas``), and the setup warns when the kernel's differ.
- It is never uploaded or published from here. It is shared as a private release asset and stays outside git
  (``.gitignore``: ``footprint_bundle*.zip``).

The second half of this module keeps the notebook free of logic. Setup, the workbook upload, the tables, the
download and the verification checks are all calls into it. It never imports ``footprint.pipeline``: the notebook
cells call the pipeline themselves, as the app does.
"""

from __future__ import annotations

import fnmatch
import gzip
import hashlib
import html
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import tomllib
import zipfile
import zlib
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from typing import Any

from footprint import __version__
from footprint.models import STUDENT_FIELDS, AssessmentResult, EvidenceItem, VendorFindings

# =========================================================================== bundle

BUNDLE_NAME = "footprint_bundle.zip"
MANIFEST_NAME = "BUNDLE_MANIFEST.json"
BUNDLE_FORMAT = "footprint-bundle/1"
VERIFIED_MARKER = ".bundle_verified"
MAX_BUNDLE_BYTES = 100_000_000
"""The bundle must stay under 100 MB (decimal), so a browser upload to Colab stays practical."""

DEFAULT_INCLUDE: tuple[str, ...] = (
    "pyproject.toml", "uv.lock", "README.md", "CLAUDE.md", "config", "prompts", "seeds", "src", "app",
    "notebooks", "data/input", "evidence", "runs", "review", "tests/gold",
)
"""Paths under the repository root that the bundle takes. A missing entry is an error."""

EXCLUDED_DIRS: frozenset[str] = frozenset({
    ".git", ".venv", "venv", "__pycache__", ".pytest_cache", ".ipynb_checkpoints", ".ruff_cache", ".mypy_cache",
    "scratch", "submission", "dist", "build", "node_modules",
})
"""Folder names that are never bundled, at any depth."""

STREAMLIT_DIR = ".streamlit"
STREAMLIT_FILES: frozenset[str] = frozenset({"config.toml"})
"""The only files a ``.streamlit`` folder contributes: the app's settings (localhost only, no usage statistics), so
the app keeps them when it runs from an extracted bundle. ``secrets.toml``, credentials and sub-folders stay out."""

EXCLUDED_DIR_PATTERNS: tuple[str, ...] = ("*.egg-info",)
EXCLUDED_PREFIXES: tuple[str, ...] = ("review/sandbox/",)
"""Demo sandboxes (git-ignored): the notebook and the app work on copies of the review stores there."""

EXCLUDED_FILE_PATTERNS: tuple[str, ...] = (
    ".env", ".env.*", "*.pyc", "*.pyo", "*.tmp", "*.partial", "~$*", "footprint_bundle*.zip", ".bundle_verified",
    ".ds_store", "thumbs.db", "desktop.ini", "secrets.toml", "*.pem", "*.key", "*.p12", "*.pfx", "id_rsa*",
    "id_ecdsa*", "id_ed25519*", ".netrc", ".pypirc", "credentials*.json", "client_secret*.json",
    "service_account*.json", "token.json",
)
"""File names (compared in lower case) that are never bundled: secrets, caches, editor and OS litter."""

ALLOWED_DOTENV_NAMES: frozenset[str] = frozenset({".env.example"})

SECRET_PATTERN = re.compile(
    rb"AIza[0-9A-Za-z_\-]{35}"                        # Google API key (classic format)
    rb"|(?<![0-9A-Za-z_\-])AQ\.[0-9A-Za-z_\-]{40,}"   # Google API key (AQ. format)
    rb"|gh[pousr]_[0-9A-Za-z]{36,}"                   # GitHub tokens
    rb"|github_pat_[0-9A-Za-z_]{22,}"                 # GitHub fine-grained tokens
)
_SECRET_PREFIXES: tuple[bytes, ...] = (b"AIza", b"AQ.", b"ghp_", b"gho_", b"ghu_", b"ghs_", b"ghr_", b"github_pat_")
SECRET_ENV_VARS: tuple[str, ...] = ("GEMINI_API_KEY", "GOOGLE_API_KEY", "GH_TOKEN", "GITHUB_TOKEN")
"""Environment variables whose values must never appear in a bundled file."""
_SECRET_NAME = re.compile(r"(^|_)(KEY|TOKEN|SECRET|PASSWORD|PASSWD|PWD)$", re.IGNORECASE)
"""A .env entry with such a name holds a secret value."""
MIN_SECRET_CHARS = 8

REPLAY_PACKAGES: dict[str, str] = {
    "charset-normalizer": "charset_normalizer", "courlan": "courlan", "htmldate": "htmldate", "justext": "justext",
    "lxml": "lxml", "protego": "protego", "pydantic": "pydantic", "pydantic-core": "pydantic_core",
    "pypdf": "pypdf", "trafilatura": "trafilatura",
}
"""Distribution -> import name of the packages whose versions decide what replay rebuilds: text extraction
(trafilatura and its parsers, the lxml fallback, pypdf, encoding and date detection), robots parsing, and pydantic
with its core, which generates the LLM response schemas whose SHA-256 is part of every LLM cache key (another
version could turn every cache lookup into a silent miss, and the cells into rules-only ones). The manifest records
the build machine's versions, so the Colab setup installs the same ones."""

ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
_ZIP_FILE_MODE = 0o100644
_ZIP_LEVEL = 6


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_excluded(arcname: str) -> bool:
    """True when the bundle never takes this repository-relative POSIX path."""
    parts = arcname.split("/")
    if any(_excluded_dir(part) for part in parts[:-1]):
        return True
    if STREAMLIT_DIR in parts[:-1]:
        return not (parts[-2] == STREAMLIT_DIR and parts[-1] in STREAMLIT_FILES)
    if any(arcname.startswith(prefix) for prefix in EXCLUDED_PREFIXES):
        return True
    name = parts[-1].lower()
    if name in ALLOWED_DOTENV_NAMES:
        return False
    return any(fnmatch.fnmatchcase(name, pattern) for pattern in EXCLUDED_FILE_PATTERNS)


def bundle_files(root: str | Path = ".", include: Sequence[str] = DEFAULT_INCLUDE, *,
                 skip: Iterable[str | Path] = ()) -> list[tuple[str, Path]]:
    """(arcname, path) for every file the bundle takes from ``root``, sorted by arcname.

    ``include`` entries are POSIX paths relative to ``root``; a missing entry raises ValueError, and so does one
    that leaves ``root``. Excluded folders are not entered, symbolic links are never followed, and the paths in
    ``skip`` (for example the output zip) are left out.
    """
    base = Path(root).resolve()
    skipped = {_same_path_key(p) for p in skip}
    found: dict[str, Path] = {}
    missing: list[str] = []
    for entry in include:
        rel = PurePosixPath(entry)
        if rel.is_absolute() or ".." in rel.parts or not rel.parts:
            raise ValueError(f"bundle include entries must be relative paths inside the root, not {entry!r}")
        path = base.joinpath(*rel.parts)
        if path.is_symlink():
            raise ValueError(f"bundle include entry {entry!r} is a symbolic link; links are never followed")
        if not path.exists():
            missing.append(entry)
            continue
        candidates = [(rel.as_posix(), path)] if path.is_file() else _walk(path, rel.as_posix())
        for arcname, file in candidates:
            if not is_excluded(arcname) and _same_path_key(file) not in skipped:
                found[arcname] = file
    if missing:
        raise ValueError(f"bundle include entries not found under {base}: {', '.join(missing)}")
    return sorted(found.items())


def _same_path_key(path: str | Path) -> str:
    return os.path.normcase(os.path.abspath(path))


def _excluded_dir(name: str) -> bool:
    return name in EXCLUDED_DIRS or any(fnmatch.fnmatchcase(name, p) for p in EXCLUDED_DIR_PATTERNS)


def _walk(folder: Path, prefix: str) -> list[tuple[str, Path]]:
    """(arcname, path) of the regular files under ``folder``; excluded folders and links are not entered."""
    files: list[tuple[str, Path]] = []
    with os.scandir(folder) as listing:
        entries = sorted(listing, key=lambda e: e.name)
    for entry in entries:
        if entry.is_symlink():
            continue
        arcname = f"{prefix}/{entry.name}"
        if entry.is_dir(follow_symlinks=False):
            if not _excluded_dir(entry.name):
                files.extend(_walk(Path(entry.path), arcname))
        elif entry.is_file(follow_symlinks=False):
            files.append((arcname, Path(entry.path)))
    return files


def parse_dotenv(path: str | Path) -> dict[str, str]:
    """KEY=VALUE lines of a .env file (comments and blank lines skipped, surrounding quotes removed).

    The values are secrets: callers must never print or log them.
    """
    p = Path(path)
    if not p.is_file():
        return {}
    values: dict[str, str] = {}
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def secret_values(root: str | Path = ".") -> list[bytes]:
    """Values that must never be bundled: SECRET_ENV_VARS from the environment and every key-, token-, secret- or
    password-named entry of ``root/.env``. Kept in memory only, never printed."""
    values = {os.environ.get(name, "") for name in SECRET_ENV_VARS}
    dotenv = parse_dotenv(Path(root) / ".env")
    values |= {value for key, value in dotenv.items() if key in SECRET_ENV_VARS or _SECRET_NAME.search(key)}
    return sorted(value.encode("utf-8") for value in values if len(value) >= MIN_SECRET_CHARS)


def _scan_blob(data: bytes, values: Sequence[bytes]) -> bool:
    if any(value in data for value in values):
        return True
    for prefix in _SECRET_PREFIXES:  # find() is fast; the pattern only runs where a token could start
        start = data.find(prefix)
        while start != -1:
            if SECRET_PATTERN.match(data, start):
                return True
            start = data.find(prefix, start + 1)
    return False


def _scan_values(data: bytes, values: Sequence[bytes]) -> bool:
    return any(value in data for value in values)


def holds_secret(data: bytes, values: Sequence[bytes] = (), *, tokens: bool = True) -> bool:
    """True when ``data`` holds one of ``values`` or, with ``tokens``, a secret-like token. Gzip data is scanned
    decompressed, and every member of a zip (a wheel, an .xlsx) is scanned too."""
    scan = _scan_blob if tokens else _scan_values
    if scan(data, values):
        return True
    if data[:2] == b"\x1f\x8b":
        try:
            return scan(gzip.decompress(data), values)
        except (OSError, EOFError, zlib.error):
            return False
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                return any(scan(archive.read(info), values) for info in archive.infolist())
        except (zipfile.BadZipFile, OSError, RuntimeError, ValueError, NotImplementedError, zlib.error):
            return False
    return False


def scan_secrets(entries: Iterable[tuple[str, bytes]], values: Sequence[bytes] = (), *,
                 allow_tokens_in: Sequence[str] = ()) -> list[str]:
    """Names of the entries that hold a secret (never the matches themselves).

    ``allow_tokens_in`` holds fnmatch patterns of entry names that are exempt from the token patterns (a public
    capture whose page embeds a third-party key, such as a Maps API key). They are still checked for ``values``.
    """
    return [name for name, data in entries
            if holds_secret(data, values, tokens=not any(fnmatch.fnmatchcase(name, p) for p in allow_tokens_in))]


def build_wheel(root: str | Path = ".", out_dir: str | Path | None = None, *, offline: bool = False) -> Path:
    """Build the footprint wheel with ``uv build --wheel`` and return its path (a temporary folder by default).

    uv builds pure-Python wheels reproducibly (fixed timestamps), so the bundle stays deterministic.
    """
    uv = shutil.which("uv")
    if uv is None:
        raise RuntimeError("uv is not on PATH: install uv, or pass wheel=<path of a built wheel> to build_bundle")
    target = Path(out_dir) if out_dir is not None else Path(tempfile.mkdtemp(prefix="footprint-wheel-"))
    target.mkdir(parents=True, exist_ok=True)
    cmd = [uv, "build", "--wheel", "--out-dir", str(target), str(Path(root).resolve())]
    if offline:
        cmd.insert(2, "--offline")
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)
    if proc.returncode != 0:
        tail = (proc.stderr or proc.stdout).strip()[-600:]
        raise RuntimeError(f"uv build failed (exit {proc.returncode}): {tail}")
    wheels = sorted(target.glob("*.whl"))
    if len(wheels) != 1:
        raise RuntimeError(f"expected one wheel in {target}, found {len(wheels)}")
    return wheels[0]


def package_versions(names: Iterable[str] = REPLAY_PACKAGES) -> dict[str, str]:
    """name -> installed version, for the names that are installed in this environment."""
    from importlib import metadata

    found: dict[str, str] = {}
    for name in names:
        try:
            found[name] = metadata.version(name)
        except metadata.PackageNotFoundError:
            continue
    return found


def llm_schema_hashes() -> dict[str, str]:
    """role -> SHA-256 of the response schema that this kernel's pydantic generates for it, as the LLM cache keys
    use them (footprint.ai.schema_sha256); empty when footprint.ai cannot be imported."""
    try:
        from footprint import ai
    except ImportError:
        return {}
    return {role: ai.schema_sha256(schema) for role, schema in (
        ("expand", ai.EXPAND_SCHEMA), ("extract", ai.EXTRACT_SCHEMA), ("triage", ai.TRIAGE_SCHEMA))}


def _project(root: Path) -> tuple[str, str]:
    """(name, version) from root/pyproject.toml, else the installed package's."""
    try:
        with open(root / "pyproject.toml", "rb") as fh:
            project = tomllib.load(fh).get("project", {})
        return str(project.get("name", "footprint")), str(project.get("version", __version__))
    except (OSError, tomllib.TOMLDecodeError):
        return "footprint", __version__


def _compressed(data: bytes) -> bool:
    """gzip or zip data (evidence blobs, the wheel, .xlsx): deflating it again gains nothing."""
    return data[:2] == b"\x1f\x8b" or data[:4] == b"PK\x03\x04"


def _zip_info(name: str, data: bytes = b"") -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(name, date_time=ZIP_EPOCH)
    info.compress_type = zipfile.ZIP_STORED if _compressed(data) else zipfile.ZIP_DEFLATED
    info.create_system = 3  # Unix, so the permission bits below read the same on every platform
    info.external_attr = _ZIP_FILE_MODE << 16
    return info


def build_bundle(out_zip: str | Path = BUNDLE_NAME, *, root: str | Path = ".",
                 include: Sequence[str] = DEFAULT_INCLUDE, wheel: bool | str | Path = True,
                 max_bytes: int = MAX_BUNDLE_BYTES, allow_tokens_in: Sequence[str] = ()) -> dict[str, Any]:
    """Write the Colab bundle and return {path, files, bytes, sha256, wheel}.

    - ``out_zip`` is relative to the current directory, like a CLI argument. Write it outside git: the default
      name is git-ignored.
    - ``wheel``: True builds it with ``uv build``. A path uses a wheel you built already. False leaves it out.
      The wheel goes in at ``dist/<wheel name>``.
    - ``allow_tokens_in``: fnmatch patterns of bundle paths exempt from the token patterns, for a public capture
      whose page embeds a third-party key (a Maps API key, say). Name single files after reviewing them. The
      values from ``.env`` and the environment are still refused everywhere.
    - ``files`` counts the bundled files, the wheel included and the manifest left out. ``bytes`` and ``sha256``
      describe the zip.

    Raises ValueError, and writes nothing, when an include entry is missing, when a file holds a secret (the
    message names the files, never the values) or when the zip would reach ``max_bytes``.
    """
    base = Path(root).resolve()
    out = Path(out_zip).resolve()
    files = bundle_files(base, include, skip=[out])
    entries: list[tuple[str, bytes]] = [(arcname, path.read_bytes()) for arcname, path in files]
    wheel_name = ""
    if wheel:
        with tempfile.TemporaryDirectory(prefix="footprint-wheel-") as tmp:
            built = build_wheel(base, tmp) if wheel is True else Path(wheel)
            if built.suffix != ".whl" or not built.is_file():
                raise ValueError(f"not a wheel file: {built}")
            wheel_name = f"dist/{built.name}"
            entries.append((wheel_name, built.read_bytes()))
    offenders = scan_secrets(entries, secret_values(base), allow_tokens_in=allow_tokens_in)
    if offenders:
        shown = ", ".join(offenders[:10]) + (f" and {len(offenders) - 10} more" if len(offenders) > 10 else "")
        raise ValueError(f"secret-like content found in {len(offenders)} file(s): {shown}. The bundle was not "
                         "written; remove the secret (or exclude the file) and build again. A public capture "
                         "that embeds a third-party key can be listed in allow_tokens_in after review.")
    name, version = _project(base)
    manifest = {
        "format": BUNDLE_FORMAT,
        "name": name,
        "version": version,
        "wheel": wheel_name,
        "replay_packages": package_versions(),
        "llm_schemas": llm_schema_hashes(),
        "include": list(include),
        "file_count": len(entries),
        "total_bytes": sum(len(data) for _, data in entries),
        "files": {arcname: sha256_bytes(data) for arcname, data in sorted(entries)},
    }
    manifest_bytes = (json.dumps(manifest, sort_keys=True, indent=1, ensure_ascii=False) + "\n").encode("utf-8")
    payload = sorted([*entries, (MANIFEST_NAME, manifest_bytes)])

    out.parent.mkdir(parents=True, exist_ok=True)
    partial = out.with_name(out.name + ".partial")
    try:
        with zipfile.ZipFile(partial, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=_ZIP_LEVEL) as archive:
            for arcname, data in payload:
                info = _zip_info(arcname, data)
                archive.writestr(info, data, compress_type=info.compress_type, compresslevel=_ZIP_LEVEL)
        size = partial.stat().st_size
        if size >= max_bytes:
            raise ValueError(f"the bundle would be {size:,} bytes, over the {max_bytes:,}-byte limit; "
                             "leave large folders out of `include`")
        os.replace(partial, out)
    finally:
        partial.unlink(missing_ok=True)
    return {"path": str(out), "files": len(entries), "bytes": size, "sha256": sha256_file(out),
            "wheel": wheel_name}


def read_manifest(path: str | Path) -> dict[str, Any]:
    """BUNDLE_MANIFEST.json of a bundle zip, or of the folder it was extracted to."""
    p = Path(path)
    if p.is_file() and zipfile.is_zipfile(p):
        with zipfile.ZipFile(p) as archive:
            return json.loads(archive.read(MANIFEST_NAME).decode("utf-8"))
    return json.loads((p / MANIFEST_NAME).read_text(encoding="utf-8"))


def verify_bundle(path: str | Path = ".") -> list[str]:
    """Check a bundle zip, or the folder it was extracted to, against its manifest.

    Returns one line per problem ("missing: <path>", "changed: <path>"); an empty list means every file matches.
    Files the manifest does not list are ignored.
    """
    p = Path(path)
    manifest = read_manifest(p)
    problems: list[str] = []
    if p.is_file():
        with zipfile.ZipFile(p) as archive:
            names = set(archive.namelist())
            for arcname, sha in sorted(manifest["files"].items()):
                if arcname not in names:
                    problems.append(f"missing: {arcname}")
                elif sha256_bytes(archive.read(arcname)) != sha:
                    problems.append(f"changed: {arcname}")
        return problems
    for arcname, sha in sorted(manifest["files"].items()):
        file = p.joinpath(*PurePosixPath(arcname).parts)
        if not file.is_file():
            problems.append(f"missing: {arcname}")
        elif sha256_file(file) != sha:
            problems.append(f"changed: {arcname}")
    return problems


# =========================================================================== notebook support

LIVE_MODULES: tuple[str, ...] = ("protego", "trafilatura", "pypdf")
"""Modules the live modes need (extra `live`)."""
AI_MODULES: tuple[str, ...] = ("google.genai",)
"""Modules live_ai needs (extra `ai`)."""
COLAB_SECRETS: dict[str, tuple[str, ...]] = {
    "replay": ("FOOTPRINT_TEAM_NAME",),
    "live_rules": ("FOOTPRINT_TEAM_NAME", "FOOTPRINT_SEC_CONTACT"),
    "live_ai": ("FOOTPRINT_TEAM_NAME", "FOOTPRINT_SEC_CONTACT", "GEMINI_API_KEY"),
}
"""Colab secrets read per mode, when present: only what the mode uses, so the Gemini key enters the kernel only
for live_ai. Secrets are never typed into a cell and never printed."""
MODES: tuple[str, ...] = tuple(COLAB_SECRETS)
PLACEHOLDER_TEAM = "Team TBD"
COLUMN_LETTERS = "LMNOPQRSTUV"
REQUIRED_SHEETS: tuple[str, ...] = ("Evidence Log", "Coverage Log")
"""Appended sheets the export must always carry (the design's 'never cut' list)."""
REVERIFY_CHECK = "Cited evidence re-verifies (capture and text SHA-256, exact quote)"


def in_colab() -> bool:
    return "google.colab" in sys.modules


@dataclass(frozen=True)
class NotebookEnv:
    """What the setup cell found. Holds no secret values: only whether a key is available."""

    root: str
    colab: bool
    mode: str
    team: str
    gemini_key: bool
    bundle: str
    footprint: str = __version__
    python: str = field(default_factory=lambda: sys.version.split()[0])
    notes: tuple[str, ...] = ()

    def lines(self) -> list[str]:
        where = "Google Colab" if self.colab else "local Jupyter"
        key = "available" if self.gemini_key else "not set (replay and live_rules need none)"
        placeholder = self.team == PLACEHOLDER_TEAM
        team = self.team + (" (placeholder: set TEAM in the parameters cell)" if placeholder else "")
        return [f"footprint {self.footprint} on Python {self.python} ({where})", f"Root: {self.root}",
                f"Mode: {self.mode}", f"Bundle: {self.bundle}", f"Gemini key: {key}", f"Column V team: {team}",
                *(f"Note: {n}" for n in self.notes)]

    def __repr__(self) -> str:
        return "\n".join(self.lines())

    def _repr_html_(self) -> str:
        return "<br>".join(html.escape(line) for line in self.lines())


def find_root(start: str | Path | None = None) -> Path:
    """The repository checkout or extracted bundle: the first folder at or above ``start`` (default: the current
    directory) that holds ``src/footprint`` and ``config``, else the folder the imported package lives in."""
    here = Path(start or Path.cwd()).resolve()
    for folder in (here, *here.parents):
        if (folder / "src" / "footprint").is_dir() and (folder / "config").is_dir():
            return folder
    package_root = Path(__file__).resolve().parents[2]
    if (package_root / "config").is_dir():
        return package_root
    raise FileNotFoundError(f"no footprint checkout or bundle at or above {here}: run the setup cell from the "
                            "repository, or upload footprint_bundle.zip on Colab")


def load_dotenv_names(path: str | Path) -> list[str]:
    """Load a .env file into os.environ (variables already set win) and return the names loaded, never values."""
    loaded = []
    for key, value in parse_dotenv(path).items():
        if key not in os.environ:
            os.environ[key] = value
            loaded.append(key)
    return sorted(loaded)


def colab_secret(name: str) -> str | None:
    """A Colab secret (the key icon in the sidebar), or None when it is missing or access was not granted."""
    try:
        from google.colab import userdata  # type: ignore[import-not-found]

        value = userdata.get(name)
    except Exception:  # noqa: BLE001 - SecretNotFoundError, NotebookAccessError, TimeoutException, ...
        return None
    return str(value) if value else None


def _missing_modules(names: Iterable[str]) -> list[str]:
    missing = []
    for name in names:
        try:
            found = importlib.util.find_spec(name) is not None
        except (ImportError, ValueError):
            found = False
        if not found:
            missing.append(name)
    return missing


def _imported_already(requirements: Sequence[str]) -> list[str]:
    """Import names of the pinned replay packages in ``requirements`` that this kernel has imported already: pip
    replaces their files, but the running kernel keeps the old version until it restarts."""
    pinned = [req.split("==", 1)[0] for req in requirements if "==" in req]
    return sorted(REPLAY_PACKAGES[name] for name in pinned
                  if name in REPLAY_PACKAGES and REPLAY_PACKAGES[name] in sys.modules)


def _pip_install(requirements: Sequence[str]) -> None:
    """Colab only: pip install (quiet) into the running kernel's environment."""
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", *requirements], check=True)
    importlib.invalidate_caches()


def _colab_requirements(root: Path, mode: str) -> tuple[list[str], str]:
    """What the Colab runtime still needs, and a note saying why ("" when nothing is missing).

    - Replay: the replay packages at the versions recorded in the bundle manifest, so text extraction matches the
      build machine and replay rebuilds identical cells.
    - Live modes: also the bundled wheel's live and ai extras.
    """
    manifest = read_manifest(root) if (root / MANIFEST_NAME).is_file() else {}
    installed = package_versions(manifest.get("replay_packages", {}))
    pins = [f"{name}=={version}" for name, version in sorted(manifest.get("replay_packages", {}).items())
            if installed.get(name) != version]
    needed = [*LIVE_MODULES, *(AI_MODULES if mode == "live_ai" else ())]
    if mode == "replay" or not _missing_modules(needed):
        return pins, (f"installed {', '.join(pins)} (the build machine's versions)" if pins else "")
    wheel = manifest.get("wheel", "")
    if not wheel:
        raise RuntimeError(f"{mode} needs the live extras, and this bundle has no wheel to install them from")
    extras = f"{root.joinpath(*PurePosixPath(wheel).parts)}[live,ai]"
    return [extras, *pins], f"installed the live and ai extras for {mode}"


def schema_drift(root: Path) -> list[str]:
    """Roles whose LLM response schema hash in this kernel differs from the one the bundle manifest records (empty
    for a repository checkout, or a manifest that records none). A drifted schema misses every cached reply."""
    if not (root / MANIFEST_NAME).is_file():
        return []
    recorded = read_manifest(root).get("llm_schemas") or {}
    current = llm_schema_hashes()
    return sorted(role for role, sha in recorded.items() if current.get(role) != sha)


def _check_bundle(root: Path) -> str:
    """Verify an extracted bundle once (a marker records the manifest that passed) and describe it."""
    manifest_file = root / MANIFEST_NAME
    if not manifest_file.is_file():
        return "repository checkout (no BUNDLE_MANIFEST.json)"
    manifest_sha = sha256_file(manifest_file)
    marker = root / VERIFIED_MARKER
    count = len(read_manifest(root)["files"])
    if marker.is_file() and marker.read_text(encoding="utf-8").strip() == manifest_sha:
        return f"verified at first setup ({count} files)"
    problems = verify_bundle(root)
    if problems:
        raise RuntimeError(f"the extracted bundle does not match its manifest ({len(problems)} problems, e.g. "
                           f"{'; '.join(problems[:5])}). Upload footprint_bundle.zip again.")
    marker.write_text(manifest_sha + "\n", encoding="utf-8", newline="\n")
    return f"verified ({count} files, SHA-256 per file)"


def notebook_setup(*, mode: str = "replay", vendors: Sequence[str] | None = None, team: str | None = None,
                   load_secrets: bool = True, root: str | Path | None = None) -> NotebookEnv:
    """Prepare the kernel for the demo notebook and describe the environment.

    - Finds the checkout or extracted bundle and makes it the working directory (relative paths such as
      ``evidence`` and ``seeds`` resolve there), and checks that the imported package can read ``config/``.
    - Verifies an extracted bundle against BUNDLE_MANIFEST.json (once).
    - ``load_secrets``: locally, reads ``.env``; on Colab, reads the Colab secrets the mode uses (COLAB_SECRETS:
      the Gemini key only for live_ai). Values go to os.environ only and are never shown.
    - Colab: installs the replay packages at the versions the bundle manifest records (identical text
      extraction, so identical cells), and for live modes the bundled wheel's live and ai extras. Locally,
      `uv sync --all-extras` provides everything.
    - Warns (a note starting "WARNING") when the kernel's LLM response schemas differ from the ones the bundle
      manifest records, since the LLM cache would then miss for them (``schema_drift``).
    - Live modes run one vendor at a time.
    """
    if mode not in MODES:
        raise ValueError(f"MODE must be one of {', '.join(MODES)}, not {mode!r}")
    if mode != "replay" and (not vendors or len(vendors) != 1):
        raise ValueError("live modes run one vendor at a time in the notebook: set VENDORS = ['V-00x']")
    colab = in_colab()
    base = Path(root).resolve() if root is not None else find_root()
    os.chdir(base)
    from footprint.criticality import DEFAULT_RUBRIC_PATH

    if not Path(DEFAULT_RUBRIC_PATH).is_file():
        raise RuntimeError(f"footprint was imported from {Path(__file__).resolve().parent}, where config/ cannot be "
                           f"found; put {base / 'src'} first on sys.path before importing footprint (see the setup "
                           "cell)")
    notes: list[str] = []
    bundle = _check_bundle(base)
    if load_secrets:
        if colab:
            for name in COLAB_SECRETS[mode]:
                value = None if os.environ.get(name) else colab_secret(name)
                if value:
                    os.environ[name] = value
        else:
            load_dotenv_names(base / ".env")
    if colab:
        requirements, why = _colab_requirements(base, mode)
        if requirements:
            stale = _imported_already(requirements)
            _pip_install(requirements)
            notes.append(why)
            if stale:
                raise RuntimeError(f"installed the build machine's versions of {', '.join(stale)}, but this session "
                                   "had imported them already and keeps the old ones: restart the session (Runtime > "
                                   "Restart session), then run the notebook again from the top (the bundle stays "
                                   "extracted)")
        if mode != "replay":
            notes.append("Colab has no Chromium for excerpt screenshots unless you run "
                         "`!python -m playwright install --with-deps chromium` first")
    drifted = schema_drift(base)
    if drifted:
        pins = read_manifest(base).get("replay_packages", {})
        wanted = ", ".join(f"{n}=={pins[n]}" for n in ("pydantic", "pydantic-core") if n in pins)
        notes.append(f"WARNING: the LLM response schemas for {', '.join(drifted)} differ from the build machine's, so "
                     "every cached Gemini reply for them misses and those cells fall back to the rules; install "
                     + (wanted or "the build machine's pydantic") + " and restart the session")
    if mode != "replay":
        needed = [*LIVE_MODULES, *(AI_MODULES if mode == "live_ai" else ())]
        if _missing_modules(needed):
            raise RuntimeError(f"{mode} needs the optional extras ({', '.join(_missing_modules(needed))} missing): "
                               "run `uv sync --all-extras` and restart the kernel")
        notes.append("live collection is polite (ToS register, robots.txt, rate limits) and takes minutes")
        if mode == "live_ai" and not os.environ.get("GEMINI_API_KEY"):
            notes.append("no Gemini key: live_ai falls back to the rules")
    resolved_team = team or os.environ.get("FOOTPRINT_TEAM_NAME") or PLACEHOLDER_TEAM
    _pandas_display()
    return NotebookEnv(root=str(base), colab=colab, mode=mode, team=resolved_team,
                       gemini_key=bool(os.environ.get("GEMINI_API_KEY")), bundle=bundle, notes=tuple(notes))


def _pandas_display() -> None:
    try:
        import pandas as pd
    except ImportError:
        return
    pd.set_option("display.max_colwidth", None)
    pd.set_option("display.max_rows", 500)


# --------------------------------------------------------------------------- workbook upload


class WorkbookInput:
    """The workbook the notebook assesses: an upload when there is one, else the default path.

    With ``ask=True`` Colab opens its file picker at once; local Jupyter shows an ipywidgets upload button, read
    whenever ``source()`` is called (so pick the file before running the next cell).
    """

    def __init__(self, default: str | Path, *, ask: bool = False) -> None:
        self.default = Path(default)
        self.uploaded_name = ""
        self._uploaded: bytes | None = None
        self.widget: Any = None
        self.message = ""
        if not ask:
            return
        if in_colab():
            from google.colab import files  # type: ignore[import-not-found]

            picked = sorted((name, data) for name, data in files.upload().items() if name.lower().endswith(".xlsx"))
            if picked:
                self.uploaded_name, self._uploaded = picked[0][0], bytes(picked[0][1])
            else:
                self.message = "No .xlsx uploaded: using the default workbook."
            return
        try:
            import ipywidgets
        except ImportError:
            self.message = "ipywidgets is not installed: using the default workbook."
            return
        self.widget = ipywidgets.FileUpload(accept=".xlsx", multiple=False, description="Workbook (.xlsx)")

    def _widget_file(self) -> tuple[str, bytes] | None:
        value = getattr(self.widget, "value", None)
        if not value:
            return None
        if isinstance(value, dict):  # ipywidgets 7: {name: {"content": bytes, ...}}
            name, info = sorted(value.items())[0]
            return str(name), bytes(info["content"])
        first = value[0]  # ipywidgets 8: ({"name": ..., "content": memoryview, ...},)
        return str(first["name"]), bytes(first["content"])

    def source(self) -> bytes | Path:
        """The uploaded bytes, else the default path (FileNotFoundError when it does not exist)."""
        picked = self._widget_file()
        if picked is not None:
            return picked[1]
        if self._uploaded is not None:
            return self._uploaded
        if not self.default.is_file():
            raise FileNotFoundError(f"no workbook uploaded and {self.default} does not exist")
        return self.default

    @property
    def name(self) -> str:
        picked = self._widget_file()
        if picked is not None:
            return picked[0]
        return self.uploaded_name or self.default.name

    def sha256(self) -> str:
        src = self.source()
        return sha256_bytes(src) if isinstance(src, bytes) else sha256_file(src)

    def describe(self) -> str:
        try:
            origin = "uploaded" if isinstance(self.source(), bytes) else f"from {self.default}"
            text = f"Workbook: {self.name} ({origin}; SHA-256 {self.sha256()[:16]}...)"
        except FileNotFoundError as exc:
            text = f"Workbook: none yet ({exc})"
        if self.widget is not None and self._widget_file() is None:
            text += ". Pick a file above to replace it before running the next cell."
        return " ".join(part for part in (self.message, text) if part)

    def __repr__(self) -> str:
        return self.describe()

    def _ipython_display_(self) -> None:
        from IPython.display import display

        if self.widget is not None:
            display(self.widget)
        print(self.describe())


def _source_of(workbook: WorkbookInput | Any) -> Any:
    return workbook.source() if isinstance(workbook, WorkbookInput) else workbook


def read_inventory(workbook: WorkbookInput | Any) -> Any:
    """Read and validate the vendor workbook: print each issue, raise ValueError when there are errors."""
    from footprint.workbook import read_workbook

    data = read_workbook(_source_of(workbook))
    for issue in data.issues:
        where = f" {issue.cell}" if issue.cell else ""
        print(f"[{issue.severity}] {issue.code}{where}: {issue.message}")
    if not data.ok:
        raise ValueError("the workbook failed validation; fix the errors above and upload it again")
    print(f'Sheet "{data.sheet_name}": {len(data.vendors)} vendors'
          + (f" (worked example {data.example.vendor_id} skipped)" if data.example else ""))
    return data


def tier_overrides() -> Any:
    """The HC1 override store (review/overrides.jsonl or $FOOTPRINT_OVERRIDES), or None when there is none."""
    from footprint.review import OverrideStore, default_overrides_path

    path = default_overrides_path()
    return OverrideStore(path) if path.exists() else None


# --------------------------------------------------------------------------- tables


def _frame(rows: list[dict[str, Any]], columns: Sequence[str]) -> Any:
    try:
        import pandas as pd
    except ImportError as exc:  # pragma: no cover - pandas ships with Colab and the `app` extra
        raise ImportError("the notebook tables need pandas: pip install pandas") from exc
    return pd.DataFrame(rows, columns=list(columns))


class TextTable:
    """Long texts (cells, rationales) as a wrapped table: HTML in Jupyter, plain text elsewhere."""

    def __init__(self, headers: Sequence[str], rows: Sequence[Sequence[Any]], title: str = "") -> None:
        self.headers = list(headers)
        self.rows = [["" if v is None else str(v) for v in row] for row in rows]
        self.title = title

    def __repr__(self) -> str:
        blocks = [self.title] if self.title else []
        for row in self.rows:
            blocks.append("\n".join(f"{h}: {v}" for h, v in zip(self.headers, row)))
        return "\n\n".join(blocks)

    def _repr_html_(self) -> str:
        cell = 'style="text-align:left;vertical-align:top;white-space:pre-wrap"'
        head = "".join(f"<th {cell}>{html.escape(h)}</th>" for h in self.headers)
        body = "".join("<tr>" + "".join(f"<td {cell}>{html.escape(v)}</td>" for v in row) + "</tr>"
                       for row in self.rows)
        caption = f"<caption>{html.escape(self.title)}</caption>" if self.title else ""
        return f"<table>{caption}<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"

    def to_frame(self) -> Any:
        return _frame([dict(zip(self.headers, row)) for row in self.rows], self.headers)


def _level(factor: Any) -> str:
    return f"{factor.level}-P" if str(factor.anchor).endswith("-P") else str(factor.level)


CRITICALITY_COLUMNS = ("Vendor ID", "Vendor", "O", "D", "P", "R", "V", "Score", "Floors", "Tier (L)", "Override",
                       "Stable under ±1 weights")


def criticality_frame(assessments: Sequence[Any], example: Any = None) -> Any:
    """One row per vendor: factor levels, score, floors and tier (column L); the V-000 example last, if given."""
    rows = []
    for a, calibration in [*((a, False) for a in assessments), *([(example, True)] if example is not None else [])]:
        c = a.criticality
        rows.append({
            "Vendor ID": a.profile.vendor_id + (" (calibration)" if calibration else ""),
            "Vendor": a.profile.name,
            **{code: _level(c.factors[code]) for code in ("O", "D", "P", "R", "V")},
            "Score": c.score,
            "Floors": ", ".join(c.floors_fired) or "-",
            "Tier (L)": c.tier.value,
            "Override": f"{c.override.tier.value} by {c.override.analyst}: {c.override.reason}" if c.override else "",
            "Stable under ±1 weights": "yes" if c.perturbation_stable else "no",
        })
    return _frame(rows, CRITICALITY_COLUMNS)


def rationale_table(assessments: Sequence[Any]) -> TextTable:
    """Column M (criticality rationale) for each vendor, wrapped."""
    from footprint.criticality import render_rationale

    rows = [(a.profile.vendor_id, a.profile.name, a.criticality.tier.value,
             render_rationale(a.profile, a.criticality)) for a in assessments]
    return TextTable(["Vendor ID", "Vendor", "Tier (L)", "Criticality rationale (M)"], rows)


DEPTH_COLUMNS = ("Vendor ID", "Tier", "Depth (N label)", "Mandatory families", "Other families", "Modifiers",
                 "Fetch budget", "Gemini calls", "Analyst minutes", "Saturation window", "Reserved for Meridian")


def depth_frame(assessments: Sequence[Any]) -> Any:
    """The depth plan per vendor (column N): mandatory families with their mode and cap, budgets and modifiers."""
    rows = []
    for a in assessments:
        plan = a.depth
        mandatory = [f"{fp.family.value} ({fp.mode}, cap {fp.cap})" for fp in plan.families if fp.mandatory]
        other = [f"{fp.family.value} ({fp.mode})" for fp in plan.families
                 if not fp.mandatory and fp.mode not in ("not_used", "not_applicable")]
        rows.append({
            "Vendor ID": plan.vendor_id, "Tier": plan.tier.value, "Depth (N label)": plan.label,
            "Mandatory families": "; ".join(mandatory) or "-", "Other families": "; ".join(other) or "-",
            "Modifiers": ", ".join(plan.modifiers) or "-", "Fetch budget": plan.discretionary_fetches,
            "Gemini calls": plan.gemini_calls, "Analyst minutes": plan.analyst_minutes,
            "Saturation window": plan.saturation_window, "Reserved for Meridian": "; ".join(plan.reserved_for_meridian),
        })
    return _frame(rows, DEPTH_COLUMNS)


def _findings(result: AssessmentResult, *, example: bool = False) -> list[tuple[VendorFindings, bool]]:
    found = [(f, False) for f in result.vendors]
    if example and result.example is not None:
        found.append((result.example, True))
    return found


def _eids(findings: VendorFindings, keys: Sequence[str]) -> str:
    """E-IDs of the items with these keys (the first 12 hex of the key while no E-ID is assigned)."""
    ids = []
    for key in keys:
        item = findings.item(key)
        if item is not None and item.evidence_id:
            ids.append(item.evidence_id)
        else:
            ids.append(key[:12])
    return ", ".join(ids)


SUMMARY_COLUMNS = ("Vendor ID", "Vendor", "Tier (L)", "AI usage (O)", "Verdict", "Likelihood", "Confidence",
                   "AI risk (S)", "Ceiling if confirmed", "Score (ARP of 18)", "Decisive evidence", "Flip condition")


def summary_frame(result: AssessmentResult) -> Any:
    """Answer first: tier, AI usage, verdict wording, risk class, ceiling and flip condition per vendor."""
    rows = []
    for f, calibration in _findings(result, example=True):
        risk = f.risk
        rows.append({
            "Vendor ID": f.vendor_id + (" (calibration)" if calibration else ""),
            "Vendor": f.profile.name,
            "Tier (L)": f.criticality.tier.value,
            "AI usage (O)": f.verdict.column_o,
            "Verdict": f.verdict.label,
            "Likelihood": f.verdict.likelihood,
            "Confidence": f.verdict.confidence,
            "AI risk (S)": risk.final_class + (" (provisional)" if risk.provisional else ""),
            "Ceiling if confirmed": risk.ceiling_class or "-",
            "Score (ARP of 18)": risk.arp,
            "Decisive evidence": _eids(f, f.verdict.decisive) or "-",
            "Flip condition": risk.flip_condition or "-",
        })
    return _frame(rows, SUMMARY_COLUMNS)


EVIDENCE_COLUMNS = ("Evidence ID", "Vendor ID", "Role", "Strength", "Tags", "Source type", "Publisher", "Published",
                    "URL", "Excerpt", "Method", "Review")


def _evidence_row(item: EvidenceItem) -> dict[str, Any]:
    return {
        "Evidence ID": item.evidence_id or item.item_key[:12], "Vendor ID": item.vendor_id, "Role": item.role,
        "Strength": item.strength, "Tags": item.tag_string, "Source type": item.source_type,
        "Publisher": item.publisher, "Published": item.published or f"retrieved {item.retrieved_at[:10]}",
        "URL": item.url, "Excerpt": item.excerpt, "Method": item.method,
        "Review": item.review_status + (f" ({item.review_reason})" if item.review_reason else ""),
    }


def evidence_frame(result: AssessmentResult, vendor_id: str | None = None, *, cited_only: bool = True) -> Any:
    """Evidence items in E-ID order: the cited ones by default, every Evidence Log row with cited_only=False."""
    rows = [_evidence_row(item) for f, _ in _findings(result) if vendor_id in (None, f.vendor_id)
            for item in (f.cited() if cited_only else f.evidence)]
    return _frame(rows, EVIDENCE_COLUMNS)


RISK_COLUMNS = ("Vendor ID", "E (exposure)", "K (decision impact)", "TP (tier)", "TG (gaps)", "ARP", "Base class",
                "Gate met", "Escalators fired", "Cap", "Final class (S)", "Ceiling", "Missing checks", "Themes")


def _input(value: int, assumed: bool) -> str:
    return f"{value} (assumed)" if assumed else str(value)


def risk_frame(result: AssessmentResult) -> Any:
    """The ordered risk steps per vendor: ARP = 2E + 2K + TP + TG, gate, escalators, cap and ceiling."""
    rows = []
    for f, calibration in _findings(result, example=True):
        r, i = f.risk, f.risk.inputs
        rows.append({
            "Vendor ID": f.vendor_id + (" (calibration)" if calibration else ""),
            "E (exposure)": _input(i.e, i.e_assumed), "K (decision impact)": _input(i.k, i.k_assumed),
            "TP (tier)": i.tp, "TG (gaps)": i.tg, "ARP": r.arp, "Base class": r.base_class,
            "Gate met": "yes" if r.gate_met else "no", "Escalators fired": ", ".join(r.escalators_fired) or "-",
            "Cap": r.cap or "-", "Final class (S)": r.final_class + (" (provisional)" if r.provisional else ""),
            "Ceiling": r.ceiling_class or "-", "Missing checks": ", ".join(i.missing_gaps) or "-",
            "Themes": ", ".join(r.themes) or "-",
        })
    return _frame(rows, RISK_COLUMNS)


COVERAGE_COLUMNS = ("Coverage ID", "Vendor ID", "Family", "Mandatory", "Status", "Collector", "Endpoint / query",
                    "Requests", "Cap", "Documents", "AI passages", "Note")


def coverage_frame(result: AssessmentResult, vendor_id: str | None = None) -> Any:
    """The Coverage Log (negative evidence): one row per family attempt, with its {vendor}-C-nn id."""
    rows = []
    for f, _ in _findings(result):
        if vendor_id not in (None, f.vendor_id):
            continue
        for n, e in enumerate(f.coverage, start=1):
            rows.append({
                "Coverage ID": f"{f.vendor_id}-C-{n:02d}", "Vendor ID": e.vendor_id, "Family": e.family.value,
                "Mandatory": "yes" if e.mandatory else "no", "Status": e.status.value, "Collector": e.collector,
                "Endpoint / query": e.endpoint, "Requests": e.requests_used, "Cap": e.cap, "Documents": e.documents,
                "AI passages": e.ai_passages, "Note": e.note,
            })
    return _frame(rows, COVERAGE_COLUMNS)


def _headers() -> dict[str, str]:
    from footprint.workbook import HEADER_ALIASES

    return {f: HEADER_ALIASES[f][0] for f in STUDENT_FIELDS}


def cells_table(result: AssessmentResult, vendor_id: str | None = None) -> TextTable:
    """Columns L-V for one vendor as they will be written (the first vendor when ``vendor_id`` is not in the run)."""
    if not result.vendors:
        return TextTable(["Column", "Header", "Text"], [], title="No vendors in this run")
    f = result.vendor(vendor_id or "") or result.vendors[0]
    headers = _headers()
    rows = [(letter, headers[name], getattr(f.cells, name)) for letter, name in zip(COLUMN_LETTERS, STUDENT_FIELDS)]
    return TextTable(["Column", "Header", "Text"], rows, title=f"{f.vendor_id} {f.profile.name}: columns L-V")


# --------------------------------------------------------------------------- run, export, download


def progress_printer() -> Callable[[str, float], None]:
    """A ``progress(message, fraction)`` callback for run_assessment that prints one line per stage."""

    def report(message: str, fraction: float) -> None:
        pct = max(0.0, min(1.0, float(fraction))) * 100
        print(f"[{pct:5.1f}%] {message}", flush=True)

    return report


def output_path(out_dir: str | Path, result: AssessmentResult) -> Path:
    """Where the exported workbook goes: <out_dir>/Meridian_Vendor_Assessment_<run_id>.xlsx (folder created)."""
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"Meridian_Vendor_Assessment_{result.run_id}.xlsx"


def offer_download(path: str | Path) -> Any:
    """Colab: start the browser download. Local Jupyter: a link to the file (and its absolute path)."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"{p} does not exist; run the export cell first")
    if in_colab():
        from google.colab import files  # type: ignore[import-not-found]

        files.download(str(p))
        return f"Downloading {p.name} ({p.stat().st_size:,} bytes)"
    print(f"Exported workbook: {p.resolve()} ({p.stat().st_size:,} bytes)")
    try:
        from IPython.display import FileLink
    except ImportError:
        return str(p.resolve())
    try:
        link = os.path.relpath(p)
    except ValueError:  # another drive (Windows): no relative link exists
        link = str(p.resolve())
    return FileLink(link)


# --------------------------------------------------------------------------- verification


def _bytes_of(source: Any) -> bytes:
    if isinstance(source, WorkbookInput):
        source = source.source()
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, os.PathLike)):
        return Path(source).read_bytes()
    position = source.tell()
    source.seek(0)
    data = source.read()
    source.seek(position)
    return data


def _norm(value: object) -> str:
    """A cell value for comparison: None and "" are the same (an untouched student cell), and characters an .xlsx
    cannot store are dropped, as the writer drops them."""
    from footprint.workbook import _storable  # the writer's own rule, so a dropped control character is no diff

    return "" if value is None else _storable(str(value))[0]


def _sheet_cells(data: bytes) -> tuple[dict[str, dict[str, object]], list[str]]:
    """vendor_id -> {field: value of L-V} and the sheet names of an exported workbook."""
    import openpyxl

    from footprint.workbook import read_workbook

    inventory = read_workbook(data)
    book = openpyxl.load_workbook(io.BytesIO(data))
    ws = book[inventory.sheet_name]
    cells = {v.vendor_id: {name: ws[f"{inventory.column_map[name]}{v.row}"].value for name in STUDENT_FIELDS}
             for v in inventory.vendors if all(name in inventory.column_map for name in STUDENT_FIELDS)}
    return cells, list(book.sheetnames)


def reverify(item: EvidenceItem, store: Any) -> list[str]:
    """The demo's Re-verify (footprint.verify.reverify_item) as a list of problems; [] means the stored raw bytes
    and text still hash to the item's SHA-256 values and the excerpt is still the exact slice [start:end]."""
    from footprint.verify import reverify_item

    outcome = reverify_item(item, store)
    return [] if outcome.ok else (list(outcome.details) or ["re-verification failed"])


def verify_export(result: AssessmentResult, xlsx_in: Any, out_path: str | Path, *, reference: Any = None,
                  store: Any = None) -> Any:
    """Integrity checks on the exported workbook, as a table (Check, Result, Detail); ``attrs["ok"]`` is True
    when nothing failed.

    1. The input is the workbook that was assessed (SHA-256).
    2. Fidelity: the provided cells, the V-000 row, styles and sheets are unchanged (workbook.check_fidelity).
    3. Every vendor's L-V cells in the file equal the assessment's cells.
    4. The Evidence Log and Coverage Log sheets are present.
    5. Every cited evidence item re-verifies against the evidence store.
    6. With ``reference`` (an earlier export or the submitted workbook): replay reproduced its L-V cells.
    """
    from footprint.capture.store import EvidenceStore
    from footprint.workbook import check_fidelity

    checks: list[tuple[str, str, str]] = []

    def add(name: str, ok: bool | None, detail: str) -> None:
        checks.append((name, "SKIP" if ok is None else "PASS" if ok else "FAIL", detail))

    original = _bytes_of(xlsx_in)
    exported = Path(out_path).read_bytes()
    digest = sha256_bytes(original)
    add("Input is the workbook that was assessed", digest == result.input_sha256, f"SHA-256 {digest[:16]}...")

    try:
        problems = check_fidelity(original, exported)
        _, original_sheets = _sheet_cells(original)
    except Exception as exc:  # noqa: BLE001 - an unreadable input fails the check instead of the cell
        problems, original_sheets = [f"cannot compare with the input ({type(exc).__name__})"], []
    add("Provided cells, V-000, styles and sheets unchanged", not problems,
        "; ".join(problems[:3]) or "no differences outside L-V and the appended sheets")

    cells, sheet_names = _sheet_cells(exported)
    mismatched = []
    for f in result.vendors:
        written = cells.get(f.vendor_id)
        if written is None:
            mismatched.append(f"{f.vendor_id}: row not found")
            continue
        for letter, name in zip(COLUMN_LETTERS, STUDENT_FIELDS):
            expected = getattr(f.cells, name)
            if name == "assessed_by" and expected is None:
                continue  # rendered at export from the team name and as_of
            if _norm(expected) != _norm(written[name]):
                mismatched.append(f"{f.vendor_id} {letter}")
    add("Exported L-V cells equal the assessment", not mismatched,
        ", ".join(mismatched[:8]) or f"{len(result.vendors)} vendors x 11 columns")

    appended = [n for n in sheet_names if n not in original_sheets]
    missing = [n for n in REQUIRED_SHEETS if n not in sheet_names]
    detail = "appended: " + (", ".join(appended) or "none")
    add("Evidence Log and Coverage Log sheets present", not missing,
        f"missing {', '.join(missing)}; {detail}" if missing else detail)

    store = store if store is not None else EvidenceStore()
    cited = [item for f in result.vendors for item in f.cited()]
    if not cited:
        add(REVERIFY_CHECK, None, "no cited items in this run")
    elif importlib.util.find_spec("footprint.verify") is None:
        add(REVERIFY_CHECK, None, "footprint.verify is missing")
    else:
        failed = []
        for item in cited:
            problems_found = reverify(item, store)
            if problems_found:
                failed.append(f"{item.evidence_id or item.item_key[:12]}: {problems_found[0]}")
        add(REVERIFY_CHECK, not failed, "; ".join([f"{len(cited) - len(failed)} of {len(cited)} items", *failed[:3]]))

    if reference is not None:
        ref_cells, _ = _sheet_cells(_bytes_of(reference))
        differ = [f"{vid} {letter}" for vid in (f.vendor_id for f in result.vendors)
                  for letter, name in zip(COLUMN_LETTERS, STUDENT_FIELDS)
                  if _norm(ref_cells.get(vid, {}).get(name)) != _norm(cells.get(vid, {}).get(name))]
        add("Replay reproduces the reference workbook's L-V cells", not differ,
            ", ".join(differ[:8]) or "identical for every vendor in this run")

    columns = ("Check", "Result", "Detail")
    frame = _frame([dict(zip(columns, check)) for check in checks], columns)
    frame.attrs["ok"] = all(outcome != "FAIL" for _, outcome, _ in checks)
    return frame
