"""Shared helpers for the footprint Streamlit demo (Outcome 05): app/streamlit_app.py and app/pages/*.py.

The pages stay thin. Everything that can be decided without a browser lives here as plain functions: input
handling, the P1 fallback, run caching, evidence filters, review records, the why trace and the export, so
tests/unit/test_app.py checks it offline. The session and sidebar helpers at the end use ``st.session_state``.

Contracts: docs/contracts_p3.md section 15 (run_assessment, rescore_vendor, export_assessment and "UI usage") and
section 10 (review records). Until footprint.pipeline provides run_assessment, a run falls back to criticality and
depth (P1, columns L-N), so the app always starts.

The app serves localhost only and publishes nothing. It records analyst decisions in a review store (by default a
sandbox copy of review/) and never contacts a vendor: live runs go through the pipeline's ToS register -> robots.txt
-> rate-limit gates and the Gemini payload guard; the live DNS check asks the two allowlisted DoH resolvers only.
"""

from __future__ import annotations

import datetime as dt
import functools
import hashlib
import html
import io
import json
import os
import re
import shutil
import tempfile
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import streamlit as st

from footprint import pipeline, review
from footprint.criticality import FACTORS, Rubric, load_rubric, render_rationale
from footprint.depth import DepthConfig, load_depth_config, render_depth_cell
from footprint.models import (
    CITED_ROLES,
    GAP_KEYS,
    STUDENT_FIELDS,
    TIER_ORDER,
    AssessmentResult,
    CoverageEntry,
    CriticalityResult,
    DepthPlan,
    EvidenceItem,
    StudentCells,
    Tier,
    TierOverride,
    ValidationIssue,
    VendorFindings,
    VendorProfile,
    WorkbookData,
)
from footprint.review import OverrideStore, ReviewRecord
from footprint.workbook import DEFAULT_LENGTH_BUDGETS, HEADER_ALIASES, check_fidelity, read_workbook

# =========================================================================== paths and labels

APP_DIR = Path(__file__).resolve().parent
REPO_ROOT = APP_DIR.parent
BUNDLED_INPUT = REPO_ROOT / "data" / "input" / "Meridian_Vendor_Input.xlsx"
EVIDENCE_DIR = REPO_ROOT / "evidence"
SEEDS_DIR = REPO_ROOT / "seeds"
RUNS_DIR = REPO_ROOT / "runs"
DOTENV_PATH = REPO_ROOT / ".env"

SANDBOX_ENV = "FOOTPRINT_APP_SANDBOX"  # another sandbox directory (tests, a second demo)
NO_DOTENV_ENV = "FOOTPRINT_APP_NO_DOTENV"  # set to skip reading .env (hermetic tests)
TEAM_ENV = "FOOTPRINT_TEAM_NAME"
GEMINI_ENV = "GEMINI_API_KEY"
OVERRIDES_FILE = "overrides.jsonl"
REVIEWS_FILE = "reviews.jsonl"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MODES: tuple[str, ...] = ("replay", "live_rules", "live_ai")
MODE_LABELS: dict[str, str] = {"replay": "Replay", "live_rules": "Live rules", "live_ai": "Live AI"}
MODE_CAPTIONS: dict[str, str] = {
    "replay": "Offline, from the frozen evidence pack and caches. Default.",
    "live_rules": "Polite live collection, rules only.",
    "live_ai": "Live collection, rules plus Gemini on public text.",
}
STORE_LABELS: dict[str, str] = {"sandbox": "Sandbox (demo copy)", "project": "Project store (review/)"}

TIER_NAMES: tuple[str, ...] = tuple(t.value for t in reversed(TIER_ORDER))  # Critical .. Low
CLASS_COLORS: dict[str, str] = {
    "Critical": "red", "High": "orange", "Medium": "blue", "Low": "green", "None identified": "gray",
}
FAMILY_NAMES: dict[str, str] = {
    "LEG": "Legal and trust pages",
    "REG": "Regulatory filings (SEC)",
    "PRD": "Product pages and newsroom",
    "JOB": "Vendor job postings",
    "DNS": "DNS records",
    "IND": "Independent corroboration",
    "HIST": "Archive history (Wayback)",
    "EXEC": "Executive channels (manual)",
}
METHOD_LABELS: dict[str, str] = {
    "rule": "Rules",
    "rule+llm_agree": "Rules + Gemini agree",
    "llm_proposed_accepted": "Gemini proposal",
    "adjudicated": "Adjudicated",
}
REVIEW_LABELS: dict[str, str] = {"unreviewed": "Unreviewed", "accepted": "Accepted", "rejected": "Rejected"}
LABEL_KEYS: tuple[str, ...] = (
    "temporal", "action_level", "sp", "rl", "locus", "u_class", "ai_type", "claim_kind", "subject", "providers",
)
GAP_LABELS: dict[str, str] = {
    "t1": "AI use in the service disclosed by an accountable source (R3, SR A/B)",
    "t2": "AI providers named by the vendor itself",
    "t3": "Data-use terms for AI in a legal document",
    "t4": "Human oversight of AI output described",
    "t5": "AI governance attestation (NIST AI RMF, ISO/IEC 42001 or an AI policy)",
    "t6": "AI incident or change notification",
}

# HC2 reason codes. The record's reason is "{code}: {wording}", plus the analyst's note.
ACCEPT_CODES: dict[str, str] = {
    "ACC-VERIFIED": "quote checked against the captured source and supports the tags",
    "ACC-SCOPE": "statement concerns the vendor's service to Meridian",
    "ACC-LABELS": "labels adjudicated against the captured source",
}
REJECT_CODES: dict[str, str] = {
    "REJ-NOT-AI": "definition test fails: automation, rules or scheduling described as AI",
    "REJ-ENTITY": "different entity or a name collision",
    "REJ-SUBJECT": "statement is not about the vendor or its service",
    "REJ-MARKETING": "aspirational or marketing wording only",
    "REJ-STALE": "superseded or out-of-date statement",
    "REJ-DUPLICATE": "duplicate of another item from the same origin",
    "REJ-CONTEXT": "excerpt is misleading out of context",
}
OTHER_CODE = "OTHER"
OTHER_WORDING = "other reason (see note)"
MIN_REASON = review.MIN_REASON_CHARS

P1_NOTE = (
    "The full assessment (footprint.pipeline.run_assessment) is not available in this build, so this run holds "
    "criticality and depth only (columns L-N)."
)
EXPORT_P1_NOTE = (
    "The full export (footprint.pipeline.export_assessment) is not available in this build, so the workbook holds "
    "criticality and depth only (columns L-N), plus the Criticality Workings and Method & Legend sheets."
)


# =========================================================================== formatting

_MD_SPECIAL = re.compile(r"([\\`*_{}\[\]()#+\-.!|<>$~])")
_SECRET = re.compile(r"AIza[0-9A-Za-z_\-]{20,}|gh[pousr]_[0-9A-Za-z]{20,}|github_pat_[0-9A-Za-z_]{20,}")
_SPACE = re.compile(r"\s+")
_LEGAL_SUFFIX = re.compile(r",?\s+(?:Inc|L\.?L\.?C|Ltd|Limited|Corp|Corporation|plc|LLP|Company|Co)\.?$",
                           re.IGNORECASE)


def md_escape(text: object) -> str:
    """Text safe to put inside st.markdown: Markdown punctuation and '$' (LaTeX) are backslash-escaped."""
    return _MD_SPECIAL.sub(r"\\\1", str(text or ""))


def clip(text: object, limit: int) -> str:
    """One line, whitespace collapsed, cut at a word boundary with an ellipsis when longer than ``limit``."""
    flat = _SPACE.sub(" ", str(text or "")).strip()
    if len(flat) <= limit:
        return flat
    cut = flat[: max(limit - 1, 1)].rsplit(" ", 1)[0].rstrip(" ,;:")
    return (cut or flat[: max(limit - 1, 1)]) + "…"


def scrub(text: object) -> str:
    """An error message with anything key-like removed (API keys and tokens never reach the screen)."""
    out = _SECRET.sub("[redacted]", str(text or ""))
    key = os.environ.get(GEMINI_ENV, "").strip()
    if len(key) >= 8:
        out = out.replace(key, "[redacted]")
    return out


def dmy(iso: str) -> str:
    """'2026-10-02' -> '02-10-2026' (the workbook's date style); other text is returned unchanged."""
    try:
        return dt.date.fromisoformat(iso[:10]).strftime("%d-%m-%Y") if iso else ""
    except ValueError:
        return iso


def today() -> str:
    """Today's date for review records (record metadata only, never decision logic)."""
    return dt.date.today().isoformat()


def short_name(name: str) -> str:
    """'Financial Statement Services, Inc. (FSSI)' -> 'FSSI'; 'Fiserv, Inc.' -> 'Fiserv'."""
    acronym = re.search(r"\(([A-Z][A-Z0-9&]{1,7})\)", name or "")
    if acronym:
        return acronym.group(1)
    base = re.sub(r"\s*\([^)]*\)", "", name or "").strip(" ,")
    while True:
        trimmed = _LEGAL_SUFFIX.sub("", base).strip(" ,")
        if trimmed == base or not trimmed:
            break
        base = trimmed
    return clip(base or name or "vendor", 32)


# =========================================================================== environment


def load_env() -> None:
    """Read .env (KEY=VALUE; existing variables win) unless FOOTPRINT_APP_NO_DOTENV is set. Values are never shown."""
    if not os.environ.get(NO_DOTENV_ENV):
        pipeline.load_dotenv(DOTENV_PATH)


def gemini_configured() -> bool:
    """True when a Gemini key is set. Only this yes/no is ever displayed."""
    return bool(os.environ.get(GEMINI_ENV, "").strip())


def default_team() -> str:
    return os.environ.get(TEAM_ENV, "").strip()


# =========================================================================== input workbook


@dataclass(frozen=True)
class InputFile:
    """The workbook under assessment: its name, raw bytes and their SHA-256 (the run cache key)."""

    name: str
    data: bytes = field(repr=False)
    sha256: str
    source: str = "upload"  # "upload" or "bundled"

    @property
    def short_sha(self) -> str:
        return self.sha256[:12]


def make_input(name: str, data: bytes, source: str = "upload") -> InputFile:
    raw = bytes(data)
    return InputFile(name=name or "workbook.xlsx", data=raw, sha256=hashlib.sha256(raw).hexdigest(), source=source)


def bundled_input(path: Path = BUNDLED_INPUT) -> InputFile:
    """The input workbook shipped with the repo (data/input/Meridian_Vendor_Input.xlsx)."""
    return make_input(path.name, path.read_bytes(), "bundled")


def read_input(inp: InputFile) -> WorkbookData:
    """Read and validate the workbook; problems come back as WorkbookData.issues, never as an exception."""
    return read_workbook(inp.data)


def issue_rows(issues: Iterable[ValidationIssue]) -> list[dict[str, str]]:
    return [{"Severity": i.severity, "Code": i.code, "Cell": i.cell, "Message": i.message} for i in issues]


def example_cells(inp: InputFile, data: WorkbookData) -> dict[str, str]:
    """The worked example's (V-000) student cells as text, by field name, for the calibration comparison."""
    if data.example is None or not data.column_map:
        return {}
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(inp.data), read_only=True, data_only=True)
    try:
        ws = wb[data.sheet_name]
        row = data.example.row
        out: dict[str, str] = {}
        for fld in STUDENT_FIELDS:
            letter = data.column_map.get(fld)
            if letter:
                value = ws[f"{letter}{row}"].value
                out[fld] = "" if value is None else str(value)
        return out
    finally:
        wb.close()


# =========================================================================== review stores


@dataclass(frozen=True)
class ReviewStores:
    """The analyst review store in use: tier and risk-input overrides plus HC2 evidence reviews."""

    kind: str  # "sandbox" or "project"
    directory: Path
    overrides: OverrideStore
    reviews: OverrideStore

    @property
    def label(self) -> str:
        return STORE_LABELS.get(self.kind, self.kind)


def project_store_paths() -> tuple[Path, Path]:
    """(overrides, reviews) of the project review store: $FOOTPRINT_OVERRIDES or review/overrides.jsonl, and
    review/reviews.jsonl."""
    return review.default_overrides_path(), review.REVIEWS_PATH


def sandbox_dir() -> Path:
    """$FOOTPRINT_APP_SANDBOX, else review/sandbox/ (git-ignored)."""
    return Path(os.environ.get(SANDBOX_ENV) or (review.REVIEW_DIR / "sandbox"))


def ensure_sandbox(directory: Path | None = None, *, reset: bool = False) -> Path:
    """Create the demo sandbox: copies of the project store's overrides.jsonl and reviews.jsonl.

    Existing sandbox files are kept, so demo decisions survive a page reload. ``reset`` restores the project
    copies (a file the project store lacks is removed from the sandbox). The project store is never written.
    """
    target_dir = Path(directory) if directory is not None else sandbox_dir()
    pairs = [(source, target_dir / name)
             for source, name in zip(project_store_paths(), (OVERRIDES_FILE, REVIEWS_FILE))]
    if any(source.resolve() == target.resolve() for source, target in pairs):
        raise ValueError(f"The sandbox {target_dir} is the project review store; set {SANDBOX_ENV} to another folder.")
    target_dir.mkdir(parents=True, exist_ok=True)
    for source, target in pairs:
        if target.exists() and not reset:
            continue
        if source.exists():
            shutil.copyfile(source, target)
        elif reset:
            target.unlink(missing_ok=True)
    return target_dir


def review_stores(kind: str = "sandbox", directory: Path | None = None) -> ReviewStores:
    """The sandbox (default; created on first use) or the project review store."""
    if kind == "project":
        overrides, reviews = project_store_paths()
        return ReviewStores("project", overrides.parent, OverrideStore(overrides), OverrideStore(reviews))
    if kind != "sandbox":
        raise ValueError(f"unknown review store {kind!r}")
    path = ensure_sandbox(directory)
    return ReviewStores("sandbox", path, OverrideStore(path / OVERRIDES_FILE), OverrideStore(path / REVIEWS_FILE))


def _require_reason(reason: str, analyst: str) -> tuple[str, str]:
    if not (analyst or "").strip():
        raise ValueError("Enter the analyst's name or initials in the sidebar first.")
    if len((reason or "").strip()) < MIN_REASON:
        raise ValueError(f"Give a reason of at least {MIN_REASON} characters.")
    return reason.strip(), analyst.strip()


def record_tier_override(store: OverrideStore, vendor_id: str, tier: str, reason: str, analyst: str, *,
                         date: str | None = None) -> ReviewRecord:
    """HC1: append a tier override (footprint.review) for one vendor."""
    reason, analyst = _require_reason(reason, analyst)
    if tier not in TIER_NAMES:
        raise ValueError(f"Unknown tier {tier!r}.")
    override = TierOverride(vendor_id=vendor_id, tier=Tier(tier), reason=reason, analyst=analyst, date=date or today())
    return store.add_tier_override(override)


def tier_override_map(store: OverrideStore | None, vendor_ids: Iterable[str]) -> dict[str, str]:
    """vendor_id -> overridden tier ('' when none): what a run was computed with, to detect stale runs."""
    out: dict[str, str] = {}
    for vid in vendor_ids:
        override = store.tier_override(vid) if store is not None else None
        out[vid] = override.tier.value if override else ""
    return out


def review_reason(code: str, note: str = "") -> str:
    """'ACC-VERIFIED: quote checked ...' plus ' - note'; OTHER needs a note of 10 or more characters."""
    note = (note or "").strip()
    if code == OTHER_CODE:
        if len(note) < MIN_REASON:
            raise ValueError(f"Explain an 'other' reason in a note of at least {MIN_REASON} characters.")
        return f"{OTHER_CODE}: {note}"
    wording = ACCEPT_CODES.get(code) or REJECT_CODES.get(code)
    if wording is None:
        raise ValueError(f"Unknown reason code {code!r}.")
    return f"{code}: {wording}" + (f" - {note}" if note else "")


def evidence_review_record(vendor_id: str, item_key: str, decision: str, code: str, note: str, analyst: str, *,
                           labels: Mapping[str, str] | None = None, date: str | None = None) -> ReviewRecord:
    """HC2: an evidence_review record (contracts section 10). ``labels`` adjudicates rule vs Gemini labels."""
    if decision not in ("accepted", "rejected"):
        raise ValueError("The decision is 'accepted' or 'rejected'.")
    if decision == "accepted" and code in REJECT_CODES:
        raise ValueError("Pick an acceptance reason (ACC-…) to accept an item.")
    if decision == "rejected" and code in ACCEPT_CODES:
        raise ValueError("Pick a rejection reason (REJ-…) to reject an item.")
    reason, analyst = _require_reason(review_reason(code, note), analyst)
    extra: dict[str, Any] = {"labels": dict(labels)} if labels else {}
    return ReviewRecord(kind="evidence_review", vendor_id=vendor_id, key=item_key, value=decision, reason=reason,
                        analyst=analyst, date=date or today(), extra=extra)


def risk_override_record(vendor_id: str, key: str, value: str, reason: str, analyst: str, *,
                         evidence: Sequence[str] = (), date: str | None = None) -> ReviewRecord:
    """A risk_input_override record: e/k take '0'..'3'; t1..t6 take 'missing' or 'closed'. Evidence (item keys)
    makes the value evidenced rather than assumed."""
    if key in ("e", "k"):
        if value not in ("0", "1", "2", "3"):
            raise ValueError(f"{key.upper()} takes 0, 1, 2 or 3.")
    elif key in GAP_KEYS:
        if value not in ("missing", "closed"):
            raise ValueError(f"Check {key} is 'missing' or 'closed'.")
    else:
        raise ValueError(f"Unknown risk input {key!r}.")
    reason, analyst = _require_reason(reason, analyst)
    extra: dict[str, Any] = {"evidence": sorted(set(evidence))} if evidence else {}
    return ReviewRecord(kind="risk_input_override", vendor_id=vendor_id, key=key, value=value, reason=reason,
                        analyst=analyst, date=date or today(), extra=extra)


# =========================================================================== criticality and depth


@functools.lru_cache(maxsize=1)
def rubric() -> Rubric:
    return load_rubric()


@functools.lru_cache(maxsize=1)
def depth_config() -> DepthConfig:
    return load_depth_config()


def assess_p1(data: WorkbookData, overrides: OverrideStore | None = None) -> tuple[list[Any], Any]:
    """P1: criticality (with any HC1 override) and depth for every vendor, plus the V-000 calibration example."""
    return pipeline.assess_profiles(data, overrides), pipeline.assess_example(data)


def rationale(profile: VendorProfile, result: CriticalityResult) -> str:
    """Column M, as the pipeline writes it."""
    return render_rationale(profile, result, rubric())


def level_label(result: CriticalityResult, code: str) -> str:
    fs = result.factors[code]
    spec = rubric().factors[code]
    anchor = spec.anchor(fs.anchor)
    return anchor.label if anchor is not None and anchor.label else spec.labels[fs.level]


def tier_summary_rows(assessments: Sequence[Any]) -> list[dict[str, Any]]:
    rows = []
    for a in assessments:
        c = a.criticality
        rows.append({
            "Vendor ID": a.profile.vendor_id,
            "Vendor": a.profile.name,
            "Tier": c.tier.value,
            "Score": c.score,
            "Tier by score": c.score_tier.value,
            "Floors": ", ".join(c.floors_fired) or "-",
            "Computed tier": c.computed_tier.value,
            "Override": f"{c.override.tier.value} ({c.override.analyst})" if c.override else "",
            "Depth": a.depth.label,
        })
    return rows


def factor_rows(result: CriticalityResult) -> list[dict[str, Any]]:
    rows = []
    for code in FACTORS:
        fs = result.factors[code]
        rows.append({
            "Factor": f"{code} · {rubric().factors[code].name}",
            "Level": f"{fs.level}/4 · {level_label(result, code)}",
            "Anchor": fs.anchor,
            "Trigger (profile words)": fs.trigger or "-",
            "Points": f"{fs.weight} × {fs.level} = {fs.points}",
            "Note": fs.note,
        })
    return rows


def floor_rows(result: CriticalityResult) -> list[dict[str, str]]:
    rows = []
    for fid in result.floors_fired:
        try:
            floor = rubric().floor(fid)
        except ValueError:
            rows.append({"Floor": fid, "Name": "", "At least": "", "Rule": ""})
            continue
        rows.append({"Floor": fid, "Name": floor.name, "At least": floor.tier.value, "Rule": floor.description})
    return rows


def depth_rows(assessments: Sequence[Any]) -> list[dict[str, Any]]:
    rows = []
    for a in assessments:
        plan: DepthPlan = a.depth
        rows.append({
            "Vendor ID": plan.vendor_id,
            "Tier": plan.tier.value,
            "Depth": plan.label,
            "Mandatory families": ", ".join(f.family.value for f in plan.families if f.mandatory),
            "Modifiers": ", ".join(plan.modifiers) or "-",
            "Fetch budget": plan.discretionary_fetches,
            "Gemini calls": plan.gemini_calls,
            "Analyst time": f"{plan.analyst_minutes} min",
        })
    return rows


def family_rows(plan: DepthPlan) -> list[dict[str, Any]]:
    return [{
        "Family": f.family.value,
        "Source": FAMILY_NAMES.get(f.family.value, f.family.value),
        "Mandatory": "Yes" if f.mandatory else "No",
        "Mode": f.mode,
        "Cap": f.cap,
        "Why": f.reason,
    } for f in plan.families]


def modifier_lines(plan: DepthPlan) -> list[str]:
    cfg = depth_config()
    out = []
    for mod in plan.modifiers:
        policy = cfg.modifiers.get(mod)
        out.append(f"{mod} ({policy.title}): {policy.note}" if policy else mod)
    return out


def depth_cell(plan: DepthPlan, coverage: Sequence[CoverageEntry] | None = None) -> str:
    """Column N as the pipeline renders it (with the coverage sentence when coverage is known)."""
    try:
        return render_depth_cell(plan, coverage if coverage else None)
    except ValueError as exc:
        return f"(column N could not be rendered: {exc})"


_COLLECTION_RUN = re.compile(r"^(?P<vid>.+)-\d{8}-[0-9a-f]{8}$")


def latest_collection_coverage(vendor_id: str, runs_dir: Path = RUNS_DIR) -> list[CoverageEntry]:
    """Coverage Log of the vendor's newest P2 collection run (runs/<vendor>-<date>-<sha8>/), or []."""
    if not runs_dir.is_dir():
        return []
    best: tuple[str, Path] | None = None
    for d in runs_dir.iterdir():
        match = _COLLECTION_RUN.match(d.name)
        if not d.is_dir() or match is None or match["vid"] != vendor_id or not (d / "coverage.jsonl").exists():
            continue
        stamp = ""
        try:
            stamp = str(json.loads((d / "manifest.json").read_text(encoding="utf-8")).get("created_at", ""))
        except (OSError, ValueError, AttributeError):
            pass
        if best is None or (stamp, d.name) > (best[0], best[1].name):
            best = (stamp, d)
    if best is None:
        return []
    try:
        return pipeline.load_run_coverage(best[1].name, runs_dir)
    except (OSError, ValueError):
        return []


def coverage_rows(entries: Iterable[CoverageEntry]) -> list[dict[str, Any]]:
    return [{
        "Family": e.family.value,
        "Mandatory": "Yes" if e.mandatory else "No",
        "Status": e.status.value,
        "Collector": e.collector,
        "Endpoint": e.endpoint,
        "Requests": f"{e.requests_used}/{e.cap}" if e.cap else str(e.requests_used),
        "Documents": e.documents,
        "AI passages": e.ai_passages,
        "Note": e.note,
    } for e in entries]


# =========================================================================== pipeline access and runs


@dataclass(frozen=True)
class PipelineApi:
    """The P3/P4 entry points of footprint.pipeline, each None until the integrator provides it."""

    run_assessment: Callable[..., Any] | None
    rescore_vendor: Callable[..., Any] | None
    export_assessment: Callable[..., Any] | None
    input_error: type[BaseException] | None

    @property
    def full(self) -> bool:
        return self.run_assessment is not None


def pipeline_api() -> PipelineApi:
    """Look the entry points up at call time, so a newer pipeline (or a test fake) is picked up without restart."""
    return PipelineApi(
        run_assessment=getattr(pipeline, "run_assessment", None),
        rescore_vendor=getattr(pipeline, "rescore_vendor", None),
        export_assessment=getattr(pipeline, "export_assessment", None),
        input_error=getattr(pipeline, "InputError", None),
    )


def run_key(sha256: str, mode: str) -> tuple[str, str]:
    """Runs are cached per (input SHA-256, mode)."""
    return (sha256, mode)


@dataclass
class RunRecord:
    """One cached run: the full AssessmentResult, or the P1 fallback (criticality and depth only)."""

    key: tuple[str, str]
    mode: str
    kind: str  # "full" or "p1"
    p1: list[Any]
    example: Any = None
    result: AssessmentResult | None = None
    store_kind: str = "sandbox"
    store_dir: str = ""
    tier_overrides: dict[str, str] = field(default_factory=dict)
    vendors: tuple[str, ...] = ()
    team: str = ""
    notes: list[str] = field(default_factory=list)
    error: str = ""
    issues: list[ValidationIssue] = field(default_factory=list)
    version: int = 0  # bumped by every rescore, so exports built from older findings are not offered

    @property
    def full(self) -> bool:
        return self.kind == "full" and self.result is not None

    def findings(self, vendor_id: str) -> VendorFindings | None:
        return self.result.vendor(vendor_id) if self.result is not None else None

    @property
    def vendor_ids(self) -> list[str]:
        if self.result is not None:
            return [f.vendor_id for f in self.result.vendors]
        return [a.profile.vendor_id for a in self.p1]


def parse_as_of(text: str) -> str | None:
    """The optional 'as of' date of a run: '' -> None (the pipeline's default); else a strict YYYY-MM-DD date."""
    text = (text or "").strip()
    if not text:
        return None
    try:
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            raise ValueError(text)
        return dt.date.fromisoformat(text).isoformat()
    except ValueError:
        raise ValueError(f"'As of' must be a date written YYYY-MM-DD, not {clip(text, 20)!r}.") from None


def execute_run(inp: InputFile, data: WorkbookData, mode: str, stores: ReviewStores, *, team: str = "",
                vendors: Sequence[str] | None = None, progress: Callable[[str, float], None] | None = None,
                api: PipelineApi | None = None, store: Any = None, as_of: str | None = None) -> RunRecord:
    """Run the full assessment when the pipeline provides it; otherwise, or when it fails, fall back to P1.

    ``as_of`` (YYYY-MM-DD) pins the run date; None lets the pipeline choose (replay: the frozen runs' date).
    Never raises for a pipeline problem: the record carries the error (and the workbook issues for an InputError)
    next to the P1 criticality and depth, so the pages keep working.
    """
    if mode not in MODES:
        raise ValueError(f"unknown mode {mode!r}")
    as_of = parse_as_of(as_of or "")
    api = api or pipeline_api()
    wanted = tuple(vendors or ())
    unknown = sorted(set(wanted) - {v.vendor_id for v in data.vendors})
    if unknown:
        raise ValueError(f"Unknown vendor id(s): {', '.join(unknown)}")
    p1, example = assess_p1(data, stores.overrides)
    if wanted:
        p1 = [a for a in p1 if a.profile.vendor_id in wanted]
    base: dict[str, Any] = dict(
        key=run_key(inp.sha256, mode), mode=mode, p1=p1, example=example, store_kind=stores.kind,
        store_dir=str(stores.directory), vendors=wanted, team=team,
        tier_overrides=tier_override_map(stores.overrides, [v.vendor_id for v in data.vendors]),
    )
    if api.run_assessment is None:
        return RunRecord(kind="p1", notes=[P1_NOTE], **base)
    if store is None:
        from footprint.capture.store import EvidenceStore

        store = EvidenceStore(EVIDENCE_DIR)
    try:
        result = api.run_assessment(
            inp.data, mode, vendors=list(wanted) or None, as_of=as_of, store=store, seeds_dir=SEEDS_DIR,
            runs_dir=RUNS_DIR, overrides=stores.overrides, reviews=stores.reviews, team=team or None,
            progress=progress,
        )
    except NotImplementedError:
        return RunRecord(kind="p1", notes=[P1_NOTE], **base)
    except Exception as exc:  # noqa: BLE001 - the app must keep working; the error is shown
        if api.input_error is not None and isinstance(exc, api.input_error):
            issues = [i for i in getattr(exc, "issues", []) if isinstance(i, ValidationIssue)]
            return RunRecord(kind="p1", error="The workbook failed validation; nothing was assessed.",
                             issues=issues, **base)
        message = f"The assessment stopped: {type(exc).__name__}: {clip(scrub(exc), 400)}"
        return RunRecord(kind="p1", error=message,
                         notes=["Showing criticality and depth only (columns L-N) for this run."], **base)
    if not isinstance(result, AssessmentResult):
        return RunRecord(kind="p1", error="The pipeline returned an unexpected result; showing criticality and "
                                          "depth only.", **base)
    return RunRecord(kind="full", result=result, **base)


def is_stale(record: RunRecord, stores: ReviewStores, data: WorkbookData) -> bool:
    """True when the review store or a tier override changed since the run (run it again to apply)."""
    if record.store_dir != str(stores.directory):
        return True
    return record.tier_overrides != tier_override_map(stores.overrides, [v.vendor_id for v in data.vendors])


def replace_findings(record: RunRecord, findings: VendorFindings) -> None:
    """Swap one vendor's findings in the cached result (after a rescore)."""
    if record.result is None:
        raise ValueError("this run has no full result")
    vendors = [findings if f.vendor_id == findings.vendor_id else f for f in record.result.vendors]
    record.result = record.result.model_copy(update={"vendors": vendors})
    record.version += 1


def rescore(record: RunRecord, vendor_id: str, stores: ReviewStores, *, team: str = "",
            api: PipelineApi | None = None) -> VendorFindings:
    """Re-apply reviews and overrides for one vendor (pipeline.rescore_vendor) and cache the new findings."""
    api = api or pipeline_api()
    if record.result is None:
        raise ValueError("this run has no full result")
    if api.rescore_vendor is None:
        raise RuntimeError("footprint.pipeline.rescore_vendor is not available in this build")
    current = record.result.vendor(vendor_id)
    if current is None:
        raise ValueError(f"{vendor_id} is not in this run")
    updated = api.rescore_vendor(current, as_of=record.result.as_of, team=team or record.team or "",
                                 overrides=stores.overrides, reviews=stores.reviews)
    replace_findings(record, updated)
    return updated


def run_summary_rows(record: RunRecord) -> list[dict[str, Any]]:
    if record.full:
        assert record.result is not None
        rows = []
        for f in record.result.vendors:
            v, r = f.verdict, f.risk
            rows.append({
                "Vendor ID": f.vendor_id,
                "Vendor": f.profile.name,
                "Tier": f.criticality.tier.value,
                "AI usage (O)": v.column_o,
                "Verdict": v.label,
                "Likelihood": v.likelihood,
                "Confidence": v.confidence,
                "AI risk (S)": risk_label(r),
                "Score": f"{r.arp}/18",
                "Cited items": len(f.cited()),
                "Items": len(f.evidence),
                "Coverage complete": "Yes" if v.coverage_complete else "No",
            })
        return rows
    return [{
        "Vendor ID": a.profile.vendor_id,
        "Vendor": a.profile.name,
        "Tier": a.criticality.tier.value,
        "Depth": a.depth.label,
    } for a in record.p1]


def risk_label(risk: Any) -> str:
    return risk.final_class + (" (provisional)" if risk.provisional else "")


def class_change(before: Any, after: Any) -> str:
    """The live class after a rescore, in words: 'AI risk High → Critical; score 12 → 14 of 18.' A class held below
    its score band by the materiality gate or the verdict cap says so, because the score alone then looks
    inconsistent with it."""
    b, a = risk_label(before), risk_label(after)
    head = f"AI risk {b} → {a}" if b != a else f"AI risk stays {a}"
    score = f"score {before.arp} → {after.arp} of 18" if before.arp != after.arp else f"score {after.arp} of 18"
    pre = after.pre_cap_class
    held = []
    if not after.gate_met and pre and after.base_class != pre:
        held.append(f"materiality gate not met: {after.base_class} lowered to {pre}")
    if after.cap in ("High", "Medium") and pre and pre != after.final_class:
        held.append(f"capped at {after.cap} by the verdict; {pre} before the cap")
    return f"{head}" + (f" ({'; '.join(held)})" if held else "") + f"; {score}."


def score_sentence(risk: Any) -> str:
    """The score in words, as column T writes it (design 2.8)."""
    i = risk.inputs
    return (f"exposure {i.e}/3, decision impact {i.k}/3, tier {i.tp}, transparency gap {i.tg} → {risk.arp} of 18 = "
            f"{risk.base_class} (exposure and decision impact count double)")


def coverage_for(vendor_id: str, record: RunRecord | None) -> list[CoverageEntry]:
    """The vendor's Coverage Log: from the full run when there is one, else from its newest collection run."""
    if record is not None and record.full:
        findings = record.findings(vendor_id)
        if findings is not None:
            return list(findings.coverage)
    return latest_collection_coverage(vendor_id)


# =========================================================================== evidence


@dataclass(frozen=True)
class EvidenceFilter:
    """Evidence page filters; an empty tuple means 'any'."""

    vendors: tuple[str, ...] = ()
    families: tuple[str, ...] = ()
    strengths: tuple[str, ...] = ()
    methods: tuple[str, ...] = ()
    roles: tuple[str, ...] = ()
    statuses: tuple[str, ...] = ()
    cited_only: bool = False
    text: str = ""


def filter_items(findings_list: Sequence[VendorFindings], flt: EvidenceFilter
                 ) -> list[tuple[VendorFindings, EvidenceItem]]:
    """Items matching every filter, vendors in run order and items in evidence (E-ID) order."""
    needle = flt.text.strip().lower()
    out = []
    for f in findings_list:
        if flt.vendors and f.vendor_id not in flt.vendors:
            continue
        for item in f.evidence:
            if flt.families and item.family.value not in flt.families:
                continue
            if flt.strengths and item.strength not in flt.strengths:
                continue
            if flt.methods and item.method not in flt.methods:
                continue
            if flt.roles and item.role not in flt.roles:
                continue
            if flt.statuses and item.review_status not in flt.statuses:
                continue
            if flt.cited_only and item.role not in CITED_ROLES:
                continue
            if needle and needle not in " ".join((item.excerpt, item.title, item.url, item.publisher,
                                                  item.evidence_id)).lower():
                continue
            out.append((f, item))
    return out


def item_ref(item: EvidenceItem) -> str:
    return item.evidence_id or item.item_key[:12]


def ref_for(findings: VendorFindings, key: str) -> str:
    item = findings.item(key)
    return item_ref(item) if item is not None else key[:12]


def refs_for(findings: VendorFindings, keys: Iterable[str]) -> str:
    return ", ".join(ref_for(findings, k) for k in keys) or "-"


def counts_label(item: EvidenceItem) -> str:
    """Whether the item counts toward the verdict, risk and cells (EvidenceItem.citable), and if not, why."""
    if item.tags.ai_type == "not_ai":
        return "no: definition-test trap"
    if item.review_status == "rejected":
        return "no: rejected"
    if item.proposed:
        return "no: Gemini proposal awaiting review"
    return "yes"


def item_option_label(item: EvidenceItem) -> str:
    return f"{item_ref(item)} · {item.strength} · {item.source_type} · {clip(item.excerpt, 70)}"


def evidence_rows(pairs: Sequence[tuple[VendorFindings, EvidenceItem]]) -> list[dict[str, Any]]:
    return [{
        "Evidence ID": item_ref(item),
        "Vendor": f.vendor_id,
        "Role": item.role,
        "Strength": item.strength,
        "Tags": item.tag_string,
        "Family": item.family.value,
        "Source": item.source_type,
        "Publisher": item.publisher,
        "Published": item.published or f"retrieved {item.retrieved_at[:10]}",
        "Method": METHOD_LABELS.get(item.method, item.method),
        "Review": REVIEW_LABELS.get(item.review_status, item.review_status),
        "Counts": counts_label(item),
        "Excerpt": clip(item.excerpt, 160),
    } for f, item in pairs]


def present_values(findings_list: Sequence[VendorFindings], attr: str, order: Sequence[str] = ()) -> list[str]:
    """Distinct values of an item attribute across the run, in ``order`` first, then sorted."""
    seen: set[str] = set()
    for f in findings_list:
        for item in f.evidence:
            value = getattr(item, attr)
            seen.add(value.value if hasattr(value, "value") else str(value))
    ordered = [v for v in order if v in seen]
    return ordered + sorted(seen - set(ordered))


@dataclass(frozen=True)
class ExcerptContext:
    """An excerpt with up to CONTEXT_CHARS of its source text on either side (contracts: UI usage)."""

    before: str
    excerpt: str
    after: str
    clipped_start: bool
    clipped_end: bool
    matches: bool  # the stored text still holds the excerpt at its offsets


CONTEXT_CHARS = 300


def excerpt_context(text: str, start: int, end: int, excerpt: str | None = None,
                    margin: int = CONTEXT_CHARS) -> ExcerptContext:
    lo, hi = max(0, start - margin), min(len(text), end + margin)
    found = text[start:end]
    return ExcerptContext(before=text[lo:start], excerpt=found, after=text[end:hi], clipped_start=lo > 0,
                          clipped_end=hi < len(text), matches=excerpt is None or found == excerpt)


def load_context(item: EvidenceItem, store: Any = None) -> ExcerptContext | None:
    """The excerpt in context from the evidence store's text copy, or None when the text is not stored."""
    if store is None:
        from footprint.capture.store import EvidenceStore

        store = EvidenceStore(EVIDENCE_DIR)
    try:
        text = store.get_text(item.doc_id)
    except (KeyError, OSError):
        return None
    return excerpt_context(text, item.start, item.end, item.excerpt)


def _html_text(text: str) -> str:
    return html.escape(text).replace("$", "&#36;").replace("\n", "<br>")


def context_html(ctx: ExcerptContext) -> str:
    """HTML for st.html: the source text with the excerpt highlighted (all text escaped)."""
    lead = "… " if ctx.clipped_start else ""
    tail = " …" if ctx.clipped_end else ""
    return (
        '<div style="font-family: Georgia, \'Source Serif Pro\', serif; font-size: 0.98rem; line-height: 1.6; '
        'padding: 0.8rem 1rem; border-left: 3px solid #0f766e; background: rgba(127,127,127,0.07); '
        'border-radius: 0 6px 6px 0; overflow-wrap: anywhere;">'
        f'<span style="opacity: 0.72">{lead}{_html_text(ctx.before)}</span>'
        '<mark style="background: rgba(250, 204, 21, 0.45); color: inherit; padding: 0 0.12em; '
        f'border-radius: 2px;">{_html_text(ctx.excerpt)}</mark>'
        f'<span style="opacity: 0.72">{_html_text(ctx.after)}{tail}</span></div>'
    )


def screenshot_file(item: EvidenceItem, root: Path = REPO_ROOT) -> Path | None:
    """The item's screenshot on disk (repo-relative path inside the repo), or None."""
    if not item.screenshot_path:
        return None
    path = Path(item.screenshot_path)
    path = (path if path.is_absolute() else root / path).resolve()
    try:
        path.relative_to(root.resolve())
    except ValueError:
        return None
    return path if path.is_file() else None


def label_rows(item: EvidenceItem) -> list[dict[str, str]]:
    """Rules vs Gemini, label by label; a V8 disagreement keeps the rule value until an analyst adjudicates."""
    keys = [k for k in LABEL_KEYS if k in item.rule_labels or k in item.llm_labels]
    keys += sorted((set(item.rule_labels) | set(item.llm_labels)) - set(keys))
    disagree = set(item.label_disagreements)
    rows = []
    for k in keys:
        rule, llm = item.rule_labels.get(k), item.llm_labels.get(k)
        if k in disagree:
            status = "disagree: rule value stands (HC2)"
        elif rule is not None and llm is not None:
            status = "agree" if rule == llm else "differ (informational)"
        else:
            status = "rules only" if llm is None else "Gemini only"
        rows.append({"Label": k, "Rules": rule if rule is not None else "-", "Gemini": llm if llm is not None else "-",
                     "Status": status})
    return rows


def reverify(item: EvidenceItem, store: Any = None) -> tuple[bool, list[str]] | None:
    """Re-verify an item against the store (footprint.verify.reverify_item): (ok, details), or None if unavailable."""
    try:
        from footprint.verify import reverify_item
    except ImportError:
        return None
    if store is None:
        from footprint.capture.store import EvidenceStore

        store = EvidenceStore(EVIDENCE_DIR)
    result = reverify_item(item, store)
    details = list(result.details) or [f"{code} failed" for code in result.failures]
    return bool(result.ok), details


def cited_unreviewed(findings: VendorFindings) -> list[EvidenceItem]:
    """Items with a cited role that no analyst has reviewed yet (the 'approve' action accepts them)."""
    return [i for i in findings.evidence if i.role in CITED_ROLES and i.review_status == "unreviewed"]


# =========================================================================== findings and risk


def why_trace(findings: VendorFindings) -> list[dict[str, str]]:
    """The why trace: verdict rule lines, then the ordered risk steps (contracts: UI usage)."""
    v, r = findings.verdict, findings.risk
    rows = [{"Stage": "AI usage", "Step": line} for line in v.trace]
    rows.append({"Stage": "AI usage", "Step": f"Decision: rule {v.rule}) gives {v.label}, written '{v.column_o}' in "
                                              f"column O; it is {v.likelihood} that AI is used, confidence "
                                              f"{v.confidence} because {v.confidence_reason}."})
    rows += [{"Stage": "AI risk", "Step": step} for step in r.steps]
    rows.append({"Stage": "AI risk", "Step": f"Result: {risk_label(r)}"
                                             + (f"; ceiling {r.ceiling_class} if confirmed" if r.ceiling_class else "")
                                             + "."})
    return rows


def risk_input_rows(findings: VendorFindings) -> list[dict[str, str]]:
    i = findings.risk.inputs
    return [
        {"Input": "E · exposure of Meridian data to AI", "Value": f"{i.e}/3",
         "Basis": "assumed (unknowns rule)" if i.e_assumed else "evidenced", "Reason": i.e_reason,
         "Evidence": refs_for(findings, i.e_items)},
        {"Input": "K · decision impact", "Value": f"{i.k}/3",
         "Basis": "assumed (unknowns rule)" if i.k_assumed else "evidenced", "Reason": i.k_reason,
         "Evidence": refs_for(findings, i.k_items)},
        {"Input": "TP · tier points", "Value": f"{i.tp}/3", "Basis": f"{findings.criticality.tier.value} tier",
         "Reason": "Critical 3, High 2, Medium 1, Low 0", "Evidence": "-"},
        {"Input": "TG · transparency gaps", "Value": f"{i.tg}/3",
         "Basis": f"{len(i.missing_gaps)} of 6 checks missing" if i.gaps else "not assessed",
         "Reason": "0-1 missing 0, 2-3 missing 1, 4-5 missing 2, all 6 missing 3", "Evidence": "-"},
    ]


def gap_rows(findings: VendorFindings) -> list[dict[str, str]]:
    i = findings.risk.inputs
    rows = []
    for g in GAP_KEYS:
        status = ("missing" if i.gaps.get(g) else "closed") if i.gaps else "not assessed"
        rows.append({"Check": g, "What": GAP_LABELS[g], "Status": status,
                     "Closed by": refs_for(findings, i.gap_items.get(g, [])), "Note": i.gap_reasons.get(g, "")})
    return rows


def gap_values(findings: VendorFindings) -> dict[str, str]:
    """Current t1..t6 states as the override editor shows them ('missing' or 'closed')."""
    gaps = findings.risk.inputs.gaps
    return {g: "missing" if gaps.get(g) else "closed" for g in GAP_KEYS}


def changed_gap_records(findings: VendorFindings, values: Mapping[str, str], reason: str, analyst: str, *,
                        date: str | None = None) -> list[ReviewRecord]:
    """One risk_input_override per check whose state the analyst changed."""
    current = gap_values(findings)
    return [risk_override_record(findings.vendor_id, g, values[g], reason, analyst, date=date)
            for g in GAP_KEYS if g in values and values[g] != current[g]]


_DEFAULT_LETTERS: dict[str, str] = dict(zip(STUDENT_FIELDS, "LMNOPQRSTUV"))


def cells_rows(cells: StudentCells, column_map: Mapping[str, str] | None = None) -> list[dict[str, Any]]:
    """Columns L-V as they will be written, with each cell's length against its budget."""
    letters = {**_DEFAULT_LETTERS, **{k: v for k, v in (column_map or {}).items() if k in _DEFAULT_LETTERS}}
    rows = []
    for fld in STUDENT_FIELDS:
        text = getattr(cells, fld)
        rows.append({"Column": letters[fld], "Field": HEADER_ALIASES[fld][0], "Text": text or "",
                     "Length": f"{len(text or '')}/{DEFAULT_LENGTH_BUDGETS[fld]}"})
    return rows


def store_review(stores: ReviewStores, record: ReviewRecord) -> ReviewRecord:
    """Append to the right file: evidence reviews to reviews.jsonl, everything else to overrides.jsonl."""
    target = stores.reviews if record.kind == "evidence_review" else stores.overrides
    return target.add(record)


# =========================================================================== export


@dataclass
class ExportOutcome:
    """The validated output workbook and what the checks found."""

    data: bytes = field(repr=False)
    file_name: str
    kind: str  # "full" (columns L-V and all sheets) or "p1" (columns L-N)
    written: list[str]
    sheets_added: list[str]
    warnings: list[str]
    fidelity: list[str]
    sheets: list[str]
    example_unchanged: bool | None
    note: str = ""  # why a full run was exported as columns L-N only

    @property
    def ok(self) -> bool:
        return not self.fidelity and self.example_unchanged is not False


def output_name(inp: InputFile, as_of: str = "") -> str:
    stem = re.sub(r"[^\w.-]+", "_", Path(inp.name).stem).strip("_") or "vendor_inventory"
    return f"{stem}_assessed_{as_of or today()}.xlsx"


def _example_row_unchanged(source: bytes, output: bytes, data: WorkbookData) -> bool | None:
    if data.example is None:
        return None
    from openpyxl import load_workbook

    def row_values(raw: bytes) -> list[str]:
        wb = load_workbook(io.BytesIO(raw), read_only=True)
        try:
            ws = wb[data.sheet_name]
            cells = next(ws.iter_rows(min_row=data.example.row, max_row=data.example.row), ())
            values = ["" if c.value is None else str(c.value) for c in cells]
            while values and not values[-1]:
                values.pop()
            return values
        finally:
            wb.close()

    return row_values(source) == row_values(output)


def _sheet_names(raw: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def export_key(inp: InputFile, record: RunRecord | None, team: str, overwrite: bool) -> tuple[Any, ...]:
    """What a built export depends on: the input, the run and its rescore version, the team and overwrite."""
    return (inp.sha256, record.key if record else None, record.version if record else 0, team, overwrite)


def export_run(record: RunRecord | None, inp: InputFile, data: WorkbookData, stores: ReviewStores, *, team: str = "",
               overwrite: bool = False, api: PipelineApi | None = None) -> ExportOutcome:
    """Build the output workbook in memory: export_assessment for a full run, else the P1 export (L-N).

    Both paths refuse to write anything that breaks a workbook rule and run the fidelity check; the check is run
    again here so the page can show its result, together with the sheet list and the V-000 row comparison. A full
    run falls back to the P1 export (with ``note`` saying so) only while export_assessment is missing or raises
    NotImplementedError; any other export failure propagates, so nothing unchecked is offered for download.
    """
    api = api or pipeline_api()
    team = team or (record.team if record is not None else "")
    note = ""
    out = b""
    report: Any = None
    if record is not None and record.full:
        assert record.result is not None
        if api.export_assessment is None:
            note = EXPORT_P1_NOTE
        else:
            buffer = io.BytesIO()
            try:
                report = api.export_assessment(record.result, inp.data, buffer, team, overwrite=overwrite)
            except NotImplementedError:
                note = EXPORT_P1_NOTE
            else:
                out, kind, as_of = buffer.getvalue(), "full", record.result.as_of
    if report is None:
        assessments, example = (record.p1, record.example) if record is not None else assess_p1(data, stores.overrides)
        with tempfile.TemporaryDirectory(prefix="footprint-export-") as tmp:
            target = Path(tmp) / "out.xlsx"
            report = pipeline.export_workbook(inp.data, target, assessments, example=example, team=team or None,
                                              overwrite=overwrite)
            out = target.read_bytes()
        kind, as_of = "p1", ""
    return ExportOutcome(
        data=out,
        file_name=output_name(inp, as_of),
        kind=kind,
        written=list(getattr(report, "written", [])),
        sheets_added=list(getattr(report, "sheets_added", [])),
        warnings=[f"{w.code} {w.cell}: {w.message}".strip() for w in getattr(report, "warnings", [])],
        fidelity=check_fidelity(inp.data, out),
        sheets=_sheet_names(out),
        example_unchanged=_example_row_unchanged(inp.data, out, data),
        note=note,
    )


# =========================================================================== live DNS check (demo step 3)


@dataclass
class DnsCheck:
    """AI-provider DNS verification tokens: as captured in the evidence pack vs a live lookup now."""

    vendor_id: str
    domains: list[str]
    captured: list[tuple[str, str]]
    captured_at: str
    live: list[tuple[str, str]]
    live_at: str
    notes: list[str] = field(default_factory=list)

    @property
    def captured_providers(self) -> list[str]:
        return sorted({p for p, _ in self.captured})

    @property
    def live_providers(self) -> list[str]:
        return sorted({p for p, _ in self.live})

    @property
    def unchanged(self) -> bool:
        return set(self.captured) == set(self.live)


def ai_tokens_in(dns_text: str) -> list[tuple[str, str]]:
    """(provider, record) pairs from the AI_TOKEN lines of a DNS document."""
    out = []
    for line in dns_text.splitlines():
        parts = line.split("\t")
        if len(parts) == 3 and parts[0] == "AI_TOKEN":
            out.append((parts[1], parts[2]))
    return out


class ReadOnlyPack:
    """The evidence pack through a read-only lens: replayed collectors read captures, and the text they render
    stays in memory, so the frozen pack is never written."""

    def __init__(self, inner: Any) -> None:
        self._inner = inner
        self._texts: dict[str, str] = {}

    def find_by_url(self, url: str) -> Any:
        return self._inner.find_by_url(url)

    def get_raw(self, capture_id: str) -> bytes:
        return self._inner.get_raw(capture_id)

    def put_text(self, text: str) -> tuple[str, str]:
        sha = hashlib.sha256(text.encode("utf-8")).hexdigest()
        self._texts[sha] = text
        return sha, ""

    def get_text(self, sha: str) -> str:
        return self._texts[sha] if sha in self._texts else self._inner.get_text(sha)

    def put_raw(self, data: bytes, **meta: Any) -> Any:
        raise RuntimeError("the evidence pack is read-only in the app")


def _dns_pass(profile: VendorProfile, seeds: dict, plan: DepthPlan, fetcher: Any, store: Any
              ) -> tuple[list[tuple[str, str]], str, list[str]]:
    from footprint.collectors.base import CollectContext
    from footprint.collectors.dns import DnsCollector

    ctx = CollectContext(profile=profile, seeds=seeds, plan=plan, fetcher=fetcher, store=store)
    result = DnsCollector().collect(ctx)
    tokens: list[tuple[str, str]] = []
    for doc in result.documents:
        try:
            tokens += ai_tokens_in(store.get_text(doc.doc_id))
        except KeyError:
            continue
    when = min((c.retrieved_at for c in result.captures if c.retrieved_at), default="")
    return sorted(set(tokens)), when, [e.note for e in result.coverage if e.note]


def dns_check(profile: VendorProfile, plan: DepthPlan, *, seeds_dir: Path = SEEDS_DIR,
              evidence_dir: Path = EVIDENCE_DIR, live_fetcher: Callable[[Any], Any] | None = None) -> DnsCheck:
    """Replay the vendor's DNS lookup from the evidence pack, then repeat it live and compare the AI tokens.

    The live pass asks only the two DoH resolvers in the ToS register (dns.google, cloudflare-dns.com) through
    LiveFetcher (ToS -> robots.txt -> rate limit) and stores its captures in a temporary store, so the frozen
    evidence pack is never changed. ``live_fetcher(store)`` builds the fetcher (tests inject a fake).
    """
    from footprint.capture.store import EvidenceStore
    from footprint.collectors.base import vendor_domains
    from footprint.net.fetcher import ReplayFetcher

    seeds = pipeline.load_seeds(profile.vendor_id, seeds_dir)
    pack = ReadOnlyPack(EvidenceStore(evidence_dir))
    captured, captured_at, notes = _dns_pass(profile, seeds, plan, ReplayFetcher(pack), pack)
    with tempfile.TemporaryDirectory(prefix="footprint-dns-") as tmp:
        live_store = EvidenceStore(tmp)
        if live_fetcher is None:
            from footprint.net.fetcher import LiveFetcher
            from footprint.net.tou import load_tou

            fetcher: Any = LiveFetcher(live_store, load_tou())
        else:
            fetcher = live_fetcher(live_store)
        try:
            live, live_at, live_notes = _dns_pass(profile, seeds, plan, fetcher, live_store)
        finally:
            session = getattr(fetcher, "session", None)
            if session is not None and hasattr(session, "close"):
                session.close()
    return DnsCheck(vendor_id=profile.vendor_id, domains=vendor_domains(profile, seeds), captured=captured,
                    captured_at=captured_at, live=live, live_at=live_at,
                    notes=[f"captured: {n}" for n in notes] + [f"live: {n}" for n in live_notes])


# =========================================================================== session state (Streamlit)

K_INIT = "fp_init"
K_INPUT = "fp_input"
K_LAST_UPLOAD = "fp_last_upload"
K_WORKBOOKS = "fp_workbooks"
K_P1 = "fp_p1_cache"
K_EXAMPLE_CELLS = "fp_example_cells"
K_RUNS = "fp_runs"
K_MODE = "fp_mode"
K_ANALYST = "fp_analyst"
K_TEAM = "fp_team"
K_STORE_KIND = "fp_store_kind"
K_FLASH = "fp_flash"
K_EXPORT = "fp_export"
K_PREV_CLASS = "fp_prev_class"
KEEP_PREFIX = "fp_kept_"


PAGE_ICON = ":material/travel_explore:"


def page_setup(title: str) -> None:
    """First call of every page: page config, session defaults and the shared sidebar.

    Each page renders the sidebar itself instead of relying on the entrypoint: because app/pages/ exists,
    Streamlit's legacy multipage mode can run a page file without app/streamlit_app.py (a deep link on the
    server's first request, and AppTest.switch_page), and the page must still be complete then.
    """
    st.set_page_config(page_title=f"{title} · footprint", page_icon=PAGE_ICON, layout="wide")
    init_session()
    render_sidebar()


def init_session() -> None:
    """Defaults for every session key the pages share (run once per script run, before any widget)."""
    ss = st.session_state
    if not ss.get(K_INIT):
        load_env()
        ss[K_INIT] = True
    ss.setdefault(K_RUNS, {})
    ss.setdefault(K_WORKBOOKS, {})
    ss.setdefault(K_P1, {})
    ss.setdefault(K_MODE, "replay")
    ss.setdefault(K_PREV_CLASS, {})
    ss.setdefault(KEEP_PREFIX + K_TEAM, default_team())


def keep(widget_key: str, default: Any, options: Sequence[Any] | None = None, *, multi: bool = False) -> str:
    """Use as ``key=keep(...)``: the widget keeps its value across page switches and reruns.

    Streamlit drops a widget's state when the page that rendered it changes; re-asserting the value through the
    Session State API (from a plain copy kept under KEEP_PREFIX) restores it. ``options`` drops a kept value that
    is no longer offered.
    """
    ss = st.session_state
    store = KEEP_PREFIX + widget_key
    if widget_key in ss:
        ss[store] = ss[widget_key]
    value = ss.get(store, default)
    if options is not None:
        if multi:
            value = [v for v in (value or []) if v in options]
        elif value not in options:
            value = default if default in options else (options[0] if options else None)
    ss[widget_key] = value
    return widget_key


def kept(widget_key: str, default: Any = None) -> Any:
    """A kept widget's current value, also on a page that does not render the widget."""
    ss = st.session_state
    return ss[widget_key] if widget_key in ss else ss.get(KEEP_PREFIX + widget_key, default)


def session_input() -> InputFile | None:
    return st.session_state.get(K_INPUT)


def take_input(uploaded: Any, use_bundled: bool) -> None:
    """Adopt a newly uploaded workbook, or the bundled one when its button was pressed."""
    ss = st.session_state
    if uploaded is not None:
        data = uploaded.getvalue()
        sha = hashlib.sha256(data).hexdigest()
        if ss.get(K_LAST_UPLOAD) != sha:
            ss[K_LAST_UPLOAD] = sha
            ss[K_INPUT] = make_input(uploaded.name, data, "upload")
            ss.pop(K_EXPORT, None)
    if use_bundled:
        ss[K_INPUT] = bundled_input()
        ss.pop(K_EXPORT, None)


def session_workbook(inp: InputFile) -> WorkbookData:
    cache = st.session_state.setdefault(K_WORKBOOKS, {})
    if inp.sha256 not in cache:
        cache[inp.sha256] = read_input(inp)
    return cache[inp.sha256]


def session_example_cells(inp: InputFile, data: WorkbookData) -> dict[str, str]:
    """example_cells, read once per workbook (opening the workbook again on every rerun costs seconds)."""
    cache = st.session_state.setdefault(K_EXAMPLE_CELLS, {})
    if inp.sha256 not in cache:
        cache[inp.sha256] = example_cells(inp, data)
    return cache[inp.sha256]


def session_stores() -> ReviewStores:
    kind = kept(K_STORE_KIND, "sandbox")
    try:
        return review_stores(kind if kind in STORE_LABELS else "sandbox")
    except ValueError as exc:
        st.error(str(exc))
        st.stop()
        raise  # unreachable: st.stop() ends the run


def session_assessments(inp: InputFile, data: WorkbookData, stores: ReviewStores) -> tuple[list[Any], Any]:
    """P1 criticality and depth, cached until the input or the overrides file changes."""
    path = stores.overrides.path
    try:
        stat = path.stat()
        stamp: tuple[Any, ...] = (stat.st_mtime_ns, stat.st_size)
    except OSError:
        stamp = ()
    key = (inp.sha256, str(path), stamp)
    cache = st.session_state.setdefault(K_P1, {})
    if key not in cache:
        cache.clear()
        cache[key] = assess_p1(data, stores.overrides)
    return cache[key]


def session_mode() -> str:
    mode = st.session_state.get(K_MODE, "replay")
    return mode if mode in MODES else "replay"


def set_mode(mode: str) -> None:
    if mode in MODES:
        st.session_state[K_MODE] = mode


def session_analyst() -> str:
    return str(kept(K_ANALYST) or "").strip()


def session_team() -> str:
    return str(kept(K_TEAM) or "").strip()


def store_run(record: RunRecord) -> None:
    st.session_state.setdefault(K_RUNS, {})[record.key] = record
    st.session_state.pop(K_EXPORT, None)


def active_run() -> RunRecord | None:
    """The cached run for the current workbook and mode, if any."""
    inp = session_input()
    if inp is None:
        return None
    return st.session_state.get(K_RUNS, {}).get(run_key(inp.sha256, session_mode()))


def flash(message: str, kind: str = "success") -> None:
    """A message shown once, at the top of the next page run (after st.rerun)."""
    st.session_state[K_FLASH] = (kind, message)


def show_flash() -> None:
    item = st.session_state.pop(K_FLASH, None)
    if item:
        kind, message = item
        {"success": st.success, "info": st.info, "warning": st.warning, "error": st.error}.get(kind, st.info)(message)


def badge(label: str, colors: Mapping[str, str] = CLASS_COLORS) -> None:
    st.badge(label, color=colors.get(label.split(" (")[0], "gray"))  # type: ignore[arg-type]


def require_input() -> tuple[InputFile, WorkbookData]:
    inp = session_input()
    if inp is None:
        st.info("Load the vendor workbook on the Assess page first.")
        st.page_link("pages/assess.py", label="Go to Assess", icon=":material/arrow_back:")
        st.stop()
    return inp, session_workbook(inp)  # type: ignore[arg-type]


def require_full_run() -> RunRecord:
    """The active full run, or guidance (and a stop) when there is none."""
    require_input()
    record = active_run()
    if record is None:
        st.info(f"No assessment has been run for this workbook in {MODE_LABELS[session_mode()]} mode yet.")
        st.page_link("pages/assess.py", label="Run it on the Assess page", icon=":material/play_arrow:")
        st.stop()
    if not record.full:
        st.info("This run holds criticality and depth only (columns L-N): evidence, verdicts and risk need the full "
                "pipeline (footprint.pipeline.run_assessment).")
        if record.error:
            st.error(record.error)
        st.page_link("pages/assess.py", label="Back to Assess", icon=":material/arrow_back:")
        st.stop()
    return record  # type: ignore[return-value]


def vendor_picker(record: RunRecord, key: str, label: str = "Vendor") -> str:
    """A vendor selectbox over the run's vendors (sheet order)."""
    ids = record.vendor_ids
    names = {f.vendor_id: f.profile.name for f in record.result.vendors} if record.result else {
        a.profile.vendor_id: a.profile.name for a in record.p1}
    return st.selectbox(label, ids, format_func=lambda v: f"{v} · {names.get(v, '')}",
                        key=keep(key, ids[0] if ids else None, ids))


def render_sidebar() -> None:
    """Shared sidebar: workbook, run, analyst, team and review store (rendered on every page)."""
    with st.sidebar:
        st.markdown("**footprint** · vendor AI footprint")
        st.caption("Runs on this machine only. Public vendor text only goes to Gemini, through the payload guard.")
        inp = session_input()
        if inp is None:
            st.caption("No workbook loaded.")
        else:
            st.caption(f"Workbook: {inp.name} · SHA-256 {inp.short_sha}…")
            record = active_run()
            if record is None:
                st.caption(f"Mode {MODE_LABELS[session_mode()]}: not run yet.")
            elif record.full and record.result is not None:
                r = record.result
                st.caption(f"Run {r.run_id} · {MODE_LABELS[record.mode]} · as of {dmy(r.as_of)}")
            else:
                st.caption(f"Mode {MODE_LABELS[record.mode]}: criticality and depth only.")
        st.divider()
        st.text_input("Analyst", key=keep(K_ANALYST, ""), placeholder="Name or initials",
                      help="Recorded with every tier override and evidence decision.")
        st.text_input("Team (column V)", key=keep(K_TEAM, default_team()), placeholder="e.g. Team Osprey",
                      help="Written in column V as 'Team <name> / <date>'.")
        kind = st.radio("Review store", list(STORE_LABELS), key=keep(K_STORE_KIND, "sandbox", list(STORE_LABELS)),
                        format_func=STORE_LABELS.get,
                        help="The sandbox is a copy of review/ for demos; the project store is the team's audit "
                             "trail.")
        if kind == "sandbox":
            if st.button("Reset sandbox", key="fp_reset_sandbox", icon=":material/restart_alt:",
                         help="Restore the sandbox from the project review store."):
                ensure_sandbox(reset=True)
                st.session_state[K_P1] = {}
                flash("Sandbox restored from the project review store.", "info")
                st.rerun()
        st.caption("Gemini key: configured" if gemini_configured()
                   else "Gemini key: not set (Live AI uses the rules only)")
