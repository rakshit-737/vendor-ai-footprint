"""Tests for footprint.workbook: reading the inventory, writing student cells, the fidelity diff.

Every test uses the real input workbook or a copy edited in tmp_path; no binary fixtures are stored.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import shutil
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Sequence
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from footprint.models import STUDENT_FIELDS, SheetSpec, StudentCells, ValidationIssue, WorkbookData
from footprint.workbook import (
    CREAM,
    DEFAULT_LENGTH_BUDGETS,
    WorkbookWriteError,
    WriteReport,
    check_fidelity,
    read_workbook,
    write_workbook,
)

INPUT = Path(__file__).resolve().parents[2] / "data" / "input" / "Meridian_Vendor_Input.xlsx"
SHEET = "Vendor Inventory"
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

VENDORS = [
    ("V-001", "AutomWorx", "www.automworx.com"),
    ("V-002", "Fiserv, Inc.", "www.fiserv.com"),
    ("V-003", "Financial Statement Services, Inc. (FSSI)", "www.fssi-ca.com"),
    ("V-004", "Terrapin Technologies, Inc.", "www.terrapintech.com"),
    ("V-005", "BNY", "www.bny.com"),
    ("V-006", "The Clearing House Payments Company L.L.C.", "www.theclearinghouse.org"),
]
VENDOR_IDS = [vendor_id for vendor_id, _, _ in VENDORS]
VENDOR_ROWS = list(range(6, 12))
COLUMN_MAP = {
    "vendor_id": "B", "name": "C", "description": "D", "service": "E", "category": "F",
    "website": "G", "business_process": "H", "operational_dependency": "I", "data_accessed": "J",
    "data_volume": "K",
    "criticality_tier": "L", "criticality_rationale": "M", "assessment_depth": "N",
    "ai_usage_detected": "O", "evidence": "P", "how_ai_used": "Q", "ai_subprocessors": "R",
    "ai_risk_class": "S", "risk_rationale": "T", "recommended_action": "U", "assessed_by": "V",
}


# --------------------------------------------------------------------------- helpers

Edit = Callable[[Workbook], object]


def edited_copy(tmp_path: Path, *edits: Edit, name: str = "edited.xlsx") -> Path:
    """The input workbook with `edits` applied, saved in tmp_path (rich text kept, as the writer does)."""
    wb = openpyxl.load_workbook(INPUT, rich_text=True)
    for edit in edits:
        edit(wb)
    path = tmp_path / name
    wb.save(path)
    return path


def set_cells(**values: object) -> Edit:
    """Edit that sets Vendor Inventory cells, e.g. set_cells(B7="V-001", G8=None)."""

    def edit(wb: Workbook) -> None:
        for coordinate, value in values.items():
            wb[SHEET][coordinate] = value

    return edit


def with_code(data: WorkbookData, code: str) -> list[ValidationIssue]:
    return [issue for issue in data.issues if issue.code == code]


def write(
    tmp_path: Path,
    cells: dict[str, StudentCells],
    sheets: Sequence[SheetSpec] = (),
    source: Path = INPUT,
    **options: object,
) -> tuple[Path, WriteReport]:
    out = tmp_path / "out.xlsx"
    report = write_workbook(source, out, cells, list(sheets), **options)
    return out, report


def refused(
    tmp_path: Path, cells: dict[str, StudentCells], source: Path = INPUT, **options: object
) -> WorkbookWriteError:
    """Assert that the write is refused and that no output file was created."""
    out = tmp_path / "refused.xlsx"
    with pytest.raises(WorkbookWriteError) as caught:
        write_workbook(source, out, cells, **options)
    assert not out.exists()
    return caught.value


def inventory(path: Path | bytes) -> openpyxl.worksheet.worksheet.Worksheet:
    return openpyxl.load_workbook(io.BytesIO(path) if isinstance(path, bytes) else path)[SHEET]


def sheet_xml(path: Path) -> ET.Element:
    """Root of xl/worksheets/sheet1.xml (openpyxl names sheet parts in tab order: sheet1 = Vendor Inventory)."""
    with zipfile.ZipFile(path) as archive:
        return ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))


def l_to_n(vendor_id: str) -> StudentCells:
    return StudentCells(
        criticality_tier="High",
        criticality_rationale=f"{vendor_id}: payments data at full scale — see “Criticality Workings”.",
        assessment_depth="Standard review – tier-driven",
    )


def every_field(vendor_id: str) -> StudentCells:
    values = {field: f"{vendor_id} {field} text" for field in STUDENT_FIELDS}
    values.update(criticality_tier="High", ai_usage_detected="Inconclusive", ai_risk_class="None identified")
    return StudentCells(**values)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# --------------------------------------------------------------------------- reading the template


def test_reads_six_vendors_in_sheet_order():
    data = read_workbook(INPUT)
    assert data.sheet_name == SHEET
    assert data.header_row == 4
    assert [(v.vendor_id, v.name, v.website) for v in data.vendors] == VENDORS
    assert [v.row for v in data.vendors] == VENDOR_ROWS
    assert not any(v.is_example for v in data.vendors)
    assert data.issues == []
    assert data.ok


def test_reads_every_provided_field():
    fiserv = read_workbook(INPUT).vendors[1]
    assert fiserv.category == "Core Platform"
    assert fiserv.business_process == "Core account processing and servicing"
    assert fiserv.operational_dependency == "Critical"
    assert fiserv.data_accessed == "Full customer master, balances and transactions for 1.4 million customers."
    assert fiserv.data_volume == "240 million transactions"
    assert fiserv.description.startswith("A global financial technology and payments provider")
    assert fiserv.service.startswith("The DNA core banking platform")
    assert fiserv.domain == "fiserv.com"


def test_worked_example_is_set_aside():
    example = read_workbook(INPUT).example
    assert example is not None
    assert (example.vendor_id, example.row, example.is_example) == ("V-000", 5, True)
    assert example.name.startswith("EXAMPLE VENDOR")


def test_column_map_covers_b_to_v():
    assert read_workbook(INPUT).column_map == COLUMN_MAP


def test_reads_paths_bytes_and_streams():
    expected = read_workbook(INPUT)
    assert read_workbook(str(INPUT)) == expected
    assert read_workbook(INPUT.read_bytes()) == expected
    assert read_workbook(io.BytesIO(INPUT.read_bytes())) == expected


def test_text_is_kept_verbatim_apart_from_surrounding_whitespace(tmp_path):
    source = edited_copy(tmp_path, set_cells(C6="  “AutomWorx” – Automation’s partner \n"))
    assert read_workbook(source).vendors[0].name == "“AutomWorx” – Automation’s partner"


@pytest.mark.parametrize(
    "vendor_id, name",
    [
        ("V-900", None),  # name still says "(fictional, do not edit)"
        ("X-1", "Example Corp"),
        ("X-2", "Larkspur (Do Not Edit)"),
        ("V-000", "Some Vendor"),
    ],
)
def test_example_row_detected_by_id_or_name(tmp_path, vendor_id, name):
    edits = {"B5": vendor_id} | ({"C5": name} if name else {})
    data = read_workbook(edited_copy(tmp_path, set_cells(**edits)))
    assert data.example is not None and data.example.vendor_id == vendor_id
    assert [v.vendor_id for v in data.vendors] == VENDOR_IDS


# --------------------------------------------------------------------------- header tolerance


def test_header_variants_map_to_the_same_columns(tmp_path):
    source = edited_copy(
        tmp_path,
        set_cells(
            D4="Vendor Descriptions",  # fuzzy (rapidfuzz), not in the alias table
            E4="Service / Product Provided to Meridian",
            G4="public  website",
            O4="AI Usage Detected",
            P4="Evidence Excerpt",
            V4="Assessed By / Date",
        ),
    )
    data = read_workbook(source)
    assert data.column_map == COLUMN_MAP
    assert data.issues == []


def test_missing_provided_header_reports_closest_match(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(K4="Throughput per year")))
    [issue] = with_code(data, "E03")
    assert issue.severity == "error"
    assert "Data Volume (annual)" in issue.message
    assert "closest match" in issue.message and "Throughput per year" in issue.message
    assert issue.cell == "K4"
    assert not data.ok
    assert "data_volume" not in data.column_map
    assert [v.data_volume for v in data.vendors] == [""] * 6


def test_missing_student_header_is_an_error(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(V4=None)))
    [issue] = with_code(data, "E03")
    assert issue.severity == "error"
    assert "Assessed By / Date" in issue.message and "closest match" in issue.message
    assert not data.ok


# --------------------------------------------------------------------------- reader issue codes


def test_e01_when_no_sheet_has_vendor_ids(tmp_path):
    wb = Workbook()
    wb.active.title = "Notes"
    wb.active["A1"] = "Nothing to see here"
    wb.save(tmp_path / "notes.xlsx")
    data = read_workbook(tmp_path / "notes.xlsx")
    assert [issue.code for issue in data.issues] == ["E01"]
    assert not data.ok and data.vendors == []


def test_e01_when_the_file_is_not_a_workbook():
    data = read_workbook(b"Vendor ID,Vendor Name\nV-001,AutomWorx\n")
    assert [issue.code for issue in data.issues] == ["E01"]


def test_sheet_found_by_vendor_id_header_when_renamed(tmp_path):
    data = read_workbook(edited_copy(tmp_path, lambda wb: setattr(wb[SHEET], "title", "Inventory 2026")))
    assert data.sheet_name == "Inventory 2026"
    assert [v.vendor_id for v in data.vendors] == VENDOR_IDS
    assert data.ok


def test_e02_when_the_sheet_has_no_vendor_id_header(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(B4="Supplier Key")))
    assert [issue.code for issue in data.issues] == ["E02"]
    assert data.sheet_name == SHEET and not data.ok


def test_e04_duplicate_vendor_ids(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(B7="V-001")))
    [issue] = with_code(data, "E04")
    assert "V-001" in issue.message and issue.cell == "B7"
    assert not data.ok


def test_e05_no_vendor_rows(tmp_path):
    blank = {f"{column}{row}": None for row in VENDOR_ROWS for column in "BCDEFGHIJK"}
    data = read_workbook(edited_copy(tmp_path, set_cells(**blank)))
    assert [issue.code for issue in data.issues if issue.severity == "error"] == ["E05"]
    assert data.vendors == [] and data.example is not None and not data.ok


def test_blank_row_does_not_hide_the_vendors_below_it(tmp_path):
    blank = {f"{column}7": None for column in "BCDEFGHIJK"}
    data = read_workbook(edited_copy(tmp_path, set_cells(**blank)))
    assert [v.vendor_id for v in data.vendors] == ["V-001", "V-003", "V-004", "V-005", "V-006"]
    [issue] = with_code(data, "E05")
    assert issue.severity == "warning" and "7" in issue.message and data.ok


def test_row_with_data_but_no_vendor_id_warns(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(B7=None)))
    assert "V-002" not in [v.vendor_id for v in data.vendors] and len(data.vendors) == 5
    [issue] = with_code(data, "E05")
    assert issue.severity == "warning" and issue.cell == "B7"


@pytest.mark.parametrize(("cell", "header", "field"), [
    ("M4", "Criticality", "criticality_rationale"),
    ("T4", "Risk", "risk_rationale"),
    ("T4", "Notes on Risk Rationale", "risk_rationale"),
])
def test_generic_header_words_do_not_fuzzy_match(tmp_path, cell, header, field):
    data = read_workbook(edited_copy(tmp_path, set_cells(**{cell: header})))
    assert field not in data.column_map
    assert any(issue.code == "E03" for issue in data.issues)


def test_e06_filled_student_cells_warn(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(L6="High", M6="Because")))
    [issue] = with_code(data, "E06")
    assert issue.severity == "warning" and issue.cell == "L6"
    assert "L6" in issue.message and "M6" in issue.message
    assert data.ok


def test_e07_vendor_without_website_warns(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(G8=None)))
    [issue] = with_code(data, "E07")
    assert (issue.severity, issue.cell) == ("warning", "G8")
    assert data.vendors[2].vendor_id == "V-003" and data.vendors[2].website == ""
    assert data.ok


def test_e08_unknown_dependency_label_warns(tmp_path):
    data = read_workbook(edited_copy(tmp_path, set_cells(I9="Very High", I8=" moderate ")))
    [issue] = with_code(data, "E08")
    assert (issue.severity, issue.cell) == ("warning", "I9")
    assert "Very High" in issue.message
    assert data.ok


# --------------------------------------------------------------------------- writing student cells


def test_writes_l_to_n_for_all_vendors_and_keeps_everything_else(tmp_path):
    values = {vendor_id: l_to_n(vendor_id) for vendor_id in VENDOR_IDS}
    log = SheetSpec(
        title="Evidence Log",
        headers=["E-ID", "Vendor", "Excerpt"],
        rows=[["E-001", "V-005", "“AI-powered” payments"]],
        note="Every cited excerpt.",
    )
    out, report = write(tmp_path, values, [log])

    assert check_fidelity(INPUT, out) == []
    ws = inventory(out)
    for row, vendor_id in zip(VENDOR_ROWS, VENDOR_IDS):
        for field in ("criticality_tier", "criticality_rationale", "assessment_depth"):
            cell = ws[f"{COLUMN_MAP[field]}{row}"]
            assert cell.value == getattr(values[vendor_id], field)
            assert cell.data_type == "s"
        assert ws[f"O{row}"].value is None
    assert sorted(report.written) == sorted(f"{column}{row}" for row in VENDOR_ROWS for column in "LMN")
    assert len(report.skipped) == 6 * 8 and "O6" in report.skipped and "L6" not in report.skipped
    assert report.warnings == []
    assert report.sheets_added == ["Evidence Log"]
    assert [v.vendor_id for v in read_workbook(out).vendors] == VENDOR_IDS


def test_col_spans_survive_writing_every_student_column(tmp_path):
    out, _ = write(tmp_path, {vendor_id: every_field(vendor_id) for vendor_id in VENDOR_IDS})
    spans = [(int(col.get("min")), int(col.get("max"))) for col in sheet_xml(out).iter(f"{MAIN}col")]
    assert spans.count((12, 15)) == 1
    assert spans.count((18, 20)) == 1
    assert [span for span in spans if span[0] in {13, 14, 15, 19, 20}] == []
    assert check_fidelity(INPUT, out) == []


def test_formula_like_text_stays_text(tmp_path):
    payloads = {"evidence": '=HYPERLINK("http://x")', "how_ai_used": "=1+1", "ai_subprocessors": "#N/A"}
    out, _ = write(tmp_path, {"V-001": StudentCells(**payloads)})
    ws = inventory(out)
    for field, text in payloads.items():
        cell = ws[f"{COLUMN_MAP[field]}6"]
        assert (cell.value, cell.data_type) == (text, "s")
    cells = {c.get("r"): c for c in sheet_xml(out).iter(f"{MAIN}c")}
    assert cells["P6"].get("t") == "inlineStr"
    assert [r for r, c in cells.items() if c.find(f"{MAIN}f") is not None] == []


def test_control_characters_are_removed_and_reported(tmp_path):
    out, report = write(tmp_path, {"V-001": StudentCells(criticality_rationale="Bell\x07 and null\x00 removed")})
    assert inventory(out)["M6"].value == "Bell and null removed"
    assert [(w.code, w.severity, w.cell) for w in report.warnings] == [("E11", "warning", "M6")]


def test_example_row_is_never_written(tmp_path):
    error = refused(tmp_path, {"V-000": StudentCells(criticality_tier="High")})
    assert error.code == "E12"
    assert "example" in str(error)


def test_unknown_vendor_is_refused(tmp_path):
    assert refused(tmp_path, {"V-999": StudentCells(criticality_tier="High")}).code == "E12"


def test_provided_columns_cannot_be_written(tmp_path):
    assert set(StudentCells.model_fields) == set(STUDENT_FIELDS)  # no field can address B-K
    # A template whose "Criticality Tier" header sits over a provided column: the cream check refuses.
    swapped = edited_copy(tmp_path, set_cells(C4="Criticality Tier", L4="Vendor Name"), name="swapped.xlsx")
    assert read_workbook(swapped).column_map["criticality_tier"] == "C"
    error = refused(tmp_path, {"V-001": StudentCells(criticality_tier="High")}, source=swapped)
    assert error.code == "E09" and error.issues[0].cell == "C6"


def test_non_cream_target_is_refused(tmp_path):
    def whiten(wb: Workbook) -> None:
        wb[SHEET]["L7"].fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")

    source = edited_copy(tmp_path, whiten, name="white.xlsx")
    error = refused(tmp_path, {"V-002": StudentCells(criticality_tier="High")}, source=source)
    assert error.code == "E09" and error.issues[0].cell == "L7"


def test_filled_cell_needs_overwrite(tmp_path):
    source = edited_copy(tmp_path, set_cells(L6="Low"), name="filled.xlsx")
    error = refused(tmp_path, {"V-001": StudentCells(criticality_tier="High")}, source=source)
    assert error.code == "E06" and error.issues[0].cell == "L6"

    out, report = write(tmp_path, {"V-001": StudentCells(criticality_tier="High")}, source=source, overwrite=True)
    assert inventory(out)["L6"].value == "High"
    assert [(w.code, w.severity, w.cell) for w in report.warnings] == [("E06", "warning", "L6")]


def test_template_errors_block_writing(tmp_path):
    source = edited_copy(tmp_path, set_cells(B7="V-001"), name="duplicate.xlsx")
    assert refused(tmp_path, {}, source=source).code == "E04"


def test_default_length_budgets_match_the_design():
    assert DEFAULT_LENGTH_BUDGETS["evidence"] == 1390
    assert DEFAULT_LENGTH_BUDGETS["how_ai_used"] == 1670
    assert DEFAULT_LENGTH_BUDGETS["recommended_action"] == 1480
    others = set(STUDENT_FIELDS) - {"evidence", "how_ai_used", "recommended_action"}
    assert {DEFAULT_LENGTH_BUDGETS[field] for field in others} == {1050}


def test_length_budget_is_enforced_not_truncated(tmp_path):
    at_budget = "x" * DEFAULT_LENGTH_BUDGETS["evidence"]
    out, _ = write(tmp_path, {"V-001": StudentCells(evidence=at_budget)})
    assert inventory(out)["P6"].value == at_budget

    error = refused(tmp_path, {"V-001": StudentCells(evidence=at_budget + "x")})
    assert error.code == "E10" and error.issues[0].cell == "P6"
    assert "1,391" in str(error) and "1,390" in str(error)


def test_length_budgets_can_be_overridden(tmp_path):
    signed = {"V-001": StudentCells(assessed_by="Team Test / 05-10-2026")}
    assert refused(tmp_path, signed, length_budgets={"assessed_by": 10}).code == "E10"
    with pytest.raises(ValueError, match="not_a_field"):
        write_workbook(INPUT, tmp_path / "out.xlsx", {}, length_budgets={"not_a_field": 10})


@pytest.mark.parametrize(
    "field, value",
    [
        ("criticality_tier", "Very High"),
        ("criticality_tier", "high"),
        ("ai_usage_detected", "Y"),
        ("ai_usage_detected", "Yes "),
        ("ai_risk_class", "None"),
        ("ai_risk_class", "Medium-High"),
    ],
)
def test_values_outside_the_allowed_set_are_refused(tmp_path, field, value):
    error = refused(tmp_path, {"V-001": StudentCells(**{field: value})})
    assert error.code == "E10" and error.issues[0].cell == f"{COLUMN_MAP[field]}6"


def test_allowed_values_are_accepted(tmp_path):
    accepted = {"criticality_tier": "Low", "ai_usage_detected": "No", "ai_risk_class": "None identified"}
    out, _ = write(tmp_path, {"V-003": StudentCells(**accepted)})
    ws = inventory(out)
    assert [ws[f"{COLUMN_MAP[field]}8"].value for field in accepted] == list(accepted.values())


def test_every_problem_is_reported_at_once(tmp_path):
    error = refused(
        tmp_path,
        {
            "V-001": StudentCells(criticality_tier="Severe", evidence="x" * 2000),
            "V-999": StudentCells(assessed_by="Team Test"),
        },
        sheets=[SheetSpec(title="Field Guide", headers=["A"])],
    )
    assert sorted(issue.code for issue in error.issues) == ["E10", "E10", "E12", "E13"]


def test_source_file_is_never_modified(tmp_path):
    source = tmp_path / "template.xlsx"
    shutil.copyfile(INPUT, source)
    before = sha256(source)
    write(tmp_path, {vendor_id: every_field(vendor_id) for vendor_id in VENDOR_IDS}, source=source)
    assert sha256(source) == before


def test_destination_must_differ_from_source(tmp_path, monkeypatch):
    source = tmp_path / "template.xlsx"
    shutil.copyfile(INPUT, source)
    before = sha256(source)
    monkeypatch.chdir(tmp_path)
    for destination in (source, str(source), "template.xlsx", Path(".") / "template.xlsx"):
        with pytest.raises(WorkbookWriteError) as caught:
            write_workbook(source, destination, {"V-001": l_to_n("V-001")})
        assert caught.value.code == "E13"
    assert sha256(source) == before


def test_bytes_in_and_stream_out(tmp_path):
    buffer = io.BytesIO()
    write_workbook(INPUT.read_bytes(), buffer, {"V-005": l_to_n("V-005")})
    assert check_fidelity(INPUT, buffer.getvalue()) == []
    assert inventory(buffer.getvalue())["L10"].value == "High"


# --------------------------------------------------------------------------- row heights


def test_row_height_grows_to_fit_the_longest_cell(tmp_path):
    out, report = write(
        tmp_path,
        {
            "V-001": StudentCells(criticality_rationale="x" * 900),  # M is 30.5 wide: 35 chars a line, 26 lines
            "V-002": StudentCells(criticality_tier="High"),  # one line: the template's 55 pt already fits
        },
    )
    ws = inventory(out)
    assert ws.row_dimensions[6].height == 26 * 12.75 + 8
    assert ws.row_dimensions[7].height == 55
    assert report.row_heights == {6: 339.5}


def test_row_height_is_never_lowered(tmp_path):
    def tall(wb: Workbook) -> None:
        wb[SHEET].row_dimensions[6].height = 300

    source = edited_copy(tmp_path, tall, name="tall.xlsx")
    out, report = write(tmp_path, {"V-001": StudentCells(criticality_tier="High")}, source=source)
    assert inventory(out).row_dimensions[6].height == 300
    assert report.row_heights == {}


def test_row_height_is_capped_at_409(tmp_path):
    sixty_lines = "\n".join(["x"] * 60)
    out, report = write(tmp_path, {"V-001": StudentCells(criticality_rationale=sixty_lines)})
    assert inventory(out).row_dimensions[6].height == 409
    assert report.row_heights == {6: 409}


def test_row_taller_than_the_cap_is_kept(tmp_path):
    def very_tall(wb: Workbook) -> None:
        wb[SHEET].row_dimensions[6].height = 420

    source = edited_copy(tmp_path, very_tall, name="very-tall.xlsx")
    out, _ = write(tmp_path, {"V-001": StudentCells(criticality_rationale="\n".join(["x"] * 60))}, source=source)
    assert inventory(out).row_dimensions[6].height == 420


# --------------------------------------------------------------------------- appended sheets


def test_appended_sheet_layout(tmp_path):
    spec = SheetSpec(
        title="Evidence Log",
        headers=["E-ID", "Vendor", "Excerpt", "Bytes"],
        rows=[["E-001", "V-005", '=cmd|" /C calc"!A0', 1234], ["E-002", "V-002", "plain", 2.5]],
        column_widths=[12, 14],
        note="Every cited excerpt, with its retrieval date and SHA-256.",
    )
    out, report = write(tmp_path, {}, [spec])
    ws = openpyxl.load_workbook(out)["Evidence Log"]

    assert ws["A1"].value == spec.note and ws["A1"].font.i
    header = ws[2]
    assert [cell.value for cell in header] == spec.headers
    for cell in header:
        assert cell.font.b and cell.alignment.wrap_text
        assert (cell.fill.fill_type, cell.fill.fgColor.rgb) == ("solid", CREAM)
    assert ws.freeze_panes == "A3"
    assert ws.auto_filter.ref == "A2:D4"
    assert (ws["C3"].value, ws["C3"].data_type) == ('=cmd|" /C calc"!A0', "s")
    assert (ws["D3"].value, ws["D4"].value) == (1234, 2.5)
    assert ws["A3"].alignment.wrap_text and ws["A3"].alignment.vertical == "top"
    assert (ws.column_dimensions["A"].width, ws.column_dimensions["B"].width) == (12, 14)
    assert 10 <= ws.column_dimensions["C"].width <= 60
    assert report.sheets_added == ["Evidence Log"]


def test_appended_sheet_without_note_starts_at_row_1(tmp_path):
    spec = SheetSpec(title="Coverage Log", headers=["Vendor", "Family"], rows=[["V-001", "DNS"]])
    out, _ = write(tmp_path, {}, [spec])
    ws = openpyxl.load_workbook(out)["Coverage Log"]
    assert [cell.value for cell in ws[1]] == ["Vendor", "Family"]
    assert ws["A1"].fill.fgColor.rgb == CREAM
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B2"


def test_appended_sheet_control_characters_are_reported(tmp_path):
    spec = SheetSpec(title="Run Info", headers=["Key", "Value"], rows=[["note", "tab\tkept, bell\x07 dropped"]])
    out, report = write(tmp_path, {}, [spec])
    assert openpyxl.load_workbook(out)["Run Info"]["B2"].value == "tab\tkept, bell dropped"
    assert [(w.code, w.cell) for w in report.warnings] == [("E11", "'Run Info'!B2")]


@pytest.mark.parametrize("title", ["Field Guide", "field guide", "VENDOR INVENTORY", "Evidence/Log", "x" * 32,
                                   "History", "history"])
def test_appended_sheet_never_replaces_or_breaks_a_sheet(tmp_path, title):
    error = refused(tmp_path, {}, sheets=[SheetSpec(title=title, headers=["A"])])
    assert error.code == "E13"


def test_two_appended_sheets_cannot_share_a_title(tmp_path):
    specs = [SheetSpec(title="Run Info", headers=["A"]), SheetSpec(title="run info", headers=["B"])]
    assert refused(tmp_path, {}, sheets=specs).code == "E13"


def test_sheet_order_active_tab_and_author(tmp_path):
    specs = [SheetSpec(title="Evidence Log", headers=["E-ID"]), SheetSpec(title="Run Info", headers=["Key"])]
    out, _ = write(tmp_path, {}, specs, last_modified_by="Team Test")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [SHEET, "Field Guide", "Points to consider", "Evidence Log", "Run Info"]
    assert wb.active.title == "Field Guide"
    assert wb.properties.lastModifiedBy == "Team Test"


# --------------------------------------------------------------------------- fidelity check


def test_round_trip_is_faithful(tmp_path):
    assert check_fidelity(INPUT, edited_copy(tmp_path)) == []


def test_flattened_rich_text_is_detected(tmp_path):
    wb = openpyxl.load_workbook(INPUT)  # default load drops the bold runs of the example's P5
    wb.save(tmp_path / "flat.xlsx")
    assert any("P5" in diff for diff in check_fidelity(INPUT, tmp_path / "flat.xlsx"))


def test_allowed_changes_pass(tmp_path):
    def allowed(wb: Workbook) -> None:
        ws = wb[SHEET]
        ws["L6"] = "High"
        ws["V11"] = "Team Test / 05-10-2026"
        ws.row_dimensions[6].height = 120
        wb.create_sheet("Evidence Log")["A1"] = "E-ID"

    assert check_fidelity(INPUT, edited_copy(tmp_path, allowed)) == []


@pytest.mark.parametrize(
    "edit, expected",
    [
        (set_cells(C6="AutomWorx Ltd"), "C6"),
        (set_cells(L5="High"), "L5"),  # the example row never changes
        (set_cells(W6="stray"), "W6"),
        (lambda wb: setattr(wb[SHEET]["G10"], "hyperlink", "http://example.org/"), "G10"),
        (lambda wb: setattr(wb[SHEET]["L6"], "fill", PatternFill(fill_type="solid", fgColor="FFFFFFFF")), "L6"),
        (lambda wb: setattr(wb[SHEET]["D6"], "number_format", "0.00"), "D6"),
        (lambda wb: setattr(wb[SHEET].column_dimensions["M"], "width", 30.54296875), "column M"),
        (lambda wb: setattr(wb[SHEET].column_dimensions["P"], "width", 50), "column P"),
        (lambda wb: setattr(wb[SHEET].row_dimensions[6], "height", 40), "row 6"),
        (lambda wb: setattr(wb[SHEET].row_dimensions[5], "height", 400), "row 5"),
        (lambda wb: wb[SHEET].unmerge_cells("B3:V3"), "merged"),
        (lambda wb: setattr(wb[SHEET], "freeze_panes", "C5"), "freeze"),
        (lambda wb: setattr(wb[SHEET].sheet_view, "zoomScale", 100), "view"),
        (lambda wb: wb.move_sheet("Points to consider", offset=-2), "order"),
        (lambda wb: setattr(wb, "active", 0), "active"),
        (lambda wb: wb.remove(wb["Points to consider"]), "Points to consider"),
        (lambda wb: setattr(wb["Field Guide"]["D4"], "value", "Changed"), "Field Guide"),
    ],
)
def test_fidelity_detects_changes(tmp_path, edit, expected):
    diffs = check_fidelity(INPUT, edited_copy(tmp_path, edit))
    assert any(expected in diff for diff in diffs), diffs


# --------------------------------------------------------------------------- reviewer regressions (P1 fixer)


@pytest.mark.parametrize("bad", ["\ud800", "￿", "￾"])
def test_unstorable_characters_are_removed_and_reported(tmp_path, bad):
    out, report = write(tmp_path, {"V-001": StudentCells(criticality_rationale=f"a{bad}b")})
    assert inventory(out)["M6"].value == "ab"
    assert [w.code for w in report.warnings] == ["E11"]
    assert [p.name for p in tmp_path.iterdir()] == ["out.xlsx"]  # no temporary file left behind


def test_unstorable_characters_in_appended_sheets_are_removed(tmp_path):
    out, report = write(tmp_path, {}, [SheetSpec(title="Run Info", headers=["k"], rows=[["x\ud800y"]])])
    assert openpyxl.load_workbook(out)["Run Info"]["A2"].value == "xy"
    assert [w.code for w in report.warnings] == ["E11"]


def test_failed_save_raises_the_real_error_and_leaves_no_temporary_file(tmp_path, monkeypatch):
    def broken(self: Workbook, filename: object) -> None:
        raise RuntimeError("serialisation failed")

    monkeypatch.setattr(Workbook, "save", broken)
    with pytest.raises(RuntimeError, match="serialisation failed"):
        write_workbook(INPUT, tmp_path / "out.xlsx", {})
    assert list(tmp_path.iterdir()) == []


def test_literal_excel_escape_is_escaped(tmp_path):
    out, _ = write(tmp_path, {"V-001": StudentCells(criticality_rationale="lit _x0041_ end")})
    with zipfile.ZipFile(out) as archive:
        xml = "".join(archive.read(name).decode() for name in archive.namelist() if name.startswith("xl/"))
    assert "lit _x005F_x0041_ end" in xml


def test_fidelity_flags_a_formula_in_a_student_cell(tmp_path):
    out, _ = write(tmp_path, {})
    wb = openpyxl.load_workbook(out, rich_text=True)
    wb[SHEET]["M6"] = "=1+1"
    wb.save(out)
    assert any("M6" in problem and "formula" in problem for problem in check_fidelity(INPUT, out))


def test_output_is_byte_reproducible_with_a_fixed_time(tmp_path, monkeypatch):
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    cells = {"V-001": l_to_n("V-001")}
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "1767225600")  # 2026-01-01T00:00:00Z
    write_workbook(INPUT, first, cells)
    monkeypatch.delenv("SOURCE_DATE_EPOCH")
    write_workbook(INPUT, second, cells, modified=dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc))
    assert sha256(first) == sha256(second)
    assert openpyxl.load_workbook(first).properties.modified.year == 2026
