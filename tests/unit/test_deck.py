"""deck/build_deck.py and deck/diagrams.py: the Team Osprey walkthrough deck (python-pptx, native shapes only).

Every test is offline: the build reads only the local workbook, config/ and the given findings, and a guard makes any
socket connection fail. Decks are written to tmp_path, never to submission/.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import socket
import sys
import zipfile
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Emu  # noqa: E402

from footprint.criticality import load_rubric, score_profile  # noqa: E402
from footprint.depth import load_depth_config  # noqa: E402
from footprint.models import SourceFamily, Tier  # noqa: E402
from footprint.workbook import read_workbook  # noqa: E402

REPO = Path(__file__).resolve().parents[2]
WORKBOOK = REPO / "data" / "input" / "Meridian_Vendor_Input.xlsx"
CONFIG = REPO / "config"
AS_OF = "2026-10-02"
TEAM = "Team Osprey"
MAIN = 13
APPENDIX_FIXED = 6  # divider + A1-A5
BRANDING = ("optiv",)  # the sponsor's name never appears in the deck, in any case


def _load(name: str, path: Path):
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    return sys.modules[name]


deck = _load("footprint_deck_build", REPO / "deck" / "build_deck.py")
samples = _load("p3_samples", REPO / "tests" / "fixtures" / "p3_samples.py")


# --------------------------------------------------------------------------- helpers


@pytest.fixture(autouse=True)
def no_network(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(*_args, **_kwargs):
        raise AssertionError("the deck build must not open network connections")

    monkeypatch.setattr(socket.socket, "connect", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.delenv("FOOTPRINT_TEAM_NAME", raising=False)


def shape_texts(shape) -> list[str]:
    out: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            out += shape_texts(child)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        out.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        out += [cell.text for row in shape.table.rows for cell in row.cells]
    return out


def slide_text(slide) -> str:
    return "\n".join(t for shape in slide.shapes for t in shape_texts(shape))


def notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def table_rows(slide, name: str) -> list[list[str]]:
    frame = next(s for s in slide.shapes if s.name == name)
    return [[cell.text for cell in row.cells] for row in frame.table.rows]


def iter_runs(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        yield slide, shape, run


def sample_assessment():
    """Two synthetic vendors (fictional Acme, V-901/V-902): one Confirmed with a trap, one Not detected."""
    item = samples.item(evidence_id="V-901-E-0001", role="Primary", providers=[])
    trap = samples.item(
        excerpt="Alerts are reviewed by an analyst before any payment is held.", evidence_id="V-901-E-0002",
        tags=samples.tags(u_class="U7", sp="S1", rl="R1", ai_type="not_ai", strength="Weak", ic=3),
        indicators=[], role="Logged")
    first = samples.findings(evidence=[item, trap], verdict=samples.verdict([item.item_key]))
    p2 = samples.profile(vendor_id="V-902", name="Beta Mailing Services, Inc. (BMS)", row=13)
    second = samples.findings(
        profile=p2, evidence=[], coverage=[samples.coverage(vendor_id="V-902")],
        verdict=samples.verdict([], rule="e", likelihood="unlikely", qualifying=[], decisive=[]),
        risk=samples.risk_result(
            samples.risk_inputs(e=0, k=0, e_items=[], k_items=[]), arp=4, base_class="Low", gate_met=False,
            cap="None identified", pre_cap_class="Low", final_class="None identified",
            flip_condition="Generative AI on Meridian data would make the class Medium."),
        actions=samples.action_plan(gap_blocks=["GOV"], questionnaire_items=["Q13"],
                                    text="Request a written no-AI attestation."))
    return samples.assessment(vendors=[first, second], manifest={
        "llm": {"calls": {"V-901": {"planned": 6, "actual": 4}}, "withheld": {"V-901": {"count": 1, "rate": 0.1}}}})


@pytest.fixture(scope="module")
def placeholder(tmp_path_factory):
    out = tmp_path_factory.mktemp("deck") / "placeholder.pptx"
    info = deck.build_deck(out, findings=None, workbook=WORKBOOK, team=TEAM, as_of=AS_OF)
    return info, Presentation(str(out))


@pytest.fixture(scope="module")
def with_findings(tmp_path_factory):
    folder = tmp_path_factory.mktemp("deck")
    path = folder / "assessment.json"
    path.write_text(sample_assessment().model_dump_json(), encoding="utf-8")
    out = folder / "findings.pptx"
    info = deck.build_deck(out, findings=path, workbook=WORKBOOK, team=TEAM)
    return info, Presentation(str(out))


# --------------------------------------------------------------------------- build and structure


def test_placeholder_build_has_13_main_slides_plus_appendix(placeholder) -> None:
    info, prs = placeholder
    assert info["main_slides"] == MAIN
    assert len(prs.slides) == info["slides"] == MAIN + APPENDIX_FIXED
    assert info["findings"] is False and info["run_id"] == ""
    assert prs.slide_width == Emu(12192000) and prs.slide_height == Emu(6858000)  # 16:9, 13.33 x 7.5 in


def test_every_slide_has_speaker_notes(placeholder, with_findings) -> None:
    for _, prs in (placeholder, with_findings):
        for index, slide in enumerate(prs.slides, start=1):
            assert len(notes(slide).strip()) >= 40, f"slide {index} has no speaker notes"


def test_slide_titles_follow_the_design_outline(placeholder) -> None:
    _, prs = placeholder
    titles = [slide.shapes.title.text if slide.shapes.title is not None else "" for slide in prs.slides]
    expected = ["Reading the public footprint", "Answer first", "Approach", "Functional design", "Architecture",
                "Source strategy", "Evidentiary weight", "AI where it adds recall", "Criticality", "Depth",
                "Findings and risk", "Live demo", "Recommendations"]
    for index, start in enumerate(expected):
        assert titles[index].startswith(start), (index + 1, titles[index])
    assert "Team Osprey".upper() in slide_text(prs.slides[0]).upper()


def test_outcome_corner_tags_on_outcome_slides(placeholder) -> None:
    _, prs = placeholder
    tags = {4: "OUTCOME 01", 5: "OUTCOME 02", 6: "OUTCOME 03", 7: "OUTCOME 03", 9: "OUTCOME 04", 10: "OUTCOME 04",
            12: "OUTCOME 05"}
    for number, tag in tags.items():
        assert tag in slide_text(prs.slides[number - 1]), number


def test_no_sponsor_branding_anywhere(placeholder, with_findings, tmp_path) -> None:
    for info, prs in (placeholder, with_findings):
        for index, slide in enumerate(prs.slides, start=1):
            text = (slide_text(slide) + "\n" + notes(slide)).lower()
            for word in BRANDING:
                assert word not in text, f"slide {index} mentions {word!r}"
        props = prs.core_properties
        meta = " ".join(str(v) for v in (props.title, props.subject, props.author, props.keywords, props.comments,
                                         props.category, props.last_modified_by)).lower()
        assert not any(word in meta for word in BRANDING)
        with zipfile.ZipFile(info["path"]) as package:
            for name in package.namelist():
                if name.endswith(".xml"):
                    assert b"optiv" not in package.read(name).lower(), name


def test_diagrams_are_native_shapes_not_pictures(placeholder, with_findings) -> None:
    for info, prs in (placeholder, with_findings):
        for index, slide in enumerate(prs.slides, start=1):
            kinds = [shape.shape_type for shape in slide.shapes]
            assert MSO_SHAPE_TYPE.PICTURE not in kinds, f"slide {index} contains a picture"
        with zipfile.ZipFile(info["path"]) as package:
            assert not [n for n in package.namelist() if n.startswith("ppt/media/")]
    _, prs = placeholder
    for number, minimum in ((4, 25), (5, 20)):
        slide = prs.slides[number - 1]
        autoshapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
        lines = [s for s in slide.shapes if s.shape_type in (MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM)
                 or s.element.tag.endswith("}cxnSp")]
        assert len(autoshapes) >= minimum, (number, len(autoshapes))
        assert len(lines) >= 6, (number, len(lines))


def test_functional_flow_has_every_component_and_checkpoint(placeholder) -> None:
    _, prs = placeholder
    names = {shape.name for shape in prs.slides[3].shapes}
    text = slide_text(prs.slides[3])
    for code in [f"C{i}" for i in range(1, 11)]:
        assert any(n.startswith(code + " ") for n in names), code
    for checkpoint in ("HC1", "HC2"):
        assert checkpoint in text
    assert "LLM audit gate" in text  # the release is gated by code, not by a third sign-off
    assert "Stop-rule loop" in text and "Evidence store" in text and "Gemini" in text
    # connectors are glued to the shapes they join, so they follow them when edited
    glued = [s for s in prs.slides[3].shapes if s.element.tag.endswith("}cxnSp")
             and s.element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}stCxn") is not None]
    assert len(glued) >= 12


def test_architecture_shows_layers_services_and_trust_boundary(placeholder) -> None:
    _, prs = placeholder
    text = slide_text(prs.slides[4])
    for word in ("CLI", "Streamlit", "Notebook", "run_assessment", "Net policy", "Collect", "Analyse", "Decide",
                 "evidence/", "SEC EDGAR", "DNS over HTTPS", "ATS APIs", "Wayback Machine", "Gemini API",
                 "Payload guard", "Trust boundary", "replay"):
        assert word in text, word
    boundary = next(s for s in prs.slides[4].shapes if s.name == "Trust boundary")
    assert boundary.line.dash_style is not None


def test_text_is_readable_and_inside_the_slide(placeholder, with_findings) -> None:
    for _, prs in (placeholder, with_findings):
        for slide, shape, run in iter_runs(prs):
            if run.text.strip() and run.font.size is not None:
                assert run.font.size.pt >= 10, (shape.name, run.text, run.font.size.pt)
        for index, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.width is None or shape.element.tag.endswith("}cxnSp"):
                    continue
                assert shape.left >= 0 and shape.top >= 0, (index, shape.name)
                assert shape.left + shape.width <= prs.slide_width + Emu(9525), (index, shape.name)
                assert shape.top + shape.height <= prs.slide_height + Emu(9525), (index, shape.name)


def test_build_is_deterministic(placeholder, tmp_path) -> None:
    first, _ = placeholder
    again = deck.build_deck(tmp_path / "again.pptx", findings=None, workbook=WORKBOOK, team=TEAM, as_of=AS_OF)
    assert again["sha256"] == first["sha256"]
    assert hashlib.sha256((tmp_path / "again.pptx").read_bytes()).hexdigest() == again["sha256"]


def test_output_must_be_pptx(tmp_path) -> None:
    with pytest.raises(ValueError):
        deck.build_deck(tmp_path / "deck.pdf", workbook=WORKBOOK, team=TEAM, as_of=AS_OF)


def test_team_name_comes_from_the_environment_when_not_given(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("FOOTPRINT_TEAM_NAME", "Team Heron")
    info = deck.build_deck(tmp_path / "t.pptx", workbook=None, as_of=AS_OF, appendix=False)
    prs = Presentation(info["path"])
    assert len(prs.slides) == MAIN
    assert "TEAM HERON" in slide_text(prs.slides[0])
    assert prs.core_properties.author == "Team Heron"


# --------------------------------------------------------------------------- live inputs


def test_slide_9_tier_table_is_computed_live_from_the_workbook(placeholder) -> None:
    if not WORKBOOK.is_file():
        pytest.skip("input workbook not present")
    _, prs = placeholder
    rows = table_rows(prs.slides[8], "Tier table")
    data = read_workbook(WORKBOOK)
    rubric = load_rubric(CONFIG / "rubric.toml")
    profiles = [data.example, *data.vendors]
    assert len(rows) == 1 + len(profiles)  # header + V-000 + six vendors
    for row, profile in zip(rows[1:], profiles):
        result = score_profile(profile, rubric)
        assert row[6] == str(result.score)
        assert row[8] == result.tier.value
        assert profile.vendor_id in row[0] or "V-000" in row[0]
    assert "payment-path floor" in slide_text(prs.slides[8])


def test_slide_6_reads_the_depth_config(placeholder) -> None:
    _, prs = placeholder
    cfg = load_depth_config(CONFIG / "depth.toml")
    rows = table_rows(prs.slides[5], "Family by tier table")
    assert [r[0].split("\n")[0] for r in rows[1:]] == [f.value for f in SourceFamily]
    leg = cfg.tiers[Tier.CRITICAL].families[SourceFamily.LEG]
    assert f"≤ {leg.cap} req." in rows[1][1]
    excluded = slide_text(prs.slides[5])
    for source in cfg.excluded_sources:
        assert source.source in excluded


def test_sr_ladder_from_sources_toml_and_fallback(tmp_path) -> None:
    ladder, source = deck.load_sr_ladder(tmp_path)
    assert source.startswith("design") and [g for g, _, _ in ladder] == ["A", "B", "C", "D", "E–F"]
    (tmp_path / "sources.toml").write_text(
        '[[source]]\nsource_type = "Annual report"\nsr = "A"\nfamily = "REG"\n'
        '[[source]]\nsource_type = "Forum post"\nsr = "E"\nfamily = "IND"\n', encoding="utf-8")
    ladder, source = deck.load_sr_ladder(tmp_path)
    assert source == "config/sources.toml"
    assert ladder[0][0] == "A" and "Annual report" in ladder[0][2]
    assert ladder[-1][0] == "E–F" and "Forum post" in ladder[-1][2]


def test_placeholders_when_no_findings(placeholder) -> None:
    _, prs = placeholder
    assert "Pending frozen run" in slide_text(prs.slides[1])
    assert "--findings" in slide_text(prs.slides[1])
    assert "V-000" in slide_text(prs.slides[10])


# --------------------------------------------------------------------------- findings


def test_findings_fill_slides_2_and_11_and_the_vendor_appendix(with_findings) -> None:
    info, prs = with_findings
    assert info["findings"] is True and info["run_id"] == "A-20261002-0000abcd"
    assert len(prs.slides) == MAIN + APPENDIX_FIXED + 2
    answer = table_rows(prs.slides[1], "Answer-first table")
    assert [r[0].split("\n")[0] for r in answer[1:]] == ["Acme Payments", "BMS"]
    assert answer[1][2].startswith("Yes · Confirmed") and answer[1][3].startswith("High")
    assert answer[2][2].startswith("No · Not detected") and answer[2][3].startswith("None identified")
    assert "V-901-E-0001" in notes(prs.slides[1])
    chain = slide_text(prs.slides[10])
    assert "V-901-E-0001" in chain and "12 of 18" in chain
    assert samples.EXCERPT[:40] in chain
    vendor = slide_text(prs.slides[MAIN + APPENDIX_FIXED])
    assert "Acme Payments (V-901)" in vendor and "V-901-E-0001" in vendor
    assert "V-901-E-0002" in vendor  # the definition-test trap, logged but never cited


def test_trap_and_strong_item_on_the_weight_slide(with_findings) -> None:
    _, prs = with_findings
    text = slide_text(prs.slides[6])
    assert "V-901-E-0002" in text and "Definition test fails" in text
    assert "V-901-E-0001" in text and "Strong" in text


def test_ai_slide_counts_come_from_the_manifest(with_findings) -> None:
    _, prs = with_findings
    text = slide_text(prs.slides[7])
    assert "4 Gemini calls" in text and "1 passage withheld" in text and "0 URLs taken from the model" in text


def test_summary_json_round_trip_and_headline(tmp_path) -> None:
    def vendor(vid, name, usage, verdict, risk, **kw):
        return {"vendor_id": vid, "name": name, "tier": "Critical", "usage": usage, "verdict": verdict, "risk": risk,
                **kw}

    def evidence(eid):
        return [{"evidence_id": eid, "role": "Primary", "strength": "Strong", "excerpt": "synthetic excerpt text"}]

    summary = {"run_id": "A-1", "mode": "replay", "as_of": AS_OF, "vendors": [
        vendor("V-001", "AutomWorx", "Inconclusive", "Inconclusive", "Medium", provisional=True, ceiling="Critical"),
        vendor("V-002", "Fiserv", "Yes", "Probable", "High", flip="Confirmation on DNA would raise it to Critical.",
               arp=13, evidence=evidence("V-002-E-0001")),
        vendor("V-003", "FSSI", "No", "Not detected", "None identified"),
        vendor("V-004", "Terrapin", "Inconclusive", "Inconclusive", "Medium", provisional=True, ceiling="High"),
        vendor("V-005", "BNY", "Yes", "Confirmed", "High", flip="External models would raise the class to Critical.",
               arp=12, evidence=evidence("V-005-E-0001")),
        vendor("V-006", "TCH", "Yes", "Probable", "High", flip="Isolated GenAI would lower the class to Medium.",
               evidence=evidence("V-006-E-0002")),
    ]}
    path = tmp_path / "summary.json"
    path.write_text(json.dumps(summary), encoding="utf-8")
    data = deck.load_findings(path)
    assert [v.name for v in data.vendors][:2] == ["AutomWorx", "Fiserv"]
    assert deck.headline(data) == (
        "AI runs in service delivery at 3 of 6 vendors (Fiserv, BNY and TCH). None reaches Critical AI risk on "
        "public evidence; AutomWorx's ceiling is Critical if confirmed, and Fiserv and BNY would become Critical on "
        "a single confirming answer.")
    assert deck.pick_chain(data).vendor_id == "V-005"
    assert deck.pick_chain(data, "V-006").vendor_id == "V-006"
    with pytest.raises(ValueError):
        deck.pick_chain(data, "V-999")


def test_load_findings_rejects_unknown_json(tmp_path) -> None:
    path = tmp_path / "bad.json"
    path.write_text(json.dumps({"vendors": [{"vendor_id": "V-001"}]}), encoding="utf-8")
    with pytest.raises(ValueError):
        deck.load_findings(path)
    path.write_text("{not json", encoding="utf-8")
    with pytest.raises(ValueError):
        deck.load_findings(path)


def test_findings_from_a_run_folder(tmp_path) -> None:
    run = tmp_path / "A-20261002-0000abcd"
    run.mkdir()
    (run / "assessment.json").write_text(sample_assessment().model_dump_json(), encoding="utf-8")
    data = deck.load_findings(run)
    assert data.run_id == "A-20261002-0000abcd" and data.mode == "replay" and data.as_of == AS_OF
    assert data.llm_calls == 4 and data.withheld == 1
    first = data.vendors[0]
    assert (first.name, first.usage, first.verdict, first.risk, first.arp) == (
        "Acme Payments", "Yes", "Confirmed", "High", 12)
    assert [e.evidence_id for e in first.evidence] == ["V-901-E-0001"]
    assert first.trap is not None and first.trap.evidence_id == "V-901-E-0002"
    assert data.vendors[1].name == "BMS" and data.vendors[1].evidence == []


def test_verdict_wording_never_repeats_a_word(with_findings) -> None:
    same = deck.DeckVendor(vendor_id="V-004", name="Terrapin", tier="High", usage="Inconclusive",
                           verdict="Inconclusive", risk="Medium")
    yes = same.model_copy(update={"usage": "Yes", "verdict": "Confirmed"})
    assert deck.verdict_words(same) == deck.verdict_words(same, sep="(") == "Inconclusive"
    assert deck.verdict_words(yes) == "Yes · Confirmed" and deck.verdict_words(yes, sep="(") == "Yes (Confirmed)"
    _, prs = with_findings
    for slide in prs.slides:
        text = slide_text(slide) + notes(slide)
        assert not re.search(r"\b(Inconclusive|Yes|No) \(\1\)", text)
    subtitle = slide_text(prs.slides[MAIN + APPENDIX_FIXED + 1])
    assert "No (Not detected) · AI risk None identified" in subtitle


def test_concentration_counts_by_basis() -> None:
    vendors = [
        deck.DeckVendor(vendor_id="V-1", name="A", tier="High", usage="Yes", risk="High",
                        providers_named=["OpenAI"], providers_dns=["Anthropic"]),
        deck.DeckVendor(vendor_id="V-2", name="B", tier="High", usage="Yes", risk="High",
                        providers_named=["OpenAI", "Google"]),
        deck.DeckVendor(vendor_id="V-3", name="C", tier="High", usage="No", risk="None identified",
                        providers_dns=["OpenAI"]),
    ]
    rows = deck.concentration(vendors)
    assert rows[0] == ("OpenAI", 2, 0, 1)
    assert ("Anthropic", 0, 0, 1) in rows and ("Google", 1, 0, 0) in rows
    assert deck.concentration_note(rows[0]) == "OpenAI: named by 2 vendors; DNS verification token only at 1 more"


def test_canonical_provider_and_short_name() -> None:
    assert deck.canonical_provider("ChatGPT") == "OpenAI"
    assert deck.canonical_provider("Amazon Bedrock") == "AWS"
    assert deck.canonical_provider("Personetics") == "Personetics"
    assert deck.short_name("Financial Statement Services, Inc. (FSSI)") == "FSSI"
    assert deck.short_name("Fiserv, Inc.") == "Fiserv"
    assert deck.short_name("The Clearing House Payments Company L.L.C.", ["TCH", "The Clearing House"]) == "TCH"


# --------------------------------------------------------------------------- checkpoints, demo script, actions


def all_text(prs) -> str:
    return "\n".join(slide_text(slide) + "\n" + notes(slide) for slide in prs.slides)


def shape_text(slide, name: str) -> str:
    return next(s for s in slide.shapes if s.name == name).text_frame.text


def demo_steps_of(prs) -> list[str]:
    return [shape_text(prs.slides[11], f"Demo text {i}") for i in range(1, 8)]


def test_only_two_human_checkpoints_and_no_release_sign_off(placeholder, with_findings) -> None:
    """The design dropped HC3 and nothing enforces a release approval, so the deck never presents one: HC1 and HC2
    are the human checkpoints (as in the README), and code gates the export (LLM audit, integrity check)."""
    for _, prs in (placeholder, with_findings):
        text = all_text(prs)
        assert "HC3" not in text
        assert not re.search(r"three (human )?checkpoints", text)
        assert "unapproved" not in text and "waits for" not in text
    _, prs = placeholder
    approach = slide_text(prs.slides[2])
    assert "two human checkpoints" in approach and "HC1" in approach and "HC2" in approach
    assert "Approve action" in notes(prs.slides[2]) and "part of HC2" in notes(prs.slides[2])
    assert "LLM audit" in notes(prs.slides[3]) and "integrity check" in notes(prs.slides[3])
    assert "Risk, then sign-off (HC2)" in demo_steps_of(prs)[4]


def test_demo_script_promises_only_what_the_run_holds(with_findings) -> None:
    """Slide 12 is generated from the run: no background screen run (the app has none), the excerpt and the trap
    the run actually holds, and no screenshot, Gemini comparison or Evidence Images when the run has none."""
    _, prs = with_findings
    steps = demo_steps_of(prs)
    text = "\n".join(steps) + "\n" + shape_text(prs.slides[11], "Watch card")
    for absent in ("background", "screen run", "screenshot", "Gemini with the rules", "Gemini and rule labels",
                   "Evidence Images", "RTP"):
        assert absent not in text, absent
    assert "Acme Payments' excerpt V-901-E-0001 in its source context" in steps[3]
    assert "reject Acme Payments' trap V-901-E-0002" in steps[3]
    assert steps[4].startswith("Risk, then sign-off (HC2)\nAcme Payments: exposure") and "Approve" in steps[4]
    assert "Evidence Log and Coverage Log" in steps[5]
    if WORKBOOK.is_file():  # step 2 comes from the live tiers: the vendor a floor lifts, with its real score
        data = read_workbook(WORKBOOK)
        rubric = load_rubric(CONFIG / "rubric.toml")
        lifted = [p for p in data.vendors
                  if score_profile(p, rubric).computed_tier.rank > score_profile(p, rubric).score_tier.rank]
        for profile in lifted[:1]:
            assert f"scores {score_profile(profile, rubric).score} and the " in steps[1]


def test_demo_script_uses_screenshots_traps_and_gemini_when_the_run_has_them(tmp_path) -> None:
    data = deck.DeckData(as_of=AS_OF, vendors=[
        deck.DeckVendor(vendor_id="V-001", name="Alpha", tier="Critical", usage="Yes", verdict="Confirmed",
                        risk="High", arp=12, e=2, k=2, tp=3, tg=1,
                        evidence=[deck.DeckEvidence(evidence_id="V-001-E-0001", excerpt="synthetic")]),
        deck.DeckVendor(vendor_id="V-002", name="Bravos", tier="High", usage="Yes", verdict="Probable",
                        risk="Medium", screenshots=3, llm_labelled=4, llm_disagree=2,
                        shot=deck.DeckEvidence(evidence_id="V-002-E-0007", excerpt="synthetic", screenshot=True)),
        deck.DeckVendor(vendor_id="V-003", name="Charlie", tier="High", usage="No", verdict="Not detected",
                        risk="None identified", trap=deck.DeckEvidence(evidence_id="V-003-E-0002")),
    ])
    info = deck.build_deck(tmp_path / "d.pptx", findings=data, workbook=None, team=TEAM, appendix=False)
    prs = Presentation(info["path"])
    steps = demo_steps_of(prs)
    assert "Bravos' excerpt V-002-E-0007 with its screenshot, then Re-verify" in steps[3]
    assert "reject Charlie's trap V-003-E-0002" in steps[3]
    assert "compare Gemini with the rules on Bravos (2 disagreements)" in steps[3]
    assert "with 3 screenshots in Evidence Images" in steps[5]
    assert "Each vendor's score, factor levels and floors" in steps[1]  # no workbook: neutral wording
    assert "Gemini and rule labels side by side" in shape_text(prs.slides[11], "Watch card")


def eight_vendors() -> deck.DeckData:
    """The replay run's shape (Fiserv Critical-tier, Yes (Probable), Medium and not provisional), plus a Low-class
    and a Low-tier None identified vendor, so every column U playbook appears."""
    def v(vid, name, tier, usage, verdict, risk, action, **kw):
        return deck.DeckVendor(vendor_id=vid, name=name, tier=tier, usage=usage, verdict=verdict, risk=risk,
                               action=action, **kw)

    q14 = "Issue questionnaire items Q1 to Q4 within 15 business days to establish whether AI is used in the service."
    return deck.DeckData(as_of=AS_OF, vendors=[
        v("V-001", "AutomWorx", "High", "Inconclusive", "Inconclusive", "Medium", q14, provisional=True,
          ceiling="Critical"),
        v("V-002", "Fiserv", "Critical", "Yes", "Probable", "Medium",
          "Issue a questionnaire within 15 business days covering which Meridian data the AI touches."),
        v("V-003", "FSSI", "High", "No", "Not detected", "None identified",
          "Obtain a written attestation from the vendor that no AI is used in the service or on Meridian data."),
        v("V-004", "Terrapin", "High", "Inconclusive", "Inconclusive", "Medium", q14, provisional=True,
          ceiling="High"),
        v("V-005", "BNY", "Critical", "Yes", "Probable", "High",
          "Issue a targeted questionnaire within 15 business days covering which Meridian data the AI touches.",
          providers_named=["Microsoft"]),
        v("V-006", "TCH", "Critical", "Yes", "Probable", "High",
          "Issue a targeted questionnaire within 15 business days covering which Meridian data the AI touches."),
        v("V-007", "Delta", "Medium", "Yes", "Probable", "Low",
          "Issue a questionnaire within 30 business days covering which Meridian data the AI touches."),
        v("V-008", "Echo", "Low", "No", "Not detected", "None identified",
          "Take no AI-specific action beyond monitoring."),
    ])


def test_slide_13_gives_every_vendor_its_thirty_day_action(tmp_path, with_findings) -> None:
    """Every vendor appears exactly once in 'Next 30 days', with its own deadline from column U: a Medium vendor
    that is not provisional (Fiserv in the replay run) is no longer dropped."""
    data = eight_vendors()
    lines = deck.thirty_day_actions(data.vendors)
    vendor_lines = [line for line in lines if not line.startswith("Register AI sub-processors")]
    for vendor in data.vendors:
        hits = [line for line in vendor_lines if re.search(rf"\b{re.escape(vendor.name)}\b", line)]
        assert len(hits) == 1, (vendor.name, hits)
    assert "Send a targeted AI questionnaire to BNY and TCH within 15 business days." in lines
    assert ("Send an AI questionnaire to Fiserv within 15 business days; Delta within 30 business days; review AI "
            "clauses at renewal.") in lines
    assert "Put questions Q1–Q4 to AutomWorx and Terrapin within 15 business days, then reclassify on the answers." \
        in lines
    assert "Ask FSSI for a written no-AI attestation and an AI change-notification clause." in lines
    assert "No AI-specific action for Echo beyond monitoring." in lines
    assert any("(Microsoft)" in line for line in lines)

    info = deck.build_deck(tmp_path / "d.pptx", findings=data, workbook=None, team=TEAM, appendix=False)
    first = shape_text(Presentation(info["path"]).slides[12], "Column text Next 30 days")
    for vendor in data.vendors:
        assert vendor.name in first, vendor.name
    _, prs = with_findings
    first = shape_text(prs.slides[12], "Column text Next 30 days")
    assert "Acme Payments" in first and "BMS" in first


@pytest.mark.parametrize("tier", ["Critical", "High", "Medium", "Low"])
@pytest.mark.parametrize("provisional", [False, True])
@pytest.mark.parametrize("risk", ["Critical", "High", "Medium", "Low", "None identified"])
def test_slide_13_groups_follow_the_column_u_playbook(tier: str, provisional: bool, risk: str) -> None:
    from types import SimpleNamespace

    from footprint.actions import playbook_key

    v = deck.DeckVendor(vendor_id="V-1", name="A", tier=tier, usage="Yes", risk=risk, provisional=provisional)
    assert deck.action_group(v) == playbook_key(SimpleNamespace(final_class=risk, provisional=provisional), tier)


def test_deadline_days_reads_column_u() -> None:
    v = deck.DeckVendor(vendor_id="V-1", name="A", tier="High", usage="Yes", risk="High",
                        action="Issue a targeted questionnaire within 30 business days covering X.")
    assert deck.deadline_days(v) == 30
    assert deck.deadline_days(v.model_copy(update={"action": "Obtain a written attestation."})) is None


def test_summary_records_screenshots_and_gemini_labels() -> None:
    shot = samples.item(evidence_id="V-901-E-0001", role="Primary", screenshot_path="evidence/screenshots/ab.png",
                        screenshot_sha256="a" * 64, llm_labels={"ai_type": "genai_llm"}, llm_model="gemini-test")
    row = deck.summarize_vendor(samples.findings(evidence=[shot], verdict=samples.verdict([shot.item_key])),
                                name="Acme")
    assert row.shot is not None and row.shot.evidence_id == "V-901-E-0001" and row.shot.screenshot
    assert row.screenshots == 1 and row.llm_labelled == 1 and row.gemini
    plain = deck.summarize_vendor(samples.findings(), name="Acme")
    assert plain.shot is None and plain.screenshots == 0 and not plain.gemini


def test_fit_size_shrinks_only_when_the_text_would_overflow() -> None:
    short = [["One short line."]]
    assert deck.fit_size(short, 2.6, 3.35) == 11.5
    long = [["word " * 60] * 5]
    assert deck.fit_size(long, 2.6, 3.35) == 10
    assert deck.wrapped_lines("aaa bbb ccc", 7) == 2 and deck.wrapped_lines("", 7) == 1


# --------------------------------------------------------------------------- CLI


def test_cli_builds_without_findings(tmp_path, capsys) -> None:
    out = tmp_path / "cli.pptx"
    code = deck.main(["--out", str(out), "--as-of", AS_OF, "--team", TEAM, "--no-appendix"])
    assert code == 0 and out.is_file()
    assert "13 slides" in capsys.readouterr().out


def test_cli_reports_a_missing_findings_file(tmp_path, capsys) -> None:
    code = deck.main(["--findings", str(tmp_path / "missing.json"), "--out", str(tmp_path / "x.pptx")])
    assert code == 2 and "not found" in capsys.readouterr().err
    assert not (tmp_path / "x.pptx").exists()
