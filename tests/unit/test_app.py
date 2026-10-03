"""Streamlit demo app (app/): offline AppTest runs of the real pages, plus unit tests of the helpers in app/ui.py.

Hermetic: no network, no Gemini, no .env (FOOTPRINT_APP_NO_DOTENV), and every review-store path points into
tmp_path (the demo sandbox via FOOTPRINT_APP_SANDBOX, the project store via FOOTPRINT_OVERRIDES and
review.REVIEWS_PATH), so review/ is never touched. The P3 entry points of footprint.pipeline are either removed
(the P1 fallback) or replaced by fakes built from tests/fixtures/p3_samples.py, so these tests hold before and after
the integrator lands run_assessment.
"""

from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import os
import subprocess
import sys
import tomllib
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip("streamlit")
pytest.importorskip("pandas")

import openpyxl  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

from footprint import pipeline, review  # noqa: E402
from footprint.capture.store import EvidenceStore  # noqa: E402
from footprint.models import (  # noqa: E402
    AssessmentResult,
    CoverageEntry,
    CoverageStatus,
    STRENGTH_ORDER,
    FetchOutcome,
    SourceFamily,
    StudentCells,
    ValidationIssue,
)
from footprint.review import OverrideStore  # noqa: E402
from footprint.workbook import WriteReport, read_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
APP = ROOT / "app" / "streamlit_app.py"
INPUT = ROOT / "data" / "input" / "Meridian_Vendor_Input.xlsx"
TIERS = {"V-001": "High", "V-002": "Critical", "V-003": "High", "V-004": "High", "V-005": "Critical",
         "V-006": "Critical"}
TIMEOUT = 120  # the first AppTest run imports pandas, openpyxl and the engine
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _load(name: str, path: Path) -> Any:
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules.setdefault(name, module)
    spec.loader.exec_module(module)
    return module


ui = _load("footprint_app_ui", ROOT / "app" / "ui.py")
S = _load("p3_samples", ROOT / "tests" / "fixtures" / "p3_samples.py")


# --------------------------------------------------------------------------- fixtures and helpers


@pytest.fixture
def hermetic(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv(ui.NO_DOTENV_ENV, "1")
    monkeypatch.setenv(ui.SANDBOX_ENV, str(tmp_path / "sandbox"))
    monkeypatch.setenv(review.OVERRIDES_ENV, str(tmp_path / "project" / "overrides.jsonl"))
    monkeypatch.setattr(review, "REVIEWS_PATH", tmp_path / "project" / "reviews.jsonl")
    monkeypatch.delenv(ui.GEMINI_ENV, raising=False)
    monkeypatch.delenv(ui.TEAM_ENV, raising=False)
    return tmp_path


@pytest.fixture
def p1_only(monkeypatch: pytest.MonkeyPatch) -> None:
    """The pipeline without its P3/P4 entry points (the build before the integrator)."""
    for name in ("run_assessment", "rescore_vendor", "export_assessment", "InputError"):
        monkeypatch.delattr(pipeline, name, raising=False)


def _findings() -> Any:
    f = S.findings()
    item = f.evidence[0].model_copy(update={"evidence_id": "V-901-E-0001", "role": "Primary"})
    return f.model_copy(update={"evidence": [item]})


class FakePipeline:
    """Stand-ins for run_assessment, rescore_vendor and export_assessment (docs/contracts_p3.md section 15)."""

    def __init__(self) -> None:
        self.calls: dict[str, Any] = {}

    def run_assessment(self, xlsx: bytes, mode: str = "replay", **kw: Any) -> AssessmentResult:
        self.calls["run"] = {"mode": mode, **kw}
        return S.assessment(input_sha256=hashlib.sha256(xlsx).hexdigest(), vendors=[_findings()])

    def rescore_vendor(self, findings: Any, *, as_of: str, team: str, overrides: Any = None,
                       reviews: Any = None) -> Any:
        self.calls["rescore"] = {"as_of": as_of, "team": team, "overrides": overrides, "reviews": reviews}
        items = []
        for item in findings.evidence:
            rec = reviews.latest("evidence_review", findings.vendor_id, item.item_key) if reviews else None
            if rec is not None:
                item = item.model_copy(update={"review_status": rec.value, "review_reason": rec.reason})
            items.append(item)
        risk = findings.risk
        e = overrides.latest("risk_input_override", findings.vendor_id, "e") if overrides else None
        if e is not None:  # a crude rescore: E3 lifts this sample to Critical
            inputs = risk.inputs.model_copy(update={"e": int(e.value)})
            risk = S.risk_result(inputs, base_class="Critical", pre_cap_class="Critical", final_class="Critical")
        return findings.model_copy(update={"evidence": items, "risk": risk})

    def export_assessment(self, result: Any, xlsx_in: bytes, out_path: Any, team: str, *,
                          overwrite: bool = False) -> WriteReport:
        self.calls["export"] = {"team": team, "run_id": result.run_id, "overwrite": overwrite}
        out_path.write(xlsx_in)
        return WriteReport(written=["L6"], sheets_added=[])


@pytest.fixture
def fake(monkeypatch: pytest.MonkeyPatch) -> FakePipeline:
    fp = FakePipeline()
    for name in ("run_assessment", "rescore_vendor", "export_assessment"):
        monkeypatch.setattr(pipeline, name, getattr(fp, name), raising=False)
    return fp


def start(**state: Any) -> AppTest:
    at = AppTest.from_file(str(APP), default_timeout=TIMEOUT)
    for key, value in state.items():
        at.session_state[key] = value
    return at.run()


def upload(at: AppTest, name: str = INPUT.name, data: bytes | None = None) -> AppTest:
    at.file_uploader(key="fp_upload").upload(name, INPUT.read_bytes() if data is None else data, XLSX)
    return at.run()


def ok(at: AppTest) -> AppTest:
    assert not at.exception, [e.value for e in at.exception]
    return at


def table(at: AppTest, column: str) -> Any:
    """The first dataframe on the page that has ``column``."""
    for df in at.dataframe:
        if column in df.value.columns:
            return df.value
    raise AssertionError(f"no dataframe with a {column!r} column")


def texts(elements: Any) -> list[str]:
    return [e.value for e in elements]


# --------------------------------------------------------------------------- configuration


def test_streamlit_config_keeps_the_app_local_private_and_quiet() -> None:
    project = tomllib.loads((ROOT / ".streamlit" / "config.toml").read_text(encoding="utf-8"))
    script = tomllib.loads((ROOT / "app" / ".streamlit" / "config.toml").read_text(encoding="utf-8"))
    assert script == project  # the script-level copy and the repo-root one never drift apart
    for cfg in (project, script):
        assert cfg["server"]["address"] == "localhost"
        assert cfg["browser"]["gatherUsageStats"] is False
        assert cfg["client"]["showErrorDetails"] == "none"
        assert cfg["server"]["runOnSave"] is False


def test_streamlit_applies_the_app_config_from_any_working_directory(tmp_path: Path) -> None:
    """A clone or an extracted bundle may start Streamlit from another directory, where no .streamlit/ exists:
    the script-level config next to app/streamlit_app.py still binds the server to localhost and turns usage
    statistics off. (Run in a fresh interpreter: Streamlit's config is process-global. Nothing is served.)"""
    probe = (
        "import json, os, sys\n"
        "from pathlib import Path\n"
        "import streamlit.config as config\n"
        "os.chdir(sys.argv[1])\n"
        "os.environ['HOME'] = os.environ['USERPROFILE'] = sys.argv[1]\n"
        "Path.home = classmethod(lambda cls: Path(sys.argv[1]))\n"
        "config._main_script_path = os.path.abspath(sys.argv[2])\n"
        "config.get_config_options(force_reparse=True)\n"
        "print(json.dumps({k: config.get_option(k) for k in ('server.address', 'browser.gatherUsageStats', "
        "'client.showErrorDetails', 'browser.serverAddress')}))\n"
    )
    env = {k: v for k, v in os.environ.items() if not k.startswith("STREAMLIT_")}
    out = subprocess.run([sys.executable, "-c", probe, str(tmp_path), str(APP)], capture_output=True, text=True,
                         timeout=120, env=env, check=True).stdout
    options = json.loads(out.strip().splitlines()[-1])
    assert not (tmp_path / ".streamlit").exists()
    assert options == {"server.address": "localhost", "browser.gatherUsageStats": False,
                       "client.showErrorDetails": "none", "browser.serverAddress": "localhost"}


# --------------------------------------------------------------------------- Assess page (AppTest)


def test_app_starts_and_asks_for_a_workbook(hermetic: Path) -> None:
    at = ok(start())
    assert texts(at.title) == ["Vendor AI footprint"]
    assert any("Upload the vendor inventory" in text for text in texts(at.info))
    assert at.text_input(key=ui.K_ANALYST).value == ""
    assert at.radio(key=ui.K_STORE_KIND).value == "sandbox"
    assert not (hermetic / "project").exists()


def test_upload_of_the_real_workbook_shows_six_vendors_and_their_tiers(hermetic: Path) -> None:
    at = ok(upload(start()))
    assert any("6 vendors" in text for text in texts(at.success))
    tiers = table(at, "Tier by score")
    assert dict(zip(tiers["Vendor ID"], tiers["Tier"])) == TIERS
    assert [tab.label.split(" · ")[0] for tab in at.tabs] == list(TIERS)
    scores = dict(zip(tiers["Vendor ID"], tiers["Score"]))
    assert scores["V-005"] == 79  # BNY: score tier High, raised to Critical by the payment-path floor
    assert not (hermetic / "project").exists()


def test_bundled_workbook_button_loads_the_same_vendors(hermetic: Path) -> None:
    at = ok(start())
    ok(at.button(key="fp_use_bundled").click().run())
    tiers = table(at, "Tier by score")
    assert dict(zip(tiers["Vendor ID"], tiers["Tier"])) == TIERS


def test_a_file_that_is_not_a_workbook_is_reported_not_assessed(hermetic: Path) -> None:
    at = ok(upload(start(), "notes.xlsx", b"not a workbook"))
    assert any("cannot be assessed" in text for text in texts(at.error))
    assert "E01" in set(table(at, "Code")["Code"])
    assert not at.tabs


def test_hc1_override_is_recorded_in_the_sandbox_only(hermetic: Path) -> None:
    at = ok(upload(start(**{ui.K_ANALYST: "RK"})))
    at.selectbox(key="fp_hc1_tier_V-003").set_value("Critical")
    at.text_area(key="fp_hc1_reason_V-003").input("Dependency confirmed as High by the business owner.")
    ok(at.button(key="fp_hc1_submit_V-003").click().run())
    records = OverrideStore(hermetic / "sandbox" / "overrides.jsonl").records("tier_override")
    assert [(r.vendor_id, r.value, r.analyst) for r in records] == [("V-003", "Critical", "RK")]
    tiers = table(at, "Tier by score")
    row = tiers[tiers["Vendor ID"] == "V-003"].iloc[0]
    assert (row["Tier"], row["Computed tier"], row["Override"]) == ("Critical", "High", "Critical (RK)")
    assert any("tier override to Critical recorded" in text for text in texts(at.success))
    assert not (hermetic / "project").exists()


def test_hc1_override_needs_an_analyst_and_a_reason(hermetic: Path) -> None:
    at = ok(upload(start()))
    at.selectbox(key="fp_hc1_tier_V-001").set_value("Critical")
    at.text_area(key="fp_hc1_reason_V-001").input("Confirmed by the owner, long enough.")
    ok(at.button(key="fp_hc1_submit_V-001").click().run())
    assert any("analyst" in text for text in texts(at.error))
    at.text_input(key=ui.K_ANALYST).input("RK")
    at.text_area(key="fp_hc1_reason_V-001").input("short")
    ok(at.button(key="fp_hc1_submit_V-001").click().run())
    assert any("at least 10 characters" in text for text in texts(at.error))
    assert not (hermetic / "sandbox" / "overrides.jsonl").exists()


def test_run_without_the_full_pipeline_falls_back_to_criticality_and_depth(hermetic: Path, p1_only: None) -> None:
    at = ok(upload(start()))
    assert any("not available in this build" in text for text in texts(at.info))
    ok(at.button(key="fp_run").click().run())
    summary = table(at, "Depth")
    assert list(summary["Vendor ID"]) == list(TIERS)
    assert list(summary["Tier"]) == list(TIERS.values())
    assert any("criticality and depth only" in text for text in texts(at.sidebar.caption))
    ok(at.switch_page("pages/evidence.py").run())
    assert any("criticality and depth only" in text for text in texts(at.info))


def test_pipeline_errors_are_scrubbed_and_the_app_keeps_working(hermetic: Path,
                                                                  monkeypatch: pytest.MonkeyPatch) -> None:
    secret = "AIza" + "S" * 35

    def broken(*args: Any, **kw: Any) -> None:
        raise RuntimeError(f"quota exceeded for key {secret}")

    monkeypatch.setattr(pipeline, "run_assessment", broken, raising=False)
    at = ok(upload(start()))
    ok(at.button(key="fp_run").click().run())
    errors = " ".join(texts(at.error)).replace("\\", "")  # the message is Markdown-escaped
    assert "The assessment stopped: RuntimeError" in errors and "[redacted]" in errors
    assert secret not in errors
    assert list(table(at, "Depth")["Vendor ID"]) == list(TIERS)


def test_an_input_error_lists_the_workbook_issues(hermetic: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    class InputError(ValueError):
        def __init__(self, issues: list[ValidationIssue]) -> None:
            super().__init__("invalid workbook")
            self.issues = issues

    def refuse(*args: Any, **kw: Any) -> None:
        raise InputError([ValidationIssue(code="E04", message='Duplicate vendor ID "V-002".', cell="B8")])

    monkeypatch.setattr(pipeline, "InputError", InputError, raising=False)
    monkeypatch.setattr(pipeline, "run_assessment", refuse, raising=False)
    at = ok(upload(start()))
    ok(at.button(key="fp_run").click().run())
    assert any("failed validation" in text for text in texts(at.error))
    assert list(table(at, "Code")["Code"]) == ["E04"]


def test_pages_without_a_run_point_back_to_assess(hermetic: Path, fake: FakePipeline) -> None:
    at = ok(start())
    for page in ("pages/evidence.py", "pages/findings.py", "pages/export.py"):
        ok(at.switch_page(page).run())
        assert any("Load the vendor workbook" in text for text in texts(at.info)), page
    ok(at.switch_page("pages/assess.py").run())
    upload(at)
    ok(at.switch_page("pages/evidence.py").run())
    assert any("No assessment has been run" in text for text in texts(at.info))


# --------------------------------------------------------------------------- full flow (fake pipeline)


def _run_full(hermetic: Path, fake: FakePipeline) -> AppTest:
    at = ok(upload(start(**{ui.K_ANALYST: "RK", ui.K_TEAM: "Team Test"})))
    ok(at.button(key="fp_run").click().run())
    return at


def test_full_run_is_cached_per_workbook_and_mode(hermetic: Path, fake: FakePipeline) -> None:
    at = _run_full(hermetic, fake)
    call = fake.calls["run"]
    assert call["mode"] == "replay" and call["vendors"] is None and call["team"] == "Team Test"
    assert call["as_of"] is None  # replay takes the date of the frozen collection runs
    assert call["overrides"].path == hermetic / "sandbox" / "overrides.jsonl"
    assert call["reviews"].path == hermetic / "sandbox" / "reviews.jsonl"
    assert Path(call["seeds_dir"]) == ROOT / "seeds" and Path(call["runs_dir"]) == ROOT / "runs"
    summary = table(at, "AI usage (O)")
    assert list(summary["Vendor ID"]) == ["V-901"] and list(summary["AI risk (S)"]) == ["High"]
    runs = at.session_state[ui.K_RUNS]
    assert list(runs) == [(hashlib.sha256(INPUT.read_bytes()).hexdigest(), "replay")]
    assert any("Run A-20261002-0000abcd" in text for text in texts(at.sidebar.caption))
    at.radio(key="fp_mode_widget").set_value("live_rules")
    ok(at.run())
    assert any("No Live rules run" in text for text in texts(at.caption))
    assert any("polite GET requests" in text for text in texts(at.warning))


def test_an_invalid_as_of_date_is_refused_before_running(hermetic: Path, fake: FakePipeline) -> None:
    at = ok(upload(start()))
    at.text_input(key="fp_as_of").input("02-10-2026")
    ok(at.button(key="fp_run").click().run())
    assert any("YYYY-MM-DD" in text.replace("\\", "") for text in texts(at.error))  # Markdown-escaped
    assert "run" not in fake.calls and not at.session_state[ui.K_RUNS]
    at.text_input(key="fp_as_of").input("2026-10-02")
    ok(at.button(key="fp_run").click().run())
    assert fake.calls["run"]["as_of"] == "2026-10-02"
    assert list(table(at, "AI usage (O)")["Vendor ID"]) == ["V-901"]


def test_evidence_review_is_recorded_and_the_vendor_rescored(hermetic: Path, fake: FakePipeline) -> None:
    at = _run_full(hermetic, fake)
    ok(at.switch_page("pages/evidence.py").run())
    assert at.text_input(key=ui.K_ANALYST).value == "RK"  # sidebar settings survive the page switch
    rows = table(at, "Evidence ID")
    assert list(rows["Evidence ID"]) == ["V-901-E-0001"] and list(rows["Role"]) == ["Primary"]
    key = _findings().evidence[0].item_key[:16]
    ok(at.radio(key=f"fp_dec_{key}").set_value("rejected").run())
    at.selectbox(key=f"fp_code_{key}_rejected").set_value("REJ-NOT-AI")
    at.text_input(key=f"fp_note_{key}").input("scheduling relabelled as AI")
    ok(at.button(key=f"fp_record_{key}").click().run())
    records = OverrideStore(hermetic / "sandbox" / "reviews.jsonl").records("evidence_review")
    assert [(r.vendor_id, r.value, r.analyst) for r in records] == [("V-901", "rejected", "RK")]
    assert records[0].reason.startswith("REJ-NOT-AI: definition test fails")
    assert records[0].reason.endswith("- scheduling relabelled as AI")
    assert fake.calls["rescore"]["as_of"] == "2026-10-02" and fake.calls["rescore"]["team"] == "Team Test"
    assert any("V-901-E-0001 rejected (REJ-NOT-AI). V-901: Yes (Confirmed); AI risk stays High; score 12 of 18."
               in text for text in texts(at.success))
    assert list(table(at, "Evidence ID")["Review"]) == ["Rejected"]
    assert not (hermetic / "project").exists()


def test_findings_show_the_why_trace_and_override_e_with_a_reason(hermetic: Path, fake: FakePipeline) -> None:
    at = _run_full(hermetic, fake)
    ok(at.switch_page("pages/findings.py").run())
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["AI usage (column O)"] == "Yes" and metrics["AI security risk (live)"] == "High"
    trace = table(at, "Stage")
    assert set(trace["Stage"]) == {"AI usage", "AI risk"}
    assert any(step.startswith("Decision: rule b) gives Confirmed") for step in trace["Step"])
    cells = table(at, "Column")
    assert list(cells["Column"]) == list("LMNOPQRSTUV")
    vid = "V-901"
    at.selectbox(key=f"fp_ov_e_value_{vid}").set_value(3)
    at.text_input(key=f"fp_ov_e_reason_{vid}").input("Model confirmed to process all payment data.")
    ok(next(b for b in at.button if b.label == "Record E").click().run())
    records = OverrideStore(hermetic / "sandbox" / "overrides.jsonl").records("risk_input_override")
    assert [(r.key, r.value, r.analyst) for r in records] == [("e", "3", "RK")]
    assert any("Live class: AI risk High → Critical; score 12 → 14 of 18." in text for text in texts(at.success))
    live = next(m for m in at.metric if m.label == "AI security risk (live)")
    assert live.value == "Critical" and live.delta == "was High"


def test_export_builds_a_checked_workbook_for_download(hermetic: Path, fake: FakePipeline) -> None:
    at = _run_full(hermetic, fake)
    ok(at.switch_page("pages/export.py").run())
    ok(at.button(key="fp_export_build").click().run())
    assert fake.calls["export"] == {"team": "Team Test", "run_id": "A-20261002-0000abcd", "overwrite": False}
    assert any("Integrity check passed" in text for text in texts(at.success))
    assert [b.label for b in at.get("download_button")] == ["Download the assessed workbook"]


def test_export_of_a_full_run_without_export_assessment_falls_back_to_l_to_n(hermetic: Path, fake: FakePipeline,
                                                                             monkeypatch: pytest.MonkeyPatch) -> None:
    at = _run_full(hermetic, fake)
    monkeypatch.delattr(pipeline, "export_assessment")
    ok(at.switch_page("pages/export.py").run())
    assert ui.EXPORT_P1_NOTE in texts(at.info)
    ok(at.button(key="fp_export_build").click().run())
    assert texts(at.info).count(ui.EXPORT_P1_NOTE) == 1
    assert any("Integrity check passed" in text for text in texts(at.success))
    assert {m.label: m.value for m in at.metric}["Workbook"] == "L-N"


# --------------------------------------------------------------------------- ui helpers: inputs and stores


def test_inputs_are_identified_by_the_sha256_of_their_bytes() -> None:
    inp = ui.make_input("x.xlsx", b"abc")
    assert inp.sha256 == hashlib.sha256(b"abc").hexdigest() and inp.short_sha == inp.sha256[:12]
    bundled = ui.bundled_input()
    assert bundled.source == "bundled" and bundled.sha256 == hashlib.sha256(INPUT.read_bytes()).hexdigest()
    assert ui.run_key(inp.sha256, "replay") == (inp.sha256, "replay")


def test_sandbox_copies_the_project_store_once_and_resets(hermetic: Path) -> None:
    project = OverrideStore(hermetic / "project" / "overrides.jsonl")
    ui.record_tier_override(project, "V-001", "Critical", "Confirmed with the business owner.", "RK")
    stores = ui.review_stores("sandbox")
    assert stores.directory == hermetic / "sandbox" and stores.label == "Sandbox (demo copy)"
    assert stores.overrides.tier_override("V-001").tier.value == "Critical"
    ui.record_tier_override(stores.overrides, "V-002", "High", "Demo override, sandbox only.", "RK")
    stores.reviews.add(ui.evidence_review_record("V-002", "k" * 64, "accepted", "ACC-VERIFIED", "", "RK"))
    assert project.tier_override("V-002") is None  # the project store never changes
    assert ui.review_stores("sandbox").overrides.tier_override("V-002") is not None  # kept until a reset
    ui.ensure_sandbox(reset=True)
    assert ui.review_stores("sandbox").overrides.tier_override("V-002") is None
    assert not (hermetic / "sandbox" / "reviews.jsonl").exists()
    assert ui.review_stores("project").overrides.path == project.path
    with pytest.raises(ValueError):
        ui.review_stores("elsewhere")


def test_review_records_check_codes_reasons_and_the_analyst() -> None:
    rec = ui.evidence_review_record("V-005", "k" * 64, "accepted", "ACC-VERIFIED", "", "RK",
                                    labels={"temporal": "in_production"}, date="2026-10-08")
    assert (rec.kind, rec.value, rec.date) == ("evidence_review", "accepted", "2026-10-08")
    assert rec.reason == "ACC-VERIFIED: " + ui.ACCEPT_CODES["ACC-VERIFIED"]
    assert rec.extra == {"labels": {"temporal": "in_production"}}
    other = ui.evidence_review_record("V-005", "k" * 64, "rejected", "OTHER", "Supplier brochure, not the vendor.",
                                      "RK")
    assert other.reason == "OTHER: Supplier brochure, not the vendor." and other.extra == {}
    for decision, code, note, analyst in [("accepted", "REJ-ENTITY", "", "RK"), ("rejected", "ACC-SCOPE", "", "RK"),
                                          ("rejected", "OTHER", "short", "RK"), ("accepted", "ACC-VERIFIED", "", " "),
                                          ("maybe", "ACC-VERIFIED", "", "RK"), ("accepted", "NOPE", "", "RK")]:
        with pytest.raises(ValueError):
            ui.evidence_review_record("V-005", "k" * 64, decision, code, note, analyst)


def test_risk_override_records_follow_the_contract() -> None:
    rec = ui.risk_override_record("V-005", "e", "3", "Customer data confirmed in scope.", "RK",
                                  evidence=["b" * 64, "a" * 64, "a" * 64])
    assert (rec.kind, rec.key, rec.value) == ("risk_input_override", "e", "3")
    assert rec.extra == {"evidence": ["a" * 64, "b" * 64]}
    assert ui.risk_override_record("V-005", "t4", "closed", "Human review described.", "RK").extra == {}
    for key, value in [("e", "4"), ("k", "high"), ("t1", "open"), ("x", "1")]:
        with pytest.raises(ValueError):
            ui.risk_override_record("V-005", key, value, "A reason that is long enough.", "RK")
    with pytest.raises(ValueError, match="at least 10"):
        ui.risk_override_record("V-005", "e", "1", "too short", "RK")
    with pytest.raises(ValueError, match="Unknown tier"):
        ui.record_tier_override(OverrideStore("unused.jsonl"), "V-001", "Huge", "Reason long enough.", "RK")


# --------------------------------------------------------------------------- ui helpers: runs


def _workbook() -> tuple[Any, Any]:
    inp = ui.bundled_input()
    return inp, read_workbook(inp.data)


def test_execute_run_without_the_pipeline_returns_criticality_and_depth(hermetic: Path, p1_only: None) -> None:
    inp, data = _workbook()
    record = ui.execute_run(inp, data, "replay", ui.review_stores("sandbox"))
    assert (record.kind, record.full, record.notes) == ("p1", False, [ui.P1_NOTE])
    assert [a.criticality.tier.value for a in record.p1] == list(TIERS.values())
    assert record.example.profile.vendor_id == "V-000"
    assert record.tier_overrides == {vid: "" for vid in TIERS}
    assert ui.run_summary_rows(record)[0]["Depth"] == record.p1[0].depth.label


def test_execute_run_passes_scope_and_stores_to_the_pipeline(hermetic: Path, fake: FakePipeline) -> None:
    inp, data = _workbook()
    stores = ui.review_stores("sandbox")
    record = ui.execute_run(inp, data, "live_rules", stores, team="Team Test", vendors=["V-005"], store=object())
    assert record.full and record.key == (inp.sha256, "live_rules") and record.vendors == ("V-005",)
    assert [a.profile.vendor_id for a in record.p1] == ["V-005"]
    assert fake.calls["run"]["vendors"] == ["V-005"] and fake.calls["run"]["mode"] == "live_rules"
    assert fake.calls["run"]["overrides"] is stores.overrides and fake.calls["run"]["reviews"] is stores.reviews
    with pytest.raises(ValueError, match="V-999"):
        ui.execute_run(inp, data, "replay", stores, vendors=["V-999"])
    with pytest.raises(ValueError):
        ui.execute_run(inp, data, "turbo", stores)


def test_parse_as_of_accepts_only_strict_iso_dates(hermetic: Path, fake: FakePipeline) -> None:
    assert ui.parse_as_of("") is None and ui.parse_as_of("  ") is None
    assert ui.parse_as_of(" 2026-10-02 ") == "2026-10-02"
    for bad in ("02-10-2026", "2026-10-2", "2026-13-01", "2026-02-30", "20261002", "today"):
        with pytest.raises(ValueError, match="YYYY-MM-DD"):
            ui.parse_as_of(bad)
    inp, data = _workbook()
    record = ui.execute_run(inp, data, "replay", ui.review_stores("sandbox"), as_of="2026-10-01", store=object())
    assert record.full and fake.calls["run"]["as_of"] == "2026-10-01"
    with pytest.raises(ValueError, match="YYYY-MM-DD"):
        ui.execute_run(inp, data, "replay", ui.review_stores("sandbox"), as_of="1 Oct", store=object())


def test_a_run_goes_stale_when_a_tier_override_or_the_store_changes(hermetic: Path, p1_only: None) -> None:
    inp, data = _workbook()
    stores = ui.review_stores("sandbox")
    record = ui.execute_run(inp, data, "replay", stores)
    assert not ui.is_stale(record, stores, data)
    ui.record_tier_override(stores.overrides, "V-004", "Critical", "Holdings data confirmed in scope.", "RK")
    assert ui.is_stale(record, stores, data)
    fresh = ui.execute_run(inp, data, "replay", stores)
    assert not ui.is_stale(fresh, stores, data) and fresh.tier_overrides["V-004"] == "Critical"
    assert ui.is_stale(fresh, ui.review_stores("project"), data)


def test_rescore_swaps_the_vendor_findings_in_the_cached_run(hermetic: Path, fake: FakePipeline) -> None:
    inp, data = _workbook()
    stores = ui.review_stores("sandbox")
    record = ui.execute_run(inp, data, "replay", stores, team="Team Test", store=object())
    stores.overrides.add(ui.risk_override_record("V-901", "e", "3", "Payment data confirmed in scope.", "RK"))
    before = ui.export_key(inp, record, "Team Test", False)
    updated = ui.rescore(record, "V-901", stores)
    assert updated.risk.final_class == "Critical" and record.findings("V-901").risk.final_class == "Critical"
    assert fake.calls["rescore"]["team"] == "Team Test"  # falls back to the team of the run
    assert record.version == 1 and ui.export_key(inp, record, "Team Test", False) != before
    with pytest.raises(ValueError):
        ui.rescore(record, "V-999", stores)


# --------------------------------------------------------------------------- ui helpers: evidence


def test_excerpt_context_and_html_escape_the_source_text() -> None:
    text = "A" * 400 + "We use <b>ML</b> for $5 payments.\n\nNext" + "B" * 400
    start = text.index("We use")
    end = start + len("We use <b>ML</b> for $5 payments.")
    ctx = ui.excerpt_context(text, start, end, text[start:end])
    assert ctx.matches and ctx.clipped_start and ctx.clipped_end
    assert len(ctx.before) == ui.CONTEXT_CHARS and ctx.after.startswith("\n\nNext")
    html = ui.context_html(ctx)
    assert "<b>ML</b>" not in html and "&lt;b&gt;ML&lt;/b&gt;" in html
    assert "$" not in html and "&#36;5" in html and "\n" not in html
    assert not ui.excerpt_context(text, start, end, "something else").matches
    short = ui.excerpt_context("abc", 0, 3)
    assert (short.before, short.excerpt, short.after) == ("", "abc", "")
    assert not short.clipped_start and not short.clipped_end


def test_load_context_reads_the_evidence_store(tmp_path: Path) -> None:
    store = EvidenceStore(tmp_path / "evidence")
    doc_id, _ = store.put_text(S.DOC_TEXT)
    item = S.item(doc_id=doc_id)
    ctx = ui.load_context(item, store)
    assert ctx is not None and ctx.matches and ctx.excerpt == S.EXCERPT
    assert ui.load_context(S.item(), store) is None  # text not stored


def _three_items() -> Any:
    base = S.findings()
    strong = base.evidence[0].model_copy(update={"evidence_id": "V-901-E-0001", "role": "Primary"})
    job = S.item(excerpt="Alerts are reviewed by an analyst before any payment is held.", family=SourceFamily.JOB,
                 tags=S.tags(strength="Moderate"), method="rule+llm_agree", evidence_id="V-901-E-0002",
                 role="Supporting", review_status="accepted")
    weak = S.item(excerpt="Instant payments", family=SourceFamily.IND, tags=S.tags(strength="Weak"),
                  method="llm_proposed_accepted", evidence_id="V-901-E-0003", title="Trade press",
                  url="https://news.example/x")
    return base.model_copy(update={"evidence": [strong, job, weak]})


def test_filter_items_combines_every_filter() -> None:
    f = _three_items()

    def ids(**kw: Any) -> list[str]:
        return [item.evidence_id for _, item in ui.filter_items([f], ui.EvidenceFilter(**kw))]

    assert ids() == ["V-901-E-0001", "V-901-E-0002", "V-901-E-0003"]
    assert ids(vendors=("V-902",)) == []
    assert ids(families=("JOB", "IND")) == ["V-901-E-0002", "V-901-E-0003"]
    assert ids(strengths=("Strong",)) == ["V-901-E-0001"]
    assert ids(methods=("llm_proposed_accepted",)) == ["V-901-E-0003"]
    assert ids(cited_only=True) == ["V-901-E-0001", "V-901-E-0002"]
    assert ids(statuses=("accepted",)) == ["V-901-E-0002"]
    assert ids(text="news.example") == ["V-901-E-0003"]
    rows = ui.evidence_rows(ui.filter_items([f], ui.EvidenceFilter()))
    assert [r["Method"] for r in rows] == ["Rules", "Rules + Gemini agree", "Gemini proposal"]
    assert [r["Counts"] for r in rows] == ["yes", "yes", "no: Gemini proposal awaiting review"]
    assert rows[0]["Tags"] == "U1 · SR:B · SP:S2 · RL:R3 · RC:T3 · IC:2 · locus=service_feature"
    assert ui.present_values([f], "strength", STRENGTH_ORDER) == ["Strong", "Moderate", "Weak"]
    assert ui.cited_unreviewed(f) == [f.evidence[0]]


def test_counts_label_explains_why_an_item_does_not_count() -> None:
    assert ui.counts_label(S.item()) == "yes"
    assert ui.counts_label(S.item(tags=S.tags(ai_type="not_ai", u_class="U7", strength="Marketing only"))) \
        == "no: definition-test trap"
    assert ui.counts_label(S.item(review_status="rejected")) == "no: rejected"
    assert ui.counts_label(S.item(method="llm_proposed_accepted")) == "no: Gemini proposal awaiting review"
    assert ui.counts_label(S.item(method="llm_proposed_accepted", review_status="accepted")) == "yes"


def test_label_rows_flag_the_v8_disagreements() -> None:
    item = S.item(rule_labels={"temporal": "in_production", "sp": "S2", "locus": "service_feature"},
                  llm_labels={"temporal": "pilot_or_beta", "sp": "S2", "subject": "vendor_product"})
    rows = {r["Label"]: r for r in ui.label_rows(item)}
    assert rows["temporal"]["Status"].startswith("disagree") and rows["temporal"]["Gemini"] == "pilot_or_beta"
    assert rows["sp"]["Status"] == "agree"
    assert rows["locus"]["Status"] == "rules only" and rows["subject"]["Status"] == "Gemini only"
    assert list(rows) == ["temporal", "sp", "locus", "subject"]


def test_screenshot_paths_stay_inside_the_repo(tmp_path: Path) -> None:
    shot = tmp_path / "evidence" / "shots" / "a.png"
    shot.parent.mkdir(parents=True)
    shot.write_bytes(b"png")
    assert ui.screenshot_file(S.item(screenshot_path="evidence/shots/a.png"), tmp_path) == shot.resolve()
    assert ui.screenshot_file(S.item(screenshot_path="../outside.png"), tmp_path) is None
    assert ui.screenshot_file(S.item(screenshot_path="evidence/shots/missing.png"), tmp_path) is None
    assert ui.screenshot_file(S.item(), tmp_path) is None


# --------------------------------------------------------------------------- ui helpers: findings and cells


def test_why_trace_and_risk_tables_cite_evidence_ids() -> None:
    f = _three_items()
    key = f.evidence[0].item_key
    f = f.model_copy(update={"verdict": S.verdict([key], trace=["a) no conflict", "b) Q at R3 with a K"])})
    trace = ui.why_trace(f)
    assert [r["Step"] for r in trace[:2]] == ["a) no conflict", "b) Q at R3 with a K"]
    assert trace[2]["Step"].startswith("Decision: rule b) gives Confirmed, written 'Yes' in column O")
    assert trace[-1] == {"Stage": "AI risk", "Step": "Result: High."}
    assert ui.score_sentence(f.risk) == ("exposure 2/3, decision impact 2/3, tier 3, transparency gap 1 → 12 of 18 = "
                                         "High (exposure and decision impact count double)")
    inputs = {r["Input"][:2]: r for r in ui.risk_input_rows(f)}
    assert inputs["E "]["Value"] == "2/3" and inputs["E "]["Evidence"] == "V-901-E-0001"
    assert inputs["TG"]["Basis"] == "3 of 6 checks missing"
    gaps = {r["Check"]: r["Status"] for r in ui.gap_rows(f)}
    assert gaps == {"t1": "closed", "t2": "missing", "t3": "missing", "t4": "closed", "t5": "closed",
                    "t6": "missing"}
    records = ui.changed_gap_records(f, {**ui.gap_values(f), "t2": "closed", "t5": "missing"},
                                     "Sub-processor list reviewed on the trust page.", "RK")
    assert [(r.key, r.value) for r in records] == [("t2", "closed"), ("t5", "missing")]


def test_class_change_reports_the_class_the_score_and_the_cap() -> None:
    high = S.risk_result()
    assert ui.class_change(high, high) == "AI risk stays High; score 12 of 18."
    critical = S.risk_result(S.risk_inputs(e=3), base_class="Critical", pre_cap_class="Critical",
                             final_class="Critical")
    assert ui.class_change(high, critical) == "AI risk High → Critical; score 12 → 14 of 18."
    capped = S.risk_result(S.risk_inputs(e=3), base_class="Critical", pre_cap_class="Critical", cap="Medium",
                           final_class="Medium", provisional=True)
    assert ui.class_change(high, capped) == ("AI risk High → Medium (provisional) (capped at Medium by the verdict; "
                                             "Critical before the cap); score 12 → 14 of 18.")
    gated = S.risk_result(S.risk_inputs(e=3), base_class="Critical", gate_met=False, pre_cap_class="Medium",
                          cap="Medium", final_class="Medium", provisional=True)
    assert ui.class_change(capped, gated) == ("AI risk stays Medium (provisional) (materiality gate not met: Critical "
                                              "lowered to Medium); score 14 of 18.")


def test_cells_rows_follow_the_column_map_and_budgets() -> None:
    rows = ui.cells_rows(StudentCells(criticality_tier="High", ai_usage_detected="Yes"), {"criticality_tier": "L"})
    assert [r["Column"] for r in rows] == list("LMNOPQRSTUV")
    assert rows[0] == {"Column": "L", "Field": "Criticality Tier", "Text": "High", "Length": "4/1050"}
    assert rows[4]["Length"] == "0/1390" and rows[10]["Field"] == "Assessed By / Date"


# --------------------------------------------------------------------------- ui helpers: export


def test_p1_export_writes_columns_l_to_n_and_passes_the_fidelity_check(hermetic: Path, p1_only: None) -> None:
    inp, data = _workbook()
    outcome = ui.export_run(None, inp, data, ui.review_stores("sandbox"), team="Team Test")
    assert outcome.kind == "p1" and outcome.ok and outcome.fidelity == [] and outcome.example_unchanged is True
    assert outcome.file_name.startswith("Meridian_Vendor_Input_assessed_") and outcome.file_name.endswith(".xlsx")
    assert sorted(outcome.written) == sorted(f"{c}{r}" for c in "LMN" for r in range(6, 12))
    assert outcome.sheets == ["Vendor Inventory", "Field Guide", "Points to consider", "Criticality Workings",
                              "Method & Legend"]
    wb = openpyxl.load_workbook(io.BytesIO(outcome.data))
    ws = wb["Vendor Inventory"]
    assert [ws[f"L{r}"].value for r in range(6, 12)] == list(TIERS.values())
    assert wb.properties.lastModifiedBy == "Team Test"


def test_full_export_uses_export_assessment(hermetic: Path, fake: FakePipeline) -> None:
    inp, data = _workbook()
    stores = ui.review_stores("sandbox")
    record = ui.execute_run(inp, data, "replay", stores, team="Team Run", store=object())
    outcome = ui.export_run(record, inp, data, stores, team="")
    assert outcome.kind == "full" and outcome.data == inp.data and outcome.ok
    assert outcome.file_name.endswith("_assessed_2026-10-02.xlsx")
    assert fake.calls["export"]["team"] == "Team Run"


def test_full_run_export_falls_back_to_l_to_n_only_while_export_assessment_is_missing(
        hermetic: Path, fake: FakePipeline, monkeypatch: pytest.MonkeyPatch) -> None:
    inp, data = _workbook()
    stores = ui.review_stores("sandbox")
    record = ui.execute_run(inp, data, "replay", stores, team="Team Run", store=object())
    monkeypatch.delattr(pipeline, "export_assessment")
    outcome = ui.export_run(record, inp, data, stores)
    assert (outcome.kind, outcome.note, outcome.ok) == ("p1", ui.EXPORT_P1_NOTE, True)
    assert sorted(outcome.written) == sorted(f"{c}{r}" for c in "LMN" for r in range(6, 12))
    assert openpyxl.load_workbook(io.BytesIO(outcome.data)).properties.lastModifiedBy == "Team Run"

    def not_yet(*args: Any, **kw: Any) -> None:
        raise NotImplementedError

    monkeypatch.setattr(pipeline, "export_assessment", not_yet, raising=False)
    assert ui.export_run(record, inp, data, stores).note == ui.EXPORT_P1_NOTE

    def broken(*args: Any, **kw: Any) -> None:
        raise ValueError("input workbook does not match the run")

    monkeypatch.setattr(pipeline, "export_assessment", broken, raising=False)
    with pytest.raises(ValueError, match="does not match"):  # a real export failure is never papered over
        ui.export_run(record, inp, data, stores)


# --------------------------------------------------------------------------- ui helpers: coverage and DNS


def _coverage_run(runs: Path, name: str, created_at: str, status: CoverageStatus) -> None:
    d = runs / name
    d.mkdir(parents=True)
    entry = CoverageEntry(vendor_id="V-777", family=SourceFamily.DNS, mandatory=True, status=status)
    (d / "coverage.jsonl").write_text(entry.model_dump_json() + "\n", encoding="utf-8")
    (d / "manifest.json").write_text(json.dumps({"created_at": created_at}), encoding="utf-8")


def test_latest_collection_coverage_uses_the_newest_run(tmp_path: Path) -> None:
    runs = tmp_path / "runs"
    _coverage_run(runs, "V-777-20261001-aaaaaaaa", "2026-10-01T10:00:00Z", CoverageStatus.ERROR)
    _coverage_run(runs, "V-777-20261002-bbbbbbbb", "2026-10-02T10:00:00Z", CoverageStatus.DONE)
    (runs / "A-20261002-cccccccc").mkdir()  # an assessment run is not a collection run
    entries = ui.latest_collection_coverage("V-777", runs)
    assert [e.status for e in entries] == [CoverageStatus.DONE]
    assert ui.coverage_rows(entries)[0]["Family"] == "DNS"
    assert ui.latest_collection_coverage("V-778", runs) == []
    assert ui.latest_collection_coverage("V-777", tmp_path / "missing") == []


def _doh(record: str) -> bytes:
    return json.dumps({"Status": 0, "Answer": [{"name": "acme.example.", "type": 16, "data": f'"{record}"'}]}).encode()


class FakeLive:
    """A live fetcher stand-in: answers the TXT queries from ``answers`` and fails everything else."""

    def __init__(self, store: EvidenceStore, answers: dict[str, bytes]) -> None:
        self.store, self.answers, self.urls = store, answers, []

    def get(self, url: str, *, vendor_id: str, family: SourceFamily, collector: str, accept: str | None = None,
            seeded: bool = False) -> FetchOutcome:
        self.urls.append(url)
        if url not in self.answers:
            return FetchOutcome(ok=False, reason="network_error")
        cap = self.store.put_raw(self.answers[url], vendor_id=vendor_id, family=family, collector=collector,
                                 url_requested=url, status=200, content_type="application/dns-json",
                                 retrieved_at="2026-10-03T09:00:00Z")
        return FetchOutcome(ok=True, capture=cap, status=200)


def test_dns_check_compares_captured_and_live_ai_tokens(tmp_path: Path) -> None:
    profile, plan = S.profile(), S.depth()
    urls = [f"{base}?name=acme.example&type=TXT" for base in ("https://dns.google/resolve",
                                                               "https://cloudflare-dns.com/dns-query")]
    pack = EvidenceStore(tmp_path / "pack")
    for url in urls:
        pack.put_raw(_doh("openai-domain-verification=dv-abc"), vendor_id=profile.vendor_id, family=SourceFamily.DNS,
                     collector="dns", url_requested=url, status=200, content_type="application/dns-json",
                     retrieved_at="2026-10-02T10:00:00Z")
    live_answer = _doh("anthropic-domain-verification-x1=abc")
    made: list[FakeLive] = []

    def live(store: EvidenceStore) -> FakeLive:
        made.append(FakeLive(store, {url: live_answer for url in urls}))
        return made[-1]

    check = ui.dns_check(profile, plan, seeds_dir=tmp_path / "no-seeds", evidence_dir=tmp_path / "pack",
                         live_fetcher=live)
    assert check.domains == ["acme.example"]
    assert check.captured_providers == ["OpenAI"] and check.captured_at == "2026-10-02T10:00:00Z"
    assert check.live_providers == ["Anthropic"] and not check.unchanged
    assert all(u.startswith(("https://dns.google/", "https://cloudflare-dns.com/")) for u in made[0].urls)
    assert len(EvidenceStore(tmp_path / "pack").captures()) == 2  # the evidence pack is never written
    assert not (tmp_path / "pack" / "text").exists()


# --------------------------------------------------------------------------- ui helpers: text


def test_text_helpers_escape_clip_scrub_and_shorten(monkeypatch: pytest.MonkeyPatch) -> None:
    assert ui.md_escape("Costs $5 *now*") == r"Costs \$5 \*now\*"
    assert ui.clip("one  two\nthree four", 12) == "one two…"
    assert ui.clip("short", 12) == "short"
    monkeypatch.setenv(ui.GEMINI_ENV, "plain-secret-value")
    message = ui.scrub("failed: AIza" + "A" * 30 + " and plain-secret-value and ghp_" + "b" * 30)
    assert "AIza" not in message and "plain-secret-value" not in message and "ghp_" not in message
    assert ui.dmy("2026-10-02") == "02-10-2026" and ui.dmy("") == "" and ui.dmy("undated") == "undated"
    assert ui.short_name("Financial Statement Services, Inc. (FSSI)") == "FSSI"
    assert ui.short_name("Fiserv, Inc.") == "Fiserv"
    assert ui.short_name("The Clearing House Payments Company L.L.C.") == "The Clearing House Payments"
