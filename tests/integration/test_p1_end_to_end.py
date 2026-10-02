"""P1 end to end: real workbook -> assess -> export -> re-read, plus CLI smoke tests."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from typer.testing import CliRunner

from footprint.cli import app
from footprint.pipeline import assess_example, assess_profiles, export_workbook
from footprint.workbook import check_fidelity, read_workbook

INPUT = Path(__file__).resolve().parents[2] / "data" / "input" / "Meridian_Vendor_Input.xlsx"
EXPECTED = ["High", "Critical", "High", "High", "Critical", "Critical"]
SHEETS = ["Vendor Inventory", "Field Guide", "Points to consider", "Criticality Workings", "Method & Legend"]


@pytest.fixture(scope="module")
def exported(tmp_path_factory: pytest.TempPathFactory) -> Path:
    data = read_workbook(INPUT)
    out = tmp_path_factory.mktemp("p1") / "out.xlsx"
    export_workbook(INPUT, out, assess_profiles(data), example=assess_example(data), team="Test Team")
    return out


def test_tiers_and_text(exported: Path) -> None:
    ws = openpyxl.load_workbook(exported)["Vendor Inventory"]
    assert [ws[f"B{r}"].value for r in range(6, 12)] == [f"V-00{i}" for i in range(1, 7)]
    assert [ws[f"L{r}"].value for r in range(6, 12)] == EXPECTED
    for r in range(6, 12):
        for col in "MN":
            text = ws[f"{col}{r}"].value
            assert text and len(text) <= 1050
        for col in "OPQRSTUV":
            assert ws[f"{col}{r}"].value in (None, "")


def test_example_row_and_fidelity(exported: Path) -> None:
    src = openpyxl.load_workbook(INPUT)["Vendor Inventory"]
    out = openpyxl.load_workbook(exported)["Vendor Inventory"]
    assert [c.value for c in out[5]] == [c.value for c in src[5]]
    assert check_fidelity(INPUT, exported) == []
    assert openpyxl.load_workbook(exported).sheetnames == SHEETS


def test_cli_criticality_and_override(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("FOOTPRINT_OVERRIDES", str(tmp_path / "review" / "overrides.jsonl"))
    runner = CliRunner()
    result = runner.invoke(app, ["criticality", str(INPUT)])
    assert result.exit_code == 0, result.output
    for i in range(1, 7):
        assert f"V-00{i}" in result.output

    result = runner.invoke(app, ["override-tier", "V-001", "Critical", "--reason", "analyst judgement for test",
                                 "--analyst", "Test Analyst"])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "review" / "overrides.jsonl").exists()

    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, ["criticality", str(INPUT), "--out", str(out)])
    assert result.exit_code == 0, result.output
    ws = openpyxl.load_workbook(out)["Vendor Inventory"]
    assert ws["L6"].value == "Critical"
    assert "override" in ws["M6"].value.lower()


def test_cli_rejects_bad_tier(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("FOOTPRINT_OVERRIDES", str(tmp_path / "overrides.jsonl"))
    result = CliRunner().invoke(app, ["override-tier", "V-001", "Huge", "--reason", "long enough reason",
                                      "--analyst", "A"])
    assert result.exit_code != 0


@pytest.fixture
def isolated(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("FOOTPRINT_OVERRIDES", str(tmp_path / "overrides.jsonl"))
    return tmp_path


def test_cli_write_errors_are_short_messages(isolated: Path) -> None:
    source = isolated / "in.xlsx"
    source.write_bytes(INPUT.read_bytes())
    result = CliRunner().invoke(app, ["criticality", str(source), "--out", str(source), "--force"])
    assert result.exit_code == 2
    assert "E13" in result.output and "Traceback" not in result.output
    assert not isinstance(result.exception, Exception) or isinstance(result.exception, SystemExit)


def test_cli_refuses_an_existing_out_file_without_force(isolated: Path) -> None:
    out = isolated / "out.xlsx"
    out.write_bytes(b"keep me")
    result = CliRunner().invoke(app, ["criticality", str(INPUT), "--out", str(out)])
    assert result.exit_code == 2 and out.read_bytes() == b"keep me"
    result = CliRunner().invoke(app, ["criticality", str(INPUT), "--out", str(out), "--force"])
    assert result.exit_code == 0, result.output
    assert openpyxl.load_workbook(out)["Vendor Inventory"]["L6"].value == "High"


def test_cli_override_store_does_not_depend_on_the_cwd(isolated: Path) -> None:
    (isolated / "review").mkdir()
    (isolated / "review" / "overrides.jsonl").write_text("not json\n", encoding="utf-8")  # ignored: not the store
    result = CliRunner().invoke(app, ["criticality", str(INPUT)])
    assert result.exit_code == 0, result.output
    assert "Overrides: none" in result.output


def test_export_uses_one_rubric_for_cells_and_sheets(tmp_path: Path) -> None:
    from footprint.criticality import load_rubric

    data = read_workbook(INPUT)
    rubric = load_rubric().model_copy(update={"version": "9.9-test"})
    with pytest.raises(ValueError, match="another rubric version"):
        export_workbook(INPUT, tmp_path / "x.xlsx", assess_profiles(data), rubric=rubric)
    out = tmp_path / "y.xlsx"
    export_workbook(INPUT, out, assess_profiles(data, rubric=rubric), rubric=rubric)
    assert "rubric v9.9-test" in openpyxl.load_workbook(out)["Vendor Inventory"]["M6"].value
