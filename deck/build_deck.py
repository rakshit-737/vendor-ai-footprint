"""Build the Team Osprey walkthrough deck: 13 slides plus an appendix, every diagram as native, editable shapes.

    uv run python deck/build_deck.py --findings runs/findings.json \
        --out submission/Team_Osprey_Vendor_AI_Footprint.pptx

Data sources:
- Slides 2 and 11 (and one appendix slide per vendor) come from a findings JSON: the AssessmentResult that
  footprint.pipeline.run_assessment writes (runs/<run_id>/assessment.json, or just the run folder), or the smaller
  deck summary (``DeckData``). Without one, those slides show clearly marked placeholders.
- Slide 9's tier table is computed live with footprint.criticality from the input workbook.
- Slides 6 and 10 read config/depth.toml and config/tou.toml, plus config/sources.toml when it exists.

The design follows docs/design.md ("Deck outline" and Appendix A section 3). Team styling only: no sponsor branding.
Confidential: the deck is written to a local file. Nothing here uploads, publishes or contacts anyone.
"""

from __future__ import annotations

import argparse
import datetime as dt
import functools
import hashlib
import importlib.util
import io
import json
import math
import os
import re
import sys
import tomllib
import uuid
import zipfile
from collections import defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from footprint.criticality import Rubric, load_rubric, score_profile
from footprint.depth import DepthConfig, load_depth_config, plan_depth
from footprint.models import (
    ROLE_ORDER,
    AssessmentResult,
    CriticalityResult,
    DepthPlan,
    SourceFamily,
    Tier,
    VendorFindings,
    VendorProfile,
)
from footprint.workbook import read_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "input" / "Meridian_Vendor_Input.xlsx"
DEFAULT_OUT = ROOT / "submission" / "Team_Osprey_Vendor_AI_Footprint.pptx"
DEFAULT_TEAM = "Team Osprey"
MAIN_SLIDES = 13


def _load_diagrams() -> Any:
    """Load the sibling diagrams.py under a private module name (deck/ is not a package)."""
    name = "footprint_deck_diagrams"
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, Path(__file__).with_name("diagrams.py"))
        if spec is None or spec.loader is None:
            raise ImportError("deck/diagrams.py not found")
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    return sys.modules[name]


D = _load_diagrams()
C = D.C
R = D.run
P = D.para

SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in (16:9)
LEFT, WIDTH = 0.55, 12.23
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid": fills and rules are ours
NAMESPACE = uuid.UUID("5f1c7a2e-8d3b-4c55-9a8e-2b6f0d4c9e17")
CLASS_RANK = {"None identified": -1, "Low": 0, "Medium": 1, "High": 2, "Critical": 3}


# =========================================================================== findings input


class DeckEvidence(BaseModel):
    """One cited excerpt as the deck shows it (never paraphrased; clipped only for display)."""

    model_config = ConfigDict(extra="ignore")

    evidence_id: str = ""
    role: str = ""
    strength: str = ""
    source_type: str = ""
    publisher: str = ""
    published: str = ""
    url: str = ""
    excerpt: str = ""
    tags: str = ""
    screenshot: bool = Field(default=False, description="the item has a captured screenshot the app can show")


class DeckVendor(BaseModel):
    """One vendor's headline results: the deck summary row."""

    model_config = ConfigDict(extra="ignore")

    vendor_id: str
    name: str
    tier: str
    score: int | None = None
    usage: str = Field(description="column O: Yes, No or Inconclusive")
    verdict: str = Field(default="", description="Confirmed, Probable, Affirmed negative, Not detected, Inconclusive")
    likelihood: str = ""
    confidence: str = ""
    confidence_reason: str = ""
    risk: str = Field(description="column S: Critical, High, Medium, Low or None identified")
    provisional: bool = False
    ceiling: str = ""
    arp: int | None = None
    e: int | None = None
    k: int | None = None
    tp: int | None = None
    tg: int | None = None
    e_assumed: bool = False
    k_assumed: bool = False
    flip: str = ""
    action: str = ""
    gap_blocks: list[str] = Field(default_factory=list)
    questionnaire: list[str] = Field(default_factory=list)
    evidence: list[DeckEvidence] = Field(default_factory=list, description="decisive items, best first")
    trap: DeckEvidence | None = None
    shot: DeckEvidence | None = Field(default=None, description="the best cited item that has a screenshot")
    screenshots: int = Field(default=0, description="cited items with a screenshot (the Evidence Images sheet)")
    providers_named: list[str] = Field(default_factory=list, description="named by the vendor (SR A/B, not DNS)")
    providers_third: list[str] = Field(default_factory=list, description="named only by third parties")
    providers_dns: list[str] = Field(default_factory=list, description="seen only as DNS verification tokens")
    llm_agree: int = 0
    llm_disagree: int = 0
    llm_proposed: int = 0
    llm_labelled: int = Field(default=0, description="items with Gemini labels shown beside the rule labels")
    cells: dict[str, str] = Field(default_factory=dict, description="workbook cells P-U, for the presenter notes")

    @property
    def gemini(self) -> bool:
        """Whether this vendor's evidence carries any Gemini output the app can compare with the rules."""
        return bool(self.llm_labelled or self.llm_agree or self.llm_disagree or self.llm_proposed)


class DeckData(BaseModel):
    """The deck summary format; also what an AssessmentResult is reduced to."""

    model_config = ConfigDict(extra="ignore")

    run_id: str = ""
    mode: str = ""
    as_of: str = ""
    vendors: list[DeckVendor] = Field(default_factory=list)
    llm_calls: int | None = None
    withheld: int | None = None


PROVIDER_CANON = {
    "openai": "OpenAI", "chatgpt": "OpenAI", "gpt": "OpenAI", "azure openai": "OpenAI",
    "anthropic": "Anthropic", "claude": "Anthropic",
    "microsoft": "Microsoft", "copilot": "Microsoft", "microsoft copilot": "Microsoft", "github copilot": "Microsoft",
    "azure": "Microsoft", "m365 copilot": "Microsoft",
    "google": "Google", "gemini": "Google", "vertex ai": "Google", "google cloud": "Google",
    "aws": "AWS", "amazon": "AWS", "amazon bedrock": "AWS", "bedrock": "AWS", "amazon web services": "AWS",
    "servicenow": "ServiceNow", "cognition": "Cognition", "devin": "Cognition", "meta": "Meta", "llama": "Meta",
    "mistral": "Mistral AI", "ibm": "IBM", "watsonx": "IBM", "ollama": "Ollama",
}


def canonical_provider(name: str) -> str:
    """'ChatGPT' -> 'OpenAI', 'Amazon Bedrock' -> 'AWS'; unknown names are kept as given."""
    key = " ".join(name.lower().split())
    if key in PROVIDER_CANON:
        return PROVIDER_CANON[key]
    for alias in sorted(PROVIDER_CANON, key=len, reverse=True):
        if key.startswith((alias + " ", alias + "-")):
            return PROVIDER_CANON[alias]
    return name.strip()


_LEGAL_SUFFIX = re.compile(r",?\s+(?:Inc|L\.?L\.?C|Ltd|Limited|Corp|Corporation|plc|LLP)\.?$", re.IGNORECASE)


@functools.lru_cache(maxsize=None)
def _seed_aliases(seeds_dir: str, vendor_id: str) -> tuple[str, ...]:
    path = Path(seeds_dir) / f"{vendor_id}.toml"
    if not path.is_file():
        return ()
    try:
        with path.open("rb") as handle:
            return tuple(str(a) for a in tomllib.load(handle).get("aliases", []))
    except (OSError, tomllib.TOMLDecodeError):
        return ()


def short_name(name: str, aliases: Sequence[str] = ()) -> str:
    """'Financial Statement Services, Inc. (FSSI)' -> 'FSSI'; 'Fiserv, Inc.' -> 'Fiserv'; a long legal name falls
    back to its shortest seeded alias ('The Clearing House Payments Company L.L.C.' -> 'TCH')."""
    acronym = re.search(r"\(([A-Z][A-Z0-9&]{1,7})\)", name)
    if acronym:
        return acronym.group(1)
    if name.upper().startswith("EXAMPLE") and "—" in name:
        name = name.split("—", 1)[1]
    base = _LEGAL_SUFFIX.sub("", re.sub(r"\s*\([^)]*\)", "", name).strip(" ,")).strip(" ,") or name.strip()
    if len(base) <= 14:
        return base
    short = sorted((a for a in aliases if len(a) <= 14), key=lambda a: (len(a), a))
    return short[0] if short else base


def _first_sentence(text: str) -> str:
    text = " ".join(text.split())
    match = re.match(r"(.+?[.!?])(?:\s|$)", text)
    return match.group(1) if match else text


def _evidence(item: Any) -> DeckEvidence:
    return DeckEvidence(evidence_id=item.evidence_id, role=item.role, strength=item.strength,
                        source_type=item.source_type, publisher=item.publisher, published=item.published,
                        url=item.url, excerpt=item.excerpt, tags=item.tag_string,
                        screenshot=bool(item.screenshot_path))


def summarize_vendor(findings: VendorFindings, *, name: str) -> DeckVendor:
    """Reduce one VendorFindings to the deck's summary row (decisive evidence first, providers by basis)."""
    by_key = {i.item_key: i for i in findings.evidence}
    decisive = [by_key[k] for k in findings.verdict.decisive if k in by_key]
    if not decisive:
        decisive = sorted(findings.cited(), key=lambda i: (ROLE_ORDER.index(i.role), i.evidence_id or i.item_key))
    named: set[str] = set()
    third: set[str] = set()
    dns: set[str] = set()
    for item in findings.evidence:
        if not item.citable:
            continue
        is_dns = item.family == SourceFamily.DNS or item.source_type.lower().startswith("dns")
        for provider in item.providers:
            canon = canonical_provider(provider)
            (dns if is_dns else named if item.tags.sr in ("A", "B") else third).add(canon)
    trap = next((i for i in findings.evidence if i.tags.ai_type == "not_ai"), None)
    cited = sorted(findings.cited(), key=lambda i: (ROLE_ORDER.index(i.role), i.evidence_id or i.item_key))
    cited_keys = {i.item_key for i in cited}
    shot = next((i for i in (*decisive, *cited) if i.screenshot_path and i.item_key in cited_keys), None)
    risk, inputs = findings.risk, findings.risk.inputs
    actions = findings.actions
    action = next((_first_sentence(t) for t in (*actions.class_playbook, actions.text,
                                                findings.cells.recommended_action or "") if t), "")
    cells = {key: value for key, value in findings.cells.model_dump().items()
             if value and key in ("evidence", "how_ai_used", "ai_subprocessors", "risk_rationale",
                                  "recommended_action")}
    return DeckVendor(
        vendor_id=findings.vendor_id, name=name, tier=findings.criticality.tier.value,
        score=findings.criticality.score, usage=findings.verdict.column_o, verdict=findings.verdict.label,
        likelihood=findings.verdict.likelihood, confidence=findings.verdict.confidence,
        confidence_reason=findings.verdict.confidence_reason, risk=risk.final_class, provisional=risk.provisional,
        ceiling=risk.ceiling_class, arp=risk.arp, e=inputs.e, k=inputs.k, tp=inputs.tp, tg=inputs.tg,
        e_assumed=inputs.e_assumed, k_assumed=inputs.k_assumed, flip=risk.flip_condition, action=action,
        gap_blocks=list(actions.gap_blocks), questionnaire=list(actions.questionnaire_items),
        evidence=[_evidence(i) for i in decisive[:3]], trap=_evidence(trap) if trap else None,
        shot=_evidence(shot) if shot else None, screenshots=sum(1 for i in cited if i.screenshot_path),
        providers_named=sorted(named, key=str.lower), providers_third=sorted(third - named, key=str.lower),
        providers_dns=sorted(dns - named - third, key=str.lower),
        llm_agree=sum(1 for i in findings.evidence if i.method == "rule+llm_agree"),
        llm_disagree=sum(1 for i in findings.evidence if i.llm_labels and i.label_disagreements),
        llm_proposed=sum(1 for i in findings.evidence if i.method == "llm_proposed_accepted"),
        llm_labelled=sum(1 for i in findings.evidence if i.llm_labels),
        cells=cells,
    )


def summarize_assessment(result: AssessmentResult | Sequence[VendorFindings], *,
                         seeds_dir: str | Path = ROOT / "seeds") -> DeckData:
    vendors = result.vendors if isinstance(result, AssessmentResult) else list(result)
    rows = [summarize_vendor(f, name=short_name(f.profile.name, _seed_aliases(str(seeds_dir), f.vendor_id)))
            for f in vendors]
    if not isinstance(result, AssessmentResult):
        return DeckData(vendors=rows)
    llm = result.manifest.get("llm", {}) if isinstance(result.manifest, dict) else {}
    calls = llm.get("calls", {}) if isinstance(llm, dict) else {}
    withheld = llm.get("withheld", {}) if isinstance(llm, dict) else {}
    return DeckData(
        run_id=result.run_id, mode=result.mode, as_of=result.as_of, vendors=rows,
        llm_calls=sum(int(v.get("actual", 0)) for v in calls.values() if isinstance(v, dict)) if calls else None,
        withheld=sum(int(v.get("count", 0)) for v in withheld.values() if isinstance(v, dict)) if withheld else None,
    )


def findings_file(path: str | Path) -> Path:
    """A findings path as given, or a run folder (runs/<run_id>/) holding assessment.json."""
    path = Path(path)
    return path / "assessment.json" if path.is_dir() else path


def load_findings(path: str | Path, *, seeds_dir: str | Path = ROOT / "seeds") -> DeckData:
    """Read an AssessmentResult JSON, a list of VendorFindings, or the deck summary format (a run folder works too:
    its assessment.json is read)."""
    path = findings_file(path)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"{path} is not valid JSON: {exc}") from exc
    try:
        if isinstance(data, dict) and "input_sha256" in data:
            return summarize_assessment(AssessmentResult.model_validate(data), seeds_dir=seeds_dir)
        items = data if isinstance(data, list) else data.get("vendors", []) if isinstance(data, dict) else []
        if items and isinstance(items[0], dict) and "profile" in items[0]:
            deck = summarize_assessment([VendorFindings.model_validate(v) for v in items], seeds_dir=seeds_dir)
            if isinstance(data, dict):
                deck = deck.model_copy(update={k: data[k] for k in ("run_id", "mode", "as_of") if k in data})
            return deck
        return DeckData.model_validate(data)
    except ValidationError as exc:
        raise ValueError(f"{path} is neither an AssessmentResult nor a deck summary: {exc.errors()[0]['msg']}") from exc


# =========================================================================== live inputs (workbook and config)


@dataclass
class Scored:
    profile: VendorProfile
    result: CriticalityResult
    plan: DepthPlan
    name: str


SR_LADDER_DEFAULT: tuple[tuple[str, str, str], ...] = (
    ("A", "Legally accountable", "SEC filings, privacy notices, DPAs, sub-processor lists"),
    ("B", "First-party", "product pages, trust pages, own job postings, DNS"),
    ("C", "Promotional", "press releases, blogs, provider stories, trade press"),
    ("D", "Syndicated", "aggregators and syndicated copies"),
    ("E–F", "Excluded", "user-generated or unknown origin"),
)
SR_NAMES = {row[0]: row[1] for row in SR_LADDER_DEFAULT}
SR_EXAMPLE_BUDGET = 44


def ladder_examples(types: Sequence[tuple[str, str]], budget: int = SR_EXAMPLE_BUDGET) -> str:
    """Short, varied examples for one SR grade: the shortest source type of each family in turn, while they fit the
    character budget, then '+N more'. ``types`` holds (source_type, family) pairs in file order."""
    groups: dict[str, list[str]] = {}
    for name, family in types:
        if name not in groups.setdefault(family, []):
            groups[family].append(name)
    for names in groups.values():
        names.sort(key=lambda n: (len(n), n))
    picked: list[str] = []
    used = 0
    while any(groups.values()):
        for family in list(groups):
            if not groups[family]:
                continue
            name = groups[family].pop(0)
            cost = len(name) + (2 if picked else 0)
            if used + cost > budget:
                groups[family] = []
                continue
            picked.append(name)
            used += cost
    total = len({name for name, _ in types})
    more = total - len(picked)
    return ", ".join(picked) + (f" +{more} more" if more > 0 else "")


def load_sr_ladder(config_dir: Path) -> tuple[list[tuple[str, str, str]], str]:
    """The SR ladder: from config/sources.toml when it lists source types with an `sr` grade, else the design's.

    sources.toml is owned by the rules module and its layout may change, so every table that has both `sr` and
    `source_type` (or `type`) keys is read, wherever it sits in the file."""
    path = config_dir / "sources.toml"
    if path.is_file():
        try:
            with path.open("rb") as handle:
                tree = tomllib.load(handle)
        except (OSError, tomllib.TOMLDecodeError):
            tree = {}
        found: dict[str, list[tuple[str, str]]] = defaultdict(list)

        def walk(node: Any) -> None:
            if isinstance(node, dict):
                kind = node.get("source_type") or node.get("type")
                grade = node.get("sr")
                if isinstance(kind, str) and isinstance(grade, str) and grade.upper() in ("A", "B", "C", "D", "E",
                                                                                         "F"):
                    found[grade.upper()].append((kind, str(node.get("family", ""))))
                for value in node.values():
                    walk(value)
            elif isinstance(node, list):
                for value in node:
                    walk(value)

        walk(tree)
        if found:
            ladder = [(g, SR_NAMES[g], ladder_examples(found[g])) for g in ("A", "B", "C", "D") if found.get(g)]
            rest = found.get("E", []) + found.get("F", [])
            if rest:
                ladder.append(("E–F", SR_NAMES["E–F"], ladder_examples(rest)))
            return ladder, "config/sources.toml"
    return list(SR_LADDER_DEFAULT), "design §2.3 (config/sources.toml not present)"


def load_terms_register(config_dir: Path) -> dict[str, Any]:
    path = config_dir / "tou.toml"
    counts = {"full": 0, "limited": 0, "none": 0}
    manual: list[str] = []
    if path.is_file():
        with path.open("rb") as handle:
            hosts = tomllib.load(handle).get("host", [])
        for host in hosts:
            kind = str(host.get("automation", "limited"))
            counts[kind] = counts.get(kind, 0) + 1
            if kind == "none":
                manual.append(str(host.get("match", "")))
    return {"counts": counts, "manual": manual, "present": path.is_file()}


@dataclass
class Context:
    team: str
    as_of: str
    data: DeckData | None
    rubric: Rubric
    depth: DepthConfig
    scored: list[Scored] = field(default_factory=list)
    example: Scored | None = None
    workbook_name: str = ""
    sr_ladder: list[tuple[str, str, str]] = field(default_factory=list)
    sr_source: str = ""
    terms: dict[str, Any] = field(default_factory=dict)
    chain: DeckVendor | None = None
    appendix: list[str] = field(default_factory=list)

    @property
    def vendors(self) -> list[DeckVendor]:
        return self.data.vendors if self.data else []


def score_workbook(path: Path, *, rubric: Rubric, depth: DepthConfig,
                   seeds_dir: Path) -> tuple[list[Scored], Scored | None]:
    """Criticality and depth for every vendor in the workbook, computed live (footprint.criticality)."""
    if not path.is_file():
        return [], None
    workbook = read_workbook(path)
    if not workbook.ok:
        return [], None

    def one(profile: VendorProfile) -> Scored:
        result = score_profile(profile, rubric)
        aliases = _seed_aliases(str(seeds_dir), profile.vendor_id)
        return Scored(profile, result, plan_depth(result, profile, config=depth), short_name(profile.name, aliases))

    return [one(p) for p in workbook.vendors], one(workbook.example) if workbook.example else None


def pick_chain(data: DeckData | None, preferred: str | None = None) -> DeckVendor | None:
    """The vendor slide 11 walks through: the given one, else a Confirmed vendor (highest score), else Probable."""
    if data is None or not data.vendors:
        return None
    if preferred:
        chosen = next((v for v in data.vendors if v.vendor_id == preferred), None)
        if chosen is None:
            raise ValueError(f"--chain-vendor {preferred} is not in the findings")
        return chosen
    order = {"Confirmed": 0, "Probable": 1}
    candidates = [(order.get(v.verdict, 2), -(v.arp or 0), i, v) for i, v in enumerate(data.vendors) if v.evidence]
    return min(candidates)[3] if candidates else data.vendors[0]


# =========================================================================== text helpers


def clip(text: str, limit: int, marker: str = "…") -> str:
    text = " ".join(str(text).split())
    if len(text) <= limit:
        return text
    cut = text[: max(limit - len(marker), 1)].rsplit(" ", 1)[0].rstrip(" ,;:.")
    return cut + marker


def quote(text: str, limit: int) -> str:
    """A display copy of a verified excerpt: never reworded; an omission is marked [...]."""
    return "“" + clip(text, limit, " […]") + "”"


def plural(count: int, one: str, many: str) -> str:
    return f"{count} {one if count == 1 else many}"


def tag_lines(tags: str) -> list[str]:
    """'U1 · SR:B · SP:S2 · RL:R3 · RC:T3 · IC:1 · locus=service_feature' -> three short display lines."""
    parts = [p.strip() for p in tags.split("·") if p.strip()]
    codes = [p for p in parts if not p.startswith("locus=")]
    locus = [p.split("=", 1)[1].replace("_", " ") for p in parts if p.startswith("locus=")]
    lines = [" · ".join(codes[i:i + 3]) for i in range(0, len(codes), 3)]
    return lines + [f"locus: {locus[0]}"] if locus else lines


def and_list(items: Sequence[str]) -> str:
    items = list(items)
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def verdict_words(v: DeckVendor, *, sep: str = " · ") -> str:
    """'Yes · Confirmed' (or 'Yes (Confirmed)' with sep='(') for column O plus its label; 'Inconclusive' alone when
    the two are the same word."""
    if not v.verdict or v.verdict == v.usage:
        return v.usage
    return f"{v.usage} ({v.verdict})" if sep == "(" else f"{v.usage}{sep}{v.verdict}"


def flips_to(text: str) -> str:
    """The class a flip condition moves to (the last class word it names), or ''."""
    found = re.findall(r"\b(Critical|High|Medium|Low)\b", text or "")
    return found[-1] if found else ""


MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December")


def month_year(as_of: str) -> str:
    day = dt.date.fromisoformat(as_of)
    return f"{MONTHS[day.month - 1]} {day.year}"


def long_date(as_of: str) -> str:
    day = dt.date.fromisoformat(as_of)
    return f"{day.day} {MONTHS[day.month - 1]} {day.year}"


def bullets(items: Iterable[str], *, size: float = 11, color: str = C.SLATE, after: float = 4,
            bullet_color: str = C.TEAL) -> list[dict[str, Any]]:
    return [P(R(text, size=size, color=color), bullet=True, space_after=after, bullet_color=bullet_color)
            for text in items]


def headline(data: DeckData) -> str:
    """Slide 2's headline, generated from the run (design §3 wording)."""
    vendors = data.vendors
    names = lambda vs: and_list([v.name for v in vs])  # noqa: E731
    yes = [v for v in vendors if v.usage == "Yes"]
    if yes:
        first = f"AI runs in service delivery at {len(yes)} of {len(vendors)} vendors ({names(yes)})."
    else:
        first = f"No vendor shows AI in service delivery on public evidence ({len(vendors)} assessed)."
    critical = [v for v in vendors if v.risk == "Critical"]
    clauses = [f"{names(critical)} {'reaches' if len(critical) == 1 else 'reach'} Critical AI risk" if critical
               else "None reaches Critical AI risk on public evidence"]
    ceilings = [v for v in vendors if v.provisional and v.ceiling == "Critical"]
    if len(ceilings) == 1:
        clauses.append(f"{ceilings[0].name}'s ceiling is Critical if confirmed")
    elif ceilings:
        clauses.append(f"the ceilings of {names(ceilings)} are Critical if confirmed")
    flips = [v for v in vendors if v.usage == "Yes" and v.risk != "Critical" and flips_to(v.flip) == "Critical"]
    if flips:
        clauses.append(f"{names(flips)} would become Critical on a single confirming answer")
    second = clauses[0] + ("; " + ", and ".join(clauses[1:]) if len(clauses) > 1 else "")
    return f"{first} {second}."


# =========================================================================== deck frame


def _apply_theme(prs: Any) -> None:
    """Team theme: Calibri for headings and body, team colours, 16:9 placeholders on the master and layouts."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root.set("name", "Team Osprey")
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", "Team Osprey")
    colours = {"dk2": C.INK, "lt2": C.SURFACE, "accent1": C.TEAL, "accent2": C.INK, "accent3": C.SEAFOAM,
               "accent4": C.SAFFRON, "accent5": "9B1C1C", "accent6": C.SLATE, "hlink": C.TEAL_DARK,
               "folHlink": C.MUTED}
    for tag, value in colours.items():
        slot = scheme.find(f"a:{tag}", ns)
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, qn("a:srgbClr"), val=value)
    fonts = root.find(".//a:fontScheme", ns)
    fonts.set("name", "Team Osprey")
    for which in ("majorFont", "minorFont"):
        fonts.find(f"a:{which}/a:latin", ns).set("typeface", "Calibri")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    scale = SLIDE_W / 9144000
    for container in (prs.slide_master, *prs.slide_layouts):
        for shape in container.placeholders:
            if shape.left is not None and shape.width is not None:
                shape.left, shape.width = int(shape.left * scale), int(shape.width * scale)
    title_style = prs.slide_master._element.find(".//" + qn("p:titleStyle"))
    if title_style is not None:
        rpr = title_style.find(qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
        if rpr is not None:
            rpr.set("sz", "3000")
            rpr.set("b", "1")


def _footer(slide: Any, ctx: Context, number: int, *, dark: bool = False) -> None:
    colour = "9FB3BA" if dark else C.MUTED
    D.label(slide, LEFT, 7.04, 9.5, 0.28, P(R(f"{ctx.team} · Vendor AI footprint · Confidential case-study material, "
                                              "not for distribution", size=10, color=colour)),
            anchor=MSO_ANCHOR.MIDDLE, name="Footer")
    box = D.label(slide, 11.98, 7.04, 0.8, 0.28, P(R(str(number), size=10, color=colour)), align=PP_ALIGN.RIGHT,
                  anchor=MSO_ANCHOR.MIDDLE, name="Slide number")
    p = box.text_frame.paragraphs[0]._p
    old = p.find(qn("a:r"))
    field_id = "{" + str(uuid.uuid5(NAMESPACE, f"slide-number-{number}")).upper() + "}"
    fld = p.makeelement(qn("a:fld"), {"id": field_id, "type": "slidenum"})
    rpr = fld.makeelement(qn("a:rPr"), {"lang": "en-GB", "sz": "1000", "dirty": "0"})
    fill = rpr.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": colour}))
    rpr.append(fill)
    text = fld.makeelement(qn("a:t"), {})
    text.text = str(number)
    fld.extend([rpr, text])
    old.addprevious(fld)
    p.remove(old)


def _notes(slide: Any, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text.strip()


def content_slide(prs: Any, ctx: Context, title: str, kicker: str, tag: str, notes: str) -> Any:
    """A light content slide: title placeholder, corner tag, kicker line, footer and speaker notes."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    heading = slide.shapes.title
    heading.left, heading.top, heading.width, heading.height = Inches(LEFT), Inches(0.34), Inches(10.4), Inches(0.7)
    D.write_text(heading.text_frame, P(R(title, size=30, bold=True, color=C.INK)), anchor=MSO_ANCHOR.BOTTOM,
                 margins=(0, 0, 0, 0))
    D.pill(slide, 11.2, 0.5, 1.58, 0.34, tag, fill=C.TEAL_TINT, color=C.TEAL_DARK, size=10, name="Corner tag")
    if kicker:
        D.label(slide, LEFT, 1.08, WIDTH, 0.55, P(R(kicker, size=15, color=C.TEAL_DARK)), name="Kicker")
    _footer(slide, ctx, len(prs.slides))
    _notes(slide, notes)
    return slide


def dark_slide(prs: Any, ctx: Context, notes: str, layout: int = 6) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = D.rgb(C.INK)
    _footer(slide, ctx, len(prs.slides), dark=True)
    _notes(slide, notes)
    return slide


def _set_cell_rule(cell: Any, colour: str, width: float = 0.75) -> None:
    """A hairline under a table cell (tcPr/lnB, inserted before the cell fill as the schema requires)."""
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("a:lnB")):
        tcpr.remove(old)
    ln = tcpr.makeelement(qn("a:lnB"), {"w": str(int(Pt(width))), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": colour}))
    ln.append(fill)
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
    D._insert_before_successor(tcpr, ln, ("a:lnTlToBr", "a:lnBlToTr", "a:cell3D", "a:noFill", "a:solidFill",
                                          "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill", "a:headers",
                                          "a:extLst"))


def table(slide: Any, x: float, y: float, widths: Sequence[float], header: Sequence[str],
          rows: Sequence[Sequence[Any]], *, row_h: float | Sequence[float] = 0.4, header_h: float = 0.34,
          size: float = 10, name: str | None = None) -> Any:
    """A native table. A cell is a string, or a dict with 'paras' (paragraph specs) and optional 'fill', 'color',
    'bold', 'align' and 'size'."""
    heights = list(row_h) if isinstance(row_h, Sequence) else [row_h] * len(rows)
    frame = slide.shapes.add_table(len(rows) + 1, len(widths), Inches(x), Inches(y), Inches(sum(widths)),
                                   Inches(header_h + sum(heights)))
    if name:
        frame.name = name
    tbl = frame.table
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = tbl._tbl.tblPr.makeelement(qn("a:tableStyleId"), {})
        tbl._tbl.tblPr.append(style)
    style.text = NO_STYLE_TABLE
    tbl.first_row, tbl.horz_banding = True, False
    for index, width in enumerate(widths):
        tbl.columns[index].width = Inches(width)
    tbl.rows[0].height = Inches(header_h)
    for index, height in enumerate(heights, start=1):
        tbl.rows[index].height = Inches(height)
    for col, text in enumerate(header):
        _cell(tbl.cell(0, col), {"paras": [P(R(text, size=size, bold=True, color=C.WHITE))], "fill": C.INK})
    for r_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            spec = value if isinstance(value, dict) else {"paras": [P(R(str(value), size=size))]}
            _cell(tbl.cell(r_index, col), spec, size=size)
            _set_cell_rule(tbl.cell(r_index, col), C.LINE)
    return frame


def _cell(cell: Any, spec: dict[str, Any], *, size: float = 10) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = D.rgb(spec.get("fill", C.WHITE))
    cell.margin_left = cell.margin_right = Inches(spec.get("pad", 0.07))
    cell.margin_top = cell.margin_bottom = Inches(0.035)
    cell.vertical_anchor = spec.get("anchor", MSO_ANCHOR.MIDDLE)
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    D.write_text(tf, spec["paras"], size=spec.get("size", size), color=spec.get("color", C.SLATE),
                 bold=spec.get("bold", False), align=spec.get("align", PP_ALIGN.LEFT), margins=(0.07, 0.035, 0.07,
                                                                                               0.035))


def card(slide: Any, x: float, y: float, w: float, h: float, title: str, body: Sequence[Any], *,
         fill: str = C.SURFACE, title_color: str = C.INK, title_size: float = 12, name: str | None = None,
         line: str | None = None) -> Any:
    """A rounded card with a bold heading and body paragraphs, anchored to the top."""
    return D.box(slide, x, y, w, h, [P(R(title, size=title_size, bold=True, color=title_color), space_after=4),
                                     *body], fill=fill, line=line, radius=0.1, anchor=MSO_ANCHOR.TOP,
                 margins=(0.14, 0.1, 0.12, 0.08), name=name or title)


def class_cell(name: str, sub: str = "", *, size: float = 11) -> dict[str, Any]:
    solid, tint, ink = D.class_colours(name)
    paras = [P(R(name, size=size, bold=True, color=ink))]
    if sub:
        paras.append(P(R(sub, size=10, color=C.SLATE)))
    return {"paras": paras, "fill": tint}


def pending_cell(text: str = "Pending frozen run") -> dict[str, Any]:
    return {"paras": [P(R(text, size=10, italic=True, color=C.MUTED))]}


# =========================================================================== slides 1-13


def s01_title(prs: Any, ctx: Context) -> None:
    notes = f"""
[0:15] Title.
Good morning. We are {ctx.team}. Meridian's third-party risk team asked a precise question: from a vendor's
public footprint alone, can we tell whether the vendor genuinely uses AI in the service it provides, rather than
only talking about AI, and what that use means for Meridian's data, operations and risk?
Everything you will see runs offline from a frozen, hashed evidence pack. The findings live in the workbook's
student columns L to V; this deck explains how they were reached.
"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = D.rgb(C.INK)
    D.box(slide, 0.55, 0.5, 0.44, 0.44, None, shape=MSO_SHAPE.OVAL, fill=None, line=C.SEAFOAM, line_w=2.25,
          name="Logo ring")
    D.box(slide, 0.67, 0.62, 0.2, 0.2, None, shape=MSO_SHAPE.OVAL, fill=C.SAFFRON, line=None, name="Logo eye")
    D.label(slide, 1.12, 0.5, 4.0, 0.44, P(R(ctx.team.upper(), size=14, bold=True, color=C.WHITE)),
            anchor=MSO_ANCHOR.MIDDLE, name="Team name")
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(0.55), Inches(1.55), Inches(8.0), Inches(2.0)
    D.write_text(title.text_frame, P(R("Reading the public footprint", size=54, bold=True, color=C.WHITE),
                                     line_spacing=0.9), anchor=MSO_ANCHOR.BOTTOM, margins=(0, 0, 0, 0))
    subtitle = slide.placeholders[1]
    subtitle.left, subtitle.top, subtitle.width, subtitle.height = (Inches(0.55), Inches(3.75), Inches(7.6),
                                                                    Inches(1.0))
    D.write_text(subtitle.text_frame, P(R("How Meridian can tell genuine vendor AI use from AI marketing, using "
                                          "public evidence only, and what that use means for its risk.", size=20,
                                          color=C.SEAFOAM)), margins=(0, 0, 0, 0))
    D.label(slide, 0.55, 4.88, 7.6, 0.35, P(R(f"Case study 1 · Vendor AI usage analysis · {month_year(ctx.as_of)}",
                                              size=14, color="C9D6DB")), name="Date line")
    outcomes = ("01 Functional design", "02 Architecture", "03 Source strategy", "04 Criticality & depth",
                "05 Working demo")
    x = 0.55
    for text in outcomes:
        w = 0.32 + 0.078 * len(text)
        D.pill(slide, x, 5.6, w, 0.36, text, fill=C.INK, color=C.WHITE, line=C.SEAFOAM, size=10,
               name=f"Outcome {text[:2]}")
        x += w + 0.12
    # radar of public signals: rings, a sweep and five blips
    cx, cy = 10.75, 3.35
    for radius, colour in ((2.05, "24525C"), (1.5, "2E6670"), (0.95, "3F8A8E"), (0.42, C.SEAFOAM)):
        D.box(slide, cx - radius, cy - radius, 2 * radius, 2 * radius, None, shape=MSO_SHAPE.OVAL, fill=None,
              line=colour, line_w=1.25, name=f"Radar ring {radius}")
    D.arrow(slide, cx, cy, cx + 2.05 * math.cos(math.radians(60)), cy - 2.05 * math.sin(math.radians(60)),
            color=C.SEAFOAM, width=1.5, tail=None, name="Radar sweep")
    blips = (("SEC filing", 0.95, 140, C.SEAFOAM), ("product page", 1.5, 60, C.SAFFRON),
             ("job posting", 2.05, 250, C.SEAFOAM), ("DNS token", 1.5, 200, C.SEAFOAM),
             ("trade press", 2.05, 100, C.SEAFOAM))
    for text, radius, angle, colour in blips:
        bx = cx + radius * math.cos(math.radians(angle))
        by = cy - radius * math.sin(math.radians(angle))
        D.box(slide, bx - 0.08, by - 0.08, 0.16, 0.16, None, shape=MSO_SHAPE.OVAL, fill=colour, line=None,
              name=f"Blip {text}")
        D.label(slide, bx + 0.14, by - 0.15, 1.3, 0.3, P(R(text, size=10, color=C.WHITE)),
                anchor=MSO_ANCHOR.MIDDLE, name=f"Blip label {text}")
    D.label(slide, 0.55, 6.45, 9.0, 0.3, P(R("Confidential: prepared for Meridian's third-party risk team. Not for "
                                             "publication or onward distribution.", size=10, color="9FB3BA")),
            name="Confidentiality line")
    _notes(slide, notes)


def s02_answer(prs: Any, ctx: Context) -> None:
    data = ctx.data
    widths = (1.6, 1.25, 1.75, 1.85, 2.9, 2.88)
    header = ("Vendor", "Criticality", "AI usage (column O)", "AI risk (column S)", "What would change it",
              "Next action for Meridian")
    rows: list[list[Any]] = []
    note_lines: list[str] = []
    if data:
        kicker = headline(data)
        for v in data.vendors:
            usage = verdict_words(v)
            wording = "; ".join(t for t in (v.likelihood, f"confidence {v.confidence.lower()}" if v.confidence else "")
                                if t)
            if v.provisional and v.ceiling:
                risk_sub = f"Provisional; ceiling {v.ceiling} if confirmed"
            elif v.provisional:
                risk_sub = "Provisional"
            else:
                risk_sub = f"score {v.arp} of 18" if v.arp is not None else ""
            rows.append([
                {"paras": [P(R(v.name, size=11, bold=True, color=C.INK)), P(R(v.vendor_id, size=10, color=C.MUTED))]},
                class_cell(v.tier, f"score {v.score}" if v.score is not None else ""),
                {"paras": [P(R(usage, size=11, bold=True, color=C.INK)), P(R(wording, size=10))]},
                class_cell(v.risk, risk_sub),
                {"paras": [P(R(clip(v.flip, 150) or "No flip condition recorded", size=10))]},
                {"paras": [P(R(clip(v.action, 150) or "See column U", size=10))]},
            ])
            ids = ", ".join(e.evidence_id for e in v.evidence if e.evidence_id) or "no decisive excerpt"
            ceiling = f", provisional with a ceiling of {v.ceiling} if confirmed" if v.provisional and v.ceiling else ""
            note_lines.append(f"- {v.vendor_id} {v.name}: {v.tier} tier; {usage} ({v.likelihood or 'n/a'}, "
                              f"confidence {v.confidence or 'n/a'}); AI risk {v.risk}{ceiling}. Evidence: {ids}. "
                              f"Flip: {v.flip or 'none'} Next: {v.action or 'see column U'}")
        source = (f"Source: run {data.run_id or 'deck summary'} ({data.mode or 'summary'}), evidence as of "
                  f"{data.as_of or ctx.as_of}. Evidence Log IDs for every vendor claim are in the speaker notes; "
                  "the workbook's cells L–V hold the full wording.")
    else:
        kicker = ("Criticality below is final. AI usage, risk and next actions are injected from the frozen run "
                  "(build with --findings).")
        for s in ctx.scored:
            rows.append([
                {"paras": [P(R(s.name, size=11, bold=True, color=C.INK)),
                           P(R(s.profile.vendor_id, size=10, color=C.MUTED))]},
                class_cell(s.result.tier.value, f"score {s.result.score}"),
                pending_cell(), pending_cell(), pending_cell(), pending_cell(),
            ])
            note_lines.append(f"- {s.profile.vendor_id} {s.name}: {s.result.tier.value} tier (score "
                              f"{s.result.score}); usage, risk and action pending the frozen run.")
        source = "Placeholder build: no findings JSON was given, so only the criticality column is filled."
    notes = "[1:00] Answer first.\n" + kicker + "\n" + "\n".join(note_lines) + (
        "\nEach row is the workbook's columns L, O and S plus the flip condition from T and the first action "
        "from U. Every vendor claim above cites Evidence Log IDs; say the ID when you quote evidence.")
    slide = content_slide(prs, ctx, "Answer first", kicker, "SUMMARY", notes)
    if rows:
        table(slide, LEFT, 1.82, widths, header, rows, row_h=0.66, size=10, name="Answer-first table")
    else:
        D.box(slide, LEFT, 2.0, WIDTH, 2.0, P(R("No workbook or findings were available to this build.", size=14,
                                                 color=C.MUTED)), fill=C.SURFACE, line=None, align=PP_ALIGN.CENTER,
              name="Answer-first placeholder")
    D.label(slide, LEFT, 6.58, WIDTH, 0.4, P(R(source, size=10, color=C.MUTED)), name="Source note")


ROLES: tuple[tuple[str, str, tuple[str, ...], str, str], ...] = (
    ("01", "Determine criticality", ("C1 Intake", "C2 Criticality"),
     "A deterministic rubric on profile columns B–K: score 0–100, floors F1–F6, a sensitivity check. "
     "OSINT never changes the tier.", "HC1 tier confirmed"),
    ("02", "Set the depth", ("C3 Depth planner",),
     "The tier and three modifiers set the mandatory source families, request caps, Gemini budget and stop "
     "rules. Every family ends with a Coverage Log status.", ""),
    ("03", "Identify the sources", ("C4 Discovery", "C5 Capture"),
     "Our own discovery, no search-engine scraping: SEC, DNS, job APIs, sitemaps, WordPress REST, Wayback. "
     "Terms, robots and rate limits before every GET.", ""),
    ("04", "Build the capability", ("C6 Extraction", "C7 Signals", "C8 Verdict"),
     "Rules and Gemini propose claims; code verifies exact quotes (V1–V9), assigns transparent tags and "
     "applies verdict rules a–f.", "HC2 evidence reviewed"),
    ("05", "Surface the risk", ("C9 Risk", "C10 Compose"),
     "ARP = 2E + 2K + TP + TG with a materiality gate, escalators and a verdict cap, then playbook actions "
     "and the L–V cells.", ""),
)
"""The brief's five roles. The checkpoint pill marks the two human checkpoints, HC1 and HC2; there is no third:
the release itself is gated by code (the LLM audit gate and the integrity check), not by an approval step."""


def s03_approach(prs: Any, ctx: Context) -> None:
    notes = """
[0:30] Approach.
The brief gives us five roles. We mapped each role onto pipeline components C1 to C10, so every step has an
owner in code and an output in the workbook.
The rule that runs through all of it: AI proposes, code verifies, analyst approves. Gemini only suggests names,
priorities and candidate claims. Code decides what counts: a quote is evidence only if it is an exact slice of a
captured, hashed page. Analysts hold two checkpoints: HC1 confirms the tier, and HC2 accepts or rejects each cited
excerpt with a reason. The Approve action on the Findings page is part of HC2: it accepts a vendor's remaining
cited excerpts in one step, after the analyst confirms reading each one. Export itself is gated by code, not by a
sign-off: the LLM audit must be clean and the integrity check must pass.
The strip at the bottom shows where each required outcome is covered in this deck.
"""
    slide = content_slide(prs, ctx, "Approach: AI proposes, code verifies, analyst approves",
                          "The brief's five roles map onto ten components and two human checkpoints; every "
                          "conclusion traces to a hashed source.", "APPROACH", notes)
    principles = (("AI proposes", "Gemini expands names, triages pages by ID and extracts candidate claims. It "
                                  "never supplies URLs or final labels.", C.TEAL, C.WHITE),
                  ("Code verifies", "Every quote must be an exact slice of a captured, hashed source (V1–V9); rules "
                                    "assign tags, verdicts and risk.", C.INK, C.WHITE),
                  ("Analyst approves", "HC1 confirms the tier; HC2 accepts or rejects each cited excerpt with a "
                                       "reason.", C.SAFFRON, C.INK))
    overlap = 0.1
    w = (WIDTH + 2 * overlap) / 3
    for index, (head, text, colour, ink) in enumerate(principles):
        x = LEFT + index * (w - overlap)
        shape = MSO_SHAPE.PENTAGON if index == 0 else MSO_SHAPE.CHEVRON
        D.box(slide, x, 1.8, w, 0.56, P(R(head, size=15, bold=True, color=ink)), shape=shape, fill=colour,
              line=None, align=PP_ALIGN.CENTER, margins=(0.3, 0, 0.3, 0), name=f"Principle {head}")
        D.label(slide, x + 0.12, 2.44, w - 0.45, 0.62, P(R(text, size=10.5, color=C.SLATE)),
                name=f"Principle text {head}")
    card_w, gap, top, height = 2.3, (WIDTH - 5 * 2.3) / 4, 3.18, 3.0
    for index, (number, role, components, how, checkpoint) in enumerate(ROLES):
        x = LEFT + index * (card_w + gap)
        D.box(slide, x, top, card_w, height, None, fill=C.SURFACE, line=None, radius=0.12, name=f"Role {number}")
        D.badge(slide, x + 0.36, top + 0.36, 0.46, number, fill=C.TEAL, size=12, name=f"Role badge {number}")
        D.label(slide, x + 0.7, top + 0.14, card_w - 0.8, 0.5, P(R(role, size=13, bold=True, color=C.INK)),
                anchor=MSO_ANCHOR.MIDDLE, name=f"Role title {number}")
        y = top + 0.75
        for comp in components:
            D.pill(slide, x + 0.16, y, card_w - 0.32, 0.27, comp, fill=C.TEAL_TINT, color=C.TEAL_DARK, size=10,
                   name=f"Role {number} component {comp}")
            y += 0.33
        D.label(slide, x + 0.16, top + 1.82, card_w - 0.3, 0.86, P(R(how, size=10, color=C.SLATE)),
                name=f"Role {number} how")
        if checkpoint:
            D.pill(slide, x + 0.16, top + height - 0.4, card_w - 0.32, 0.28, checkpoint, fill=C.SAFFRON,
                   color=C.INK, size=10, name=f"Role {number} checkpoint")
    D.label(slide, LEFT, 6.4, WIDTH, 0.4, P(
        R("Where each outcome is covered:  ", size=11, bold=True, color=C.INK),
        R("O1 functional design, slide 4 · O2 architecture, 5 · O3 sources and weight, 6–8 · "
          "O4 criticality and depth, 9–10 · O5 working demo, 12", size=11, color=C.SLATE)),
        name="Outcome map")


def s04_functional(prs: Any, ctx: Context) -> None:
    notes = """
[1:00] Outcome 01: functional design.
Read the top row left to right. C1 takes the uploaded workbook, tolerates header typos and maps the cream student
cells; it skips the fictional V-000 row. C2 rates criticality from the profile alone, and HC1 is the analyst
confirming that tier; an override needs a written reason. C3 turns the tier into a depth plan: which source
families are mandatory, their request caps, the Gemini budget and the stop rules.
C4 discovers candidate sources with our own OSINT: seeds, DNS, sitemaps, SEC, job-board APIs and Wayback. Gemini
may expand names and triage pages, by ID only. C5 captures compliantly: terms register, then robots.txt, then a
per-host rate limit, then a GET. Raw bytes land in a content-addressed store with SHA-256 and UTC time.
The bottom row reads right to left. C6 extracts dated text and lexicon passages. C7 runs the rules tagger and
Gemini claim extraction side by side; every claim must pass gates V1 to V9 before it becomes an excerpt with tags.
The stop-rule loop sends work back to discovery until usage, sub-processors and data use are answered, the search
saturates or the budget runs out. HC2 is the analyst reviewing each cited excerpt. C8 applies verdict rules a to f,
C9 scores risk and picks actions, and C10 composes the cells and logs.
C10 always freezes evidence first and then renders the cells. The export refuses a run whose LLM audit log has
problems, proves with the integrity check that only the student cells changed, and an offline replay then proves
it reproduces the same cells. Those release gates are code, not a third sign-off: the human checkpoints are HC1
and HC2.
"""
    slide = content_slide(prs, ctx, "Functional design: from OSINT collection to risk output",
                          "Ten components run once per vendor. Gemini sits behind a payload guard; analysts gate the "
                          "tier (HC1) and the evidence (HC2).", "OUTCOME 01", notes)
    D.functional_flow(slide, top=1.82)


def s05_architecture(prs: Any, ctx: Context) -> None:
    notes = """
[1:00] Outcome 02: architecture.
One engine sits behind three interfaces: the typer CLI, the Streamlit app and the Colab or Jupyter notebook. All
three call a single entry point, run_assessment, with a mode switch. Replay is the default: it reads the frozen
evidence pack and the LLM cache and makes no network calls, so the demo cannot be broken by Wi-Fi or quota.
live_rules collects live without Gemini; live_ai adds Gemini through the cache.
Inside the engine, net policy is the only way out to the internet: the Live or Replay fetcher, the terms-of-use
register, robots.txt via Protego with an honest user agent, and per-host rate limits. Collectors for DNS, sites,
WordPress, SEC, three ATS APIs, Wayback, providers and seeds sit on top of it, and manual captures enter through
an import command. Analysis extracts text, runs the rules tagger and the Gemini layer, verifies with V1 to V9 and
clusters corroboration. Decisions are made by the verdict, risk and action modules, and compose writes the cells.
The stores are versioned config and prompts, the evidence store, the review records and the run manifests.
The dashed line is the trust boundary. Profile fields, tiers, verdicts and findings never cross it. The only thing
that leaves is public vendor text that has passed the payload guard, and the audit log records each call without
its text. The demo is allowlisted for one live DNS lookup on top of replay.
"""
    slide = content_slide(prs, ctx, "Architecture: layers, stores and the trust boundary",
                          "One engine behind three interfaces; every request passes the terms, robots and rate-limit "
                          "gate, and only guarded public text crosses to Gemini.", "OUTCOME 02", notes)
    D.architecture(slide, top=1.8)


MODE_WORDS = {"full": "Full", "privacy_and_subprocessors": "Privacy", "screen": "Screen", "one_query": "1 query",
              "keyword": "Keyword", "homepage": "Homepage", "on_lead": "On a lead", "manual": "Manual",
              "not_used": "—", "not_applicable": "n/a"}
FAMILY_NAMES = {"LEG": "Legal & trust", "REG": "SEC registrants", "PRD": "Product & news", "JOB": "Own job postings",
                "DNS": "DNS records", "IND": "Corroboration", "HIST": "Archive history", "EXEC": "Executives"}
TIER_COLUMNS = (Tier.CRITICAL, Tier.HIGH, Tier.MEDIUM, Tier.LOW)
REASON_LIMIT = 104


def thinning_reason(cfg: DepthConfig, family: SourceFamily) -> tuple[Tier | None, str]:
    """The first tier below Critical where the family's mode changes, and that tier's reason from depth.toml; None
    and the Low tier's reason when the mode is the same at every tier."""
    critical = cfg.tiers[Tier.CRITICAL].families[family]
    for tier in (Tier.HIGH, Tier.MEDIUM, Tier.LOW):
        policy = cfg.tiers[tier].families[family]
        if policy.mode != critical.mode:
            return tier, policy.reason
    return None, cfg.tiers[Tier.LOW].families[family].reason


def short_reason(reason: str, limit: int = REASON_LIMIT) -> str:
    """The 'why' part of a depth.toml reason (after its colon), cut at a clause boundary to fit two table lines."""
    text = reason.split(": ", 1)[1] if ": " in reason else reason
    text = text[:1].upper() + text[1:]
    if len(text) <= limit:
        return text
    head = text.split("; ", 1)[0]
    if 40 <= len(head) <= limit:
        return head.rstrip(".") + "."
    if ", " in text[:limit]:
        cut = text[:limit].rsplit(", ", 1)[0]
        if len(cut) >= 40:
            return cut.rstrip(".") + "."
    return clip(text, limit)


def s06_sources(prs: Any, ctx: Context) -> None:
    cfg = ctx.depth
    rows = []
    reasons = []
    for family in SourceFamily:
        cells: list[Any] = [{"paras": [P(R(family.value, size=11, bold=True, color=C.INK)),
                                       P(R(FAMILY_NAMES.get(family.value, family.value), size=10))]}]
        for tier in TIER_COLUMNS:
            policy = cfg.tiers[tier].families[family]
            mode = MODE_WORDS.get(policy.mode, policy.mode.replace("_", " "))
            paras = [P(R(mode, size=10, bold=policy.mandatory, color=C.INK if policy.mandatory else C.MUTED))]
            if policy.cap:
                paras.append(P(R(f"≤ {policy.cap} req.", size=10, color=C.SLATE if policy.mandatory else C.MUTED)))
            cells.append({"paras": paras, "fill": C.TEAL_TINT if policy.mandatory else C.WHITE,
                          "align": PP_ALIGN.CENTER, "pad": 0.03})
        tier, reason = thinning_reason(cfg, family)
        cells.append({"paras": [P(R(short_reason(reason), size=10))]})
        reasons.append(f"- {family.value} ({f'{tier.value} and below' if tier else 'same mode at every tier'}): "
                       f"{reason}")
        rows.append(cells)
    excluded = [(x.source, x.reason) for x in cfg.excluded_sources]
    counts = ctx.terms.get("counts", {})
    manual = ctx.terms.get("manual", [])
    notes = "\n".join([
        "[1:00] Outcome 03: source strategy.",
        "Read the table by row. Shaded cells are mandatory: the family must end with a Coverage Log status, and that "
        "status is our negative evidence. The number is the cap on automated requests per vendor. REG applies only "
        "to SEC registrants; for the others the SEC collector logs it not applicable, with proof of the EDGAR check.",
        "Why each family thins below Critical (config/depth.toml):", *reasons,
        f"Source reliability comes only from the source type ({ctx.sr_source}), never from the model: "
        + "; ".join(f"{g} {n.lower()}: {e}" for g, n, e in ctx.sr_ladder) + ".",
        f"Terms register (config/tou.toml): {counts.get('full', 0)} hosts allow automated GETs within limits, "
        f"{counts.get('limited', 0)} allow seeded URLs only, {counts.get('none', 0)} are manual only "
        f"({', '.join(manual) or 'none'}).",
        "Excluded by design, with reasons:", *[f"- {s}: {r}" for s, r in excluded],
    ])
    slide = content_slide(prs, ctx, "Source strategy: what each tier reads, and why it thins",
                          "Shaded families are mandatory and must end with a Coverage Log status; reliability (SR) "
                          "comes from the source type, never from the model.", "OUTCOME 03", notes)
    table_w = 7.9
    widths = (1.3, 0.72, 0.72, 0.72, 0.72, table_w - 1.3 - 4 * 0.72)
    table(slide, LEFT, 1.8, widths, ("Family", "Critical", "High", "Medium", "Low", "Why it thins below Critical"),
          rows, row_h=0.42, size=10, name="Family by tier table")
    # SR ladder, most reliable on the left
    ly = 5.6
    D.label(slide, LEFT, ly, table_w, 0.28, P(R("Source reliability (SR) ladder  ", size=12, bold=True, color=C.INK),
                                              R("set by the source type, never by the model", size=10,
                                                color=C.MUTED)), name="SR header")
    fills = (("0A5E5D", C.WHITE), ("0E7C7B", C.WHITE), ("B7DCD9", C.INK), ("DCEDEB", C.INK), ("EDF1F4", C.SLATE))
    rung_w = (table_w - 4 * 0.08) / 5
    for index, (grade, name, examples) in enumerate(ctx.sr_ladder[:5]):
        fill, ink = fills[min(index, 4)]
        D.box(slide, LEFT + index * (rung_w + 0.08), ly + 0.32, rung_w, 0.92, [
            P(R(grade, size=14, bold=True, color=ink), R("  " + name, size=10, bold=True, color=ink)),
            P(R(examples, size=10, color=ink), space_before=2),
        ], fill=fill, line=None, radius=0.08, anchor=MSO_ANCHOR.TOP, margins=(0.08, 0.05, 0.06, 0.03),
            name=f"SR {grade}")
    # exclusions, with their full reasons
    rx, rw = 8.75, LEFT + WIDTH - 8.75
    D.label(slide, rx, 1.8, rw, 0.3, P(R("Excluded by design", size=12, bold=True, color=C.INK)),
            name="Exclusions header")
    paras: list[Any] = []
    for source, reason in excluded:
        paras.append(P(R(source, size=10, bold=True, color=C.INK)))
        paras.append(P(R(reason, size=10), space_after=4))
    D.label(slide, rx, 2.14, rw, 3.6, paras, name="Exclusions list")
    # terms register by host class
    D.label(slide, rx, 5.6, rw, 0.28, P(R("Terms register by host class", size=12, bold=True, color=C.INK)),
            name="Terms header")
    chip_w = (rw - 2 * 0.1) / 3
    for index, (key, text) in enumerate((("full", "automated, within limits"), ("limited", "seeded URLs only"),
                                         ("none", "manual capture only"))):
        D.box(slide, rx + index * (chip_w + 0.1), 5.92, chip_w, 0.92, [
            P(R(str(counts.get(key, 0)), size=18, bold=True, color=C.TEAL_DARK if key != "none" else "8A6A00")),
            P(R(text, size=10)),
        ], fill=C.SURFACE if key != "none" else C.SAFFRON_TINT, line=None, radius=0.08, align=PP_ALIGN.CENTER,
            margins=(0.05, 0.04, 0.05, 0.04), name=f"Terms {key}")


def s07_weight(prs: Any, ctx: Context) -> None:
    trap = next((v.trap for v in ctx.vendors if v.trap and v.vendor_id == "V-003"), None) or next(
        (v.trap for v in ctx.vendors if v.trap), None)
    trap_vendor = next((v for v in ctx.vendors if v.trap is not None and v.trap == trap), None)
    strong_vendor = ctx.chain
    cited = list(strong_vendor.evidence) if strong_vendor else []
    # With no Strong item in the run (every Yes is Probable), show the strongest cited item instead of the
    # scouting placeholder, so the card never claims a strength or verdict the run does not have.
    strong = next((e for e in cited if e.strength == "Strong"), None) or next(
        (e for e in cited if e.strength == "Moderate"), None)
    notes = ["[1:00] Outcome 03: evidentiary weight.",
             "Every verified excerpt carries a tag card. U is the signal class, from U1 (AI in the exact service) to "
             "U8 (a limiting statement). SR is source reliability, set by the source type. SP is specificity from "
             "genuine-use indicators G1 to G11 against marketing indicators M1 to M7. RL is relevance to the exact "
             "service, computed locally. RC is recency against the as-of date. IC is corroboration.",
             "Four tests run in order: the EU AI Act definition, the subject, the use-state, then specificity.",
             "The strength label is the first rule that matches, in a fixed order. A DNS verification token is "
             "therefore always 'Context - relationship only', never evidence of use."]
    if trap and trap_vendor:
        notes.append(f"Trap shown: {trap_vendor.vendor_id} {trap_vendor.name}, {trap.evidence_id or 'logged item'}, "
                     f"{trap.source_type}: \"{trap.excerpt}\"")
    if strong and strong_vendor:
        notes.append(f"{strong.strength} item shown:{strong_vendor.vendor_id} {strong_vendor.name}, {strong.evidence_id}, "
                     f"{strong.source_type} ({strong.tags}): \"{strong.excerpt}\"")
    if not (trap and strong):
        notes.append("Placeholder examples come from scouting; the build with --findings swaps in the verified "
                     "excerpts and their Evidence Log IDs.")
    slide = content_slide(prs, ctx, "Evidentiary weight: tags, four tests, one label order",
                          "Each excerpt carries a transparent tag card; four ordered tests separate genuine use from "
                          "marketing, and the first matching strength label wins.", "OUTCOME 03", "\n".join(notes))
    # tag card
    D.box(slide, LEFT, 1.8, 5.95, 2.08, None, fill=C.SURFACE, line=None, radius=0.12, name="Tag card")
    D.label(slide, LEFT + 0.18, 1.9, 5.6, 0.3, P(R("Tag card on every excerpt", size=12, bold=True, color=C.INK)),
            name="Tag card title")
    sample = strong.tags if strong else "U1 · SR:B · SP:S2 · RL:R3 · RC:T3 · IC:1 · locus=service_feature"
    D.box(slide, LEFT + 0.18, 2.24, 5.6, 0.34, P(R(sample, size=10.5, bold=True, color=C.TEAL_DARK, font=C.MONO)),
          fill=C.WHITE, line=C.LINE, radius=0.08, align=PP_ALIGN.CENTER, margins=(0.04, 0, 0.04, 0),
          name="Tag string")
    tiles = (("U", "signal class: U1 exact service … U8 limiting"), ("SR", "source reliability A–F, by type"),
             ("SP", "specificity S0–S3 from G/M indicators"), ("RL", "relevance R0–R3, computed locally"),
             ("RC", "recency T0–T3 against the as-of date"), ("IC", "corroboration 1 (independent) … 6"))
    for index, (code, text) in enumerate(tiles):
        col, row = index % 3, index // 3
        x, y = LEFT + 0.18 + col * 1.88, 2.7 + row * 0.56
        D.box(slide, x, y, 1.8, 0.5, P(R(code + "  ", size=11, bold=True, color=C.TEAL_DARK), R(text, size=10)),
              fill=C.WHITE, line=None, radius=0.06, margins=(0.07, 0.02, 0.05, 0.02), name=f"Tag tile {code}")
    # four tests
    tx, tw = 6.8, LEFT + WIDTH - 6.8
    D.label(slide, tx, 1.8, tw, 0.3, P(R("Genuine use or marketing? Four tests, in order", size=12, bold=True,
                                         color=C.INK)), name="Tests title")
    tests = (("Definition", "EU AI Act Art. 3(1): scheduling, RPA, rules engines, templated composition, IMb and BI "
                            "are not AI unless ML or an LLM is evidenced"),
             ("Subject", "the sentence names the vendor, an alias or a product; first-party pages imply the product"),
             ("Use-state", "deployed cues ('live', 'uses', 'reviewed by AI', a percentage) beat 'will', 'plan', "
                           "'exploring'"),
             ("Specificity", "G1–G11 genuine-use indicators against M1–M7 marketing indicators give S3, S2, S1 "
                             "or S0"))
    for index, (name, text) in enumerate(tests):
        y = 2.16 + index * 0.43
        D.badge(slide, tx + 0.18, y + 0.19, 0.34, str(index + 1), fill=C.TEAL, size=11, name=f"Test {index + 1}")
        D.label(slide, tx + 0.45, y, tw - 0.45, 0.42, P(R(name + ": ", size=10.5, bold=True, color=C.INK),
                                                        R(text, size=10)), anchor=MSO_ANCHOR.MIDDLE,
                name=f"Test text {index + 1}")
    # strength label order
    D.label(slide, LEFT, 4.06, 5.95, 0.3, P(R("Strength label: the first rule that matches wins", size=12,
                                              bold=True, color=C.INK)), name="Labels title")
    labels = (("Negative", "U8: a limiting or negative statement", "861818", "FBE4E4"),
              ("Strong", "a qualifying signal (Q) at R3, the exact service", C.WHITE, C.TEAL_DARK),
              ("Moderate", "a Q at R2, or an independent corroborating signal (K)", C.WHITE, C.TEAL),
              ("Context", "DNS token, platform supplier or inferred affiliate: not scored", C.TEAL_DARK, C.TEAL_TINT),
              ("Marketing only", "U7, or S0 (aspiration and commentary)", "6B5000", "FFF4D1"),
              ("Weak", "anything else that verified (S1 or R1)", C.SLATE, C.SURFACE))
    for index, (name, rule, ink, fill) in enumerate(labels):
        y = 4.42 + index * 0.37
        D.label(slide, LEFT, y, 0.3, 0.32, P(R(str(index + 1), size=11, bold=True, color=C.MUTED)),
                anchor=MSO_ANCHOR.MIDDLE, name=f"Label rank {index + 1}")
        D.box(slide, LEFT + 0.3, y, 1.55, 0.32, P(R(name, size=10.5, bold=True, color=ink)), fill=fill, line=None,
              radius=0.16, align=PP_ALIGN.CENTER, margins=(0.03, 0, 0.03, 0), name=f"Label {name}")
        D.label(slide, LEFT + 1.98, y, 4.0, 0.32, P(R(rule, size=10)), anchor=MSO_ANCHOR.MIDDLE,
                name=f"Label rule {name}")
    D.label(slide, LEFT, 6.65, 6.0, 0.3, P(R("Q = SR A/B · SP ≥ S2 · RL ≥ R2 · RC ≥ T1 · U1–U4 · a delivery "
                                             "locus",
                                             size=10, color=C.MUTED)), name="Q definition")
    # trap vs strong
    cx, cw = 6.8, LEFT + WIDTH - 6.8
    if trap and trap_vendor:
        trap_text = [P(R(f"{trap_vendor.name} · {trap.source_type} · {trap.evidence_id or 'Evidence Log'}", size=10,
                         bold=True, color=C.INK)), P(R(quote(trap.excerpt, 150), size=10, italic=True))]
        verdict_line = ("Definition test fails: no ML or LLM is evidenced → kept in the Evidence Log as a trap, "
                        "never cited.")
    else:
        trap_text = [P(R("FSSI · scouting example", size=10, bold=True, color=C.INK)),
                     P(R("'Intelligent inserting', Intelligent Mail barcodes and an 'intelligent tools' press release "
                         "sound like AI.", size=10, italic=True))]
        verdict_line = ("Definition test fails: print-and-mail automation, no ML or LLM evidenced → kept in the "
                        "Evidence Log as a trap, never cited.")
    trap_text.append(P(R(verdict_line, size=10, color=C.SLATE), space_before=3))
    card(slide, cx, 4.06, cw, 1.3, "Trap · rejected", trap_text, fill="FDF0F0", title_color="861818",
         title_size=11, name="Trap card")
    if strong and strong_vendor:
        strong_text = [P(R(f"{strong_vendor.name} · {strong.source_type} · {strong.evidence_id}", size=10, bold=True,
                           color=C.INK)), P(R(quote(strong.excerpt, 150), size=10, italic=True))]
        tags_line = f"{strong.tags} → {strong.strength}; {strong.role or 'Primary'} source for " \
                    f"{strong_vendor.usage} ({strong_vendor.verdict})."
    else:
        strong_text = [P(R("BNY · scouting example", size=10, bold=True, color=C.INK)),
                       P(R("The instant-payments product page ties AI-enabled anomaly detection to its RTP network "
                           "connection.", size=10, italic=True))]
        tags_line = "U1 · SR B · RL R3 → Strong: the Primary source for a Yes (Confirmed) verdict."
    strong_text.append(P(R(tags_line, size=10, color=C.SLATE), space_before=3))
    card(slide, cx, 5.5, cw, 1.3, f"{strong.strength if strong else 'Strong'} · cited", strong_text, fill=C.TEAL_TINT, title_color=C.TEAL_DARK,
         title_size=11, name="Strong card")


GATES: tuple[tuple[str, str, str], ...] = (
    ("V1", "passage ID exists", "blocking"), ("V2", "quote is an exact slice of the source", "blocking"),
    ("V3", "no ellipsis; 25–600 characters", "blocking"), ("V4", "indicator spans lie in the quote", "repair"),
    ("V5", "providers named verbatim", "repair"), ("V6", "entity guard passes", "blocking"),
    ("V7", "relevance and locus set locally", "local"), ("V8", "rules decide the final labels", "local"),
    ("V9", "duplicates removed", "blocking"),
)


def s08_ai(prs: Any, ctx: Context) -> None:
    data = ctx.data
    if data:
        agree = sum(v.llm_agree for v in data.vendors)
        disagree = sum(v.llm_disagree for v in data.vendors)
        proposed = sum(v.llm_proposed for v in data.vendors)
        calls = plural(data.llm_calls, "Gemini call", "Gemini calls") + " · " if data.llm_calls is not None else ""
        withheld = (plural(data.withheld, "passage", "passages") + " withheld by the guard · "
                    if data.withheld is not None else "")
        stats = [R("Gemini vs rules (frozen run): ", size=11, bold=True, color=C.INK),
                 R(f"{calls}{withheld}{plural(agree, 'item agrees', 'items agree')} · "
                   f"{disagree} {'disagrees and goes' if disagree == 1 else 'disagree and go'} to the HC2 queue · "
                   f"{plural(proposed, 'Gemini-only proposal waits', 'Gemini-only proposals wait')} for an analyst · "
                   "0 URLs taken from the model", size=11)]
    else:
        stats = [R("Gemini vs rules: ", size=11, bold=True, color=C.INK),
                 R("agreement, disagreement and proposal counts are injected from the frozen run. URLs taken from the "
                   "model: zero by construction (the response schema has no URL field).", size=11)]
    notes = """
[0:45] AI where it adds recall, code where it adds trust.
Gemini does three jobs. Expand turns a vendor's public name and navigation text into names to search for. Triage
ranks pages, by ID only, so it never invents a URL. Extract labels public passages as structured JSON claims. If
Gemini is unavailable, the keyless NullLLM path does the same jobs with rules, and the rest of the pipeline is
unchanged.
Before any call, the payload guard withholds passages that contain Meridian's name, our team name or distinctive
profile wording; it redacts emails, phone numbers and listed names, and it skips hosts whose terms or robots.txt
refuse AI processing. Withheld passages still go to the rules.
Every claim then passes nine gates. Five reject it: unknown passage, a quote that is not an exact slice of the
captured text, an inserted ellipsis or a bad length, a failed entity guard, or a duplicate. Two repair it by
dropping an unverified indicator or provider name. Two keep decisions in code: relevance and locus are computed
locally, and the rules decide the final labels while Gemini's labels are shown beside them for review.
""" + ("".join(text for text, _ in stats))
    slide = content_slide(prs, ctx, "AI where it adds recall, code where it adds trust",
                          "Gemini proposes names, priorities and claims; deterministic code decides what is evidence, "
                          "and nothing Meridian-internal reaches the model.", "OUTCOME 03", notes)
    rows = [
        ["Expand", "public vendor name, homepage title and meta, navigation text",
         "names only (aliases, products, providers, affiliates) as search keys, never evidence", "≈ 6"],
        ["Triage", "rows of id · path · title · date · family, plus job titles",
         "IDs with a priority and a reason code; never URLs", "≈ 20"],
        ["Extract", "passages as [P#] (kind, date) text, 12 per call",
         "ClaimBatch JSON: verbatim quote, kind, subject, AI type, use-state, action level, indicators",
         "≈ 70–100"],
        ["Keyless", "seeds, page n-grams, URL slug and title rules",
         "NullLLM: rules-only extraction through the same verification and tags", "0"],
    ]
    cells = [[{"paras": [P(R(r[0], size=10.5, bold=True, color=C.INK))]}, r[1], r[2],
              {"paras": [P(R(r[3], size=10))], "align": PP_ALIGN.CENTER}] for r in rows]
    table(slide, LEFT, 1.8, (1.0, 2.05, 2.35, 0.75), ("Role", "Input (public only)", "Output", "Calls"), cells,
          row_h=0.56, size=10, name="AI roles table")
    D.label(slide, LEFT, 4.42, 6.15, 0.42, P(R("gemini-3.5-flash-lite (fallback 3.1-flash-lite) · JSON schema output · "
                                               "seed 1234 · minimal thinking · no temperature · tools off · cached",
                                               size=10, color=C.MUTED)), name="Model settings")
    gx, guard_y, guard_h = 7.05, 1.8, 1.8
    side_y = guard_y + guard_h / 2 - 0.475
    passages = D.box(slide, gx, side_y, 1.35, 0.95, [P(R("Public passages", size=10.5, bold=True, color=C.INK)),
                                                     P(R("vendor and provider pages", size=10))], fill=C.SURFACE,
                     line=C.LINE, align=PP_ALIGN.CENTER, name="Public passages")
    guard = D.box(slide, gx + 1.65, guard_y, 2.45, guard_h, [
        P(R("Payload guard", size=11, bold=True, color=C.INK), space_after=3),
        P(R("blocks 'Meridian' and the team name", size=10), bullet=True, indent=0.12),
        P(R("blocks profile wording (columns D, E, H, J, K as 6-word shingles)", size=10), bullet=True, indent=0.12),
        P(R("redacts emails, phones and listed names", size=10), bullet=True, indent=0.12),
        P(R("skips hosts that refuse AI processing (terms, Google-Extended, Content-Signal)", size=10), bullet=True,
          indent=0.12),
    ], fill=C.SAFFRON_TINT, line=C.SAFFRON, line_w=1.5, anchor=MSO_ANCHOR.TOP, margins=(0.1, 0.08, 0.08, 0.05),
        name="Payload guard")
    gemini = D.box(slide, gx + 4.38, side_y, 1.35, 0.95, [P(R("Gemini", size=10.5, bold=True, color=C.INK)),
                                                          P(R("claims as JSON, unverified", size=10))],
                   fill=C.WHITE, line=C.TEAL, line_w=1.25, dash=D.DASH, align=PP_ALIGN.CENTER, name="Gemini")
    withheld = D.box(slide, gx + 1.65, guard_y + guard_h + 0.22, 2.45, 0.55,
                     P(R("Withheld → rules only; audit log without text", size=10, color=C.INK)), fill=C.SURFACE,
                     line=C.LINE, align=PP_ALIGN.CENTER, name="Withheld")
    D.connect(slide, passages, guard, D.RIGHT, D.LEFT, name="passages to guard")
    D.connect(slide, guard, gemini, D.RIGHT, D.LEFT, name="guard to Gemini")
    D.connect(slide, guard, withheld, D.BOTTOM, D.TOP, name="guard to withheld")
    D.label(slide, LEFT, 4.92, WIDTH, 0.3, P(R("Every Gemini claim must pass nine gates before it can become evidence",
                                               size=12, bold=True, color=C.INK)), name="Gates title")
    chip_w = (WIDTH - 8 * 0.1) / 9
    for index, (code, text, kind) in enumerate(GATES):
        fill, ink = C.GATE[kind]
        D.box(slide, LEFT + index * (chip_w + 0.1), 5.27, chip_w, 0.78,
              [P(R(code, size=11, bold=True, color=ink)), P(R(text, size=10, color=C.SLATE))], fill=fill, line=None,
              radius=0.08, align=PP_ALIGN.CENTER, margins=(0.05, 0.03, 0.05, 0.03), name=f"Gate {code}")
    legend = (("blocking", "rejects the claim"), ("repair", "drops the unverified span or name"),
              ("local", "decided by code, not the model"))
    x = LEFT
    for kind, text in legend:
        fill, ink = C.GATE[kind]
        D.box(slide, x, 6.16, 0.22, 0.18, None, fill=fill, line=ink, line_w=0.75, radius=0.03,
              name=f"Gate legend {kind}")
        D.label(slide, x + 0.3, 6.1, 3.2, 0.3, P(R(text, size=10)), anchor=MSO_ANCHOR.MIDDLE,
                name=f"Gate legend text {kind}")
        x += 3.2
    D.label(slide, LEFT, 6.45, WIDTH, 0.42, P(*stats), name="Comparison numbers")


def _floor_text(floor: Any) -> str:
    parts = []
    for cond in floor.conditions:
        if cond.anchor:
            parts.append(cond.anchor)
        elif cond.min == 4:
            parts.append(f"{cond.factor} = 4")
        else:
            parts.append(f"{cond.factor} ≥ {cond.min}")
    target = floor.tier.value if floor.tier == Tier.CRITICAL else f"at least {floor.tier.value}"
    return f"{' and '.join(parts)} → {target}"


def s09_criticality(prs: Any, ctx: Context) -> None:
    rubric = ctx.rubric
    scored = ([ctx.example] if ctx.example else []) + ctx.scored
    lifted = [s for s in ctx.scored if s.result.computed_tier.rank > s.result.score_tier.rank]
    stable = sum(1 for s in scored if s.result.perturbation_stable)
    counts = defaultdict(int)
    for s in ctx.scored:
        counts[s.result.tier.value] += 1
    if ctx.scored:
        tiers = ", ".join(f"{counts[t.value]} {t.value}" for t in (Tier.CRITICAL, Tier.HIGH, Tier.MEDIUM, Tier.LOW)
                          if counts[t.value])
        lift = "; ".join(f"{s.name}'s score of {s.result.score} is lifted to {s.result.tier.value} by the "
                         f"{rubric.floor(_lifting_floor(s, rubric)).name}" for s in lifted)
        kicker = f"{tiers}. " + (lift + "; " if lift else "") + \
            f"{stable} of {len(scored)} tiers hold under every ±1 weight change."
    else:
        kicker = "The rubric scores profile fields only; the tier table appears when the input workbook is present."
    notes = ["[1:15] Outcome 04: criticality.",
             "Criticality comes first and from the profile alone, columns B to K, so OSINT can never change a tier. "
             "The rubric is anchored on OCC Bulletin 2023-17: an activity is critical when the third party's failure "
             "would cause significant risk, customer impact or operational impact.",
             f"Score = {' + '.join(f'{rubric.weights[f]}{f}' for f in ('O', 'D', 'P', 'R', 'V'))}, out of 100. "
             f"Critical from {rubric.score_tiers[Tier.CRITICAL]}, High from {rubric.score_tiers[Tier.HIGH]}, "
             f"Medium from {rubric.score_tiers[Tier.MEDIUM]}. Floors can only lift a tier.",
             "Per vendor (computed live by footprint.criticality, rubric v" + rubric.version + "):"]
    for s in scored:
        r = s.result
        notes.append(f"- {s.profile.vendor_id} {s.name}: " + " ".join(
            f"{c}{r.factors[c].level}" + ("P" if r.factors[c].anchor == "D3-P" else "")
            for c in ('O', 'D', 'P', 'R', 'V')) + f", score {r.score}, floors "
            f"{', '.join(r.floors_fired) or 'none'} → {r.tier.value}. Sensitivity: {' '.join(r.sensitivity)}")
    notes.append("V-000 is the fictional worked example: the rubric reproduces its Critical rating, which is our "
                 "calibration check. Any analyst override at HC1 needs a reason, and any rule change bumps the "
                 "version.")
    slide = content_slide(prs, ctx, "Criticality: a profile-only rubric with floors", kicker, "OUTCOME 04",
                          "\n".join(notes))
    w = rubric.weights
    card_h = 1.12
    card(slide, LEFT, 1.8, 4.05, card_h, "Score (0–100)", [
        P(R(f"{w['O']}·O + {w['D']}·D + {w['P']}·P + {w['R']}·R + {w['V']}·V", size=15, bold=True, color=C.TEAL_DARK),
          space_after=3),
        P(R("O dependency · D data · P payment flow · R regulation · V volume", size=10)),
        P(R(f"Critical ≥ {rubric.score_tiers[Tier.CRITICAL]} · High ≥ {rubric.score_tiers[Tier.HIGH]} · "
            f"Medium ≥ {rubric.score_tiers[Tier.MEDIUM]} · Low below", size=10, bold=True, color=C.INK)),
    ], name="Formula card")
    floors = list(rubric.floors)
    half = math.ceil(len(floors) / 2)
    D.box(slide, 4.8, 1.8, 5.2, card_h, None, fill=C.SURFACE, line=None, radius=0.1, name="Floors card")
    D.label(slide, 4.94, 1.88, 5.0, 0.28, P(R("Floors only ever lift a tier", size=12, bold=True, color=C.INK)),
            name="Floors title")
    for column, chunk in enumerate((floors[:half], floors[half:])):
        D.label(slide, 4.94 + column * 2.55, 2.2, 2.5, 0.7, [
            P(R(f"{f.id} ", size=10, bold=True, color=C.TEAL_DARK), R(_floor_text(f), size=10), space_after=1)
            for f in chunk], name=f"Floors column {column + 1}")
    card(slide, 10.2, 1.8, LEFT + WIDTH - 10.2, card_h, "Basis", [
        P(R("OCC Bulletin 2023-17 anchors", size=10)),
        P(R("Profile columns B–K only; OSINT never moves a tier", size=10)),
        P(R("Overrides need a reason (HC1)", size=10)),
    ], name="Basis card")
    if not scored:
        D.box(slide, LEFT, 3.25, WIDTH, 1.5, P(R("Input workbook not found: the live tier table appears when "
                                                 "data/input/Meridian_Vendor_Input.xlsx is present.", size=12,
                                                 color=C.MUTED)), fill=C.SURFACE, line=None, align=PP_ALIGN.CENTER,
              name="Tier table placeholder")
        return
    rows = []
    heights = []
    for s in scored:
        r = s.result
        example = s.profile.is_example
        levels = []
        for c in ("O", "D", "P", "R", "V"):
            text = f"{r.factors[c].level}P" if r.factors[c].anchor == "D3-P" else str(r.factors[c].level)
            levels.append({"paras": [P(R(text, size=10.5, bold=True, color=C.INK))], "align": PP_ALIGN.CENTER})
        sensitivity = " ".join(r.sensitivity)
        name = "V-000 worked example" if example else s.name
        rows.append([
            {"paras": [P(R(name, size=10.5, bold=True, color=C.INK)),
                       P(R("calibration, fictional" if example else s.profile.vendor_id, size=10, color=C.MUTED))],
             "fill": C.SURFACE if example else C.WHITE},
            *levels,
            {"paras": [P(R(str(r.score), size=10.5, bold=True, color=C.INK))], "align": PP_ALIGN.CENTER},
            {"paras": [P(R(", ".join(r.floors_fired) or "—", size=10))]},
            class_cell(r.tier.value, size=10.5),
            {"paras": [P(R(sensitivity, size=10))]},
        ])
        heights.append(0.56 if len(sensitivity) > 165 else 0.42)
    table(slide, LEFT, 3.06, (1.75, 0.4, 0.45, 0.4, 0.4, 0.4, 0.62, 1.15, 0.95, 5.71),
          ("Vendor", "O", "D", "P", "R", "V", "Score", "Floors", "Tier", "Sensitivity (rendered by the engine)"),
          rows, row_h=heights, size=10, name="Tier table")


def _lifting_floor(s: Scored, rubric: Rubric) -> str:
    """The floor that sets the final tier above the score tier (the first one at that tier)."""
    for floor_id in s.result.floors_fired:
        if rubric.floor(floor_id).tier == s.result.computed_tier:
            return floor_id
    return s.result.floors_fired[0]


def s10_depth(prs: Any, ctx: Context) -> None:
    cfg = ctx.depth
    by_tier: dict[Tier, list[Scored]] = defaultdict(list)
    for s in ctx.scored:
        by_tier[s.result.tier].append(s)
    notes = ["[0:45] Outcome 04: depth.",
             "The tier decides how far we look, never the volume of material a vendor publishes. Mandatory families "
             "run to completion or to their caps and always end with a Coverage Log status; the discretionary budget "
             "only covers extra work such as snowball rounds.",
             "Per tier (config/depth.toml):"]
    for tier in (Tier.CRITICAL, Tier.HIGH, Tier.MEDIUM, Tier.LOW):
        policy = cfg.tiers[tier]
        mandatory = [f.value for f, p in policy.families.items() if p.mandatory]
        notes.append(f"- {tier.value} → '{policy.label}': {', '.join(mandatory)}; {policy.discretionary_fetches} "
                     f"discretionary fetches, {policy.gemini_calls} Gemini calls, {policy.analyst_minutes} analyst "
                     f"minutes; saturation window {policy.saturation_window}. Vendors: "
                     f"{', '.join(s.name for s in by_tier[tier]) or 'none in this inventory'}.")
    notes += [f"- Modifier {m}: {mod.note}" for m, mod in cfg.modifiers.items()]
    notes.append("Steps reserved for Meridian, never done by us: " + "; ".join(
        cfg.tiers[Tier.CRITICAL].reserved_for_meridian) + ".")
    slide = content_slide(prs, ctx, "Depth: the tier decides how far we look",
                          "Critical vendors get every family and the largest budgets; below Critical, families thin "
                          "and the questionnaire takes over.", "OUTCOME 04", "\n".join(notes))
    step_h, step_gap = 1.08, 0.2
    for index, tier in enumerate((Tier.CRITICAL, Tier.HIGH, Tier.MEDIUM, Tier.LOW)):
        policy = cfg.tiers[tier]
        indent = index * 0.32
        x, y, w = LEFT + indent, 1.82 + index * (step_h + step_gap), 7.6 - indent
        solid, tint, ink = D.class_colours(tier.value)
        D.box(slide, x, y, w, step_h, None, fill=tint, line=None, radius=0.1, name=f"Depth step {tier.value}")
        D.box(slide, x + 0.12, y + 0.12, 1.05, 0.34, P(R(tier.value, size=11, bold=True, color=C.WHITE)),
              fill=solid, line=None, radius=0.17, align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0),
              name=f"Depth tier {tier.value}")
        hours = policy.analyst_minutes / 60
        budget = (f"{policy.discretionary_fetches} extra fetches · {policy.gemini_calls} Gemini calls · "
                  f"{hours:g} h analyst")
        D.label(slide, x + 1.3, y + 0.1, w - 1.4, 0.38, P(R(policy.label, size=12, bold=True, color=ink),
                                                          R("   " + budget, size=10, color=C.SLATE)),
                anchor=MSO_ANCHOR.MIDDLE, name=f"Depth label {tier.value}")
        fx = x + 0.12
        for family in SourceFamily:
            fp = policy.families[family]
            if not fp.mandatory:
                continue
            text = family.value + ("*" if family == SourceFamily.REG else "")
            D.pill(slide, fx, y + 0.6, 0.62, 0.3, text, fill=C.WHITE, color=ink, size=10, name=f"{tier.value} {text}")
            fx += 0.68
        vendors = by_tier[tier]
        if vendors:
            words = [s.name + (" (" + ", ".join(s.plan.modifiers) + ")" if s.plan.modifiers else "") for s in vendors]
            vendor_text = [R("Vendors: ", size=10, bold=True, color=C.INK), R(", ".join(words), size=10)]
        else:
            vendor_text = [R("No vendor in this inventory at this tier", size=10, italic=True, color=C.MUTED)]
        D.label(slide, fx + 0.1, y + 0.56, x + w - fx - 0.2, 0.42, P(*vendor_text), anchor=MSO_ANCHOR.MIDDLE,
                name=f"Depth vendors {tier.value}")
    rx, rw = 8.5, LEFT + WIDTH - 8.5
    critical, high = cfg.tiers[Tier.CRITICAL], cfg.tiers[Tier.HIGH]
    card(slide, rx, 1.82, rw, 1.62, "Stop when", bullets([
        "usage, sub-processor and data-use questions are answered or logged as gaps",
        f"no new S ≥ 2, R ≥ 2 cluster in the last {critical.saturation_window} actions (Critical) or "
        f"{high.saturation_window} (High)",
        "the discretionary budget is spent; 3 R1 items in a row stop a family",
    ], size=10, after=2), name="Stop rules card")
    card(slide, rx, 3.56, rw, 1.12, "Modifiers", [
        P(R(f"{m} ", size=10, bold=True, color=C.TEAL_DARK), R(mod.title, size=10), space_after=2)
        for m, mod in cfg.modifiers.items()], name="Modifiers card")
    manual = ctx.terms.get("manual", [])
    bare = [h for h in manual if h.count(".") == 1]
    hosts = and_list(bare or manual) + (f" (+{len(manual) - len(bare)} more)" if bare and len(manual) > len(bare)
                                        else "")
    card(slide, rx, 4.8, rw, 0.82, "Manual capture only", [
        P(R(f"{hosts}: analyst capture only (none this run); never sent to Gemini"
            if manual else
            "Hosts whose terms bar automation are captured by an analyst", size=10))], name="Manual card")
    card(slide, rx, 5.74, rw, 1.06, "Reserved for Meridian (non-OSINT)", [
        P(R("; ".join(critical.reserved_for_meridian), size=10))], fill=C.SAFFRON_TINT, name="Reserved card")


def s11_findings(prs: Any, ctx: Context) -> None:
    v = ctx.chain
    data = ctx.data
    if v and data:
        top = v.evidence[0] if v.evidence else None
        basis = f"a {top.strength.lower()}-strength {top.source_type} excerpt" if top else "the evidence"
        kicker = (f"{v.name}: {verdict_words(v, sep='(')} from {basis}; exposure {v.e}/3, decision impact {v.k}/3, "
                  f"tier {v.tp}, transparency gap {v.tg} → {v.arp} of 18 = {v.risk}"
                  + (" (provisional)." if v.provisional else "."))
    else:
        kicker = ("One excerpt to one risk class and action, shown on the fictional V-000 calibration until the frozen "
                  "run is injected.")
    notes = ["[1:15] Findings and risk."]
    slide = content_slide(prs, ctx, "Findings and risk: from excerpt to action", kicker, "FINDINGS", "")
    # the chain
    widths = (2.85, 1.85, 1.5, 1.6, 1.3, 2.13)  # the flip condition gets room for about four lines
    gap = (WIDTH - sum(widths)) / 5
    y, h = 1.85, 1.45
    if v and data:
        top = v.evidence[0] if v.evidence else None
        source = (clip(f"{top.source_type}, {top.publisher}, {top.published or 'dated by retrieval'}", 48)
                  if top else "")
        steps = [
            ("Excerpt · " + (top.evidence_id if top else "none"),
             [P(R(quote(top.excerpt, 120) if top else "No decisive excerpt", size=10, italic=True)),
              P(R(source, size=10, color=C.MUTED), space_before=2)]),
            ("Tags", [*[P(R(line, size=10, color=C.INK)) for line in (tag_lines(top.tags) if top else ["—"])],
                      P(R(f"→ {top.strength}" if top else "", size=10, bold=True, color=C.TEAL_DARK),
                        space_before=2)]),
            ("Verdict", [P(R(verdict_words(v), size=11, bold=True, color=C.INK)),
                         P(R(f"{v.likelihood}; confidence {v.confidence.lower()}", size=10))]),
            ("Risk inputs", [P(R(f"exposure {v.e}/3{' (assumed)' if v.e_assumed else ''}", size=10)),
                             P(R(f"decision impact {v.k}/3{' (assumed)' if v.k_assumed else ''}", size=10)),
                             P(R(f"tier {v.tp} · gap {v.tg}", size=10))]),
            ("Score", [P(R(f"{v.arp} of 18", size=16, bold=True, color=C.INK)),
                       P(R(v.risk + (" (provisional)" if v.provisional else ""), size=11, bold=True,
                           color=D.class_colours(v.risk)[2]))]),
            ("Flips if", [P(R(clip(v.flip, 110) or "no single answer moves the class", size=10)),
                          P(R(f"→ {flips_to(v.flip)}" if flips_to(v.flip) else "", size=10, bold=True,
                              color=D.class_colours(flips_to(v.flip))[2]), space_before=2)]),
        ]
        excerpt = (f"Excerpt {top.evidence_id} ({top.source_type}, {top.publisher}): \"{top.excerpt}\" Tags "
                   f"{top.tags} give {top.strength}." if top else "No decisive excerpt.")
        notes.append(f"Walk the chain for {v.vendor_id} {v.name}. {excerpt}"
                     f" Verdict {verdict_words(v, sep='(')}, {v.likelihood}; confidence {v.confidence} because "
                     f"{v.confidence_reason or 'see column T'}. Exposure {v.e}, decision impact {v.k}, tier points "
                     f"{v.tp}, transparency gap {v.tg}: ARP = 2E + 2K + TP + TG = {v.arp} of 18 → {v.risk}. "
                     f"Flip condition: {v.flip.rstrip('.') if v.flip else 'none'}.")
    else:
        steps = [
            ("Excerpt", [P(R("V-000's worked example cites a vendor executive's post describing an in-house AI "
                             "assistant used by servicing staff.", size=10, italic=True))]),
            ("Tags", [P(R("U1 · SR:B · RL:R3", size=10, color=C.INK)),
                      P(R("→ Strong", size=10, bold=True, color=C.TEAL_DARK), space_before=2)]),
            ("Verdict", [P(R("Yes · Confirmed", size=11, bold=True, color=C.INK)), P(R("no cap applies", size=10))]),
            ("Risk inputs", [P(R("exposure 3/3", size=10)), P(R("decision impact 3/3", size=10)),
                             P(R("tier 3 · gap 2", size=10))]),
            ("Score", [P(R("17 of 18", size=16, bold=True, color=C.INK)),
                       P(R("Critical", size=11, bold=True, color=D.class_colours("Critical")[2]))]),
            ("Matches", [P(R("the example's own Critical rating: our risk engine's calibration test", size=10))]),
        ]
        notes.append("Placeholder build: the chain shows the calibration of the fictional V-000 example, E3 + K3 + TP3 "
                     "+ TG2 = 17 of 18 = Critical with a Confirmed verdict and no cap. The frozen run replaces it with "
                     "the strongest real chain, normally BNY.")
    x = LEFT
    previous = None
    for (title, body), w in zip(steps, widths):
        shape = card(slide, x, y, w, h, title, body, fill=C.TEAL_TINT if title == "Score" else C.SURFACE,
                     title_color=C.TEAL_DARK, title_size=10.5, name=f"Chain {title}")
        if previous is not None:
            D.connect(slide, previous, shape, D.RIGHT, D.LEFT, width=1.5, name=f"chain to {title}")
        previous = shape
        x += w + gap
    # per-vendor verdicts
    y2 = 3.5
    if data:
        rows = []
        for vendor in data.vendors:
            e = vendor.evidence[0] if vendor.evidence else None
            evidence = ([R(f"{e.evidence_id} ", size=10, bold=True, color=C.TEAL_DARK),
                         R(f"{e.source_type}, {e.publisher}: ", size=10), R(quote(e.excerpt, 80), size=10, italic=True)]
                        if e else [R("No decisive excerpt; negative findings in the Coverage Log", size=10,
                                     italic=True, color=C.MUTED)])
            rows.append([
                {"paras": [P(R(vendor.name, size=10.5, bold=True, color=C.INK))]},
                {"paras": [P(R(verdict_words(vendor), size=10, bold=True, color=C.INK))]},
                {"paras": [P(*evidence)]},
                class_cell(vendor.risk, "provisional" if vendor.provisional else "", size=10.5),
            ])
            notes.append(f"- {vendor.vendor_id} {vendor.name}: {verdict_words(vendor, sep='(')}; risk {vendor.risk}"
                         f"; strongest excerpt {e.evidence_id + ' ' + e.source_type if e else 'none'}.")
        table(slide, LEFT, y2, (1.25, 1.55, 4.1, 1.05), ("Vendor", "Verdict", "Strongest evidence", "AI risk"), rows,
              row_h=0.5, size=10, name="Verdict table")
    elif ctx.scored:
        rows = [[{"paras": [P(R(s.name, size=10.5, bold=True, color=C.INK))]}, pending_cell(),
                 pending_cell("Decisive excerpt and Evidence Log ID from the frozen run"), pending_cell("Pending")]
                for s in ctx.scored]
        table(slide, LEFT, y2, (1.25, 1.55, 4.1, 1.05), ("Vendor", "Verdict", "Strongest evidence", "AI risk"), rows,
              row_h=0.5, size=10, name="Verdict table")
        notes.append("Per-vendor verdicts, strongest excerpts and classes are injected from the frozen run.")
    else:
        D.box(slide, LEFT, y2, 7.95, 1.2, P(R("Per-vendor verdicts are injected from the frozen run (build with "
                                               "--findings).", size=11, color=C.MUTED)),
              fill=C.SURFACE, line=None, radius=0.12, name="Verdict placeholder")
    # fourth-party concentration
    cx, cw = 8.75, LEFT + WIDTH - 8.75
    D.label(slide, cx, y2, cw, 0.3, P(R("Fourth-party AI concentration", size=12, bold=True, color=C.INK)),
            name="Concentration title")
    rows_c = concentration(ctx.vendors)
    if rows_c:
        top_n = rows_c[:6]
        most = max(sum(c[1:]) for c in top_n) or 1
        name_w = 1.1
        unit = min((cw - name_w - 0.1) / max(most, 1), 0.9)
        for index, row in enumerate(top_n):
            provider, named, third, dns = row
            ry = y2 + 0.42 + index * 0.4
            D.label(slide, cx, ry, name_w, 0.3, P(R(provider, size=10.5, bold=True, color=C.INK)),
                    anchor=MSO_ANCHOR.MIDDLE, name=f"Provider {provider}")
            bx = cx + name_w + 0.05
            for count, fill, line, ink, basis in ((named, C.TEAL, None, C.WHITE, "named"),
                                                  (third, C.SEAFOAM, None, C.INK, "third parties"),
                                                  (dns, C.WHITE, C.TEAL, C.TEAL_DARK, "DNS only")):
                if count:  # each segment is labelled with its vendor count
                    D.box(slide, bx, ry + 0.03, unit * count - 0.04, 0.24,
                          P(R(str(count), size=10, bold=True, color=ink)), fill=fill, line=line, radius=0.03,
                          align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0), name=f"Bar {provider} {basis}")
                    bx += unit * count
            notes.append("- " + concentration_note(row))
        D.label(slide, cx, y2 + 0.5 + len(top_n) * 0.4, cw, 0.5, [
            P(R("Vendors per provider:  ", size=10, bold=True, color=C.INK),
              R("■ ", size=10, color=C.TEAL), R("named by the vendor   ", size=10),
              R("■ ", size=10, color=C.SEAFOAM), R("third parties only", size=10)),
            P(R("□ ", size=10, color=C.TEAL), R("DNS verification token only: a relationship, not use", size=10))],
            name="Concentration legend")
    else:
        card(slide, cx, y2 + 0.4, cw, 2.2, "Generated from the frozen run", [
            P(R("Each AI provider is counted per vendor and split by basis:", size=10.5)),
            *bullets(["named by the vendor itself (SR A/B sources)", "named only by third parties",
                      "a DNS verification token only: a relationship, not use"], size=10.5, after=2),
            P(R("For example: “OpenAI: named by N vendors; DNS token only at M more.”", size=10.5, italic=True,
                color=C.MUTED), space_before=4),
        ], name="Concentration placeholder")
        notes.append("Fourth-party concentration is generated from the frozen run, split by basis.")
    _notes(slide, "\n".join(notes))


def concentration(vendors: Sequence[DeckVendor]) -> list[tuple[str, int, int, int]]:
    """(provider, vendors naming it, vendors where only third parties name it, vendors with a DNS token only)."""
    counts: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
    for v in vendors:
        for index, names in enumerate((v.providers_named, v.providers_third, v.providers_dns)):
            for name in names:
                counts[name][index] += 1
    return sorted(((p, *c) for p, c in counts.items()), key=lambda r: (-sum(r[1:]), -r[1], r[0].lower()))


def concentration_note(row: tuple[str, int, int, int]) -> str:
    """'OpenAI: named by 2 vendors; DNS verification token only at 1 more' (design §3 wording, split by basis)."""
    provider, *counts = row
    clauses: list[str] = []
    for label, count in zip(("named by", "named only by third parties at", "DNS verification token only at"), counts):
        if count:
            clauses.append(f"{label} {count} more" if clauses else f"{label} {plural(count, 'vendor', 'vendors')}")
    return f"{provider}: " + "; ".join(clauses)


DEMO_TIMES: tuple[str, ...] = ("0:15", "0:30", "0:25", "0:40", "0:30", "0:30", "0:10")
"""Slide 12's time slots, one per step (3:00 in all)."""


def possessive(name: str) -> str:
    return name + ("'" if name.endswith("s") else "'s")


def _ref(item: DeckEvidence) -> str:
    return f" {item.evidence_id}" if item.evidence_id else ""


def _demo_order(ctx: Context) -> list[DeckVendor]:
    """The vendors in the order the demo prefers them: the slide 11 chain vendor first, then the run's order."""
    chain = ctx.chain
    return ([chain] if chain else []) + [v for v in ctx.vendors if chain is None or v.vendor_id != chain.vendor_id]


def demo_criticality(ctx: Context) -> str:
    """Step 2, computed live from the workbook: the vendor a floor lifts above its score tier, and the vendor (not
    lifted) whose sensitivity line explains why it stays in its tier. Neutral wording when neither exists."""
    lifted = [s for s in ctx.scored if s.result.computed_tier.rank > s.result.score_tier.rank]
    lifted_ids = {s.profile.vendor_id for s in lifted}
    parts: list[str] = []
    if lifted:
        s = max(lifted, key=lambda x: x.result.score)
        floor = ctx.rubric.floor(_lifting_floor(s, ctx.rubric)).name
        parts.append(f"{s.name} scores {s.result.score} and the {floor} makes it {s.result.tier.value}")
    near = [s for s in ctx.scored if s.profile.vendor_id not in lifted_ids
            and any(line.startswith("Would become") for line in s.result.sensitivity)]
    if near:
        s = max(near, key=lambda x: x.result.score)
        parts.append(f"{possessive(s.name)} sensitivity line shows why it stays {s.result.tier.value}")
    if not parts:
        return ("Each vendor's score, factor levels and floors, with its sensitivity line; confirm the tier or "
                "override it with a reason.")
    return "; ".join(parts) + "."


def demo_evidence(ctx: Context) -> str:
    """Step 4 (HC2), from what the run holds: an excerpt with its screenshot when one exists (else in its source
    context), the definition-test trap the run logged, and Gemini against the rules only when the run has Gemini
    output. Nothing is promised that the run or the app cannot show."""
    order = _demo_order(ctx)
    if not order:
        return ("Open a decisive excerpt in its source context and Re-verify it against its capture; accept or reject "
                "it with a reason code.")
    shot = next((v for v in order if v.shot), None)
    lead = next((v for v in order if v.evidence), None)
    if shot is not None and shot.shot is not None:
        parts = [f"{possessive(shot.name)} excerpt{_ref(shot.shot)} with its screenshot, then Re-verify"]
    elif lead is not None:
        parts = [f"{possessive(lead.name)} excerpt{_ref(lead.evidence[0])} in its source context, then Re-verify"]
    else:
        parts = ["an excerpt in its source context, then Re-verify"]
    trap = next((v for v in order if v.trap), None)
    if trap is not None and trap.trap is not None:
        parts.append(f"reject {possessive(trap.name)} trap{_ref(trap.trap)}, which fails the definition test")
    else:
        parts.append("accept or reject an excerpt with a reason code")
    gemini = [v for v in order if v.gemini]
    if gemini:
        g = max(gemini, key=lambda v: v.llm_disagree)  # the first in demo order on a tie
        extra = f" ({plural(g.llm_disagree, 'disagreement', 'disagreements')})" if g.llm_disagree else ""
        parts.append(f"compare Gemini with the rules on {g.name}{extra}")
    return "; ".join(parts) + "."


def demo_export(ctx: Context) -> str:
    """Step 6: what the exported workbook shows; screenshots only when the run has some."""
    shots = sum(v.screenshots for v in ctx.vendors)
    images = f", with {plural(shots, 'screenshot', 'screenshots')} in Evidence Images" if shots else ""
    return (f"Build and open the workbook: V-000 untouched, Evidence Log and Coverage Log appended{images}; the "
            "integrity check passes.")


def demo_steps(ctx: Context) -> list[tuple[str, str, str]]:
    """Slide 12's script as (title, text, time), generated from the run and the live tiers so that a presenter
    following it never reaches for a feature, excerpt, screenshot, trap or Gemini label that is not there."""
    v = ctx.chain
    chain = (f"{v.name}: exposure {v.e}/3, decision impact {v.k}/3, tier {v.tp}, transparency gap {v.tg} → {v.arp} "
             f"of 18 = {v.risk}" if v and v.arp is not None else
             "The chain vendor's score in words (exposure, decision impact, tier, transparency gap → n of 18 = class)")
    steps = (
        ("Upload the workbook", "Header typos are tolerated; V-000 is recognised as the worked example and is never "
                                "written."),
        ("Criticality and HC1", demo_criticality(ctx)),
        ("Replay, then one live check", "Replay the frozen run offline, then the allowlisted live DNS lookup against "
                                        "the captured record (Evidence page)."),
        ("Evidence review (HC2)", demo_evidence(ctx)),
        ("Risk, then sign-off (HC2)", f"{chain}, with its flip condition; then Approve accepts its remaining cited "
                                      "excerpts."),
        ("Export and check", demo_export(ctx)),
        ("Spare time", "The V-000 worked example on Findings & Risk: a calibration check the export never writes."),
    )
    return [(title, text, time) for (title, text), time in zip(steps, DEMO_TIMES, strict=True)]


def s12_demo(prs: Any, ctx: Context) -> None:
    steps = demo_steps(ctx)
    notes = ["[3:00] Outcome 05: live demo. Review actions in steps 4 and 5 use a sandbox copy of review/, so the demo "
             "never changes the submitted decisions. HC1 and HC2 are the only human checkpoints; the export is gated "
             "by code (the LLM audit and the integrity check)."]
    notes += [f"{i}. ({time}) {title}: {text}" for i, (title, text, time) in enumerate(steps, start=1)]
    notes.append("Rebuild this deck from the final frozen run, then rehearse these steps in the app before presenting.")
    notes.append("If the app fails: Plan B is the same engine in local Jupyter; Plan C is the recorded run.")
    slide = content_slide(prs, ctx, "Live demo: upload, analyse, classify",
                          "Offline replay from the frozen evidence pack, under 60 seconds from upload to export, plus "
                          "one allowlisted live DNS lookup.", "OUTCOME 05", "\n".join(notes))
    top, step_h = 1.85, 0.68
    D.arrow(slide, LEFT + 0.25, top + 0.3, LEFT + 0.25, top + 6 * step_h + 0.3, color=C.LINE, width=2, tail=None,
            name="Timeline spine")
    for index, (title, text, time) in enumerate(steps):
        y = top + index * step_h
        D.badge(slide, LEFT + 0.25, y + 0.3, 0.44, str(index + 1), fill=C.TEAL, size=12, name=f"Demo step {index + 1}")
        D.label(slide, LEFT + 0.65, y + 0.02, 6.6, 0.62, [P(R(title, size=12, bold=True, color=C.INK)),
                                                         P(R(text, size=10.5 if len(text) <= 180 else 10))],
                name=f"Demo text {index + 1}")
        D.label(slide, 7.55, y + 0.05, 0.75, 0.32, P(R(time, size=12, bold=True, color=C.TEAL_DARK)),
                align=PP_ALIGN.RIGHT, name=f"Demo time {index + 1}")
    seconds = sum(int(t.split(":")[0]) * 60 + int(t.split(":")[1]) for _, _, t in steps)
    D.label(slide, 6.3, top + 7 * step_h + 0.02, 2.0, 0.3, P(R(f"total {seconds // 60}:{seconds % 60:02d}", size=11,
                                                                bold=True, color=C.INK)), align=PP_ALIGN.RIGHT,
            name="Demo total")
    rx, rw = 8.75, LEFT + WIDTH - 8.75
    card(slide, rx, 1.85, rw, 1.56, "What the brief asks the demo to show", bullets(
        ["upload Meridian's vendor list", "analyse every vendor from public evidence",
         "classify the AI security risk per vendor"], size=11.5, after=3), name="Brief card")
    card(slide, rx, 3.57, rw, 1.56, "Safety net", bullets(
        ["default: offline replay, no network or quota", "Plan B: the same engine in local Jupyter",
         "Plan C: a recorded run of these steps"], size=11.5, after=3), fill=C.SAFFRON_TINT, name="Safety card")
    labels = ("Gemini and rule labels side by side" if any(v.gemini for v in ctx.vendors) else
              "rules set every final tag; Gemini only proposes")
    card(slide, rx, 5.29, rw, 1.56, "Watch for", bullets(
        ["each excerpt re-verifies against its SHA-256 capture", labels,
         "V-000 and every provided cell untouched"], size=11.5, after=3), name="Watch card")


def action_group(v: DeckVendor) -> str:
    """The column U playbook behind a summary row: the same rule as footprint.actions.playbook_key (a test keeps the
    two equal), so slide 13 groups vendors exactly as their workbook actions do."""
    if v.risk == "None identified":
        return "none_high" if v.tier in ("Critical", "High") else "none_low"
    if v.provisional:
        return "provisional"
    if v.risk in ("Critical", "High"):
        return v.risk.lower()
    return "standard"


_DEADLINE = re.compile(r"\bwithin (\d+) business days\b")


def deadline_days(v: DeckVendor) -> int | None:
    """The questionnaire deadline (business days) that the vendor's column U action states, if any."""
    match = _DEADLINE.search(v.action or v.cells.get("recommended_action", ""))
    return int(match.group(1)) if match else None


def _who_by_deadline(vendors: Sequence[DeckVendor]) -> str:
    """'BNY and TCH within 15 business days', or 'BNY within 15 business days; Acme within 30 business days' when the
    deadlines differ (each vendor's own, from its column U action)."""
    groups: dict[int | None, list[str]] = {}
    for v in vendors:
        groups.setdefault(deadline_days(v), []).append(v.name)
    parts = [and_list(names) + (f" within {days} business days" if days else "") for days, names in
             sorted(groups.items(), key=lambda kv: (kv[0] is None, kv[0] or 0))]
    return "; ".join(parts)


ACTION_LINES: dict[str, str] = {
    "urgent": "Send a targeted AI questionnaire to {who}.",
    "standard": "Send an AI questionnaire to {who}; review AI clauses at renewal.",
    "provisional": "Put questions Q1–Q4 to {who}, then reclassify on the answers.",
    "none_high": "Ask {who} for a written no-AI attestation and an AI change-notification clause.",
    "none_low": "No AI-specific action for {who} beyond monitoring.",
}
"""Slide 13's first-column lines, one per playbook group (critical and high share the targeted questionnaire)."""


def thirty_day_actions(vendors: Sequence[DeckVendor]) -> list[str]:
    """Slide 13's 'Next 30 days': every vendor appears in exactly one line, grouped by its column U playbook and
    with its own questionnaire deadline, plus the sub-processor register line."""
    groups: dict[str, list[DeckVendor]] = defaultdict(list)
    for v in vendors:
        key = action_group(v)
        groups["urgent" if key in ("critical", "high") else key].append(v)
    lines: list[str] = []
    for key in ("urgent", "standard", "provisional"):
        if groups[key]:
            lines.append(ACTION_LINES[key].format(who=_who_by_deadline(groups[key])))
    if not groups["urgent"]:
        lines.insert(0, "No vendor needs a targeted AI questionnaire.")
    providers = sorted({p for x in vendors for p in x.providers_named}, key=str.lower)
    shown = providers if len(providers) <= 7 else [*providers[:6], f"{len(providers) - 6} more"]
    lines.append("Register AI sub-processors on Meridian's inventory" + (f" ({and_list(shown)})" if shown else "")
                 + ", flagging unresolved providers.")
    for key in ("none_high", "none_low"):
        if groups[key]:
            lines.append(ACTION_LINES[key].format(who=and_list([v.name for v in groups[key]])))
    return lines


def wrapped_lines(text: str, chars: int) -> int:
    """How many lines ``text`` takes when wrapped greedily at word boundaries to ``chars`` characters a line."""
    lines, used = 1, 0
    for word in text.split():
        need = len(word) if used == 0 else used + 1 + len(word)
        if need <= chars or used == 0:
            used = need
        else:
            lines, used = lines + 1, len(word)
    return lines


def fit_size(columns: Sequence[Sequence[str]], width: float, height: float, *,
             sizes: Sequence[float] = (11.5, 11, 10.5, 10), after: float = 6) -> float:
    """The largest font size at which every column of bullets fits its text box. An estimate: Calibri averages
    about 0.47 em a character; lines are 1.2 em apart; the bullet indent comes off the width."""
    for size in sizes:
        chars = max(1, int((width - 0.25) / (size / 72 * 0.47)))
        if all(sum(wrapped_lines(text, chars) for text in items) * size * 1.2 / 72 + len(items) * after / 72
               <= height for items in columns):
            return size
    return sizes[-1]


def s13_recommendations(prs: Any, ctx: Context) -> None:
    vendors = ctx.vendors
    if vendors:
        first = thirty_day_actions(vendors)
    else:
        first = ["Send targeted AI questionnaires within 15 business days to Critical- and High-risk vendors.",
                 "Put Q1–Q4 to provisional vendors, then reclassify on the answers.",
                 "Register AI sub-processors on Meridian's inventory, flagging unresolved providers."]
    later = ["Add AI clauses at the next contract review: training limits, sub-processor change notice, "
             "human-oversight thresholds, attestation rights.",
             "Escalate vendors that do not answer to the Third-Party Risk Committee.",
             "Re-scan annually or on material change; replay keeps runs comparable."]
    controls = ["Private repository; nothing published; no vendor contacted; GET requests only.",
                "Terms register, robots.txt and rate limits before every request; barred hosts "
                "logged as gaps, not crawled.",
                "Only public vendor text reaches Gemini, through the payload guard, with an audit log.",
                "Replay reproduces identical L–V cells from the SHA-256 evidence pack."]
    as_of = ctx.data.as_of if ctx.data and ctx.data.as_of else ctx.as_of
    limits = ["Absence of public evidence is not proof of no AI: a 'No (Not detected)' still needs an attestation.",
              "Public claims can overstate (marketing) or lag (unannounced tools).",
              "A DNS token proves a relationship, not processing of Meridian data.",
              "Contracts, internal tools and data flows are invisible; the questionnaire closes the gap.",
              f"A point-in-time snapshot, as of {long_date(as_of)}."]
    bottom_line = ("Every rating in the workbook traces to a captured, hashed source; whatever public evidence cannot "
                   "settle goes to Meridian's questionnaire, never to a guess.")
    notes = ["[0:30] Recommendations, controls and limits.", "Next 30 days, all actions by Meridian:", *first]
    if vendors:
        notes.append("Each vendor's first column U action (the cell holds the full playbook):")
        notes += [f"- {v.vendor_id} {v.name}: {v.action}" for v in vendors if v.action]
    notes += ["30 to 90 days:", *later, "Controls we practised:", *controls, "Limits of OSINT:", *limits,
              "Bottom line: " + bottom_line, "Close: thank you; questions."]
    slide = content_slide(prs, ctx, "Recommendations, controls and limits",
                          "What Meridian should do next, how we kept the work confidential and reproducible, and what "
                          "public evidence cannot tell us.", "NEXT STEPS", "\n".join(notes))
    columns = (("Next 30 days", first, C.TEAL_TINT, "1"), ("30–90 days", later, C.TEAL_TINT, "2"),
               ("Controls we practised", controls, C.SURFACE, "✓"), ("Limits of OSINT", limits, C.SAFFRON_TINT, "!"))
    w = (WIDTH - 3 * 0.2) / 4
    size = fit_size([items for _, items, _, _ in columns], w - 0.3, 3.35)
    for index, (title, items, fill, mark) in enumerate(columns):
        x = LEFT + index * (w + 0.2)
        D.box(slide, x, 1.85, w, 4.2, None, fill=fill, line=None, radius=0.12, name=f"Column {title}")
        D.badge(slide, x + 0.36, 2.2, 0.42, mark, fill=C.INK if index < 3 else "8A6A00", size=12,
                name=f"Column badge {title}")
        D.label(slide, x + 0.68, 2.0, w - 0.78, 0.42, P(R(title, size=13, bold=True, color=C.INK)),
                anchor=MSO_ANCHOR.MIDDLE, name=f"Column title {title}")
        D.label(slide, x + 0.16, 2.62, w - 0.3, 3.35, bullets(items, size=size, after=6),
                name=f"Column text {title}")
    D.box(slide, LEFT, 6.2, WIDTH, 0.62, P(R("Bottom line  ", size=12, bold=True, color=C.SAFFRON),
                                           R(bottom_line, size=12, color=C.WHITE)),
          fill=C.INK, line=None, radius=0.1, margins=(0.2, 0.04, 0.2, 0.04), name="Bottom line")


# =========================================================================== appendix


def a00_divider(prs: Any, ctx: Context) -> None:
    items = ["A1  Criticality rubric anchors", "A2  Tag legend", "A3  Verdict rules and wording",
             "A4  Risk anchors, gate, escalators and cap", "A5  Workbook fidelity and assumptions"]
    if ctx.vendors:
        items.append(f"A6–A{5 + len(ctx.vendors)}  One slide per vendor")
    slide = dark_slide(prs, ctx, "Appendix divider. Reference material for questions; not presented in the 13-minute "
                                 "walkthrough.")
    D.label(slide, LEFT, 1.6, 8.0, 1.0, P(R("Appendix", size=44, bold=True, color=C.WHITE)), name="Appendix title")
    D.label(slide, LEFT, 2.7, 8.0, 3.5, [P(R(t, size=16, color=C.SEAFOAM), space_after=6) for t in items],
            name="Appendix contents")


def a01_rubric(prs: Any, ctx: Context) -> None:
    rubric = ctx.rubric
    rows = []
    for code in ("O", "D", "P", "R", "V"):
        spec = rubric.factors[code]
        cells: list[Any] = [{"paras": [P(R(f"{code} · {spec.name}", size=10.5, bold=True, color=C.INK)),
                                       P(R(f"weight ×{rubric.weights[code]}", size=10, color=C.MUTED))]}]
        for level in (4, 3, 2, 1, 0):
            anchors = [a for a in spec.anchors if a.level == level]
            paras = []
            for anchor in anchors:
                label = anchor.label or spec.labels[level]
                paras.append(P(R(f"{anchor.id} {label}: ", size=10, bold=True, color=C.INK),
                               R(clip(anchor.description, 38 if len(anchors) > 1 else 72), size=10)))
            cells.append({"paras": paras or [P(R("—", size=10, color=C.MUTED))]})
        rows.append(cells)
    notes = (f"Appendix A1: rubric v{rubric.version} anchors (config/rubric.toml). The highest level whose anchor "
             "matches wins; when two adjacent levels fit, the higher is taken and logged. Each factor keeps its "
             "verbatim trigger phrase in the Criticality Workings sheet. Descriptions are shortened here; the sheet "
             "and the config file hold them in full.")
    slide = content_slide(prs, ctx, "A1 · Criticality rubric anchors",
                          "Five factors scored 0–4 from profile text; the trigger phrase behind every level is kept "
                          "for audit.", "APPENDIX", notes)
    table(slide, LEFT, 1.8, (1.83, 2.08, 2.08, 2.08, 2.08, 2.08), ("Factor", "4", "3", "2", "1", "0"), rows,
          row_h=0.72, size=10, name="Rubric anchors table")


TAG_LEGEND: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("U · signal class", ("U1 AI in the exact service", "U2 AI in the platform or an add-on",
                          "U3 AI in operations touching the service", "U4 named AI provider or relationship",
                          "U5 building AI capability", "U6 AI governance", "U7 marketing claim",
                          "U8 negative or limiting statement")),
    ("SR · source reliability", ("A legally accountable", "B first-party accountable",
                                 "C promotional or second-party", "D aggregators, syndicated copies",
                                 "E–F user-generated or unknown origin (excluded)")),
    ("SP · specificity", ("S3 (G1, G5, G6 or G7) with (G2 or G3)", "S2 G2, G4, G5, G8 or G11 tied to a named "
                                                                  "product, no M4",
                          "S1 marketing only, or legal or filing text alone", "S0 aspiration or commentary dominates")),
    ("RL · relevance", ("R3 exact-service term in the claim, its heading or title",
                        "R2 product family, company-wide delivery practice, DNS relationship",
                        "R1 other business line, inferred affiliate, platform supplier", "R0 commentary")),
    ("RC · recency (against as-of)", ("T3 12 months or less", "T2 24 months or less", "T1 36 months or less",
                                      "T0 older or undated")),
    ("IC · corroboration", ("1 corroborated by an independent K", "2 consistent, uncorroborated A/B",
                            "3 S0 or S1 only", "4 uncorroborated C/D at S2+ (doubtful)",
                            "5 contradicted by an A/B limiting statement", "6 cannot be judged")),
)


def a02_tags(prs: Any, ctx: Context) -> None:
    notes = ("Appendix A2: tag legend (design §2.3; IC 4 is our documented interpretation, the Admiralty 'doubtful'). "
             "Genuine-use indicators: G1 named model or provider, G2 named feature with observable behaviour, G3 data "
             "flow, retention, training or human review, G4 release note for a GA feature, G5 developer artifact, G6 "
             "legal artifact, G7 filing tying AI to products, G8 job posting with operational AI duties, G11 "
             "quantified outcome. Marketing indicators: M1 buzzword, M2 aspiration or capability wording, M3 industry "
             "commentary, M4 automation relabelled as AI, M5 superlative, M6 no product named, M7 logo-only "
             "partnership.")
    slide = content_slide(prs, ctx, "A2 · Tag legend", "Six tags per excerpt; the locus and the G/M indicators behind "
                                                       "SP are listed in the notes and the Method & Legend sheet.",
                          "APPENDIX", notes)
    w, h = (WIDTH - 2 * 0.2) / 3, 2.42
    for index, (title, lines) in enumerate(TAG_LEGEND):
        col, row = index % 3, index // 3
        card(slide, LEFT + col * (w + 0.2), 1.8 + row * (h + 0.2), w, h, title,
             [P(R(line.split(" ", 1)[0] + " ", size=11, bold=True, color=C.TEAL_DARK), R(line.split(" ", 1)[1],
                                                                                         size=11), space_after=1)
              for line in lines], name=f"Legend {title}")


def a03_verdict(prs: Any, ctx: Context) -> None:
    rules = (
        ("a", "Conflict: a qualifying signal (Q) is contradicted by an A/B limiting statement dated the same or later",
         "Inconclusive", "roughly even chance or unlikely"),
        ("b", "Confirmed: a Q at R3 plus an independent corroborating signal (K)", "Yes",
         "very likely; almost certain with 2+ independent SR A sources"),
        ("c", "Probable: any Q at R2 or above, or two independent K", "Yes", "likely"),
        ("d", "Affirmed negative: an A/B limiting statement covers the service, no Q or K", "No", "very unlikely"),
        ("e", "Not detected: every mandatory family complete and no S1+, R2+, T1+, U1–U4/U7 item", "No", "unlikely"),
        ("f", "Otherwise: marketing, capability, relationship or affiliate only, or incomplete coverage",
         "Inconclusive", "roughly even chance or unlikely"),
    )
    rows = [[{"paras": [P(R(r, size=12, bold=True, color=C.TEAL_DARK))], "align": PP_ALIGN.CENTER},
             cond, {"paras": [P(R(col, size=10.5, bold=True, color=C.INK))]}, like] for r, cond, col, like in rules]
    notes = ("Appendix A3: verdict rules a to f, first match wins (design §2.7). Likelihood follows ICD 203 and is "
             "written in its own sentence in column T; confidence is a separate sentence. Q = verified, SR A/B, "
             "SP S2+, RL R2+, RC T1+, class U1-U4, delivery locus. K = an independent cluster, SR A-C, SP S2+, RL R2+, "
             "RC T1+. Independent = different origin cluster and a different publisher or family.")
    slide = content_slide(prs, ctx, "A3 · Verdict rules and wording",
                          "Rules a–f run in order and the first match wins; likelihood and confidence are written as "
                          "separate sentences.", "APPENDIX", notes)
    table(slide, LEFT, 1.8, (0.55, 6.3, 1.45, 3.93), ("Rule", "Condition", "Column O", "ICD 203 likelihood"), rows,
          row_h=0.5, size=10.5, name="Verdict rules table")
    card(slide, LEFT, 5.35, WIDTH, 1.42, "Confidence (a separate sentence in column T)", bullets([
        "Low: a conflict, incomplete coverage of a mandatory family, or SR D–F sources only",
        "High: two or more consistent, independent A/B sources and complete coverage",
        "Moderate: one A/B source, two or more C sources, or complete coverage with nothing found",
    ], size=10.5, after=1), name="Confidence card")


def a04_risk(prs: Any, ctx: Context) -> None:
    notes = ("Appendix A4: risk anchors (design §2.8). Order is fixed: score, base class, materiality gate, "
             "escalator floors, verdict cap last. Assumed values (marked a) never satisfy the gate. Calibration: "
             "V-000 = E3 + K3 + TP3 + TG2 = 17, gate met, Confirmed so no cap → Critical.")
    slide = content_slide(prs, ctx, "A4 · Risk anchors, gate, escalators and cap",
                          "ARP = 2E + 2K + TP + TG (0–18): exposure and decision impact count double.", "APPENDIX",
                          notes)
    rows = [
        ["service feature or add-on", "D = 4, D3-P, or a named external model processes the data",
         "D3; or an R2 item (E3 only if confirmed)", "otherwise"],
        ["delivery operations", "the excerpt says customer data or credentials are processed",
         "AI works on tickets, cases, incidents, transactions, logs", "otherwise"],
        ["SDLC", "—", "production data is named", "otherwise"],
        ["corporate internal", "—", "—", "always E0"],
    ]
    table(slide, LEFT, 1.8, (1.9, 2.35, 2.35, 1.2), ("E: where the AI sits", "E3 (needs an R3 item)", "E2", "E1"),
          rows, row_h=0.5, size=10, name="Exposure table")
    card(slide, 8.6, 1.8, LEFT + WIDTH - 8.6, 2.36, "K, TP and TG", bullets([
        "K3 action without per-case review on customers, funds or production; K2 human-reviewed; K1 advisory",
        "TP: Critical 3, High 2, Medium 1, Low 0",
        "TG: six disclosure checks; missing 0–1 → 0, 2–3 → 1, 4–5 → 2, 6 → 3",
    ], size=11, after=3), name="K TP TG card")
    card(slide, LEFT, 4.4, 3.9, 2.4, "Classes and gate", bullets([
        "Critical 14–18 · High 10–13 · Medium 6–9 · Low 0–5",
        "High or above needs evidenced E ≥ 2 or K ≥ 2; otherwise the class is Medium",
        "Assumed inputs never satisfy the gate",
    ], size=11, after=3), name="Classes card")
    card(slide, 4.65, 4.4, 4.0, 2.4, "Escalators (High floor; Yes verdicts only)", bullets([
        "X1 provider named only by a third party, E ≥ 2", "X2 training or retention without opt-out, E ≥ 2",
        "X3 agentic AI on production without approval", "X4 K3 on credit, access or payment holds",
        "X5 AI incident in 24 months", "X6 single model provider behind a Critical vendor",
    ], size=10.5, after=1), name="Escalators card")
    card(slide, 8.85, 4.4, LEFT + WIDTH - 8.85, 2.4, "Verdict cap (always last)", bullets([
        "Probable: at most High", "Inconclusive: at most Medium, 'Provisional', with the ceiling if confirmed",
        "No: 'None identified'",
    ], size=11, after=3), fill=C.SAFFRON_TINT, name="Cap card")


def a05_fidelity(prs: Any, ctx: Context) -> None:
    notes = ("Appendix A5: workbook fidelity and the interpretations we documented (docs/contracts_p3.md §17). The "
             "writer refuses rather than bends a rule; the fidelity gate compares every provided cell, style, merge, "
             "column definition and sheet before export.")
    slide = content_slide(prs, ctx, "A5 · Workbook fidelity and assumptions",
                          "The writer touches only the cream student cells; everything else is proven unchanged.",
                          "APPENDIX", notes)
    card(slide, LEFT, 1.8, 6.0, 4.95, "Workbook rules", bullets([
        "Writes only cream cells L–V of vendor rows 6–11; never the V-000 example row or a provided cell",
        "Never reads or writes column widths for M–O or S–T: L–O and R–T are single column spans, and touching "
        "them makes Excel offer a repair",
        "Every value is forced to text, so nothing can become a formula; per-column length budgets apply",
        "Appends Evidence Log, Coverage Log, Criticality Workings, Method & Legend, Evidence Images and Run Info",
        "Fidelity gate: a semantic diff of every provided cell, style, merge, link and sheet, plus a manual open in "
        "Excel",
    ], size=12.5, after=8), name="Fidelity card")
    card(slide, 6.78, 1.8, LEFT + WIDTH - 6.78, 4.95, "Interpretations we documented", bullets([
        "IC 4 is undefined in the design: we use the Admiralty 'doubtful' (uncorroborated C/D)",
        "'M2/M3 dominate' (S0) means planned use, or aspiration without any G2, G3, G4, G8 or G11",
        "A corroborating signal must be class U1–U4: it corroborates use",
        "Definition-test traps are kept as never-citable items so reviewers can see them rejected",
        "Dates in P and V are DD-MM-YYYY, as in the V-000 example",
    ], size=12.5, after=8), fill=C.TEAL_TINT, name="Assumptions card")


def a06_vendor(prs: Any, ctx: Context, v: DeckVendor, number: int) -> None:
    ceiling = f", provisional; ceiling {v.ceiling} if confirmed" if v.provisional and v.ceiling else (
        ", provisional" if v.provisional else "")
    ids = ", ".join(e.evidence_id for e in v.evidence if e.evidence_id) or "none (negative findings only)"
    notes = [f"Appendix A{number}: {v.vendor_id} {v.name}. {v.tier} tier"
             + (f" (criticality score {v.score})" if v.score is not None else "") + ".",
             f"AI usage: {verdict_words(v, sep='(')}; it is {v.likelihood or 'not stated how likely it is'} that the "
             f"vendor uses AI in the service; confidence {(v.confidence or 'unrated').lower()}"
             + (f" because {v.confidence_reason}." if v.confidence_reason else "."),
             f"AI risk: {v.risk}{ceiling}; ARP {v.arp if v.arp is not None else 'n/a'} of 18 (E {v.e}, K {v.k}, "
             f"TP {v.tp}, TG {v.tg}). Flip: {v.flip or 'none recorded'}",
             f"Evidence Log IDs: {ids}." + (f" Trap kept in the log, never cited: {v.trap.evidence_id}."
                                            if v.trap and v.trap.evidence_id else ""),
             f"Next action: {v.action or 'see column U'}"]
    for key, label in (("evidence", "P"), ("how_ai_used", "Q"), ("ai_subprocessors", "R"), ("risk_rationale", "T"),
                       ("recommended_action", "U")):
        if v.cells.get(key):
            notes.append(f"Column {label}: {v.cells[key]}")
    slide = content_slide(prs, ctx, f"A{number} · {v.name} ({v.vendor_id})",
                          f"{v.tier} tier · {verdict_words(v, sep='(')} · AI risk {v.risk}{ceiling}", "APPENDIX",
                          "\n".join(notes))
    items = v.evidence[:3]
    body = []
    for e in items:
        body.append(P(R(f"{e.evidence_id} · {e.role or 'Cited'} · {e.strength}", size=10, bold=True,
                        color=C.TEAL_DARK), space_before=4))
        body.append(P(R(f"{e.source_type}, {e.publisher}, {e.published or 'dated by retrieval'}", size=10,
                        color=C.MUTED)))
        body.append(P(R(quote(e.excerpt, 230), size=10.5, italic=True)))
    if not body:
        body = [P(R("No decisive excerpt; see the Coverage Log for the negative findings.", size=10.5, italic=True,
                    color=C.MUTED))]
    # the workbook's own wording for columns Q and R fills the room the excerpts leave
    room = {0: 700, 1: 560, 2: 380}.get(len(items), 0)
    for key, heading in (("how_ai_used", "How AI is used (column Q)"), ("ai_subprocessors", "AI sub-processors "
                                                                                             "(column R)")):
        text = v.cells.get(key, "")
        if text and room >= 120:
            limit = min(room, 360 if key == "how_ai_used" else 200)
            body.append(P(R(heading, size=10, bold=True, color=C.INK), space_before=8))
            body.append(P(R(clip(text, limit), size=10.5)))
            room -= limit
    card(slide, LEFT, 1.8, 7.3, 3.2, "Decisive evidence", body, name="Vendor evidence card")
    reasoning = [P(R(f"It is {v.likelihood or 'not stated how likely it is'} that the vendor uses AI in the service. "
                     f"Confidence is {(v.confidence or 'unrated').lower()}"
                     + (f" because {v.confidence_reason}." if v.confidence_reason else "."), size=10.5))]
    providers = [("named by the vendor", v.providers_named), ("named only by third parties", v.providers_third),
                 ("DNS token only", v.providers_dns)]
    listed = "; ".join(f"{label}: {', '.join(names)}" for label, names in providers if names)
    reasoning.append(P(R("AI providers: " + (listed or "none named publicly"), size=10.5), space_before=4))
    if v.trap and v.trap.evidence_id:
        reasoning.append(P(R(f"Definition-test trap kept in the log, never cited: {v.trap.evidence_id}.", size=10.5,
                             color=C.MUTED), space_before=4))
    card(slide, LEFT, 5.15, 7.3, 1.65, "Why this verdict", reasoning, fill=C.TEAL_TINT, name="Vendor reasoning")
    inputs = [
        P(R(f"exposure {v.e}/3{' (assumed)' if v.e_assumed else ''} · decision impact {v.k}/3"
            f"{' (assumed)' if v.k_assumed else ''}", size=10.5)),
        P(R(f"tier points {v.tp} · transparency gap {v.tg} → {v.arp} of 18", size=10.5)),
        P(R(f"{v.risk}" + (" (provisional)" if v.provisional else ""), size=12, bold=True,
            color=D.class_colours(v.risk)[2]), space_before=3),
    ]
    card(slide, 8.1, 1.8, LEFT + WIDTH - 8.1, 1.75, "Risk", inputs, fill=D.class_colours(v.risk)[1], name="Vendor risk")
    card(slide, 8.1, 3.7, LEFT + WIDTH - 8.1, 1.35, "Flip condition", [P(R(clip(v.flip, 200) or "None recorded",
                                                                            size=10.5))], name="Vendor flip")
    actions = [P(R(clip(v.action, 160) or "See column U", size=10.5))]
    if v.questionnaire or v.gap_blocks:
        actions.append(P(R("Gap blocks " + (", ".join(v.gap_blocks) or "none") + " · questions " +
                           (", ".join(v.questionnaire) or "none"), size=10, color=C.MUTED), space_before=3))
    card(slide, 8.1, 5.2, LEFT + WIDTH - 8.1, 1.55, "Next action for Meridian", actions, fill=C.SAFFRON_TINT,
         name="Vendor action")


# =========================================================================== build


def _core_properties(prs: Any, ctx: Context) -> dt.datetime:
    when = dt.datetime.fromisoformat(ctx.as_of + "T09:00:00")
    props = prs.core_properties
    props.title = "Reading the public footprint"
    props.subject = "Vendor AI usage analysis: case study 1 walkthrough"
    props.author = ctx.team
    props.last_modified_by = ctx.team
    props.keywords = "vendor AI usage; OSINT; third-party risk; confidential"
    props.category = "Confidential"
    props.comments = "Confidential case-study material. Do not publish or distribute."
    props.revision = 1
    props.created = when
    props.modified = when
    return when


def _normalised_zip(data: bytes, when: dt.datetime) -> bytes:
    """Rewrite the package with fixed timestamps and attributes, so the same inputs give the same bytes."""
    stamp = (max(when.year, 1980), when.month, when.day, 0, 0, 0)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            item = zipfile.ZipInfo(info.filename, date_time=stamp)
            item.compress_type = zipfile.ZIP_DEFLATED
            item.create_system = 0
            item.external_attr = 0
            dst.writestr(item, src.read(info.filename))
    return out.getvalue()


def build_deck(out: str | Path = DEFAULT_OUT, *, findings: str | Path | DeckData | AssessmentResult | None = None,
               workbook: str | Path | None = DEFAULT_WORKBOOK, config_dir: str | Path = ROOT / "config",
               seeds_dir: str | Path = ROOT / "seeds", team: str | None = None, as_of: str | None = None,
               chain_vendor: str | None = None, appendix: bool = True) -> dict[str, Any]:
    """Build the deck and write it to ``out`` (a local .pptx; never uploaded). Returns a summary dict:
    {path, slides, main_slides, appendix_slides, findings, run_id, sha256}."""
    out = Path(out)
    if out.suffix.lower() != ".pptx":
        raise ValueError(f"the output must be a .pptx file, not {out.name}")
    config_dir, seeds_dir = Path(config_dir), Path(seeds_dir)
    if isinstance(findings, AssessmentResult):
        data: DeckData | None = summarize_assessment(findings, seeds_dir=seeds_dir)
    elif isinstance(findings, DeckData) or findings is None:
        data = findings
    else:
        data = load_findings(findings, seeds_dir=seeds_dir)
    rubric_path = config_dir / "rubric.toml"
    depth_path = config_dir / "depth.toml"
    rubric = load_rubric(rubric_path if rubric_path.is_file() else None)
    depth = load_depth_config(depth_path if depth_path.is_file() else None)
    when_text = as_of or (data.as_of if data and data.as_of else dt.date.today().isoformat())
    dt.date.fromisoformat(when_text)
    ctx = Context(team=team or os.environ.get("FOOTPRINT_TEAM_NAME") or DEFAULT_TEAM, as_of=when_text, data=data,
                  rubric=rubric, depth=depth)
    if workbook:
        ctx.scored, ctx.example = score_workbook(Path(workbook), rubric=rubric, depth=depth, seeds_dir=seeds_dir)
        ctx.workbook_name = Path(workbook).name
    ctx.sr_ladder, ctx.sr_source = load_sr_ladder(config_dir)
    ctx.terms = load_terms_register(config_dir)
    ctx.chain = pick_chain(data, chain_vendor)

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _apply_theme(prs)
    for build in (s01_title, s02_answer, s03_approach, s04_functional, s05_architecture, s06_sources, s07_weight,
                  s08_ai, s09_criticality, s10_depth, s11_findings, s12_demo, s13_recommendations):
        build(prs, ctx)
    if appendix:
        for build in (a00_divider, a01_rubric, a02_tags, a03_verdict, a04_risk, a05_fidelity):
            build(prs, ctx)
        for index, vendor in enumerate(ctx.vendors, start=6):
            a06_vendor(prs, ctx, vendor, index)
    when = _core_properties(prs, ctx)
    buffer = io.BytesIO()
    prs.save(buffer)
    payload = _normalised_zip(buffer.getvalue(), when)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(payload)
    return {
        "path": str(out), "slides": len(prs.slides), "main_slides": MAIN_SLIDES,
        "appendix_slides": len(prs.slides) - MAIN_SLIDES, "findings": data is not None,
        "run_id": data.run_id if data else "", "sha256": hashlib.sha256(payload).hexdigest(),
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the Team Osprey walkthrough deck (local .pptx only).")
    parser.add_argument("--findings", type=Path,
                        help="AssessmentResult JSON, a runs/<run_id>/ folder, or a deck summary JSON (optional)")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output .pptx path")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help="input workbook for the tier table")
    parser.add_argument("--config-dir", type=Path, default=ROOT / "config")
    parser.add_argument("--seeds-dir", type=Path, default=ROOT / "seeds")
    parser.add_argument("--team", help=f"team name (default: FOOTPRINT_TEAM_NAME or {DEFAULT_TEAM!r})")
    parser.add_argument("--as-of", help="date shown on the title slide, YYYY-MM-DD (default: the run's as_of)")
    parser.add_argument("--chain-vendor", help="vendor ID for the slide 11 chain (default: strongest Confirmed)")
    parser.add_argument("--no-appendix", action="store_true", help="build the 13 main slides only")
    args = parser.parse_args(argv)
    if args.findings is not None and not findings_file(args.findings).is_file():
        print(f"error: findings file not found: {findings_file(args.findings)}", file=sys.stderr)
        return 2
    try:
        info = build_deck(args.out, findings=args.findings, workbook=args.workbook, config_dir=args.config_dir,
                          seeds_dir=args.seeds_dir, team=args.team, as_of=args.as_of,
                          chain_vendor=args.chain_vendor, appendix=not args.no_appendix)
    except (ValueError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    source = f"findings {info['run_id'] or 'summary'}" if info["findings"] else "no findings (placeholders on 2 and 11)"
    print(f"Wrote {info['path']}: {info['slides']} slides ({info['main_slides']} main + {info['appendix_slides']} "
          f"appendix); {source}; sha256 {info['sha256'][:12]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
